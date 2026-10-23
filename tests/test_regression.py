"""Mandatory regression tests TEST 01–20 (directive §13). Names carry the test number for the results register."""
import copy
import random
import shutil
import time
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import C, CC, O, RE, history, new_obs
from cdf_gco_crms.analytics.dashboard_analytics import Analytics
from cdf_gco_crms.config import design_tokens as T
from cdf_gco_crms.config import round_schema as S
from cdf_gco_crms.config import workbook_settings as W
from cdf_gco_crms.excel.workbook_builder import build_workbook
from cdf_gco_crms.ingestion.workbook_reader import parse_workbook
from cdf_gco_crms.model import lifecycle as L
from cdf_gco_crms.model.entities import DashboardViewState, RoundAssessment
from cdf_gco_crms.qa import inspect as I
from cdf_gco_crms.qa import lo_engine
from cdf_gco_crms.qa.recalc_verify import check_view
from cdf_gco_crms.reconciliation.historical import compare, merge_returned, snapshot, view_sheet_edits
from cdf_gco_crms.rounds.round_expansion import RoundInput, apply_round
from cdf_gco_crms.rounds.round_parser import SchemaError, discover

needs_lo = pytest.mark.skipif(not lo_engine.available() or not lo_engine.sockets_ok(),
                              reason="LibreOffice not available")


def cached(path):
    return load_workbook(path, data_only=True)


def kpis(path):
    d = cached(path)[W.DASHBOARD_SHEET]
    return [d.cell(row=14, column=c).value for c in (2, 4, 6, 8, 10, 12)]


# ------------------------------------------------------------------------ TEST 01
def test_01_round1_ten_observations(tmp_path):
    h = history([[O]] * 10, rounds=1, outdir=tmp_path)
    p = h.files[1]
    pw = parse_workbook(p)
    ws = pw.data_ws
    assert ws.title == "ملاحظات الجولة الأولى"
    assert pw.schema.blocks == {1: {S.ROLE_CLOSURE: 6, S.ROLE_RESPONSE: 7, S.ROLE_IMPL: 8}}
    assert [o.reader_no for o in pw.state.ordered()] == [f"{i:02d}" for i in range(1, 11)]
    for r in range(2, 12):
        assert ws.cell(row=r, column=6).value == S.OPEN
        assert ws.cell(row=r, column=7).value == S.PH_TEXT
        assert ws.cell(row=r, column=8).value == S.PH_STATUS
        assert ws.cell(row=r, column=6).fill.fgColor.rgb[-6:] == T.ROUND1_OPEN_FILL
    assert kpis(p) == [10, 0, 0, 0, 10, 0]
    d = cached(p)[W.DASHBOARD_SHEET]
    assert "بانتظار الإفادة الأولى" in d["B18"].value
    hs = cached(p)[W.HELPER_SHEET]
    hdr = {hs.cell(row=1, column=c).value: c for c in range(1, hs.max_column + 1)}
    reasons = {hs.cell(row=r, column=hdr["obs.reason"]).value for r in range(2, 12)}
    assert reasons == {S.REASON_AWAIT_FIRST}            # not "overdue" or "failed" (TEST A)


# ------------------------------------------------------------------------ TEST 02
def test_02_round2_mixed_states(tmp_path):
    traj = [[O, C]] * 3 + [[O, CC]] * 2 + [[O, RE]] * 3 + [[O, O]] * 2
    h = history(traj, rounds=2, outdir=tmp_path)
    p = h.files[2]
    assert kpis(p)[:5] == [10, 3, 2, 3, 2]
    assert abs(kpis(p)[5] - 0.3) < 1e-9                  # conditional not counted as closed
    pw = parse_workbook(p)
    b = pw.schema.blocks[2]
    for o in pw.state.ordered():
        row = pw.rows[o.internal_id]
        if o.value(2) == C:
            assert pw.data_ws.cell(row=row, column=b[S.ROLE_IMPL]).value == S.IMPL_PRIOR_CLOSED
            assert pw.data_ws.cell(row=row, column=b[S.ROLE_RESPONSE]).value == S.closed_guidance("الثانية")
        else:
            assert pw.data_ws.cell(row=row, column=b[S.ROLE_RESPONSE]).value == S.PH_TEXT
            assert pw.data_ws.cell(row=row, column=b[S.ROLE_IMPL]).value == S.PH_STATUS
    an = Analytics(pw.state, pw.rows, round_cols=pw.schema.blocks, data_sheet=pw.data_ws.title)
    labels = {a.reader: a.long[2].change for a in an.obs}
    assert labels["01"] == S.CHANGE_CLOSED and labels["04"] == S.CHANGE_COND and labels["09"] == S.CHANGE_NO_PROGRESS


# ------------------------------------------------------------------------ TEST 03
def test_03_round3_new_and_reopened(tmp_path):
    traj = [[O, C, RE], [O, RE, C], [O, O, C], [O, C, C], [O, RE, RE]]
    h = history(traj, {3: [(new_obs(1), [O])]}, rounds=3, outdir=tmp_path)
    pw = parse_workbook(h.files[3])
    st = pw.state
    new = [o for o in st.observations if o.birth_round == 3]
    assert len(new) == 1 and new[0].internal_id == "GCO-TST-R3-001" and new[0].reader_no == "06"
    for r in (1, 2):
        c = pw.data_ws.cell(row=pw.rows[new[0].internal_id], column=pw.schema.col(r, S.ROLE_CLOSURE))
        assert c.value == S.NOT_RAISED and c.fill.fgColor.rgb[-6:] == T.CLOSED_ROW_GRAY
    an = Analytics(st, pw.rows, round_cols=pw.schema.blocks, data_sheet=pw.data_ws.title)
    t = an.trend[3]
    assert (t.issued, t.closed, t.reopen_event, t.closed_event, t.new) == (6, 3, 1, 2, 1)
    assert an.trend[2].issued == 5                        # new row not in the round-2 denominator


# ------------------------------------------------------------------------ TEST 04
def test_04_five_round_history(five_round):
    h, traj, out = five_round
    st = h.states[5]
    an = Analytics(st, {o.internal_id: 2 + i for i, o in enumerate(st.ordered())},
                   round_cols={}, data_sheet="x")
    o2 = next(a for a in an.obs if a.reader == "02")
    assert [o2.long[r].value for r in range(1, 6)] == [O, C, C, RE, C]
    assert [o2.long[r].change for r in range(1, 6)] == [S.CHANGE_NEW, S.CHANGE_CLOSED, S.CHANGE_STAYED_CLOSED,
                                                       S.CHANGE_REOPENED, S.CHANGE_CLOSED]
    assert o2.long[5].lcr == 5 and o2.long[3].lcr == 2
    o5 = next(a for a in an.obs if a.reader == "05")
    assert [o5.long[r].stalled for r in range(1, 6)] == [0, 0, 1, 1, 1]
    assert an.trend[5].issued == 12 and an.trend[4].issued == 12 and an.trend[3].issued == 11
    for r in range(1, 6):
        t = an.trend[r]
        assert t.closed + t.cond + t.req + t.open == t.issued


# ------------------------------------------------------------------------ TEST 05
def test_05_ten_rounds(tmp_path):
    rng = random.Random(5)
    from cdf_gco_crms.regression.synthetic import random_trajectories
    traj = random_trajectories(15, 10, rng)
    h = history(traj, rounds=10, outdir=tmp_path, build_rounds={10})
    pw = parse_workbook(h.files[10])
    assert pw.schema.rounds == list(range(1, 11))
    assert pw.schema.last_col == 8 + 4 * 9
    assert pw.data_ws.title == "الملاحظات حتى الجولة العاشرة"
    assert pw.data_ws.cell(row=1, column=pw.schema.col(10, S.ROLE_CLOSURE)).value == "حالة الإغلاق – الجولة العاشرة"


# ------------------------------------------------------------------------ TEST 06
def test_06_twenty_rounds_scalability(tmp_path):
    rng = random.Random(6)
    from cdf_gco_crms.regression.synthetic import random_trajectories
    traj = random_trajectories(40, 20, rng)
    t0 = time.time()
    h = history(traj, rounds=20, outdir=tmp_path, build_rounds={20})
    elapsed = time.time() - t0
    pw = parse_workbook(h.files[20])
    assert pw.schema.rounds == list(range(1, 21))
    assert pw.data_ws.title == "الملاحظات حتى الجولة 20"
    assert pw.schema.last_col == 8 + 4 * 19
    hs = cached(h.files[20])[W.HELPER_SHEET]
    hdr = {hs.cell(row=1, column=c).value: c for c in range(1, hs.max_column + 1)}
    assert [hs.cell(row=r + 1, column=hdr["trend.label"]).value for r in (11, 20)] == ["الجولة 11", "الجولة 20"]
    assert elapsed < 120


# ------------------------------------------------------------------------ TEST 07
def test_07_closed_row_preservation(five_round):
    h, traj, out = five_round
    for r in range(2, 6):
        before, after = snapshot(h.states[r - 1]), snapshot(h.states[r])
        assert compare(before, after, rounds=set(range(1, r)), owner_rounds=set(range(1, r - 1))) == []
    st = h.states[5]
    o1 = next(o for o in st.observations if o.reader_no == "01")
    for r in (3, 4, 5):
        assert o1.value(r) == C
        assert str(o1.cells[r].feedback) == S.carried_closed_feedback("الثانية")
        assert o1.cells[r].impl == S.IMPL_PRIOR_CLOSED


# ------------------------------------------------------------------------ TEST 08
def test_08_management_response_integrity(five_round, tmp_path):
    h, traj, out = five_round
    p4 = h.files[4]
    returned = parse_workbook(p4).state
    # the owner answers round 4 and, against the rules, edits an earlier response and a GCO cell
    victim = next(o for o in returned.observations if o.value(4) not in ("", C))
    victim.cells[4].response = "أُضيف البند (5-3) بعد التعديل."
    victim.cells[4].impl = S.IMPL_DONE
    late = next(o for o in returned.observations if o.cells.get(2) and o.value(1) == O)
    late.cells[1].response = "تعديل لاحق على إفادة الجولة الأولى"
    tampered = next(o for o in returned.observations if o.value(4) == C)
    tampered.cells[4].feedback = "نص عدّلته الإدارة"
    delivered = parse_workbook(p4).state
    merged, rep = merge_returned(delivered, returned)
    assert len(rep.gco_cell_edits) == 1 and len(rep.late_owner_edits) == 1 and rep.owner_inputs >= 1
    m = merged.by_id()
    assert str(m[tampered.internal_id].cells[4].feedback) != "نص عدّلته الإدارة"        # GCO cell restored
    assert m[late.internal_id].cells[1].response == "تعديل لاحق على إفادة الجولة الأولى"  # owner text kept
    t = traj  # next round
    assessments = [RoundAssessment(o.internal_id, 5, (C if o.value(4) == C else RE), "تحقق في النص.")
                   for o in merged.ordered() if o.value(4) != C]
    st5, _ = apply_round(merged, RoundInput(5, None, None, assessments))
    p5 = build_workbook(st5, tmp_path / "r5.xlsx").path
    after = parse_workbook(p5).state
    keys = {k for k in snapshot(merged) if k[1].endswith((".response", ".impl")) and not k[1].startswith("R5")}
    sa, sb = snapshot(merged), snapshot(after)
    assert all(sa[k][0] == sb[k][0] for k in keys)
    # text typed on the dashboard is found (Channel X, SC-61); a changed selector is not input (SC-62)
    ret = tmp_path / "returned_r4.xlsx"
    wb = load_workbook(p4)
    dws = wb[W.DASHBOARD_SHEET]
    free = next(c for row in dws.iter_rows(min_row=40) for c in row if c.value in (None, "")
                and c.protection.locked is not False and c.coordinate not in dws.merged_cells)
    free.value = "ملاحظة كتبتها الإدارة على اللوحة"
    sel = next(c for row in dws.iter_rows() for c in row if c.protection.locked is False)
    sel.value = "الجولة الثانية"
    wb.save(ret)
    edits, missing = view_sheet_edits(p4, ret)
    assert [e.coordinate for e in edits] == [free.coordinate] and not missing


# ------------------------------------------------------------------------ TEST 09
def test_09_duplicate_round_prevention(five_round, tmp_path):
    h, traj, out = five_round
    st4 = h.states[4]
    st5 = h.states[5]
    rin = RoundInput(5, None, None, [RoundAssessment(o.internal_id, 5, o.value(5), o.cells[5].feedback)
                                     for o in st5.ordered() if o.birth_round <= 5 and o.value(4) != C
                                     and o.birth_round < 5])
    again, rep = apply_round(st5, rin)
    assert rep.action == "NO_OP" and again is st5
    bad = copy.deepcopy(rin)
    bad.assessments[0].status = S.OPEN if bad.assessments[0].status != S.OPEN else S.CLOSED
    with pytest.raises(SchemaError):
        apply_round(st5, bad)
    with pytest.raises(SchemaError):
        apply_round(st4, RoundInput(6, None, None, []))       # gap
    # a sheet carrying a duplicated block is rejected by discovery
    wb = load_workbook(h.files[5])
    ws = wb[S.data_sheet_name(5)]
    pw = parse_workbook(h.files[5])
    for role, col in pw.schema.blocks[4].items():
        ws.cell(row=1, column=pw.schema.last_col + 1 + list(pw.schema.blocks[4]).index(role)).value = \
            ws.cell(row=1, column=col).value
    with pytest.raises(SchemaError):
        discover(ws)
    # rebuilding the same round twice gives identical data sheets
    a = build_workbook(st5, tmp_path / "a.xlsx").path
    b = build_workbook(st5, tmp_path / "b.xlsx").path
    assert snapshot(parse_workbook(a).state) == snapshot(parse_workbook(b).state)
    assert parse_workbook(a).schema.blocks == parse_workbook(b).schema.blocks


# ------------------------------------------------------------------------ TEST 10
@needs_lo
def test_10_formula_reconciliation(five_round, tmp_path):
    h, traj, out = five_round
    st = h.states[5]
    views = [DashboardViewState(5, S.SCOPE_ALL, "#03", 4), DashboardViewState(3, S.SCOPE_ACTION, "#12", 1),
             DashboardViewState(5, S.SCOPE_REOPENED, "#02", 2)]
    for vs in views:
        vc = check_view(st, vs, tmp_path)
        assert vc.ok, (vc.view, vc.mismatches[:5], vc.errors[:5])


# ------------------------------------------------------------------------ TEST 11
@needs_lo
def test_11_navigation_after_owner_sort(five_round, tmp_path):
    h, traj, out = five_round
    p = h.files[5]
    pw = parse_workbook(p)
    # every cached link resolves to the row holding its Internal_ID
    hs = cached(p)[W.HELPER_SHEET]
    hdr = {hs.cell(row=1, column=c).value: c for c in range(1, hs.max_column + 1)}
    for k in range(1, len(pw.state.observations) + 1):
        sheet, coord = I.hyperlink_target(hs.cell(row=k + 1, column=hdr["obs.link"]).value)
        row = int("".join(ch for ch in coord if ch.isdigit()))
        assert pw.data_ws.cell(row=row, column=2).value == hs.cell(row=k + 1, column=hdr["obs.id"]).value
    # the owner sorts the data sheet in reverse order; after recalculation the links follow the IDs (MATCH)
    q = tmp_path / "sorted.xlsx"
    shutil.copy(p, q)
    wb = load_workbook(q, rich_text=True)
    ws = wb[pw.data_ws.title]
    rows = [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)] for r in range(2, ws.max_row + 1)]
    for i, vals in enumerate(reversed(rows)):
        for c, v in enumerate(vals, start=1):
            ws.cell(row=2 + i, column=c).value = v
    wb.save(q)
    assert lo_engine.recalc(q).ok
    v = load_workbook(q, data_only=True)
    hs2, ws2 = v[W.HELPER_SHEET], v[pw.data_ws.title]
    for k in range(1, len(pw.state.observations) + 1):
        sheet, coord = I.hyperlink_target(hs2.cell(row=k + 1, column=hdr["obs.link"]).value)
        row = int("".join(ch for ch in coord if ch.isdigit()))
        assert ws2.cell(row=row, column=2).value == hs2.cell(row=k + 1, column=hdr["obs.id"]).value


# ------------------------------------------------------------------------ TEST 12
def test_12_word_excel_sync(five_round, tmp_path):
    from cdf_gco_crms.word.word_sync import build_review_docx, reconcile
    h, traj, out = five_round
    for r in (1, 3, 5):
        d = build_review_docx(h.states[r], tmp_path / f"r{r}.docx")
        rep = reconcile(d, h.files[r])
        assert rep.ok, rep.problems
    victim = next(o for o in h.states[5].ordered() if o.value(4) not in ("", C))
    d = build_review_docx(h.states[5], tmp_path / "broken.docx", break_ids={victim.internal_id})
    assert not reconcile(d, h.files[5]).ok


def test_12b_word_reply_threads(five_round, tmp_path):
    """G-SYNC-9 (v7.12): round-n follow-ups are native Word reply threads (w15 commentsExtended)."""
    from lxml import etree
    from cdf_gco_crms.word import word_sync as WS
    h, traj, out = five_round
    d = WS.build_review_docx(h.states[3], tmp_path / "r3.docx")
    comments = WS.extract_comments(d)
    by_cid = {c.cid: c for c in comments}
    follow = [c for c in comments if c.follow_up_round == 3]
    expected = {o.internal_id for o in h.states[3].ordered() if o.birth_round < 3 and (o.value(2) != C or o.value(3) != C)}
    assert follow and {c.internal_id for c in follow} == expected
    for c in follow:                                   # every round-3 follow-up replies to its round-1 comment
        parent = by_cid[c.parent_cid]
        assert parent.internal_id == c.internal_id and parent.follow_up_round is None
        assert parent.reader_no == c.reader_no and parent.parent_cid is None
        assert c.done == parent.done == (c.status == C)
    for c in comments:                                 # birth comments of round 3 are top level
        if c.follow_up_round is None:
            assert c.parent_cid is None
    rep = WS.reconcile(d, h.files[3])
    assert rep.ok, rep.problems
    assert rep.checks["G-SYNC-9 follow-ups threaded"] is True
    assert WS.reconcile(WS.build_review_docx(h.states[1], tmp_path / "r1.docx"), h.files[1]).checks[
        "G-SYNC-9 follow-ups threaded"] is True
    with zipfile.ZipFile(d) as z:                      # package structure
        names = z.namelist()
        assert WS.COMMENTS_EX_PART in names and WS.COMMENTS_IDS_PART in names
        croot = etree.fromstring(z.read("word/comments.xml"))
        ct = z.read("[Content_Types].xml").decode()
        rels = z.read("word/_rels/document.xml.rels").decode()
        ex = etree.fromstring(z.read(WS.COMMENTS_EX_PART))
    pids = [c.find(f"{{{WS.WNS}}}p").get(f"{{{WS.W14NS}}}paraId") for c in croot.findall(f"{{{WS.WNS}}}comment")]
    assert len(pids) == len(comments) == len(set(pids)) and all(len(x) == 8 and int(x, 16) < 0x80000000 for x in pids)
    assert "w14" in croot.get(f"{{{WS.MCNS}}}Ignorable", "").split() and croot.nsmap["w14"] == WS.W14NS
    assert f'PartName="/{WS.COMMENTS_EX_PART}" ContentType="{WS.COMMENTS_EX_CT}"' in ct
    assert f'Type="{WS.COMMENTS_EX_REL}" Target="commentsExtended.xml"' in rels
    parents = {e.get(f"{{{WS.W15NS}}}paraId"): e.get(f"{{{WS.W15NS}}}paraIdParent") for e in ex}
    assert set(parents) == set(pids) and all(p in pids for p in parents.values() if p)
    assert sum(1 for p in parents.values() if p) == sum(1 for c in comments if c.follow_up_round)
    if lo_engine.available() and lo_engine.sockets_ok():  # the threaded file opens in LibreOffice
        pdf = lo_engine.to_pdf(d, tmp_path / "pdf")
        assert pdf is not None and pdf.stat().st_size > 0


# ------------------------------------------------------------------------ TEST 13
def test_13_v710_migration(five_round, tmp_path):
    from cdf_gco_crms.migration.legacy_v710 import write_v710
    from cdf_gco_crms.migration.migrate import migrate
    h, traj, out = five_round
    src = write_v710(h.states[4], tmp_path / "legacy" / "doc.xlsx")
    sha = __import__("hashlib").sha256(src.read_bytes()).hexdigest()
    assert W.LEGACY_PRINT_VIEW_SHEET in load_workbook(src).sheetnames
    rep = migrate(src, tmp_path / "out")
    assert rep.status == "PASSED", rep.to_json()
    assert __import__("hashlib").sha256(src.read_bytes()).hexdigest() == sha        # source untouched
    assert Path(rep.backup).exists() and rep.cells_compared > 0 and not rep.cell_differences
    assert rep.removed_sheets == [W.LEGACY_PRINT_VIEW_SHEET] and rep.no_round_added and rep.ids_preserved
    with pytest.raises(Exception):
        migrate(src, src.parent)                                                     # never in place


# ------------------------------------------------------------------------ TEST 14
def test_14_metadata(five_round, tmp_path):
    from cdf_gco_crms.word.word_sync import build_review_docx
    h, traj, out = five_round
    for r, p in h.files.items():
        pkg = I.package(p)
        assert pkg.core["creator"] == W.AUTHOR and pkg.core["lastModifiedBy"] == W.AUTHOR
        assert pkg.app.get("Company") == W.COMPANY
        assert pkg.core["title"] == W.workbook_title(h.states[r].document.name, r)
        props = " ".join(v for k, v in I.all_text_parts(p).items() if k.startswith("docProps/")).lower()
        assert not any(x in props for x in W.PROHIBITED_PROPERTY_VALUES)
        assert p.stem == W.file_stem(h.states[r].document.name, r)
    d = build_review_docx(h.states[5], tmp_path / "d.docx")
    from docx import Document
    cp = Document(d).core_properties
    assert cp.author == W.AUTHOR and cp.last_modified_by == W.AUTHOR
    with zipfile.ZipFile(d) as z:
        props = (z.read("docProps/core.xml") + z.read("docProps/app.xml")).decode().lower()
    assert not any(x in props for x in W.PROHIBITED_PROPERTY_VALUES) and W.COMPANY in props


# ------------------------------------------------------------------------ TEST 15
def test_15_information_boundary(five_round, tmp_path):
    from cdf_gco_crms.excel import helper_writer as HW
    h, traj, out = five_round
    allowed = set(HW.META_KEYS) | set(HW.LONG_COLS) | set(HW.OBS_COLS) | set(HW.TREND_COLS)
    forbidden = {"severity", "integrity", "hash", "sha256", "closure_path", "lens", "decision_id", "evidence"}
    assert not {c for c in allowed if any(f in c.lower() for f in forbidden)}
    texts = I.all_text_parts(h.files[5])
    body = " ".join(v for k, v in texts.items() if k.startswith("xl/"))
    import re
    for pat in (r"VERIFIED_TEXT|GCO_REASSESSMENT|ROUTED_RESIDUAL|RISK_ACCEPTED", r"\bI-\d\d\b", r"\b[0-9a-f]{64}\b",
                r"\bCRITICAL\b"):
        assert not re.search(pat, body), pat
    d = cached(h.files[5])[W.DASHBOARD_SHEET]
    assert not any(isinstance(c.value, str) and "GCO-" in c.value for row in d.iter_rows() for c in row)


# ------------------------------------------------------------------------ TEST 16
@needs_lo
def test_16_visual_render_and_layout(five_round, tmp_path):
    h, traj, out = five_round
    p = h.files[5]
    pdf = lo_engine.to_pdf(p, tmp_path)
    assert pdf is not None and pdf.stat().st_size > 10000
    wb = load_workbook(p)
    fonts = set()
    for ws in wb.worksheets:
        fonts |= I.cell_fonts(ws)
    assert fonts <= {"Calibri"}
    d = wb[W.DASHBOARD_SHEET]
    assert d.sheet_view.rightToLeft and not d.sheet_view.showGridLines
    assert I.min_font_size(d) >= T.MIN_FONT
    from cdf_gco_crms.qa.ux_gates import _chart_rects
    for c1, r1, c2, r2 in _chart_rects(p, W.DASHBOARD_SHEET):
        assert all(d.cell(row=r, column=c).value in (None, "") for r in range(r1, r2 + 1) for c in range(c1, c2 + 1))


# ------------------------------------------------------------------------ TEST 17
def test_17_no_macro_dependency(five_round):
    h, traj, out = five_round
    with zipfile.ZipFile(h.files[5]) as z:
        names = z.namelist()
        ct = z.read("[Content_Types].xml").decode()
    assert not any("vbaProject" in n or "activeX" in n or "externalLink" in n or "connections" in n for n in names)
    assert "macroEnabled" not in ct and h.files[5].suffix == ".xlsx"


# ------------------------------------------------------------------------ TEST 18
def test_18_dashboard_refresh(five_round):
    h, traj, out = five_round
    for r in (2, 3, 4, 5):
        d = cached(h.files[r])[W.DASHBOARD_SHEET]
        assert d["B6"].value == S.round_label(r) and d["B10"].value == S.round_label(r)
        hs = cached(h.files[r])[W.HELPER_SHEET]
        hdr = {hs.cell(row=1, column=c).value: c for c in range(1, hs.max_column + 1)}
        assert [hs.cell(row=i + 1, column=hdr["trend.r"]).value for i in range(1, r + 1)] == list(range(1, r + 1))
        st = h.states[r]
        assert kpis(h.files[r])[0] == sum(1 for o in st.observations if o.birth_round <= r)
        pkg = I.package(h.files[r])
        assert all(c.endswith(f"${r + 1}") for ch in pkg.charts if ch.kinds[0] != "barChart/bar" for c in ch.category_refs)


# ------------------------------------------------------------------------ TEST 19
def test_19_historical_round_selection(five_round):
    h, traj, out = five_round
    st = h.states[5]
    rows = {o.internal_id: 2 + i for i, o in enumerate(st.ordered())}
    an = Analytics(st, rows, round_cols={}, data_sheet="x")
    late = next(a for a in an.obs if a.birth == 4)
    v = an.view(2, S.SCOPE_ALL, late.reader_label, 1)
    assert late not in v.table and late not in v.queue
    assert v.t.issued == 10 and len(v.table) == 10
    assert v.detail()["head"].endswith(S.NOT_RAISED)
    assert all(a.birth <= 2 for a in v.table)


@needs_lo
def test_19b_historical_selection_in_spreadsheet(five_round, tmp_path):
    h, traj, out = five_round
    vc = check_view(h.states[5], DashboardViewState(2, S.SCOPE_ALL, "#12", 1), tmp_path)
    assert vc.ok, vc.mismatches[:5]


# ------------------------------------------------------------------------ TEST 20
def test_20_closure_sustainability(five_round):
    h, traj, out = five_round
    st = h.states[5]
    o = next(x for x in st.observations if x.reader_no == "02")              # C, reopened r4, reclosed r5
    vals = {r: o.value(r) for r in range(1, 6)}
    assert L.latest_closure_round(vals, 1, 5) == 5 and L.latest_closure_round(vals, 1, 3) == 2
    assert o.cells[5].response == S.closed_guidance("الخامسة")               # 10.2: latest closure ordinal
    assert o.value(2) == C and o.value(3) == C                                # historical closure preserved
    an = Analytics(st, {x.internal_id: 2 + i for i, x in enumerate(st.ordered())}, round_cols={}, data_sheet="x")
    a = next(x for x in an.obs if x.reader == "02")
    assert a.long[4].reopen_state == 1 and a.long[5].reopen_state == 0 and a.long[5].ever == 1
    t5 = an.trend[5]
    assert t5.ever >= t5.closed and t5.reopen_state == sum(1 for x in an.obs if x.long[5].reopen_state)
    o3 = next(x for x in st.observations if x.reader_no == "03")
    assert [o3.value(r) for r in range(1, 6)] == [O, C, C, RE, RE]


# ---------------------------------------------------------------- v7.12 historical-validation regressions (TEST 25–27)
def test_25_reopened_row_gets_followup(five_round, tmp_path):
    """16.1 / AMB-164: a row closed in round 2 and reopened in round 4 carries a round-4 follow-up, threaded under its
    first comment, and the reconciliation counts it among the follow-ups of the round."""
    from cdf_gco_crms.word import word_sync as WS
    h, traj, out = five_round
    st = h.states[4]
    o3 = next(x for x in st.observations if x.reader_no == "03")             # [O, C, C, RE, RE]
    assert o3.value(3) == C and o3.value(4) == RE
    d = WS.build_review_docx(st, tmp_path / "r4.docx")
    comments = WS.extract_comments(d)
    by_cid = {c.cid: c for c in comments}
    fu = [c for c in comments if c.internal_id == o3.internal_id and c.follow_up_round == 4]
    assert len(fu) == 1 and fu[0].status == RE
    root = by_cid[fu[0].parent_cid]
    assert root.internal_id == o3.internal_id and root.follow_up_round is None
    # no follow-up for a row that stays closed through round 4
    stays = next(x for x in st.observations if x.value(3) == C and x.value(4) == C)
    assert not [c for c in comments if c.internal_id == stays.internal_id and c.follow_up_round == 4]
    rep = WS.reconcile(d, h.files[4], round_no=4)
    assert rep.ok and rep.checks["G-SYNC-1 count"] and rep.checks["G-SYNC-9 follow-ups threaded"], rep.problems


def test_26_historical_response_follows_selected_round(five_round):
    """10.4A.6: with an earlier round selected, the detail panel shows the latest owner response up to that round."""
    h, traj, out = five_round
    st = h.states[5]
    an = Analytics(st, {x.internal_id: 2 + i for i, x in enumerate(st.ordered())}, round_cols={}, data_sheet="x")
    checked = 0
    for a in an.obs:
        filled = [r for r in range(1, 6) if a.long[r].resp_filled]
        if len(filled) < 2:
            continue
        for S_ in range(1, 6):
            if not a.long[S_].exists:                                       # born after the selected round
                continue
            v = an.view(S_, S.SCOPE_ALL, a.reader_label, max(1, S_ - 1))
            resp = v.detail().get("resp", "")
            expect_round = max([r for r in filled if r <= S_], default=0)
            if expect_round == 0:
                assert resp == S.NO_RESPONSE, (a.reader, S_, resp)
            else:
                assert resp.endswith(f"({S.round_label(expect_round)})"), (a.reader, S_, resp)
                assert a.long[expect_round].resp in resp
            for later in range(S_ + 1, 6):
                if later in filled and a.long[later].resp != a.long[expect_round].resp if expect_round else True:
                    assert f"({S.round_label(later)})" not in resp
            checked += 1
    assert checked >= 10


def test_27_closure_round_is_the_actual_closure_event(five_round):
    """10.2: a row closed in round 2 and still closed in round 3 reports round 2; a reclosure reports its own round."""
    h, traj, out = five_round
    st = h.states[5]
    an = Analytics(st, {x.internal_id: 2 + i for i, x in enumerate(st.ordered())}, round_cols={}, data_sheet="x")
    o2 = next(x for x in an.obs if x.reader == "02")                          # [O, C, C, RE, C]
    assert (o2.long[2].lcr, o2.long[3].lcr, o2.long[4].lcr, o2.long[5].lcr) == (2, 2, "", 5)
    o3 = next(x for x in an.obs if x.reader == "03")                          # [O, C, C, RE, RE]
    assert (o3.long[2].lcr, o3.long[3].lcr, o3.long[4].lcr, o3.long[5].lcr) == (2, 2, "", "")
    for a in an.obs:                                                          # general law for every row
        for r in range(2, 6):
            if a.long[r].value == C and a.long[r - 1].value == C:
                assert a.long[r].lcr == a.long[r - 1].lcr, (a.reader, r)
            if a.long[r].value == C and a.long[r - 1].value not in (C, ""):
                assert a.long[r].lcr == r, (a.reader, r)
