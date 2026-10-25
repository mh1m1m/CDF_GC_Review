"""AI review of uploaded documents (CDF_DOCUMENT_UPLOADS=PUBLIC): queue, worker pipeline, UPL- cases, outputs.

The model is the explicit MOCK executor (CDF_REVIEW_EXECUTOR=mock); nothing here claims live inference. Inputs are
the fictional Arabic memo (built at test time, document_fixtures.py) and documents generated in the test.
"""
import importlib
import io
import sys
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("pdfplumber")

from document_fixtures import memo_pdf  # noqa: E402

TOKEN = "r" * 48


def headers():
    return {"Authorization": "Bearer " + TOKEN}


@pytest.fixture()
def api(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    monkeypatch.setenv("CDF_CALIBRATION_MODE", "SYNTHETIC_LOCAL_DEV")
    monkeypatch.setenv("CDF_CALIBRATION_DEV_TOKEN", TOKEN)
    monkeypatch.setenv("CDF_CALIBRATION_DB", str(tmp_path / "public-demo" / "review.sqlite"))
    monkeypatch.setenv("CDF_DOCUMENT_UPLOADS", "PUBLIC")
    monkeypatch.setenv("CDF_REVIEW_EXECUTOR", "mock")
    monkeypatch.setenv("CDF_REVIEW_RUNNER", "inline")
    sys.modules.pop("apps.api.main", None)
    module = importlib.import_module("apps.api.main")
    with TestClient(module.app) as client:
        yield client, module
    module.store.close()
    sys.modules.pop("apps.api.main", None)


def upload(client, name, content):
    response = client.post("/api/documents", params={"filename": name}, content=content,
                           headers={**headers(), "Content-Type": "application/octet-stream"})
    assert response.status_code == 201, response.text
    return response.json()


def review(client, document_id):
    response = client.post(f"/api/documents/{document_id}/review", headers=headers())
    assert response.status_code == 202, response.text
    return response.json()


def _word(paragraphs):
    from docx import Document

    word = Document()
    for text in paragraphs:
        word.add_paragraph(text)
    out = io.BytesIO()
    word.save(out)
    return out.getvalue()


def test_pdf_review_end_to_end(api):
    client, _ = api
    document = upload(client, "مذكرة تفاهم.pdf", memo_pdf())
    assert document["kind"] == "pdf" and document["summary"]["text_layer"] is True
    done = review(client, document["document_id"])
    assert done["status"] == "DONE", done
    detail = done["detail"]
    assert detail["outcome"] == "REVIEWED" and detail["executor"] == "mock" and detail["live_calls"] == 0
    assert detail["observations"] == detail["adopted"] >= 2
    case_id = done["case_id"]
    assert case_id.startswith("UPL-")
    names = {o["filename"]: o for o in detail["outputs"]}
    assert all(n.startswith("INTERNAL_DRAFT__") for n in names)
    assert "INTERNAL_DRAFT__مذكرة تفاهم - ملاحظات الجولة 1.docx" in names
    assert "INTERNAL_DRAFT__مذكرة تفاهم - الأصل مع الملاحظات.pdf" in names
    assert "INTERNAL_DRAFT__مذكرة تفاهم - سجل المراجعة.xlsx" in names
    assert "INTERNAL_DRAFT__مذكرة تفاهم - لوحة المتابعة.xlsx" in names
    assert detail["placement"]["placed"] == detail["observations"]

    listed = {d["document_id"]: d for d in client.get("/api/documents", headers=headers()).json()}
    commented = listed[names["INTERNAL_DRAFT__مذكرة تفاهم - ملاحظات الجولة 1.docx"]["document_id"]]
    assert commented["summary"]["review_comments"] == detail["observations"]
    tracker = listed[names["INTERNAL_DRAFT__مذكرة تفاهم - سجل المراجعة.xlsx"]["document_id"]]
    assert tracker["summary"]["review_workbook"] is True and tracker["summary"]["observations"] == detail["observations"]
    assert tracker["derived_from"] == document["document_id"] and tracker["action"] == "review"

    # the case opens in the existing screens: case list, calibration view and executive dashboard
    assert case_id in [c["case_id"] for c in client.get("/api/cases", headers=headers()).json()]
    dashboard = client.get(f"/api/cases/{case_id}/intelligence", headers=headers())
    assert dashboard.status_code == 200, dashboard.text
    data = dashboard.json()
    assert data["uploaded_document"] is True and data["synthetic"] is False
    assert data["limitations"][0] == "UPLOADED_DOCUMENT_WORKING_COPY"
    assert data["observation_count"] == detail["observations"]
    export = client.get(f"/api/cases/{case_id}/intelligence/export/xlsx", headers=headers())
    assert export.status_code == 200 and export.content[:2] == b"PK"
    # synthetic job runner is refused on an uploaded case
    refused = client.post("/api/operations/cdf_review_document", headers=headers(),
                          json={"case_id": case_id, "executor": "mock"})
    assert refused.status_code == 422 and "UPLOADED_CASES_ARE_REVIEWED_FROM_THE_DOCUMENTS_PAGE" in refused.text
    status = client.get("/api/document-reviews", headers=headers()).json()
    assert status["latest"][document["document_id"]]["status"] == "DONE"


def test_word_and_spreadsheet_reviews(api):
    from openpyxl import Workbook, load_workbook

    client, _ = api
    word = upload(client, "سياسة.docx", _word(["سياسة تجريبية للمشتريات",
                                               "يعتمد مدير الإدارة جميع العقود دون حد مالي أو مراجعة مستقلة للعروض.",
                                               "يجوز تجاوز إجراءات المنافسة في الحالات العاجلة دون توثيق المبررات."]))
    done = review(client, word["document_id"])
    assert done["status"] == "DONE" and done["detail"]["placement"] is None, done
    assert done["detail"]["comments"]["anchored"] == done["detail"]["observations"]

    book = Workbook()
    sheet = book.active
    sheet.title = "الميزانية"
    sheet.append(["البند", "المبلغ", "الجهة المعتمدة"])
    sheet.append(["عقد استشارات إدارية لمدة ثلاث سنوات", 450000, "غير محددة في الخطة"])
    sheet.append(["برنامج تدريب الموظفين الجدد على الأنظمة", 20000, "إدارة الموارد البشرية"])
    out = io.BytesIO()
    book.save(out)
    sheet_doc = upload(client, "الميزانية.xlsx", out.getvalue())
    assert sheet_doc["summary"]["review_workbook"] is False
    done = review(client, sheet_doc["document_id"])
    assert done["status"] == "DONE", done
    names = {o["filename"]: o["document_id"] for o in done["detail"]["outputs"]}
    annotated = names["INTERNAL_DRAFT__الميزانية - الأصل مع الملاحظات.xlsx"]
    content = client.get(f"/api/documents/{annotated}/content", headers=headers()).content
    commented_cells = [c for row in load_workbook(io.BytesIO(content))["الميزانية"].iter_rows() for c in row if c.comment]
    assert commented_cells and all("GCO-UPL-R1-" in c.comment.text for c in commented_cells)


def test_queue_guards(api):
    client, module = api
    word = upload(client, "a.docx", _word(["نص قصير جداً"]))
    # a document whose text gives the drafting agent nothing to quote: no case is created
    done = review(client, word["document_id"])
    assert done["status"] == "DONE" and done["detail"]["outcome"] == "NO_OBSERVATIONS_ADOPTED" and not done["case_id"]

    generated = upload(client, "memo.pdf", memo_pdf())
    first = review(client, generated["document_id"])
    output = first["detail"]["outputs"][0]["document_id"]
    again = client.post(f"/api/documents/{output}/review", headers=headers())
    assert again.status_code == 422 and "GENERATED_FILES_ARE_NOT_REVIEWED" in again.text
    tracker = next(o["document_id"] for o in first["detail"]["outputs"] if "سجل المراجعة" in o["filename"])
    refused = client.post(f"/api/documents/{tracker}/review", headers=headers())
    assert refused.status_code == 422 and "REVIEW_TRACKER_USE_NEXT_ROUND_TOOLS" in refused.text
    missing = client.get("/api/document-reviews/REV-" + "0" * 32, headers=headers())
    assert missing.status_code == 403


def test_pdf_with_script_is_refused(api):
    from pypdf import PdfWriter

    client, _ = api
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    out = io.BytesIO()
    writer.write(out)
    blank = out.getvalue()
    response = client.post("/api/documents", params={"filename": "blank.pdf"}, content=blank,
                           headers={**headers(), "Content-Type": "application/octet-stream"})
    assert response.status_code == 201 and response.json()["summary"]["text_layer"] is False
    pdf = blank.replace(b"%%EOF", b"/JavaScript (app.alert(1))\n%%EOF")
    response = client.post("/api/documents", params={"filename": "x.pdf"}, content=pdf,
                           headers={**headers(), "Content-Type": "application/octet-stream"})
    assert response.status_code == 422 and "PDF_ACTIVE_CONTENT" in response.text
    response = client.post("/api/documents", params={"filename": "fake.pdf"}, content=b"not a pdf",
                           headers={**headers(), "Content-Type": "application/octet-stream"})
    assert response.status_code == 422 and "NOT_A_VALID_PDF_FILE" in response.text


def test_failed_provider_is_recorded(api, monkeypatch):
    client, _ = api
    monkeypatch.setenv("CDF_REVIEW_EXECUTOR", "anthropic")
    monkeypatch.setenv("CDF_ENABLE_PROVIDER_API", "1")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_AUTH_TOKEN", raising=False)
    document = upload(client, "memo.pdf", memo_pdf())
    done = review(client, document["document_id"])
    assert done["status"] == "FAILED" and done["error"] == "ProviderFailure" and not done["case_id"]


def test_provider_carries_data_authorization(monkeypatch):
    from cdf_gco_crms.execution.providers import select_executor

    monkeypatch.setenv("CDF_ENABLE_PROVIDER_API", "1")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-not-a-key")
    monkeypatch.delenv("CDF_AI_DATA_AUTHORIZATION", raising=False)
    assert select_executor("anthropic", timeout=5).config.data_authorization == "SYNTHETIC_ONLY"
    monkeypatch.setenv("CDF_AI_DATA_AUTHORIZATION", "INSTITUTIONAL_AUTHORIZED:owner-decision-test")
    assert select_executor("anthropic", timeout=5).config.institutional_data_authorized is True


def test_uploaded_cases_close_when_uploads_are_switched_off(tmp_path, monkeypatch):
    from cdf_gco_crms.calibration.store import AccessDenied, CalibrationStore, ReviewerContext
    from cdf_gco_crms.documents.review import ReviewQueue, process_next
    from cdf_gco_crms.execution.service import ApplicationService
    from cdf_gco_crms.institutional.contracts import InstitutionalAccessDenied
    from cdf_gco_crms.institutional.isolation import bind_public_database
    from cdf_gco_crms.interfaces.hosted_documents import DocumentLibrary

    monkeypatch.setenv("CDF_DOCUMENT_UPLOADS", "PUBLIC")
    monkeypatch.setenv("CDF_REVIEW_EXECUTOR", "mock")
    store = CalibrationStore(tmp_path / "public-demo" / "switch.sqlite")
    app = ApplicationService(store=store, token="")
    library, queue = DocumentLibrary(store), ReviewQueue(store)
    user = ReviewerContext("SYN-HOSTED-A", frozenset({"CALIBRATION_ADMIN", "CALIBRATION_REVIEWER"}), True)
    document = library.upload(user.reviewer_id, "m.pdf", memo_pdf())
    queue.queue(user, document)
    case_id = process_next(app, library, queue, "t")["case_id"]
    assert app.intelligence.dashboard(case_id, user)["uploaded_document"] is True
    other = ReviewerContext("SYN-HOSTED-B", frozenset({"CALIBRATION_ADMIN", "CALIBRATION_REVIEWER"}), True)
    with pytest.raises(AccessDenied):
        app.intelligence.dashboard(case_id, other)
    monkeypatch.setenv("CDF_DOCUMENT_UPLOADS", "")
    with pytest.raises(AccessDenied):
        app.intelligence.dashboard(case_id, user)
    with pytest.raises(InstitutionalAccessDenied):
        bind_public_database(store)
    store.close()


def test_worker_refuses_without_owner_switch(monkeypatch):
    from cdf_gco_crms.documents.worker import build

    monkeypatch.delenv("CDF_DOCUMENT_UPLOADS", raising=False)
    with pytest.raises(SystemExit):
        build()
