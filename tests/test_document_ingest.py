"""Reading uploaded documents under review (Word, PDF, Excel) and pinning comments back onto PDF/Excel originals.

The PDF is a fictional Arabic memo exported from Word by LibreOffice at test time (document_fixtures.py).
"""
import io
from pathlib import Path

import pytest

from cdf_gco_crms.documents.annotate import Comment, annotate_pdf, annotate_xlsx, find_block, key
from cdf_gco_crms.documents.ingest import extract, to_docx, visual_to_logical
from document_fixtures import memo_lines, memo_pdf_path

pytest.importorskip("pdfplumber")


def test_pdf_arabic_text_in_logical_order(tmp_path):
    source = extract(memo_pdf_path(tmp_path), "pdf")
    assert source.pages == 1 and source.blocks and not source.warnings
    joined = " ".join(b.text for b in source.blocks)
    # Exact glyph mappings and paragraph spacing vary with LibreOffice and font versions; reading order, numbers,
    # brackets and ligatures must not: every line of the memo is found, in logical order, after loose normalisation.
    for line in memo_lines():
        assert key(line) in key(joined), line
    assert "(3):" in joined and "120000" in joined and "KYC" in joined and "2024." in joined
    for block in source.blocks:
        assert block.page == 1 and block.box and block.lines
    assert source.blocks[0].location == "صفحة 1، فقرة 1"


def test_visual_order_repair_keeps_numbers_and_ligatures():
    assert visual_to_logical(list(":)3(")) == "(3):"
    assert visual_to_logical(list(".2024")) == "2024."
    assert visual_to_logical(["ل", "لا", "خ"]) == "خلال"          # a ligature glyph stays one unit
    assert visual_to_logical(list("000021")) == "000021"


def test_docx_blocks_and_word_rendition(tmp_path):
    from docx import Document

    word = Document()
    for text in ("الفقرة الأولى من المستند.", "", "يجوز للطرف الثاني رفض طلبات التدقيق."):
        word.add_paragraph(text)
    word.add_table(rows=1, cols=1).cell(0, 0).text = "خلية جدول"
    path = tmp_path / "doc.docx"
    word.save(path)
    source = extract(path, "docx")
    assert [b.text for b in source.blocks] == ["الفقرة الأولى من المستند.", "يجوز للطرف الثاني رفض طلبات التدقيق.",
                                                 "خلية جدول"]
    assert source.blocks[2].location.startswith("جدول")
    pdf_source = extract(memo_pdf_path(tmp_path), "pdf")
    out = to_docx(pdf_source, "مذكرة", tmp_path / "rendition.docx")
    paragraphs = [p.text for p in Document(out).paragraphs]
    assert all(b.text in paragraphs for b in pdf_source.blocks)


def test_pdf_annotation_places_and_refuses_to_guess(tmp_path):
    from pypdf import PdfReader

    pdf = memo_pdf_path(tmp_path)
    source = extract(pdf, "pdf")
    comments = [Comment("GCO-1", "تلتزم الجهة بدفع 120000 ريال", "الملاحظة: التزام مالي. التوصية: توثيق الحد."),
                Comment("GCO-2", "رفض طلبات التدقيق", "حفظ حق المراجعة."),
                Comment("GCO-3", "عبارة لا ترد في المستند", "x")]
    out = tmp_path / "annotated.pdf"
    report = annotate_pdf(pdf, source, comments, out)
    assert [p["ref"] for p in report.placed] == ["GCO-1", "GCO-2"] and report.unplaced == ["GCO-3"]
    annots = [a.get_object() for a in PdfReader(out).pages[0]["/Annots"]]
    assert sorted(a["/Subtype"] for a in annots) == ["/Highlight", "/Highlight", "/Text", "/Text"]
    assert any("GCO-1" in a["/Contents"] for a in annots)
    # page text is untouched
    assert extract(out, "pdf").text == source.text


def test_spreadsheet_under_review(tmp_path):
    from openpyxl import Workbook, load_workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "الميزانية"
    sheet.append(["البند", "المبلغ", "الجهة المعتمدة"])
    sheet.append(["عقد استشارات", 450000, "غير محدد"])
    path = tmp_path / "budget.xlsx"
    book.save(path)
    source = extract(path, "xlsx")
    assert source.blocks[1].text == "عقد استشارات | 450000 | غير محدد"
    assert source.blocks[1].location == "ورقة الميزانية، صف 2"
    out = tmp_path / "commented.xlsx"
    report = annotate_xlsx(path, source, [Comment("GCO-9", "عقد استشارات | 450000", "جهة الاعتماد غير محددة.")], out)
    assert report.summary()["placed"] == 1
    assert "GCO-9" in load_workbook(out)["الميزانية"]["A2"].comment.text


def test_loose_matching_key(tmp_path):
    assert key("إعداد  الطلبِ") == key("اعداد الطلب")
    source = extract(memo_pdf_path(tmp_path), "pdf")
    found = find_block(source, "يوماً من توقيع الاتفاق")              # tanween in the quote, possibly absent in the PDF
    assert found is not None and "توقيع الاتفاق" in found.text
    assert find_block(source, "") is None


def _glyph(text, x0, x1, top=100.0, bottom=111.0, size=11.0):
    return {"text": text, "x0": x0, "x1": x1, "top": top, "bottom": bottom, "size": size, "upright": True}


def test_glyph_lines_from_real_world_layouts():
    """Layouts seen across LibreOffice/font versions, rebuilt without any font: Naskh splits the lam-alef ligature
    into two characters at one position (in either order), Latin and digit glyphs sit on a different box than the
    Arabic ones, and a line may carry no space glyph between words."""
    from cdf_gco_crms.documents.ingest import _glyph_lines, _line_from_glyphs

    # visual (left-to-right) glyphs of the RTL line "خلال 2024 KYC" with Latin on a shifted baseline
    chars = [
        _glyph("K", 10, 15, 103, 112), _glyph("Y", 15, 20, 103, 112), _glyph("C", 20, 25, 103, 112),
        _glyph(" ", 25, 28),
        _glyph("2", 28, 33, 102, 113), _glyph("0", 33, 38, 102, 113), _glyph("2", 38, 43, 102, 113),
        _glyph("4", 43, 48, 102, 113),
        _glyph(" ", 48, 51),
        _glyph("ل", 51, 56),                                                  # final lam (leftmost)
        _glyph("ل", 56, 62), _glyph("ا", 56, 56),                             # lam-alef: lam + zero-width alef
        _glyph("خ", 62, 70),
    ]
    lines = _glyph_lines(chars)
    assert len(lines) == 1
    assert _line_from_glyphs(lines[0], rtl=True) == "خلال 2024 KYC"
    # the same ligature with its two characters in the other order
    chars[10], chars[11] = _glyph("ا", 56, 56), _glyph("ل", 56, 62)
    assert _line_from_glyphs(_glyph_lines(chars)[0], rtl=True) == "خلال 2024 KYC"
    # no space glyph between two words: a wide gap still separates them
    words = [_glyph("ب", 10, 16), _glyph("ا", 16, 19), _glyph("ب", 30, 36), _glyph("ل", 36, 40)]
    assert _line_from_glyphs(_glyph_lines(words)[0], rtl=True) == "لب اب"
