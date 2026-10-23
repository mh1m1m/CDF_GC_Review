"""TEST 32 — production end-to-end case (patches V712-034, 035, 044, 055).

Synthetic first review on a realistic reviewed document → the owner returns Word and Excel (replies, a
tracked insertion, a silent edit, workbook responses) → rounds 2, 3 and 5 with a reopening and a conditional
closure. At every step the artefacts are reopened and cross-reconciled; the pre-issue challenge of 19A runs
before any round-n output and nothing of its record reaches the requester files. No fixture-specific shortcut:
the round inputs are ordinary RoundInput objects and the Word file is the owner's own document.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import C, CC, O, RE, base_obs
from test_word_lifecycle import build_reviewed_doc, gco_comments, simulate_return, quote_of
from cdf_gco_crms.assurance import challenge as CH
from cdf_gco_crms.assurance.evidence import EvidenceRegistry
from cdf_gco_crms.assurance.personas import RecipientPersona
from cdf_gco_crms.config import round_schema as S
from cdf_gco_crms.config import workbook_settings as W
from cdf_gco_crms.excel.workbook_builder import build_workbook
from cdf_gco_crms.ingestion.workbook_reader import parse_workbook
from cdf_gco_crms.model.entities import GText, RoundAssessment
from cdf_gco_crms.qa import lo_engine
from cdf_gco_crms.reconciliation.historical import merge_returned
from cdf_gco_crms.references.index import ReferenceIndex
from cdf_gco_crms.regression import synthetic as SY
from cdf_gco_crms.rounds.round_expansion import NewObservation, RoundInput, apply_round, build_round1
from cdf_gco_crms.word import document_comments as DC
from cdf_gco_crms.word.word_sync import reconcile

DEPT = "الحوكمة والامتثال"
N = 8
# trajectories (index 0 = round 1): closure, reopening in round 3 (RE after C), conditional close in round 5
TRAJ = [
    [O, C, C, C, C],          # 01 closed in round 2 and stays closed
    [O, C, RE, RE, C],        # 02 closed in round 2, REOPENED in round 3, reclosed in round 5
    [O, CC, CC, CC, CC],      # 03 conditional from round 2 (condition pending)
    [O, RE, RE, O, CC],       # 04 requires edit, then open, CONDITIONAL CLOSE in round 5
    [O, O, RE, C, C],         # 05 open, requires edit, closed in round 4
    [O, C, C, C, C],          # 06
    [O, RE, C, C, C],         # 07
    [O, O, O, RE, RE],        # 08 stalled
]


def _fb(text):
    return GText([(text, False)])


def _assess(state, r, values):
    """RoundInput for round r from the trajectories: every row not مغلقة at the start of the round is assessed."""
    rows = []
    for o in state.ordered():
        k = int(o.reader_no) - 1
        if k >= len(values):
            continue
        v = values[k][r - 1]
        prev = o.value(r - 1)
        if prev == C and v == C:
            continue                                      # carried forward by 10.2
        fb = {C: "عُدّل النص وفق التوصية؛ تُغلق الملاحظة.", CC: "عُدّل النص؛ شرط الإغلاق: اعتماد الصيغة من صاحب الصلاحية.",
              RE: "التعديل جزئي؛ المطلوب: استكمال الصياغة.", O: "لم يرد ما يفيد بالتعديل؛ المطلوب: بيان الإجراء."}[v]
        basis = "VERIFIED_TEXT" if v == C else None
        rows.append(RoundAssessment(o.internal_id, r, v, _fb(fb), basis))
    return rows


@pytest.fixture(scope="module")
def case(tmp_path_factory):
    d = tmp_path_factory.mktemp("e2e")
    doc = SY.fixture_document("ETE")
    doc.owning_department = DEPT
    st1 = build_round1(doc, base_obs(N), date="2026-01-15", document_version="الإصدار 0.1")
    return {"dir": d, "states": {1: st1}, "files": {}, "docx": {}, "reports": {}}


def _signed_decision(decision_id, observation_id, outcome, priv):
    """A reviewer decision signed with an ephemeral SYNTHETIC key (no CDF key or approval is impersonated)."""
    import base64
    from dataclasses import replace
    from cdf_gco_crms.assurance.objection_engine import HumanDecision
    d = HumanDecision(decision_id, observation_id, outcome, "م. مراجع الحوكمة (تجريبي)", "GCO reviewer",
                      f"DEC-{decision_id}", "2026-10-23", f"E-{decision_id}", "VERIFIED_HUMAN_RECORD", "SYNTHETIC_TEST_KEY", "")
    return replace(d, signature_b64=base64.b64encode(priv.sign(d.signature_payload())).decode())


def _owner_adds_clause(src, dst, text):
    """The owner's next version: the delivered file with one paragraph added before the section properties."""
    import zipfile
    from lxml import etree
    with zipfile.ZipFile(src) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    root = etree.fromstring(parts[DC.DOC_PART])
    body = root.find(DC.w("body"))
    p = etree.Element(DC.w("p")); r = etree.SubElement(p, DC.w("r")); etree.SubElement(r, DC.w("t")).text = f"المادة (9): {text}."
    sect = body.find(DC.w("sectPr"))
    (sect.addprevious if sect is not None else body.append)(p)
    parts[DC.DOC_PART] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return dst


def _issue(case, r, state, *, prior_docx=None, src_docx=None):
    """Pre-issue challenge (19A) → reviewer decisions on every blocker → Excel → Word → cross-reconciliation.

    An unresolved blocker refuses the issue; in rounds n ≥ 2 an objection against an ISSUED observation always
    escalates to the GCO reviewer (17.7), so the issue waits for a verified decision record (19A.9)."""
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    d = case["dir"]
    idx = ReferenceIndex.load()
    reg = EvidenceRegistry()
    reg.add_document_text(src_docx or case["docx"]["base"])
    persona = RecipientPersona.from_index(idx, DEPT, reg)
    rec = CH.run_challenge(state, reg, persona)
    assert rec.verify_chains() and rec.independence == CH.INDEPENDENCE
    blockers = rec.blockers()
    if r == 1:
        assert blockers == [], blockers                               # a fresh draft must leave the challenge clean
    else:
        assert all(b["disposition"] == "ESCALATED_TO_GCO_REVIEWER" for b in blockers), blockers
        priv = Ed25519PrivateKey.generate()
        trust = {"SYNTHETIC_TEST_KEY": priv.public_key().public_bytes(serialization.Encoding.Raw, serialization.PublicFormat.Raw)}
        # the reviewer withholds a row whose clause the owner changed (its evidence is gone) and upholds the rest on
        # the bound evidence; an upheld outcome without verified evidence is refused by the engine
        decisions = [_signed_decision(f"R{r}-{i:02d}", b["internal_id"],
                                      "WITHHELD_FOR_EVIDENCE" if b["kind"] == "CLAUSE_NOT_FOUND" else "UPHELD_UNCHANGED", priv)
                     for i, b in enumerate(blockers, 1)]
        res = CH.apply_human_decisions(rec, decisions, trust, registry=reg)
        assert res["rejected"] == [] and res["remaining_blockers"] == [], res
        case.setdefault("decisions", {})[r] = len(decisions)
    (d / f"challenge_r{r}.json").write_text(rec.to_json(), encoding="utf-8")   # internal record, never shipped
    xlsx = d / f"r{r}.xlsx"
    build_workbook(state, xlsx)
    docx = d / f"r{r}.docx"
    rep = DC.comment_document(src_docx or case["docx"]["base"], state, docx, round_no=r, prior_docx=prior_docx)
    assert rep.ok, rep.summary()
    assert DC.check_package(docx, xlsx, state) == []
    sync = reconcile(docx, xlsx, round_no=r)
    assert sync.ok, sync.checks
    assert CH.confidentiality_scan([xlsx, docx]) == []
    case["states"][r], case["files"][r], case["docx"][r], case["reports"][r] = state, xlsx, docx, rep
    return xlsx, docx


def test_e2e_round1_issue(case):
    st1 = case["states"][1]
    case["docx"]["base"] = build_reviewed_doc(st1, case["dir"] / "base.docx")
    xlsx, docx = _issue(case, 1, st1)
    pw = parse_workbook(xlsx)
    assert len(pw.state.observations) == N and pw.state.current_round == 1
    assert len(gco_comments(docx)) == N


def test_e2e_owner_return_round1(case):
    """The owner answers in both files; the engine classifies every change before round 2 is drafted."""
    d = case["dir"]
    st1, xlsx1, docx1 = case["states"][1], case["files"][1], case["docx"][1]
    # Excel: owner fills responses; also tampers one GCO cell (column E) — restored and reported
    from cdf_gco_crms.ingestion.workbook_reader import find_data_sheet
    wb = load_workbook(xlsx1, rich_text=True)
    ws = find_data_sheet(wb)
    ws["G2"] = "تم تعديل البند وفق التوصية (ص.3)."; ws["H2"] = S.IMPL_DONE     # round-1 block: F closure, G response, H impl
    ws["G3"] = "قيد التعديل."; ws["H3"] = S.IMPL_PARTIAL
    ws["E4"] = "نص مُعدّل من الإدارة"                             # tampering with a GCO cell
    ret_xlsx = d / "r1_returned.xlsx"; wb.save(ret_xlsx)
    delivered, returned = parse_workbook(xlsx1), parse_workbook(ret_xlsx)
    merged, intake = merge_returned(delivered.state, returned.state)
    assert intake.owner_inputs >= 2
    assert [x.internal_id for x in intake.gco_cell_edits] == [st1.ordered()[2].internal_id]
    assert merged.ordered()[2].comment == st1.ordered()[2].comment   # GCO text restored
    assert str(merged.ordered()[0].cells[1].response).startswith("تم تعديل البند")
    # Word: a reply, a tracked insertion and a SILENT edit of the passage of observation 02
    o2 = st1.ordered()[1]
    q2 = quote_of(o2)
    ret_docx = simulate_return(docx1, d / "r1_returned.docx", owner_comment_on=quote_of(st1.ordered()[0])[:30],
                               add_ins_on=quote_of(st1.ordered()[4])[:30], edit=(q2[:25], q2[:10], "نص مُغيّر "))
    diff = DC.diff_returned(docx1, ret_docx, st1)
    s = diff.summary()
    assert s["owner_comments"] == 1 and s["tracked_insertions"] >= 1 and s["silent_edits"] >= 1
    assert not s["missing_gco_comments"]
    case["returned"] = {"xlsx": ret_xlsx, "docx": ret_docx, "merged": merged, "diff": s}


def test_e2e_round2_from_return(case):
    merged = case["returned"]["merged"]
    rin = RoundInput(2, "2026-02-15", "الإصدار 0.2", _assess(merged, 2, TRAJ))
    st2, rep = apply_round(merged, rin)
    assert rep.action == "APPENDED" and len(rep.assessed) == N
    xlsx, docx = _issue(case, 2, st2, src_docx=case["returned"]["docx"], prior_docx=case["docx"][1])
    pw = parse_workbook(xlsx)
    assert pw.state.current_round == 2
    assert str(pw.state.ordered()[0].cells[1].response).startswith("تم تعديل البند")   # owner input survives
    r2 = case["reports"][2]
    assert len(r2.replies_added) == N and r2.roots_missing == []          # every follow-up is a native reply


def test_e2e_round3_reopening(case):
    st2 = case["states"][2]
    # owner returns round-2 Word: changes the text of observation 01 (closed in round 2) silently → 16.1
    o1 = st2.ordered()[0]
    q1 = quote_of(o1)
    ret = simulate_return(case["docx"][2], case["dir"] / "r2_returned.docx", edit=(q1[:25], q1[:8], "تعديل "))
    findings = {f.reader_no: f.status for f in DC.closed_text_integrity(st2, ret, 3)}
    assert findings.get(o1.reader_no) == DC.CHANGED
    assert findings.get(st2.ordered()[5].reader_no) == DC.INTACT
    # GCO reopens 02 (RE after C) and keeps 01 closed only because the text is intact — here it is not, so 01
    # is reopened too (16.1); the trajectory of 01 is overridden for this round
    rows = _assess(st2, 3, TRAJ)
    rows.append(RoundAssessment(o1.internal_id, 3, RE, _fb("تغيّر النص الذي أُغلقت عليه الملاحظة؛ أُعيد فتحها (16.1)."), None))
    st3, rep = apply_round(st2, RoundInput(3, "2026-03-15", "الإصدار 0.3", rows))
    assert set(rep.reopened) >= {o1.internal_id, st2.ordered()[1].internal_id}
    xlsx, docx = _issue(case, 3, st3, src_docx=ret, prior_docx=case["docx"][2])
    an = parse_workbook(xlsx).state
    assert an.ordered()[1].value(3) == RE and an.ordered()[0].value(3) == RE
    # reopened rows carry a round-3 follow-up reply (G-SYNC-9)
    refs = {(c.internal_id, c.follow_up_round) for c in gco_comments(docx)}
    assert (o1.internal_id, 3) in refs and (st2.ordered()[1].internal_id, 3) in refs
    reopened_r3 = [o.reader_no for o in an.ordered() if o.value(2) == C and o.value(3) != C]
    assert reopened_r3 == ["01", "02"]


def test_e2e_rounds4_5_conditional_close_and_new_observation(case):
    st3 = case["states"][3]
    st4, _ = apply_round(st3, RoundInput(4, "2026-04-15", "الإصدار 0.4", _assess(st3, 4, [t[:4] for t in TRAJ])))
    rows5 = _assess(st4, 5, TRAJ)
    # 01 was reopened in round 3 outside its trajectory: close it in round 5
    o1 = st4.ordered()[0]
    rows5 = [r for r in rows5 if r.internal_id != o1.internal_id] + [RoundAssessment(o1.internal_id, 5, C, _fb("أُعيد النص إلى صيغته المعتمدة؛ تُغلق."), "VERIFIED_TEXT")]
    new = [NewObservation("ص.9، المادة (9) – البند (9-1)", "نص أضيف في الإصدار 0.5 …", _fb("نوصي بربط البند بموافقة صاحب الصلاحية."),
                          reason="أُضيف البند في الإصدار 0.5.", title="ربط البند الجديد بموافقة صاحب الصلاحية")]
    st5, rep = apply_round(st4, RoundInput(5, "2026-05-15", "الإصدار 0.5", rows5, new))
    assert len(rep.new_ids) == 1
    # the owner's 0.5 document is the round-3 file with the new clause appended (threads of rounds 1–3 intact)
    base5 = _owner_adds_clause(case["docx"][3], case["dir"] / "base5.docx", quote_of(st5.ordered()[-1]))
    xlsx, docx = _issue(case, 5, st5, src_docx=base5, prior_docx=case["docx"][3])
    an = parse_workbook(xlsx).state
    by = {o.reader_no: o for o in an.ordered()}
    assert by["04"].value(5) == CC and by["02"].value(5) == C and by["01"].value(5) == C
    assert by["09"].birth_round == 5 and by["09"].title_status == S.TITLE_PROPOSED
    assert all(by[f"0{k}"].value(r) == TRAJ[k - 1][r - 1] for k in (3, 5, 6, 7, 8) for r in (1, 2, 4, 5))
    rec = json.loads((case["dir"] / "challenge_r5.json").read_text(encoding="utf-8"))
    assert rec["chains_verified"] and rec["blockers"] == [] and rec["independence"] == "SINGLE_PROCESS_RULE_BASED"
    assert len(rec["human_decisions"]) == case["decisions"].get(5, 0)
    # the requester files carry nothing of the challenge record or the reviewer workflow
    assert CH.confidentiality_scan([xlsx, docx]) == []


def test_e2e_every_round_recalculates_and_cross_reconciles(case):
    from cdf_gco_crms.qa.gates import GateContext, run_dash_gates
    files = case["files"]
    assert sorted(files) == [1, 2, 3, 5]
    for r in (3, 5):
        ctx = GateContext(xlsx=files[r], previous_xlsx=files[2 if r == 3 else 3], docx=case["docx"][r])
        res = {g.gate: g.status for g in run_dash_gates(ctx)}
        assert res["G-DASH-12"] == "PASS", res
        assert res["G-DASH-02"] == "PASS" and res["G-DASH-01"] == "PASS"


def test_e2e_blocker_is_cleared_only_by_a_verified_human_decision(case, tmp_path):
    """V712-043/044: a quote the document does not carry → WITHHELD_FOR_EVIDENCE blocks the issue; a boolean or an
    unsigned record changes nothing; a verified, signed reviewer decision clears it (19A.9)."""
    import base64
    from dataclasses import replace
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    from cdf_gco_crms.assurance.objection_engine import HumanDecision
    st1 = case["states"][1]
    victim = st1.ordered()[3].internal_id
    doc = build_reviewed_doc(st1, tmp_path / "short.docx", omit=(victim,))
    reg = EvidenceRegistry(); reg.add_document_text(doc)
    persona = RecipientPersona.from_index(ReferenceIndex.load(), DEPT, reg)
    rec = CH.run_challenge(st1, reg, persona)
    blk = rec.blockers()
    assert blk and {b["internal_id"] for b in blk} == {victim} and "CLAUSE_NOT_FOUND" in {b["kind"] for b in blk}
    res = CH.apply_human_decisions(rec, [True, {"approved": True}])
    assert len(res["rejected"]) == 2 and res["remaining_blockers"]
    unsigned = HumanDecision("D-1", victim, "WITHHELD_FOR_EVIDENCE", "م. مراجع (تجريبي)", "GCO reviewer", "DEC-1",
                             "2026-10-23", "E-1", "VERIFIED_HUMAN_RECORD", "SYNTHETIC_TEST_KEY", "")
    res = CH.apply_human_decisions(rec, [unsigned], {})
    assert res["rejected"] and res["remaining_blockers"]
    priv = Ed25519PrivateKey.generate()                       # ephemeral synthetic key: no CDF key is impersonated
    signed = replace(unsigned, signature_b64=base64.b64encode(priv.sign(unsigned.signature_payload())).decode())
    trust = {"SYNTHETIC_TEST_KEY": priv.public_key().public_bytes(serialization.Encoding.Raw, serialization.PublicFormat.Raw)}
    res = CH.apply_human_decisions(rec, [signed], trust)
    assert res["applied"] and res["remaining_blockers"] == []
    assert rec.to_dict()["human_decisions"][-1]["reviewer"] == "م. مراجع (تجريبي)"


@pytest.mark.skipif(not lo_engine.available() or not lo_engine.sockets_ok(), reason="LibreOffice not available")
def test_e2e_round5_word_renders_and_workbook_recalculates(case):
    from cdf_gco_crms.qa.recalc_verify import check_view, standard_views
    pdf = lo_engine.to_pdf(case["docx"][5], case["dir"] / "pdf")
    assert pdf is not None and pdf.stat().st_size > 1000
    st5 = case["states"][5]
    vc = check_view(st5, standard_views(st5)[0], case["dir"] / "verify")
    assert vc.ok, vc.message
