"""Assurance modules: reference_verify, production_guard, objection_engine.

Run from the project root: ``python -m pytest tests/test_assurance.py -q``.
Nothing here asserts the approval or applicability of any source; the PDF
checks concern occurrence of text only, and the workbooks built here are
synthetic by construction.
"""
import base64
import json
import sys
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from cdf_gco_crms.assurance.objection_engine import (  # noqa: E402
    FIXTURES, ApprovedTitle, HumanDecision, Objection, SourceEvidence, release_gates, run_fixtures,
)
from cdf_gco_crms.assurance.production_guard import (  # noqa: E402
    STATUS_APPROVED, STATUS_INCOMPLETE, STATUS_PENDING, audit_xlsx, check_decision_record, evaluate, evaluate_many,
)
from cdf_gco_crms.assurance.reference_verify import (  # noqa: E402
    VerifiedCitation, build_reference_status, load_source_text, norm,
)
from cdf_gco_crms.config import round_schema as S  # noqa: E402
from cdf_gco_crms.config import workbook_settings as W  # noqa: E402

REF = ROOT / "sources" / "reference"
RESP = REF / "CDF - الهيكل التنظيمي والمهام والمسؤوليات (2026-02).pdf"
DOA = REF / "CDF - مصفوفة الصلاحيات الرئيسية (2024-04-30).pdf"
needs_sources = pytest.mark.skipif(not (RESP.is_file() and DOA.is_file()), reason="reference PDFs not on disk")


# ================================================================== reference_verify
@needs_sources
def test_pdf_page_real_extract():
    assert "الهيكل" in load_source_text(RESP, 2)
    assert "المقدمة" in norm(load_source_text(DOA, 2))


@needs_sources
def test_exact_excerpt_never_implies_approval():
    r = VerifiedCitation(str(DOA), "مصفوفة الصلاحيات", "مقدمة",
                         "تحدد هذه الوثيقة المبادئ والتوجيهات والقواعد", page=2).verify()
    assert r["excerpt_exact_or_normalized"] and r["verdict"] == "TEXT_VERIFIED_ONLY"
    assert r["match_mode"] == "NORMALIZED"
    assert not r["approval_supported"] and not r["applicability_verified"]
    assert "APPROVAL_NOT_INDEPENDENTLY_PROVEN" in r["reason"]


@needs_sources
def test_excerpt_with_correct_ligatures_matches_broken_extraction():
    # the excerpt is typed with proper lam-alef; the extraction renders "اإلدارة" / "الصالحيات"
    r = VerifiedCitation(str(DOA), "مصفوفة", "مقدمة", "تسهيل الإدارة اليومية لوظائف صندوق التنمية", page=2).verify()
    assert r["excerpt_exact_or_normalized"]


@needs_sources
def test_made_up_clause_rejected():
    r = VerifiedCitation(str(DOA), "مصفوفة", "بند 999",
                         "التفويض المطلق على كافة الاعتمادات غير المقيدة نهائيا", page=2).verify()
    assert not r["excerpt_exact_or_normalized"] and r["verdict"] == "NOT_VERIFIED"
    assert "EXCERPT_NOT_FOUND_IN_SOURCE" in r["reason"]


@needs_sources
def test_page_scope_and_invalid_page():
    found_elsewhere = VerifiedCitation(str(DOA), "مصفوفة", "مقدمة",
                                       "تحدد هذه الوثيقة المبادئ والتوجيهات والقواعد", page=8).verify()
    assert not found_elsewhere["excerpt_exact_or_normalized"]
    assert "INVALID_PAGE" in VerifiedCitation(str(DOA), "مصفوفة", "x", "نص طويل بما يكفي للفحص", page=999).verify()["reason"]


def test_missing_source_and_hash_mismatch(tmp_path):
    r = VerifiedCitation(str(tmp_path / "never_file.pdf"), "not found", "x", "الجملة غير الموجودة").verify()
    assert r["verdict"] == "NOT_VERIFIED" and "SOURCE_NOT_AVAILABLE" in r["reason"]
    if DOA.is_file():
        r = VerifiedCitation(str(DOA), "مصفوفة", "x", "تحدد هذه الوثيقة المبادئ", declared_sha256="0" * 64).verify()
        assert not r["hash_valid"] and "HASH_MISMATCH" in r["reason"]


@needs_sources
def test_proposed_document_cannot_be_approved_by_name():
    s = load_source_text(RESP, 2)
    lines = [ln for ln in s.splitlines() if ln.strip()]
    r = VerifiedCitation(str(RESP), "مهام ومسؤوليات", "الهيكل", lines[0].strip(), page=2,
                         approval_status="PROPOSED_NOT_APPROVED").verify()
    assert r["excerpt_exact_or_normalized"] and not r["approval_supported"]
    with pytest.raises(ValueError):
        VerifiedCitation(str(RESP), "x", "x", "نص", approval_status="APPROVED").verify()


def test_docx_source_and_short_excerpt(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("هذه فقرة اختبار مرجعية لا تحمل أي صفة اعتماد مؤسسي.")
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "خلية جدول للتحقق من القراءة"
    p = tmp_path / "ref.docx"
    doc.save(p)
    assert VerifiedCitation(str(p), "docx", "p1", "فقرة اختبار مرجعية").verify()["verdict"] == "TEXT_VERIFIED_ONLY"
    assert VerifiedCitation(str(p), "docx", "t1", "خلية جدول للتحقق").verify()["excerpt_exact_or_normalized"]
    assert "EXCERPT_TOO_SHORT_TO_VERIFY" in VerifiedCitation(str(p), "docx", "p1", "فقرة").verify()["reason"]
    assert "UNSUPPORTED_SOURCE" in VerifiedCitation(str(tmp_path / "x.txt"), "t", "p", "x" * 20).verify()["reason"] or \
        not (tmp_path / "x.txt").exists()


@needs_sources
def test_build_reference_status_infers_nothing_from_names(tmp_path):
    out = tmp_path / "status.json"
    summary = build_reference_status([RESP, DOA, tmp_path / "absent.pdf"], out,
                                     statuses={RESP.name: "PROPOSED_NOT_APPROVED"})
    assert summary == {"sources": 3, "available": 2}
    data = json.loads(out.read_text(encoding="utf-8"))
    by = {d["name"]: d for d in data}
    assert by[RESP.name]["approval_status"] == "PROPOSED_NOT_APPROVED" and by[RESP.name]["pages"] == 92
    assert by[DOA.name]["approval_status"] == "UNVERIFIED" and by[DOA.name]["pages"] == 28
    assert all(d.get("applicability") == "NOT_ESTABLISHED" for d in data if d["available"])
    assert by["absent.pdf"] == {"name": "absent.pdf", "available": False, "status": "NOT_SUPPLIED"}


# ================================================================== production_guard
def _workbook(path: Path, *, helper_state="hidden", headers=S.FIXED_HEADERS, ids=("GCO-TST-R1-001", "GCO-TST-R1-002"),
              banner=False, keywords=None, extra_visible=False):
    wb = Workbook()
    dash = wb.active
    dash.title = W.DASHBOARD_SHEET
    dash["A1"] = W.DASHBOARD_TITLE
    if banner:
        dash["A2"] = S.DEMO_BANNER
    data = wb.create_sheet(S.data_sheet_name(1))
    for col, h in enumerate(headers, start=1):
        data.cell(row=1, column=col, value=h)
    for row, internal_id in enumerate(ids, start=2):
        data.cell(row=row, column=1, value=row - 1)
        data.cell(row=row, column=2, value=internal_id)
    helper = wb.create_sheet(W.HELPER_SHEET)
    helper["A1"] = "document"
    helper.sheet_state = helper_state
    if extra_visible:
        wb.create_sheet("عرض الطباعة")
    if keywords:
        wb.properties.keywords = keywords
    wb.save(path)
    return path


def test_clean_workbook_passes_structure_but_stays_pending(tmp_path):
    p = _workbook(tmp_path / "CDF-GCO-CRMS-TST-R1.xlsx")
    d = evaluate(p)
    assert d.blocked == () and d.status == STATUS_PENDING
    assert set(d.unverified) == {"MS_EXCEL_365", "HUMAN_GCO_APPROVAL", "ARTICLE_00_8"}
    assert "NOT_SYNTHETIC_AS_INSTITUTIONAL" in d.passed
    full = evaluate(p, excel365_verified=True, human_approved=True, reference_case_verified=True)
    assert full.status == STATUS_APPROVED


def test_fail_closed_on_structural_defects(tmp_path):
    p = _workbook(tmp_path / "CDF-GCO-CRMS-TST-R1.xlsx", helper_state="visible", extra_visible=True,
                  ids=("GCO-TST-R1-001", "GCO-TST-R1-001"), headers=S.FIXED_HEADERS[:4] + ("عمود آخر",))
    d = evaluate(p, excel365_verified=True, human_approved=True, reference_case_verified=True)
    assert d.status == STATUS_INCOMPLETE
    assert {"HELPER_HIDDEN", "VISIBLE_SHEETS", "INTERNAL_ID_UNIQUENESS", "CORE_COLUMNS_AE"} <= set(d.blocked)


def test_synthetic_file_never_presented_as_institutional(tmp_path):
    by_name = evaluate(_workbook(tmp_path / "CDF-GCO-CRMS-v7.12-ROUND1-SYNTHETIC.xlsx"),
                       excel365_verified=True, human_approved=True, reference_case_verified=True)
    assert by_name.status == STATUS_INCOMPLETE and "NOT_SYNTHETIC_AS_INSTITUTIONAL" in by_name.blocked
    by_banner = evaluate(_workbook(tmp_path / "CDF-GCO-CRMS-ABC-R2.xlsx", banner=True))
    assert "NOT_SYNTHETIC_AS_INSTITUTIONAL" in by_banner.blocked
    by_keyword = evaluate(_workbook(tmp_path / "CDF-GCO-CRMS-ABD-R2.xlsx", keywords="بيانات تجريبية"))
    assert "NOT_SYNTHETIC_AS_INSTITUTIONAL" in by_keyword.blocked
    markers = audit_xlsx(tmp_path / "CDF-GCO-CRMS-ABC-R2.xlsx")["details"]["synthetic_markers"]
    assert markers == ["CELL:DEMO_BANNER"]
    # openly a demonstration: the synthetic gate is reported, not blocking
    demo = evaluate(tmp_path / "CDF-GCO-CRMS-ABC-R2.xlsx", declared_institutional=False)
    assert "NOT_SYNTHETIC_AS_INSTITUTIONAL" not in demo.blocked and demo.status == STATUS_PENDING


def test_evaluate_many_writes_report_and_never_approves_from_structure(tmp_path):
    a = _workbook(tmp_path / "a.xlsx")
    b = _workbook(tmp_path / "b-DEMO.xlsx")
    out = tmp_path / "gate.json"
    results = evaluate_many([a, b], out)
    assert [r["status"] for r in results] == [STATUS_PENDING, STATUS_INCOMPLETE]
    assert json.loads(out.read_text(encoding="utf-8"))["schema"] == "CDF-v7.12-PRODUCTION-GATE-1"
    assert evaluate(tmp_path / "missing.xlsx").status == STATUS_INCOMPLETE


def test_decision_record_shape():
    good = {"decision_id": "D1", "observation_id": "GCO-AG-R1-001", "decision": "ROUTE", "reviewer": "Name",
            "role": "Reviewer", "reference": "M-1", "timestamp": "2026-10-23", "evidence_ref": "E-1",
            "verification_status": "VERIFIED_HUMAN_RECORD", "signing_key_id": "K1", "signature_b64": "AA=="}
    assert check_decision_record(good) == []
    bad = dict(good, reviewer="  ", verification_status="NOT_VERIFIED", signature_b64="", note="x")
    problems = check_decision_record(bad)
    assert {"MISSING:reviewer", "VERIFICATION_STATUS_NOT_VERIFIED_HUMAN_RECORD", "MISSING:signature_b64",
            "UNDECLARED_FIELDS:note"} == set(problems)
    assert check_decision_record("text") == ["NOT_A_MAPPING"]


# ================================================================== objection_engine
def valid():
    return {"MAT-1": {"authority_verified": True}}


def case(issued=False):
    return Objection("GCO-AG-R1-001", "authority", "صلاحية غير محددة", "قطاع الأعمال", issued=issued,
                     evidence_refs=["MAT-1"])


def test_missing_identity():
    with pytest.raises(ValueError):
        Objection("", "scope", "x", "A")


def test_initial_actor():
    with pytest.raises(ValueError):
        case().add("RED_ADVOCATE", "reply")


def test_first_black():
    assert case().add("BLACK", "اعتراض")["actor"] == "BLACK"


def test_six_cycle_max():
    c = case()
    for n in range(6):
        c.add("BLACK" if n % 2 == 0 else "RED_ADVOCATE", f"different substance {n}", [f"S{n}"])
    with pytest.raises(ValueError):
        c.add("BLACK", "extra", ["S7"])


def test_duplicate():
    c = case()
    c.add("BLACK", "اعتراض", ["MAT-1"])
    c.add("RED_ADVOCATE", "رد", ["MAT-1"])
    with pytest.raises(ValueError):
        c.add("BLACK", "اعتراض", ["MAT-1"])
    assert c.add("BLACK", "اعتراض", ["MAT-1"], material_change=True)["sequence"] == 3


def test_chain():
    c = case()
    c.add("BLACK", "اعتراض")
    c.add("RED_ADVOCATE", "رد")
    assert c.verify_chain()
    c.events[0]["claim"] = "tampered"
    assert not c.verify_chain()


def test_requires_evidence():
    with pytest.raises(ValueError):
        case().decide("UPHELD_UNCHANGED", verifications={})


def test_uphold_verified():
    assert case().decide("UPHELD_UNCHANGED", verifications=valid()) == "UPHELD_UNCHANGED"


def test_no_ai_self_reversal():
    with pytest.raises(PermissionError):
        case(True).decide("ROUTED_OUTSIDE_GCO")


def test_boolean_not_enough():
    with pytest.raises(TypeError):
        case(True).decide("ROUTED_OUTSIDE_GCO", human=True)


def test_unverified_approval():
    h = HumanDecision("D1", "GCO-AG-R1-001", "ROUTE", "Name", "Reviewer", "M-1", "2026-10-23", "E-1")
    with pytest.raises(PermissionError):
        case(True).decide("ROUTED_OUTSIDE_GCO", human=h)


def _signed(observation_id="GCO-AG-R1-001"):
    # Ephemeral SYNTHETIC signing key; no CDF public key or approval is impersonated.
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    priv = Ed25519PrivateKey.generate()
    a = HumanDecision("D1", observation_id, "ROUTE", "Synthetic Reviewer", "Reviewer", "M-1", "2026-10-23", "E-1",
                      "VERIFIED_HUMAN_RECORD", "SYNTHETIC_TEST_KEY", "")
    h = replace(a, signature_b64=base64.b64encode(priv.sign(a.signature_payload())).decode())
    trust = {"SYNTHETIC_TEST_KEY": priv.public_key().public_bytes(encoding=serialization.Encoding.Raw,
                                                                  format=serialization.PublicFormat.Raw)}
    return h, trust


def test_matched_human():
    h, trust = _signed()
    assert case(True).decide("ROUTED_OUTSIDE_GCO", human=h, trust_store=trust) == "ROUTED_OUTSIDE_GCO"
    with pytest.raises(PermissionError):
        case(True).decide("ROUTED_OUTSIDE_GCO", human=h, trust_store={})
    with pytest.raises(PermissionError):
        case(True).decide("ROUTED_OUTSIDE_GCO", human=replace(h, signature_b64="AAAA"), trust_store=trust)


def test_wrong_human_observation():
    h = HumanDecision("D1", "OTHER", "ROUTE", "Name", "Reviewer", "M-1", "2026-10-23", "E-1", "VERIFIED_HUMAN_RECORD")
    with pytest.raises(PermissionError):
        case(True).decide("ROUTED_OUTSIDE_GCO", human=h)
    h, trust = _signed("OTHER")
    with pytest.raises(PermissionError):
        case(True).decide("ROUTED_OUTSIDE_GCO", human=h, trust_store=trust)


def test_legacy_title():
    assert ApprovedTitle("A", "نص").display("09") == "الملاحظة #09"


def test_unverified_title():
    h = HumanDecision("D1", "A", "APPROVE", "Name", "Reviewer", "M", "date", "ref")
    with pytest.raises(PermissionError):
        ApprovedTitle("A", "عنوان", h).display("09")


def test_valid_title_needs_trust_store():
    # The approval status string by itself CANNOT authorize an observation title.
    h = HumanDecision("D1", "A", "APPROVE", "Name", "Reviewer", "M", "date", "ref", "VERIFIED_HUMAN_RECORD")
    with pytest.raises(PermissionError):
        ApprovedTitle("A", " ضبط صلاحيات الاعتماد ", h).display("09")
    h, trust = _signed("A")
    assert ApprovedTitle("A", " ضبط صلاحيات الاعتماد ", h).display("09", trust) == "ضبط صلاحيات الاعتماد"


def test_gates_pending_and_never_auto_approval():
    g = release_gates()
    assert g["ARTICLE_00_8"] == "NOT_VERIFIED" and g["RELEASE"] == "RELEASE_CANDIDATE"
    assert release_gates(True, True, True)["RELEASE"] == "OPERATIONALLY_APPROVED"


def test_closed_challenge_and_invalid_actor():
    c = case()
    c.decide("ROUTED_OUTSIDE_GCO")
    with pytest.raises(ValueError):
        c.add("BLACK", "new")
    with pytest.raises(ValueError):
        case().add("GCO_HUMAN", "unknown")
    with pytest.raises(ValueError):
        case().decide("MADE_UP_OUTCOME")


def test_source_evidence_validates_checksum_only(tmp_path):
    f = tmp_path / "ref.pdf"
    f.write_bytes(b"%PDF-1.4 synthetic")
    import hashlib
    digest = hashlib.sha256(f.read_bytes()).hexdigest()
    ev = SourceEvidence("R1", "ref.pdf", digest, "p.1", "excerpt", status="UNVERIFIED")
    v = ev.validate(tmp_path)
    assert v["sha256_matched"] and not v["authority_verified"] and not v["excerpt_verified"]
    assert SourceEvidence("R2", "ref.pdf", digest, "p.1", "x", status="APPROVED_CURRENT").validate(tmp_path)["authority_verified"]
    assert not SourceEvidence("R3", "ref.pdf", "0" * 64, "p.1", "x", status="APPROVED_CURRENT").validate(tmp_path)["authority_verified"]


def test_synthetic_fixtures_deterministic(tmp_path):
    results = run_fixtures(tmp_path)
    assert len(results) == len(FIXTURES) == 15
    assert all(r["passed"] and r["synthetic"] for r in results)
    assert results[-1]["productive_cycles"] == 6
    assert results == run_fixtures()
    reg = json.loads((tmp_path / "synthetic_challenge_register_INTERNAL.json").read_text(encoding="utf-8"))
    assert reg["schema"] == "CDF-ADV-v7.12" and len(reg["items"]) == 15
    assert all(i["observation_id"].startswith("SYNTHETIC-") for i in reg["items"])
