"""Mode isolation uses fictional private records and ephemeral authorization only."""
from io import BytesIO
import hashlib
import json
from pathlib import Path

import pymupdf
import pytest
from PIL import Image
from openpyxl import load_workbook
from cdf_gco_crms.calibration.service import CalibrationService
from cdf_gco_crms.calibration.store import CalibrationStore, ReviewerContext
from cdf_gco_crms.assurance.evidence import EvidenceRegistry
from cdf_gco_crms.intelligence.exports import export_dashboard
from cdf_gco_crms.institutional.contracts import AccessContext, InstitutionalAccessDenied
from cdf_gco_crms.institutional.isolation import require_public_runtime, bind_public_database
from cdf_gco_crms.institutional.providers import PublicSyntheticProvider, PrivateInstitutionalProvider, ReferenceLookupService, VersionRegistry
from cdf_gco_crms.institutional.workspace import PrivateWorkspace


class TestOnlyAuthorizer:
    __test__=False
    def __init__(self, permissions):self.permissions=set(permissions)
    def authorize(self, context, permission, resource):
        if context.session_handle!='SYN-TEST-SESSION' or permission not in self.permissions:
            raise InstitutionalAccessDenied('SYNTHETIC_TEST_ACCESS_DENIED')
        return ReviewerContext('SYN-PRIVATE-REVIEWER',frozenset({'CALIBRATION_ADMIN','CALIBRATION_REVIEWER'}),True)


@pytest.fixture
def private_registry(tmp_path):
    root=tmp_path/'private';root.mkdir();(root/'originals').mkdir()
    document=pymupdf.open();page=document.new_page();page.insert_text((72,72),'Synthetic authority approval');document.save(root/'originals/source.pdf');document.close()
    digest=hashlib.sha256((root/'originals/source.pdf').read_bytes()).hexdigest()
    source={'id':'SYN-PRIVATE-SOURCE','url':'https://example.org/private-fictional','sha256':digest,'version':'SYN-v1',
        'approval_status':'UNKNOWN','effective_date':'NOT_VERIFIED','supersession_checked':False,'current_authority_confirmed':False,
        'publication_status':'RESTRICTED_TEST_FIXTURE','extraction_status':'CANDIDATE','storage_location':'originals/source.pdf'}
    authority={'id':'SYN-PRIVATE-A','source_id':source['id'],'page':1,'number':'1','title':'Synthetic authority approval','category':'SYNTHETIC',
        'roles':{'Synthetic preparer':'الإعداد','Synthetic reviewer':'المراجعة','Synthetic approver':'الاعتماد'},'qualifiers':{},'approval_status':'UNKNOWN'}
    unit={'id':'SYN-PRIVATE-U','source_id':source['id'],'page':1,'name':'Synthetic authority approval','sector':'Fictional sector',
        'tasks':['Synthetic responsibility'],'internal_contacts':['Fictional liaison'],'external_contacts':[],'approval_status':'PROPOSED_NOT_APPROVED'}
    (root/'sources.json').write_text(json.dumps({'sources':[source]}));(root/'versions.json').write_text(json.dumps({'versions':[source]}))
    (root/'candidates.json').write_text(json.dumps({'sources':[source],'authorities':[authority],'units':[unit],'synthetic':False}))
    image=BytesIO();Image.new('RGB',(16,8),'white').save(image,format='PNG');(root/'originals/brand.png').write_bytes(image.getvalue())
    brand={'variant_id':'SYN-PRIVATE-BRAND','source_document':'SYN-PRIVATE-SOURCE','original_asset_version':'SYN-v1',
        'integrity_sha256':hashlib.sha256(image.getvalue()).hexdigest(),'languages':['ar'],'orientation':'HORIZONTAL',
        'intended_use':'AUTHORIZED_DEVELOPMENT','color_restrictions':'UNCHANGED','background_restrictions':'LIGHT',
        'approval_status':'NOT_VERIFIED','publication_status':'RESTRICTED','storage_location':'originals/brand.png'}
    (root/'brands.json').write_text(json.dumps({'assets':[brand]}))
    return root


def context(case='PRIVATE-CASE'):return AccessContext('SYN-TEST-SESSION',case)


def test_public_provider_is_independent_and_cannot_route_private_records(private_registry,monkeypatch):
    monkeypatch.setenv('CDF_INSTITUTIONAL_STORAGE',str(private_registry))
    provider=PublicSyntheticProvider();data=provider.catalog()
    assert data['synthetic'] and all(s['id'].startswith('SYN-') for s in data['sources'])
    assert 'SYN-PRIVATE-A' not in json.dumps(data)
    for key in ['SYN-PRIVATE-A','../private/candidates.json','SYN-PRIVATE-U']:
        with pytest.raises(InstitutionalAccessDenied):provider.private_record(key)
    with pytest.raises(InstitutionalAccessDenied):provider.asset('SYN-PRIVATE-BRAND')
    assert provider.asset('HORIZONTAL_LIGHT_BACKGROUND')
    with pytest.raises(TypeError):PublicSyntheticProvider(private_registry)


def test_private_default_denial_precedes_any_file_read(tmp_path,monkeypatch):
    def forbidden(*args,**kwargs):raise AssertionError('Storage touched before authorization')
    provider=PrivateInstitutionalProvider(tmp_path/'nonexistent')
    monkeypatch.setattr(Path,'read_text',forbidden)
    with pytest.raises(InstitutionalAccessDenied,match='AUTHENTICATION'):provider.catalog(context())
    with pytest.raises(InstitutionalAccessDenied):provider.asset('SYN',context())


def test_private_lookup_separates_approval_from_recommendation_and_unknowns(private_registry):
    provider=PrivateInstitutionalProvider(private_registry,TestOnlyAuthorizer(['reference.read']))
    lookup=ReferenceLookupService(provider)
    item=lookup.authorities('authority',context(),role='الاعتماد')[0]
    assert item['approver']==['Synthetic approver'] and item['reviewer']==['Synthetic reviewer']
    assert item['recommender'] is None and item['requester'] is None
    assert item['financial_threshold'] is None and not item['synthetic']
    unit=lookup.responsibilities('responsibility',context())[0]
    assert unit['reporting_relationship'] is None and unit['process_owner'] is None
    assert unit['documented_internal_interfaces']==['Fictional liaison']
    data=provider.catalog(context());assert data['currency'][0]['status']=='NOT_ESTABLISHED'
    assert not data['currency'][0]['binding_authority_granted']


def test_private_source_integrity_version_conflicts_and_storage_escape(private_registry):
    provider=PrivateInstitutionalProvider(private_registry,TestOnlyAuthorizer(['reference.read','brand.read']))
    sources=json.loads((private_registry/'sources.json').read_text())['sources']
    altered={**sources[0],'sha256':'a'*64}
    assert VersionRegistry.conflicts([*sources,altered])[0]['kind']=='SAME_VERSION_DIFFERENT_CONTENT'
    (private_registry/'versions.json').write_text(json.dumps({'versions':[*sources,altered]}))
    with pytest.raises(InstitutionalAccessDenied,match='VERSION_CONFLICT'):provider.catalog(context())
    (private_registry/'versions.json').write_text(json.dumps({'versions':sources}))
    (private_registry/'originals/source.pdf').write_bytes(b'TAMPERED SYNTHETIC FIXTURE')
    with pytest.raises(InstitutionalAccessDenied,match='INTEGRITY'):provider.catalog(context())
    with pytest.raises(InstitutionalAccessDenied,match='ESCAPE'):provider._path('../original-outside')


def test_private_export_and_evidence_use_same_scoped_authorization(private_registry,tmp_path):
    permissions=['case.import','case.read','case.export','case.calibrate','reference.read','reference.attach','brand.read']
    auth=TestOnlyAuthorizer(permissions);provider=PrivateInstitutionalProvider(private_registry,auth)
    store=CalibrationStore(tmp_path/'private.sqlite');workspace=PrivateWorkspace(CalibrationService(store),provider)
    workspace.import_case('PRIVATE-CASE','a'*64,[{'observation_id':'PRIVATE-O','original_text':'Synthetic private fixture'}],context())
    digest=store._case('PRIVATE-CASE')['evidence_hash']
    workspace.calibrate('PRIVATE-CASE','PRIVATE-O',{'decision':'KEEP_UNCHANGED','rationale':'Fictional authorized reviewer decision','expected_evidence_hash':digest,'status':'SUBMITTED'},context())
    data=workspace.dashboard('PRIVATE-CASE',context());assert not data['synthetic']
    assert next(m for m in data['metrics'] if m['id']=='calibration')['numerator']==1
    assert all(not m['synthetic'] for m in data['metrics'])
    with pytest.raises(PermissionError,match='AUTHORIZED_WORKSPACE'):export_dashboard(data,'xlsx')
    workbook=load_workbook(BytesIO(workspace.export('PRIVATE-CASE','xlsx',context(),brand_asset_id='SYN-PRIVATE-BRAND')))
    assert workbook['Provenance']['B3'].value=='PRIVATE INSTITUTIONAL'
    assert workbook['Provenance']['B2'].value=='NOT_AUTHORIZED'
    registry=EvidenceRegistry();refs=workspace.retrieve_into_evidence('PRIVATE-O','authority',registry,context())
    assert refs and all(x.approval_status=='UNVERIFIED' for x in registry.basis('PRIVATE-O'))
    auth.permissions.remove('case.export')
    with pytest.raises(InstitutionalAccessDenied):workspace.export('PRIVATE-CASE','docx',context(),brand_asset_id='SYN-PRIVATE-BRAND')
    with pytest.raises(InstitutionalAccessDenied,match='SCOPE'):workspace.dashboard('ANOTHER-CASE',context())
    with pytest.raises(InstitutionalAccessDenied,match='MODE_MISMATCH'):bind_public_database(store)
    store.close()


def test_public_mode_flags_never_activate_institutional_http_or_private_database(tmp_path):
    with pytest.raises(InstitutionalAccessDenied,match='ACTIVATION_DISABLED'):
        require_public_runtime('PRIVATE_INSTITUTIONAL',tmp_path/'public.sqlite')
    with pytest.raises(InstitutionalAccessDenied,match='PRIVATE_STORAGE'):
        require_public_runtime('PUBLIC_SYNTHETIC_DEMO',tmp_path/'private'/'case.sqlite')


def test_private_retrieval_runs_before_existing_agent_roles_and_cannot_grant_issue(private_registry,tmp_path):
    from docx import Document as WordDocument
    from cdf_gco_crms.model.entities import Document,GText
    from cdf_gco_crms.rounds.round_expansion import build_round1,NewObservation
    from cdf_gco_crms.references.synthetic_index import synthetic_reference_index
    from cdf_gco_crms.agents.mat_profile import load_profile,synthetic_fixture_path
    auth=TestOnlyAuthorizer(['case.import','case.read','case.run','reference.read','reference.attach'])
    provider=PrivateInstitutionalProvider(private_registry,auth)
    store=CalibrationStore(tmp_path/'review.sqlite');workspace=PrivateWorkspace(CalibrationService(store),provider)
    path=tmp_path/'fictional.docx';word=WordDocument();word.add_paragraph('Synthetic authority approval');word.save(path)
    workspace.import_case('PRIVATE-CASE',hashlib.sha256(path.read_bytes()).hexdigest(),[{'observation_id':'PRIVATE-O'}],context())
    state=build_round1(Document(name='Fictional private fixture',code='SYN',is_demo=False),
        [NewObservation('Paragraph 1','Synthetic authority approval',GText.plain('Synthetic authority approval'))],
        date='2026-10-24',document_version='SYN-v1')
    state.observations[0].internal_id='PRIVATE-O'
    execution=workspace.execute_review(state,path,context(),index=synthetic_reference_index(),reference_queries={'PRIVATE-O':'authority'},
        mat_profile=load_profile(synthetic_fixture_path()),gate_dir=tmp_path/'gate')
    policies=[e for e in execution.registry.basis('PRIVATE-O') if e.kind=='POLICY']
    assert policies and all(e.approval_status=='UNVERIFIED' and e.applicability=='NOT_ESTABLISHED' for e in policies)
    assert execution.mat_state and execution.final_review
    assert execution.trace_summary.get('actual_live_inference_count',0)==0
    with pytest.raises(InstitutionalAccessDenied):
        workspace.execute_review(state,path,context(),index=synthetic_reference_index(),reference_queries={},runner=object(),gate_dir=tmp_path/'gate')
    path.write_bytes(b'changed fictional document')
    with pytest.raises(InstitutionalAccessDenied,match='SOURCE_VERSION'):
        workspace.execute_review(state,path,context(),index=synthetic_reference_index(),reference_queries={},gate_dir=tmp_path/'gate')
    store.close()


def test_future_or_superseded_sources_never_gain_current_authority():
    from cdf_gco_crms.institutional.providers import currency
    source={'id':'SYN-CURRENCY','version':'SYN-v2','approval_status':'APPROVED_CURRENT',
        'effective_date':'2999-01-01','supersession_checked':True,'source_confidence':'INSTITUTIONALLY_VERIFIED_RECORD','superseded_by':'SYN-v3'}
    result=currency(source)
    assert 'SOURCE_NOT_YET_EFFECTIVE' in result['reasons'] and 'SOURCE_SUPERSEDED' in result['reasons']
    assert not result['binding_authority_granted']


def test_version_history_requires_explicit_reconciled_supersession():
    old={'id':'SYN-VERSION','version':'SYN-v1','sha256':'a'*64}
    new={'id':'SYN-VERSION','version':'SYN-v2','sha256':'b'*64}
    assert VersionRegistry.conflicts([old,new])
    reconciled={**old,'supersession_verified':True,'superseded_by':'SYN-v2'}
    assert VersionRegistry.conflicts([reconciled,new])==[]
    assert VersionRegistry.active([reconciled,new])==[new]
    collision={**reconciled,'version':'SYN-v2'}
    assert VersionRegistry.conflicts([collision,new])[0]['kind']=='SAME_VERSION_DIFFERENT_CONTENT'


def test_private_authorizer_cannot_grant_access_without_a_resolved_authenticated_principal(tmp_path):
    class InvalidAuthorizer:
        def __init__(self,value):self.value=value
        def authorize(self,*args):return self.value
    for value in (None,ReviewerContext('SYN-UNAUTHENTICATED',frozenset({'CALIBRATION_REVIEWER'}),False)):
        provider=PrivateInstitutionalProvider(tmp_path/'does-not-exist',InvalidAuthorizer(value))
        with pytest.raises(InstitutionalAccessDenied,match='PRINCIPAL_NOT_RESOLVED'):
            provider.catalog(context())
        with pytest.raises(InstitutionalAccessDenied,match='PRINCIPAL_NOT_RESOLVED'):
            provider.asset('SYN-PRIVATE-BRAND',context())
