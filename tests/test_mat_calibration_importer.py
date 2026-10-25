"""Private Word import contract tested exclusively on fictional DOCX documents."""
from pathlib import Path

import pytest
from docx import Document

from cdf_gco_crms.calibration.importer import build_case_from_docx, import_private_pair
from cdf_gco_crms.calibration.store import CalibrationStore, ReviewerContext, CalibrationError


def _doc(path, texts):
    d=Document()
    p=d.add_paragraph("SYNTHETIC TEST ONLY — clause concerning response period.")
    p2=d.add_paragraph("SYNTHETIC TEST ONLY — other clause.")
    if texts[0] is not None:
        d.add_comment(p.runs[0], text=texts[0], author="Fictional Team", initials="ST")
    if texts[1] is not None:
        d.add_comment(p2.runs[0], text=texts[1], author="Fictional Team", initials="ST")
    d.save(path)
    return path


def test_synthetic_pair_import_never_infers_withdrawal(tmp_path):
    a=_doc(tmp_path/"original.docx", ("Please clarify time limit", "Audit right risk"))
    b=_doc(tmp_path/"reviewed.docx", ("Response within one business day", None))
    report=build_case_from_docx(a,b)
    assert report["original_comment_objects"]==2
    assert report["reviewed_comment_objects"]==1
    assert report["reviewed_nonempty_comments"]==1
    assert len(report["observations"])>=2
    assert all(x["documented_withdrawal"] is False for x in report["observations"])
    assert any(x["observed_change"]=="REMOVED_FROM_VERSION" for x in report["observations"])
    store=CalibrationStore(tmp_path/"private"/"calibration.db")
    admin=ReviewerContext("SYNTHETIC_ADMIN",frozenset({"CALIBRATION_ADMIN","CALIBRATION_REVIEWER"}),True)
    receipt=import_private_pair(store,"SYNTHETIC-PAIR",a,b,admin)
    assert receipt["observations"]==len(report["observations"])
    assert store.verify_chain("SYNTHETIC-PAIR",admin)
    store.close()


def test_untrusted_docx_input_refused(tmp_path):
    fake=tmp_path/"original.docx"
    fake.write_text("this is not a DOCX archive")
    other=_doc(tmp_path/"real.docx",("fictional",None))
    with pytest.raises(CalibrationError,match="INVALID_DOCX_PACKAGE_SIGNATURE"):
        build_case_from_docx(fake,other)
    with pytest.raises(CalibrationError,match="SOURCE_VERSIONS_MUST_DIFFER"):
        build_case_from_docx(other,other)
