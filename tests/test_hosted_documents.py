"""Uploaded Word/Excel documents (CDF_DOCUMENT_UPLOADS=PUBLIC): safety refusals, owner isolation, engine tools.

Every file here is generated from the public synthetic fixtures; no institutional document is used.
"""
import importlib
import io
import sys
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("httpx")

from docx import Document  # noqa: E402

from cdf_gco_crms.model.entities import plain_text  # noqa: E402

TOKEN = "d" * 48


def headers():
    return {"Authorization": "Bearer " + TOKEN}


@pytest.fixture(scope="module")
def files(tmp_path_factory):
    from cdf_gco_crms.ingestion.workbook_reader import parse_workbook
    from cdf_gco_crms.regression import synthetic as SY

    directory = tmp_path_factory.mktemp("history")
    SY.demo_history(directory, rounds=1)
    workbook = next(directory.glob("*.xlsx"))
    state = parse_workbook(workbook).state
    word = Document()
    word.add_paragraph("وثيقة تجريبية عامة لاختبار رفع المستندات")
    for observation in state.ordered():
        word.add_paragraph(plain_text(observation.quote).replace("…", "").strip() + ".")
    buffer = io.BytesIO()
    word.save(buffer)
    return {"xlsx": workbook.read_bytes(), "docx": buffer.getvalue(), "observations": len(state.observations)}


def _client(tmp_path, monkeypatch, uploads):
    from fastapi.testclient import TestClient

    monkeypatch.setenv("CDF_CALIBRATION_MODE", "SYNTHETIC_LOCAL_DEV")
    monkeypatch.setenv("CDF_CALIBRATION_DEV_TOKEN", TOKEN)
    monkeypatch.setenv("CDF_CALIBRATION_DB", str(tmp_path / "public-demo" / "uploads.sqlite"))
    if uploads:
        monkeypatch.setenv("CDF_DOCUMENT_UPLOADS", "PUBLIC")
    else:
        monkeypatch.delenv("CDF_DOCUMENT_UPLOADS", raising=False)
    sys.modules.pop("apps.api.main", None)
    module = importlib.import_module("apps.api.main")
    return TestClient(module.app), module


@pytest.fixture()
def api(tmp_path, monkeypatch):
    client, module = _client(tmp_path, monkeypatch, True)
    with client:
        yield client, module
    module.store.close()
    sys.modules.pop("apps.api.main", None)


def upload(client, name, content):
    return client.post("/api/documents", params={"filename": name}, content=content,
                       headers={**headers(), "Content-Type": "application/octet-stream"})


def test_uploads_are_off_unless_declared_public(tmp_path, monkeypatch):
    client, module = _client(tmp_path, monkeypatch, False)
    with client:
        assert client.get("/api/health").json()["document_uploads"] == "OFF"
        assert client.get("/api/documents", headers=headers()).status_code == 404
    module.store.close()
    sys.modules.pop("apps.api.main", None)


def test_sign_in_required(api):
    client, _ = api
    assert client.get("/api/health").json()["document_uploads"] == "PUBLIC"
    assert client.get("/api/documents").status_code == 401
    response = client.post("/api/documents", params={"filename": "a.docx"}, content=b"x")
    assert response.status_code == 401


def test_upload_list_download_delete(api, files):
    client, _ = api
    created = upload(client, "مذكرة تجريبية.xlsx", files["xlsx"])
    assert created.status_code == 201, created.text
    meta = created.json()
    assert meta["kind"] == "xlsx" and meta["filename"] == "مذكرة تجريبية.xlsx"
    assert meta["summary"]["review_workbook"] is True
    assert meta["summary"]["observations"] == files["observations"]
    listed = client.get("/api/documents", headers=headers()).json()
    assert [d["document_id"] for d in listed] == [meta["document_id"]]
    content = client.get(f"/api/documents/{meta['document_id']}/content", headers=headers())
    assert content.status_code == 200 and content.content == files["xlsx"]
    assert "filename*=UTF-8''" in content.headers["content-disposition"]
    assert client.delete(f"/api/documents/{meta['document_id']}", headers=headers()).status_code == 204
    assert client.get("/api/documents", headers=headers()).json() == []


@pytest.mark.parametrize("name,content,code", [
    ("notes.txt", b"plain", "ONLY_DOCX_OR_XLSX_ACCEPTED"),
    ("broken.docx", b"not a zip at all", "UNSAFE_FILE_REFUSED"),
    ("empty.docx", b"", "EMPTY_FILE"),
])
def test_refused_files(api, name, content, code):
    client, _ = api
    response = upload(client, name, content)
    assert response.status_code == 422 and code in response.text


def _repack(original: bytes, extra: dict[str, bytes] | None = None, replace: dict[str, bytes] | None = None) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info, (replace or {}).get(info.filename, source.read(info.filename)))
        for name, data in (extra or {}).items():
            target.writestr(name, data)
    return out.getvalue()


def test_macro_and_dtd_packages_refused(api, files):
    client, _ = api
    macro = _repack(files["docx"], extra={"word/vbaProject.bin": b"\x00" * 10})
    assert "UNSAFE_FILE_REFUSED" in upload(client, "macro.docx", macro).text
    with zipfile.ZipFile(io.BytesIO(files["docx"])) as package:
        body = package.read("word/document.xml")
    xxe = body.replace(b"?>", b'?><!DOCTYPE w [<!ENTITY x SYSTEM "file:///etc/passwd">]>', 1)
    response = upload(client, "xxe.docx", _repack(files["docx"], replace={"word/document.xml": xxe}))
    assert response.status_code == 422 and "DTD_IN_word/document.xml" in response.text


def test_too_large_refused(api):
    client, module = api
    response = upload(client, "big.docx", b"0" * (module.hosted_documents.MAX_BYTES + 1))
    assert response.status_code == 413


def test_word_summary(api, files):
    client, _ = api
    meta = upload(client, "doc.docx", files["docx"]).json()
    summary = meta["summary"]
    assert summary["paragraphs"] == files["observations"] + 1 and summary["comments"] == 0
    assert summary["tracked_changes"] == 0 and summary["package"]["safe"] is True


def test_engine_tools(api, files):
    client, _ = api
    workbook = upload(client, "review.xlsx", files["xlsx"]).json()["document_id"]
    document = upload(client, "owner.docx", files["docx"]).json()["document_id"]

    def act(source, action, other=None):
        response = client.post(f"/api/documents/{source}/actions/{action}", headers=headers(),
                               params={"other": other} if other else None)
        assert response.status_code == 200, response.text
        return response.json()

    qa = act(workbook, "qa")
    assert qa["output"] is None and sum(qa["result"]["counts"].values()) == len(qa["result"]["gates"]) > 0
    report = act(workbook, "word_report")
    assert report["result"]["word_matches_workbook"] is True
    assert report["output"]["filename"].startswith("INTERNAL_DRAFT__") and report["output"]["kind"] == "docx"
    assert report["output"]["derived_from"] == workbook
    rebuilt = act(workbook, "excel_rebuild")
    assert rebuilt["output"]["summary"]["observations"] == files["observations"]

    commented = act(document, "comment", workbook)
    assert commented["result"]["anchored"] == files["observations"] and commented["result"]["ok"] is True
    output = commented["output"]
    assert output["summary"]["comments"] == files["observations"]
    assert output["summary"]["review_comments"] == files["observations"]
    # the pair may be given either way round
    assert act(workbook, "comment", document)["result"]["anchored"] == files["observations"]

    compared = act(output["document_id"], "compare", document)
    assert compared["result"]["summary"]["missing_gco_comments"]       # the returned copy has no review comments

    names = {d["filename"] for d in client.get("/api/documents", headers=headers()).json()}
    assert len(names) >= 5


def test_wrong_inputs_for_tools(api, files):
    client, _ = api
    document = upload(client, "owner.docx", files["docx"]).json()["document_id"]
    workbook = upload(client, "review.xlsx", files["xlsx"]).json()["document_id"]
    for source, action, other, code in [
        (document, "qa", None, "REVIEW_WORKBOOK_REQUIRED"),
        (document, "comment", None, "SECOND_DOCUMENT_REQUIRED"),
        (document, "comment", document, "WORD_DOCUMENT_AND_REVIEW_WORKBOOK_REQUIRED"),
        (workbook, "compare", document, "TWO_WORD_DOCUMENTS_REQUIRED"),
        (document, "delete_everything", None, "UNKNOWN_DOCUMENT_ACTION"),
        (document, "compare", "DOC-" + "0" * 32, "DOCUMENT_NOT_FOUND"),
    ]:
        response = client.post(f"/api/documents/{source}/actions/{action}", headers=headers(),
                               params={"other": other} if other else None)
        assert response.status_code in (403, 422) and code in response.text, (action, response.text)


def test_owner_isolation(tmp_path, files):
    from cdf_gco_crms.calibration.store import AccessDenied, CalibrationStore
    from cdf_gco_crms.interfaces.hosted_documents import DocumentLibrary

    store = CalibrationStore(Path(tmp_path) / "iso.sqlite")
    library = DocumentLibrary(store)
    mine = library.upload("SYN-HOSTED-A", "a.docx", files["docx"])
    assert library.list("SYN-HOSTED-B") == []
    for call in (library.get, library.content, library.delete):
        with pytest.raises(AccessDenied):
            call("SYN-HOSTED-B", mine["document_id"])
    with pytest.raises(AccessDenied):
        library.prepare("SYN-HOSTED-B", mine["document_id"], "qa")
    log = store.db.execute("SELECT event FROM uploaded_document_log").fetchall()
    assert [r[0] for r in log] == ["CREATED"]
    store.close()


def test_filename_cleaning():
    from cdf_gco_crms.calibration.store import CalibrationError
    from cdf_gco_crms.interfaces.hosted_documents import clean_filename

    assert clean_filename("../../etc/مذكرة\x00 نهائية.DOCX") == ("مذكرة نهائية.docx", "docx")
    assert clean_filename("C:\\Users\\x\\file.xlsx") == ("file.xlsx", "xlsx")
    for bad in ("file.xlsm", "file", "file.docx.exe"):
        with pytest.raises(CalibrationError):
            clean_filename(bad)
