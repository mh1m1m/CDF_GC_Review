"""TEST 24 — brand library of v7.12 (Annex D.11, D.12; gates G-UX-01, G-UX-04)."""
from __future__ import annotations

import re
from pathlib import Path

from cdf_gco_crms.branding import brand_tokens as BT
from cdf_gco_crms.branding.contrast import contrast_ratio
from cdf_gco_crms.branding.guide_page_check import PRINTED_PAGE_12, check_page12, page_surface
from cdf_gco_crms.config import design_tokens as T
from cdf_gco_crms.config import round_schema as S
from cdf_gco_crms.regression import registers as RG

ROOT = Path(__file__).resolve().parents[1]
PAGE12 = ROOT / "sources" / "v7.12" / "guide_page12.png"


def test_page12_printed_hex_matches_render():
    checks = check_page12(PAGE12)
    assert len(checks) == 7 and all(c.equal for c in checks), [(c.name, c.printed, c.sampled) for c in checks]
    assert page_surface(PAGE12) == T.PAGE_SURFACE


def test_every_brand_token_is_printed_on_page12_or_sampled():
    printed = {hx for _, hx in PRINTED_PAGE_12}
    for t in BT.BRAND_TOKENS:
        assert t.verification in (BT.HEX_PRINTED, BT.RGB_SPECIFIED)
        if t.verification == BT.HEX_PRINTED and t.hex not in (T.PAGE_SURFACE, T.WHITE):
            assert t.hex in printed, t


def test_status_fills_equal_the_issued_register():
    """AMB-161: the dashboard shows a status in the very fill of every issued register (D.2)."""
    assert T.STATUS_FILL == {S.CLOSED: "E2EFDA", S.CLOSED_COND: "DDEBF7", S.REQUIRES_EDIT: "FFF2CC", S.OPEN: "FCE4D6"}
    assert {t.hex for t in BT.SEMANTIC_FILLS} == set(T.STATUS_FILL.values())
    for k, v in (("closed", S.CLOSED), ("cond", S.CLOSED_COND), ("req", S.REQUIRES_EDIT), ("open", S.OPEN)):
        assert T.CARD_CAP[k] == T.STATUS_FILL[v]
    assert T.ROUND1_OPEN_FILL == "F2F2F2"
    reg = ROOT / "sources" / "reference" / "issued_register_round4.xlsx"
    if reg.exists():                                                  # the register supplied by the protocol owner
        from openpyxl import load_workbook
        ws = load_workbook(reg).worksheets[0]
        fills = {}
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in T.STATUS_FILL and c.fill.fill_type == "solid":
                    fills.setdefault(c.value.strip(), set()).add(c.fill.fgColor.rgb[-6:])
        assert fills == {k: {v} for k, v in T.STATUS_FILL.items()}


def test_derived_tones_reach_aa_on_white_and_are_not_brand():
    brand = {t.hex for t in BT.BRAND_TOKENS}
    for v, hx in T.STATUS_TEXT.items():
        assert contrast_ratio(hx, T.WHITE) >= 4.5, (v, hx)
        assert hx not in brand and hx not in T.STATUS_FILL.values()
    assert {t.hex for t in BT.DERIVED_TONES} == set(T.STATUS_TEXT.values())


def test_ink_on_every_brand_fill_and_caps():
    for hx in list(T.STATUS_FILL.values()) + list(T.CARD_CAP.values()) + [T.TITLE_BAND, T.ACCENT, T.LABEL_FILL]:
        assert contrast_ratio(T.BRAND_INK, hx) >= 4.5, hx
    assert contrast_ratio(T.MUTED_TEXT, T.TITLE_BAND) >= 4.5
    assert contrast_ratio(T.MUTED_TEXT, T.PAGE_SURFACE) >= 4.5
    assert contrast_ratio(T.LINK, T.PAGE_SURFACE) >= 4.5


def test_contrast_register_has_no_fail():
    rows = RG.contrast_rows()
    assert rows and not [r for r in rows if r[4] == "FAIL"]


def test_no_unverified_claim_and_superseded_values_recorded():
    assert BT.UNVERIFIED_CLAIMS == ()
    sup = {t.hex for t in BT.SUPERSEDED_CONVERSIONS}
    assert sup == {"D8D3CB", "A7B0D9", "FED186", "F69271"}
    for name in dir(T):
        v = getattr(T, name)
        vals = v.values() if isinstance(v, dict) else [v]
        for x in vals:
            if isinstance(x, str) and re.fullmatch(r"[0-9A-F]{6}", x):
                assert x not in sup, (name, x)


def test_type_scale_floor():
    for name in dir(T):
        if name.startswith("SZ_"):
            assert getattr(T, name) >= T.MIN_FONT
    assert (T.SZ_TITLE, T.SZ_SECTION, T.SZ_CARD_VALUE, T.SZ_DOC) == (24, 16, 32, 14)
