"""TEST 28 — title workflow (10.4A.16, 19A.9; patches V712-011/012/013/043)."""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from cdf_gco_crms.config import round_schema as S
from cdf_gco_crms.config import workbook_settings as W
from cdf_gco_crms.excel.workbook_builder import build_workbook
from cdf_gco_crms.ingestion.workbook_reader import parse_workbook
from cdf_gco_crms.model.entities import Observation, RoundCells, TitleEvidence
from cdf_gco_crms.observations import text_derivation as TD
from cdf_gco_crms.regression import synthetic as SY
from cdf_gco_crms.rounds.round_expansion import (NewObservation, TitleDecision, build_round1, finalize_titles,
                                                 mark_titles_issued)

EV = TitleEvidence("م. مراجع الحوكمة", "2026-10-23", "GCO-DEC-2026-017")


def _state(n=4, with_titles=True):
    obs = []
    for i in range(n):
        o = SY.topic_obs(i)
        if with_titles:
            o = NewObservation(o.location, o.quote, o.comment, title=f"عنوان مقترح رقم {i + 1}")
        obs.append(o)
    return build_round1(SY.fixture_document("TTL"), obs, date="2026-10-23")


def test_boolean_is_not_a_decision():
    with pytest.raises(ValueError):
        Observation("GCO-TTL-R1-001", "01", 1, "ص.1", "نص", "تعليق.", cells={1: RoundCells(closure=S.OPEN)},
                    title="عنوان", title_status=S.TITLE_APPROVED)      # approved without evidence
    with pytest.raises(ValueError):
        Observation("GCO-TTL-R1-001", "01", 1, "ص.1", "نص", "تعليق.", cells={1: RoundCells(closure=S.OPEN)},
                    title="عنوان", title_status=S.TITLE_APPROVED,
                    title_evidence=TitleEvidence("", "2026-10-23", "X"))
    assert not TitleEvidence("م", "23/10/2026", "X").complete           # a date must be ISO
    assert EV.complete


def test_proposed_title_is_never_shown_as_approved():
    st = _state()
    o = st.observations[0]
    assert o.title_status == S.TITLE_PROPOSED
    shown = TD.display_title(o.title, o.title_status, "نوصي بتقييد الصلاحية.")
    assert shown.startswith(S.TITLE_FALLBACK_TAG) and o.title not in shown
    assert TD.display_title(None, S.TITLE_NONE, "") == S.TITLE_FALLBACK_TAG


def test_finalize_titles_requires_human_proof_and_keeps_ids():
    st = _state()
    ids = [o.internal_id for o in st.observations]
    good = TitleDecision(ids[0], "APPROVE", EV.reviewer, EV.date, EV.decision_ref)
    bad = TitleDecision(ids[1], "APPROVE", "", "", "")                   # no reviewer: refused
    reject = TitleDecision(ids[2], "REJECT", EV.reviewer, EV.date, "GCO-DEC-2026-018")
    unknown = TitleDecision("GCO-TTL-R1-099", "APPROVE", EV.reviewer, EV.date, "X")
    new, rep = finalize_titles(st, [good, bad, reject, unknown])
    assert rep.approved == [ids[0]] and rep.rejected == [ids[2]]
    assert len(rep.refused) == 2 and "19A.9" in rep.refused[0]
    assert [o.internal_id for o in new.observations] == ids             # IDs untouched (V712-012)
    assert [o.title_status for o in st.observations] == [S.TITLE_PROPOSED] * 4   # input state not mutated
    by = new.by_id()
    assert by[ids[0]].title_status == S.TITLE_APPROVED and by[ids[0]].title_evidence == EV
    assert by[ids[1]].title_status == S.TITLE_PROPOSED
    assert by[ids[2]].title_status == S.TITLE_DRAFT
    # issued titles are immutable except by a recorded withdrawal
    assert mark_titles_issued(new) == 1
    new2, rep2 = finalize_titles(new, [TitleDecision(ids[0], "APPROVE", EV.reviewer, EV.date, "Y", title="آخر")])
    assert rep2.refused and new2.by_id()[ids[0]].title == "عنوان مقترح رقم 1"
    new3, rep3 = finalize_titles(new, [TitleDecision(ids[0], "WITHDRAW", EV.reviewer, EV.date, "GCO-DEC-2026-019")])
    assert rep3.withdrawn == [ids[0]] and new3.by_id()[ids[0]].title_status == S.TITLE_NONE


def test_title_survives_rebuild_and_dashboard_shows_only_approved(tmp_path):
    st = _state()
    ids = [o.internal_id for o in st.observations]
    st, _ = finalize_titles(st, [TitleDecision(ids[1], "APPROVE", EV.reviewer, EV.date, EV.decision_ref)])
    mark_titles_issued(st)
    out = tmp_path / "t.xlsx"
    build_workbook(st, out)
    pw = parse_workbook(out)
    by = pw.state.by_id()
    assert by[ids[1]].title_status == S.TITLE_ISSUED and by[ids[1]].title_evidence == EV
    assert by[ids[0]].title_status == S.TITLE_PROPOSED and by[ids[0]].title == "عنوان مقترح رقم 1"
    # the dashboard table is title-led: header row carries the seven columns of 10.4A.11
    ws = load_workbook(out)[W.DASHBOARD_SHEET]
    heads = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "عنوان الملاحظة" in heads and "ملخص الملاحظة" not in heads and "الموضع عند الإصدار" in heads
    hs = load_workbook(out)[W.HELPER_SHEET]
    hdr = {hs.cell(row=1, column=c).value: c for c in range(1, hs.max_column + 1)}
    disp = [hs.cell(row=r, column=hdr["obs.title_disp"]).value for r in range(2, 6)]
    assert disp[1] == "عنوان مقترح رقم 2"
    assert all(d.startswith(S.TITLE_FALLBACK_TAG) for i, d in enumerate(disp) if i != 1)


def test_tampered_approval_is_demoted_on_read(tmp_path):
    st = _state()
    ids = [o.internal_id for o in st.observations]
    st, _ = finalize_titles(st, [TitleDecision(ids[0], "APPROVE", EV.reviewer, EV.date, EV.decision_ref)])
    out = tmp_path / "t.xlsx"
    build_workbook(st, out)
    wb = load_workbook(out)
    hs = wb[W.HELPER_SHEET]
    hdr = {hs.cell(row=1, column=c).value: c for c in range(1, hs.max_column + 1)}
    hs.cell(row=3, column=hdr["obs.title_status"]).value = S.TITLE_APPROVED    # forged status, no evidence
    wb.save(out)
    pw = parse_workbook(out)
    assert pw.state.by_id()[ids[1]].title_status == S.TITLE_PROPOSED
    assert any("demoted to PROPOSED" in w for w in pw.warnings)


def test_seven_columns_in_three_and_five_round_workbooks(five_round):
    """V712-048: code and protocol agree on the table definition for 3- and 5-round workbooks."""
    h = five_round[0]
    for r in (3, 5):
        ws = load_workbook(h.files[r])[W.DASHBOARD_SHEET]
        for row in ws.iter_rows():
            vals = [c.value for c in row if isinstance(c.value, str)]
            if "عنوان الملاحظة" in vals and "التغير" in vals:
                assert vals[:2] == ["الملاحظة #", "عنوان الملاحظة"] and vals[-3:] == ["التغير", "الإجراء المطلوب", "الانتقال"]
                break
        else:
            raise AssertionError(f"round {r}: title-led table header not found")
        assert "ملخص الملاحظة" not in {c.value for row in ws.iter_rows() for c in row}
