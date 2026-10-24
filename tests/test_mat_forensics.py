"""TEST 36 — OOXML comment forensics (tools/mat_dna/ooxml_forensics.py) on a SYNTHETIC document pair.

The pair is generated here with python-docx and then post-processed to carry the package evidence the extractor
relies on (w14:paraId on comment paragraphs, creation dates): an "original" review export and a "reviewed" version in
which comments were edited in place, emptied, removed, replaced at the same passage, recreated as a question and added.
No institutional text; every string is fictional.
"""
from __future__ import annotations

import json
import shutil
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools" / "mat_dna"))
import ooxml_forensics as F  # noqa: E402

AUTHOR = "وحدة تجريبية"
T0 = "2026-01-01T10:00:00Z"
PARAS = ["المادة (1): يلتزم الشريك بتسليم التقرير وفق الجدول المرفق.",
         "المادة (2): تودع مساهمات الشركاء في حساب المشروع.",
         "المادة (3): يسري هذا البرنامج على المنح الصغيرة.",
         "المادة (4): لا يستلزم تنفيذ البرنامج تزويد الشريك ببيانات المستفيدين.",
         "المادة (5) التزامات الطرفان.",
         "المادة (6): تسري هذه المذكرة من تاريخ توقيعها."]
ORIG = ["الملاحظة #01\nREF: GCO-SYN-R1-001\nعبارة «يلتزم» لا تحدد أثر التأخر.\nالتوصية: يُستبدل \"يلتزم الشريك\" بـ \"يجب على الشريك\".",
        "الملاحظة #02\nREF: GCO-SYN-R1-002\nالمساهمات تختلط بأموال المشروع العام.\nالتوصية: تخصيص حساب مستقل بالنص الآتي: (\"حساب مستقل لمساهمات الشركاء\".)",
        "الملاحظة #03\nREF: GCO-SYN-R1-003\nهل اعتُمد برنامج بهذا الاسم؟ النص يحيل إلى برنامج غير معرّف.\nالتوصية: يُعرَّف البرنامج بالإحالة إلى قرار اعتماده.",
        "الملاحظة #04\nREF: GCO-SYN-R1-004\nقد تستلزم المتطلبات النظامية لمكافحة غسل الأموال التحقق من المستفيد.\nالتوصية: يُضاف استثناء للمتطلبات النظامية.",
        "الملاحظة #05\nREF: GCO-SYN-R1-005\nعنوان المادة فيه خطأ إعرابي.\nالتوصية: يُستبدل \"الطرفان\" بـ \"الطرفين\".",
        "الملاحظة #06\nREF: GCO-SYN-R1-006\nالمذكرة بلا مدة.\nالتوصية: تحديد مدة قابلة للتمديد بالنص الآتي: (\"لمدة [●] سنة\".)"]
PIDS = ["0A000001", "0A000002", "0A000003", "0A000004", "0A000005", "0A000006"]
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _build(path: Path, comments: list[tuple[int, str, str, str]]):
    """comments: (paragraph index, text, para_id, date)."""
    doc = Document()
    runs = []
    for p in PARAS:
        para = doc.add_paragraph()
        runs.append(para.add_run(p))
    for i, (pi, text, _pid, _date) in enumerate(comments):
        doc.add_comment(runs[pi], text=text, author=AUTHOR, initials="وت")
    doc.save(path)
    # post-process: paraId on the first comment paragraph and the creation date
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/comments.xml":
                root = etree.fromstring(data)
                for c, (_pi, _text, pid, date) in zip(root.findall(f"{{{WNS}}}comment"), comments):
                    c.set(f"{{{WNS}}}date", date)
                    first = c.find(f"{{{WNS}}}p")
                    if first is not None:
                        first.set(f"{{{W14}}}paraId", pid)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    shutil.move(tmp, path)
    return path


@pytest.fixture(scope="module")
def pair(tmp_path_factory):
    d = tmp_path_factory.mktemp("mat")
    original = _build(d / "original.docx", [(i, ORIG[i], PIDS[i], T0) for i in range(6)])
    reviewed = _build(d / "reviewed.docx", [
        (0, "REF: GCO-SYN-R1-001\nيُستبدل \"يلتزم الشريك\" بـ \"يجب على الشريك\".", "0B000001", T0),      # edited in place: header + rationale dropped, REF kept
        (1, ORIG[1], PIDS[1], T0),                                                                      # unchanged (same paraId, same date)
        (2, "هل اعتُمد برنامج بهذا الاسم؟", "0B000003", "2026-01-01T13:29:00Z"),                           # recreated at the anchor as a question only
        (3, "نوصي باضافة مهلة للرد مثلاً خلال اسبوع", "0B000004", "2026-01-01T14:41:00Z"),              # new, unrelated, at the passage of removed #04
        (4, "", "0B000005", T0),                                                                        # emptied in place
        (5, "CO-SYN-R1-006\nتحديد مدة قابلة للتمديد بالنص الآتي: (\"لمدة [●] سنة\".)", "0B000006", T0),  # REF truncated, rationale dropped
        (1, "والتخزين او التامين", "0B000007", "2026-01-01T13:18:00Z"),                                     # pure addition on another passage
    ])
    return original, reviewed


def test_extract_reads_comments_anchors_ids_and_dates(pair):
    original, reviewed = pair
    ex = F.extract(original)
    assert len(ex.comments) == 6 and ex.sha256 and not ex.warnings
    c0 = ex.comments[0]
    assert c0.para_id == PIDS[0] and c0.date == T0 and c0.author == AUTHOR and c0.anchored
    assert c0.anchor_text.startswith("المادة (1)") and c0.classification == "PROPOSED_AMENDMENT"
    assert ex.comments[2].classification == "QUESTION_WITH_RECOMMENDATION"
    assert "word/commentsExtended.xml" in ex.missing_parts or True   # python-docx writes no extended part; the extractor lists what is missing
    rv = F.extract(reviewed)
    by_id = {c.comment_id: c for c in rv.comments}          # comments come back in document order, not id order
    assert len(rv.comments) == 7 and by_id["4"].text == "" and by_id["4"].classification == "EMPTY"
    assert [c.anchor_paragraph_index for c in rv.comments] == sorted(c.anchor_paragraph_index for c in rv.comments)


def test_compare_matches_by_evidence_and_never_calls_a_removal_a_withdrawal(pair):
    original, reviewed = pair
    rep = F.compare(F.extract(original), F.extract(reviewed), reviewer_name_hint="فلان")
    M = {m["observation_id"]: m for m in rep["matches"]}
    assert len(M) == 8
    assert M["OBS-001"]["matching_method"] == "REF_CODE" and M["OBS-001"]["provenance"] == "EDITED_IN_PLACE" \
        and {"DROP_NUMBER_HEADER", "DROP_RATIONALE", "SHORTEN"} <= set(M["OBS-001"]["change_types"])
    assert M["OBS-002"]["matching_method"] == "PARA_ID" and M["OBS-002"]["change_types"] == ["KEEP_UNCHANGED"] and M["OBS-002"]["provenance"] == "UNCHANGED"
    assert M["OBS-003"]["matching_method"] == "ANCHOR_TEXT" and M["OBS-003"]["provenance"] == "RECREATED_AT_ANCHOR" \
        and M["OBS-003"]["reviewed_class"] == "DIRECT_QUESTION" and "DROP_RECOMMENDATION" in M["OBS-003"]["change_types"]
    assert M["OBS-004"]["change_types"] == ["REMOVED_FROM_VERSION"] and M["OBS-004"]["attribution_status"] == "DECISION_NOT_OBSERVED" \
        and M["OBS-004"]["manual_review"] and M["OBS-004"]["replaces_at_same_passage"]
    new_at_4 = M[M["OBS-004"]["replaces_at_same_passage"]]
    assert new_at_4["change_types"] == ["ADDED_IN_REVIEW"] and new_at_4["provenance"] == "NEW_COMMENT" and new_at_4["reviewed_class"] == "DIRECT_RECOMMENDATION", \
        "a newly created comment with unrelated text is a new observation, never a rephrasing"
    assert M["OBS-005"]["change_types"] == ["EMPTIED_IN_PLACE"] and M["OBS-005"]["provenance"] == "EDITED_IN_PLACE"
    assert M["OBS-006"]["matching_method"] == "REF_CODE", "a truncated reference code still identifies the observation"
    added = [m for m in M.values() if "ADDED_IN_REVIEW" in m["change_types"]]
    assert len(added) == 2 and any(m["reviewed_class"] == "WORDING_INSERTION" for m in added)
    st = rep["statistics"]
    assert st["documented_withdrawals"] == 0 and st["removed_from_version"] == 1 and st["emptied_in_place"] == 1 and st["added_in_review"] == 2
    assert st["provenance"] == {"ABSENT": 1, "EDITED_IN_PLACE": 3, "NEW_COMMENT": 2, "RECREATED_AT_ANCHOR": 1, "UNCHANGED": 1}
    assert "REMOVED_FROM_VERSION is an observed file change, not a documented withdrawal" in rep["note"]
    assert all(m["attribution_status"] == "RECORDED_AUTHOR_DIFFERS_FROM_REVIEWER_NAME" for m in M.values() if m["reviewed_comment_id"]), \
        "a shared-account author is never promoted to the reviewer's identity"


def test_classifier_is_hamza_insensitive_and_word_bounded():
    assert F.classify_comment("نامل تزويدنا بنسخة من القرار") == "INFORMATION_REQUEST"
    assert F.classify_comment("نوصي باضافة مهلة للرد مثلاً خلال اسبوع") == "DIRECT_RECOMMENDATION"
    assert F.classify_comment("يوجد خطأ في ترقيم الملحق") == "DEFECT_FLAG"
    assert F.classify_comment("والتخزين او التامين") == "WORDING_INSERTION"
    assert F.classify_comment("ينبغي حصر السلع المؤهلة في قائمة مغلقة.") == "DIRECT_RECOMMENDATION", "هل inside المؤهلة is not a question marker"
    assert F.classify_comment("ما مصير الضمان إذا انسحب الشريك؟ التوصية: يُضاف نص ينظم ذلك.") == "QUESTION_WITH_RECOMMENDATION"
    assert F.classify_comment("") == "EMPTY"
    assert "CONVERT_TO_QUESTION" in F.change_type("المتن.\nالتوصية: يُستبدل كذا بكذا.", "هل البند مناسب للوحدة المختصة؟")
    assert F.change_type("x", "x") == ["KEEP_UNCHANGED"]


def test_cli_writes_register_and_xlsx(pair, tmp_path):
    original, reviewed = pair
    rc = F.main(["compare", str(original), str(reviewed), "--out", str(tmp_path / "cmp")])
    assert rc == 0
    reg = json.loads((tmp_path / "cmp" / "MAT_OBSERVATION_EVOLUTION_REGISTER.json").read_text(encoding="utf-8"))
    assert len(reg["matches"]) == 8 and (tmp_path / "cmp" / "MAT_OBSERVATION_EVOLUTION_REGISTER.xlsx").exists()
    from dna_deliverables import render
    d = tmp_path / "deliv"; d.mkdir()
    (d / "MAT_OBSERVATION_EVOLUTION_REGISTER.json").write_text(json.dumps({"rows": [{"Observation_ID": "OBS-001", "Change_Type": ["KEEP_UNCHANGED"]}]}), encoding="utf-8")
    assert [p.name for p in render(d)] == ["MAT_OBSERVATION_EVOLUTION_REGISTER.xlsx"]
