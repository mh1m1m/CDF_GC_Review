"""Word lifecycle on REAL documents (V712-028..036): comment_document, diff_returned, closed_text_integrity,
check_package, atomic_output. The reviewed document is built with python-docx from the fixture quotes: some
inside a table, one quote split over two runs with different formatting, one partly inside a raw w:ins (with a
w:del beside it), one written with a lam-alef ligature and doubled spaces, and one pre-existing comment by
another author."""
import hashlib
import json
import re
import shutil
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from lxml import etree

from conftest import C, CC, O, RE
from cdf_gco_crms.config import round_schema as S
from cdf_gco_crms.model.entities import plain_text
from cdf_gco_crms.qa import lo_engine
from cdf_gco_crms.word import document_comments as DC
from cdf_gco_crms.word import word_sync as WS

needs_lo = pytest.mark.skipif(not lo_engine.available() or not lo_engine.sockets_ok(),
                              reason="LibreOffice not available")
OWNER = "إدارة الشراكات"
W = DC.w
SPLIT_K, TABLE_K, FOREIGN_K, INS_K, LIG_K = 1, 2, 3, 5, 7      # positions in state.ordered()
INS_TEXT, DEL_TEXT = None, "نص محذوف بتغيير متعقب"


def quote_of(o) -> str:
    return plain_text(o.quote).replace("…", "").strip()


def build_reviewed_doc(state, path, *, omit=(), foreign=True):
    """A realistic owner document holding the quoted passages of every observation of ``state``."""
    global INS_TEXT
    doc = Document()
    doc.add_paragraph("مذكرة تفاهم بين صندوق التنمية الثقافي وجهة شريكة (نسخة تجريبية)")
    doc.add_paragraph("تمهيد: حيث يرغب الطرفان في التعاون، فقد اتفقا على ما يلي.")
    for k, o in enumerate(state.ordered()):
        if o.internal_id in omit:
            continue
        q = quote_of(o)
        if k == SPLIT_K:                                   # split over two runs, second one bold
            p = doc.add_paragraph()
            cut = q.rfind(" ", 0, len(q) // 2) + 1
            p.add_run(f"المادة ({k + 1}): {q[:cut]}")
            p.add_run(q[cut:] + ".").bold = True
        elif k == TABLE_K:                                 # inside a table cell
            t = doc.add_table(rows=1, cols=2)
            t.cell(0, 0).text = "البند"
            t.cell(0, 1).text = f"{q}."
        elif k == INS_K:                                   # second half inserted with a tracked change
            p = doc.add_paragraph()
            cut = q.rfind(" ", 0, len(q) // 2) + 1
            p.add_run(q[:cut])
            INS_TEXT = q[cut:]
            p._p.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="901" w:author="{OWNER}" w:date="2026-03-01T09:00:00Z">'
                                  f'<w:r><w:t xml:space="preserve">{INS_TEXT}</w:t></w:r></w:ins>'))
            p._p.append(parse_xml(f'<w:del {nsdecls("w")} w:id="902" w:author="{OWNER}" w:date="2026-03-01T09:00:00Z">'
                                  f'<w:r><w:delText>{DEL_TEXT}</w:delText></w:r></w:del>'))
        elif k == LIG_K:                                   # lam-alef ligature and doubled spaces
            doc.add_paragraph(q.replace("لا ", "ﻻ ", 1).replace(" ", "  ") + ".")
        else:
            p = doc.add_paragraph(f"المادة ({k + 1}): {q}.")
            if k == FOREIGN_K and foreign:                 # a comment by another author, kept untouched
                doc.add_comment(p.runs, text="يرجى مراجعة الصياغة مع الشؤون القانونية.", author=OWNER, initials="ش")
    doc.add_paragraph("تحرر من نسختين بيد كل طرف نسخة للعمل بموجبها.")
    doc.save(path)
    return path


def parts_of(path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {i.filename: z.read(i.filename) for i in z.infolist()}


def c14n(el) -> bytes:
    return etree.tostring(el, method="c14n", exclusive=True)


def gco_comments(path):
    return [c for c in WS.extract_comments(path) if c.author == WS.GCO_AUTHOR and c.internal_id]


def simulate_return(src, dst, *, edit=None, delete_ref=None, owner_comment_on=None, add_ins_on=None,
                    alter_ref=None, move_ref=None, delete_para_of=None):
    """What an owner might do to the delivered file: silent edit, delete a GCO comment, add a comment, add a
    tracked insertion, alter a GCO comment, move a passage, delete a paragraph."""
    parts = parts_of(src)
    droot = etree.fromstring(parts[DC.DOC_PART])
    croot = etree.fromstring(parts[DC.COMMENTS_PART])
    exroot = etree.fromstring(parts[DC.COMMENTS_EX_PART]) if DC.COMMENTS_EX_PART in parts else None
    doc = DC.DocText(droot)
    by_ref = {(c.internal_id, c.follow_up_round): c for c in gco_comments(src)}
    body = droot.find(W("body"))

    def para_with(text):
        return next(p for p in doc.paras if text in p.text)

    if edit:                                               # (paragraph holding this text, old, new)
        within, old, new = edit
        t = next(t for t in para_with(within).el.iter(W("t")) if old in (t.text or ""))
        t.text = t.text.replace(old, new)
    if delete_ref:
        cid = by_ref[(delete_ref, None)].cid
        for el in list(droot.iter(W("commentRangeStart"), W("commentRangeEnd"), W("commentReference"))):
            if el.get(W("id")) == cid:
                target = el.getparent() if el.tag == W("commentReference") else el
                target.getparent().remove(target)
        c = next(c for c in croot.findall(W("comment")) if c.get(W("id")) == cid)
        pid = c.find(W("p")).get(f"{{{WS.W14NS}}}paraId")
        croot.remove(c)
        if exroot is not None and pid:
            for e in list(exroot):
                if (e.get(f"{{{WS.W15NS}}}paraId") or "").upper() == pid.upper():
                    exroot.remove(e)
    if owner_comment_on:
        p = para_with(owner_comment_on)
        cid = max(int(c.get(W("id"))) for c in croot.findall(W("comment"))) + 1
        runs = [a.run for a in p.atoms if a.text]
        runs[0].addprevious(DC._marker(W("commentRangeStart"), cid))
        runs[-1].addnext(DC._reference_run(cid))
        runs[-1].addnext(DC._marker(W("commentRangeEnd"), cid))
        croot.append(parse_xml(f'<w:comment {nsdecls("w")} w:id="{cid}" w:author="{OWNER}" w:initials="ش" '
                               f'w:date="2026-04-01T08:00:00Z"><w:p><w:r><w:t>تم التعديل وفق الملاحظة.</w:t></w:r></w:p></w:comment>'))
    if add_ins_on:
        p = para_with(add_ins_on)
        p.el.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="950" w:author="{OWNER}" w:date="2026-04-01T08:00:00Z">'
                              f'<w:r><w:t xml:space="preserve"> ويُشترط لذلك موافقة صاحب الصلاحية.</w:t></w:r></w:ins>'))
    if alter_ref:
        cid = by_ref[(alter_ref, None)].cid
        c = next(c for c in croot.findall(W("comment")) if c.get(W("id")) == cid)
        t = c.findall(W("p"))[2].find(f".//{W('t')}")
        t.text = "نص عُدّل من جهة أخرى: " + (t.text or "")
    if move_ref:
        cid = by_ref[(move_ref, None)].cid
        p = doc.paras[doc.ranges[cid].para_indices[0]]
        new_p = etree.Element(W("p"))
        r = etree.SubElement(new_p, W("r"))
        etree.SubElement(r, W("t")).text = p.text
        sect = body.find(W("sectPr"))
        (sect.addprevious if sect is not None else body.append)(new_p)
        for a in p.atoms:
            for t in a.run.findall(W("t")):
                t.text = "نص بديل حل محل الفقرة الأصلية"
                break
            for t in a.run.findall(W("t"))[1:]:
                t.text = ""
    if delete_para_of:
        p = para_with(delete_para_of)
        p.el.getparent().remove(p.el)
    parts[DC.DOC_PART] = etree.tostring(droot, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts[DC.COMMENTS_PART] = etree.tostring(croot, xml_declaration=True, encoding="UTF-8", standalone=True)
    if exroot is not None:
        parts[DC.COMMENTS_EX_PART] = etree.tostring(exroot, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return dst


@pytest.fixture(scope="module")
def life(five_round, tmp_path_factory):
    h, traj, out = five_round
    d = tmp_path_factory.mktemp("lifecycle")
    base = build_reviewed_doc(h.states[5], d / "base.docx")
    r1 = d / "r1.docx"
    rep1 = DC.comment_document(base, h.states[1], r1, round_no=1)
    r3 = d / "r3.docx"
    rep3 = DC.comment_document(r1, h.states[3], r3, round_no=3, prior_docx=r1)
    return SimpleNamespace(h=h, dir=d, base=base, r1=r1, rep1=rep1, r3=r3, rep3=rep3)


def obs(life, k):
    return life.h.states[5].ordered()[k]


# ------------------------------------------------------------------------ round 1 on the real document
def test_lifecycle_round1_every_ref_present(life):
    st = life.h.states[1]
    ids = {o.internal_id for o in st.observations}
    got = {c.internal_id for c in gco_comments(life.r1) if c.follow_up_round is None}
    assert got == ids
    assert life.rep1.unanchored == [] and set(life.rep1.anchored) == ids and life.rep1.roots_missing == []
    assert set(life.rep1.new_comments) == ids and life.rep1.ok
    assert life.rep1.anchor_kind[obs(life, LIG_K).internal_id] == "normalised"
    assert all(life.rep1.anchor_kind[o.internal_id] == "exact" for k, o in enumerate(st.ordered()) if k != LIG_K)
    Document(life.r1)                                                   # python-docx reopens the package
    assert json.loads(life.rep1.to_json())["round"] == 1 and life.rep1.summary()["anchored"] == len(ids)


def test_lifecycle_comment_shape_08_2(life):
    st = life.h.states[1]
    by_id = {c.internal_id: c for c in gco_comments(life.r1)}
    croot = etree.fromstring(parts_of(life.r1)[DC.COMMENTS_PART])
    for o in st.ordered():
        c = by_id[o.internal_id]
        assert c.lines[0] == f"الملاحظة #{o.reader_no}" and c.reader_no == o.reader_no
        assert c.ref_line_index == 1 and c.ref_run_ok and c.lines[1] == f"REF: {o.internal_id}"
        assert "\n".join(c.lines[2:]) == plain_text(o.comment).strip()
    el = next(c for c in croot.findall(W("comment")) if c.get(W("id")) == by_id[obs(life, SPLIT_K).internal_id].cid)
    assert el.get(W("author")) == WS.GCO_AUTHOR and el.get(W("initials")) == "حوكمة"
    bold = [r for r in el.iter(W("r")) if r.find(f"{W('rPr')}/{W('b')}") is not None]
    assert any("".join(t.text or "" for t in r.iter(W("t"))) == "التوصية:" for r in bold)   # bold label kept
    fonts = {f.get(W("ascii")) for f in el.iter(W("rFonts"))}
    assert fonts == {"Calibri"} and el.find(W("p")).get(f"{{{WS.W14NS}}}paraId")


def test_lifecycle_preserves_foreign_comment_and_tracked_changes(life):
    before, after = parts_of(life.base), parts_of(life.r1)
    ins_bytes = re.search(rb"<w:ins .*?</w:ins>", before[DC.DOC_PART], re.S).group(0)
    del_bytes = re.search(rb"<w:del .*?</w:del>", before[DC.DOC_PART], re.S).group(0)
    assert ins_bytes in after[DC.DOC_PART] and del_bytes in after[DC.DOC_PART]       # byte for byte
    b_ins = [c14n(e) for e in etree.fromstring(before[DC.DOC_PART]).iter(W("ins"), W("del"))]
    a_ins = [c14n(e) for e in etree.fromstring(after[DC.DOC_PART]).iter(W("ins"), W("del"))]
    assert b_ins == a_ins and len(b_ins) == 2
    foreign_b = [c14n(c) for c in etree.fromstring(before[DC.COMMENTS_PART]).findall(W("comment")) if c.get(W("author")) == OWNER]
    foreign_a = [c14n(c) for c in etree.fromstring(after[DC.COMMENTS_PART]).findall(W("comment")) if c.get(W("author")) == OWNER]
    assert foreign_b == foreign_a and len(foreign_b) == 1
    foreign = next(c for c in WS.extract_comments(life.r1) if c.author == OWNER)
    assert foreign.lines == ["يرجى مراجعة الصياغة مع الشؤون القانونية."]
    untouched = {n for n in before if n not in DC.TOUCHED_PARTS}
    assert untouched <= set(life.rep1.parts_preserved)
    assert all(hashlib.sha256(before[n]).hexdigest() == hashlib.sha256(after[n]).hexdigest() for n in untouched)
    assert set(life.rep1.parts_changed) <= set(DC.TOUCHED_PARTS) and "word/styles.xml" in life.rep1.parts_preserved
    assert {"word/numbering.xml", "docProps/core.xml", "word/settings.xml"} <= set(life.rep1.parts_preserved)


def test_lifecycle_document_text_unchanged(life):
    b = DC.DocText(etree.fromstring(parts_of(life.base)[DC.DOC_PART]))
    a = DC.DocText(etree.fromstring(parts_of(life.r1)[DC.DOC_PART]))
    assert [p.text for p in b.paras] == [p.text for p in a.paras]
    assert [t.text for t in b.root.iter(W("t"))] == [t.text for t in a.root.iter(W("t"))]
    assert DEL_TEXT not in "".join(p.text for p in a.paras) and INS_TEXT in "".join(p.text for p in a.paras)


def test_lifecycle_anchors_enclose_quotes(life):
    st = life.h.states[1]
    doc = DC.DocText(etree.fromstring(parts_of(life.r1)[DC.DOC_PART]))
    by_id = {c.internal_id: c for c in gco_comments(life.r1)}
    for o in st.ordered():
        rng = doc.ranges[by_id[o.internal_id].cid]
        assert rng.complete and rng.ref_run is not None
        assert DC._passage_in_text(plain_text(o.quote), rng.text), o.internal_id
    split = doc.ranges[by_id[obs(life, SPLIT_K).internal_id].cid]
    bolds = {a.run.find(f"{W('rPr')}/{W('b')}") is not None for a in split.atoms if a.text}
    assert len([a for a in split.atoms if a.text]) == 2 and bolds == {True, False}
    table = doc.ranges[by_id[obs(life, TABLE_K).internal_id].cid]
    assert all(doc.paras[i].in_table for i in table.para_indices)
    ins = doc.ranges[by_id[obs(life, INS_K).internal_id].cid]
    assert INS_TEXT in ins.text and ins.end.getparent().tag == W("p") and ins.start.getparent().tag == W("p")
    lig = doc.ranges[by_id[obs(life, LIG_K).internal_id].cid]
    assert "ﻻ" in lig.text


def test_lifecycle_reconcile_round1(life):
    rep = WS.reconcile(life.r1, life.h.files[1])
    assert rep.ok, rep.problems
    assert DC.check_package(life.r1, life.h.files[1], life.h.states[1]) == []


def test_lifecycle_unanchored_quote_is_flagged(five_round, tmp_path):
    h, traj, out = five_round
    st = h.states[1]
    victim = st.ordered()[8]
    base = build_reviewed_doc(st, tmp_path / "no9.docx", omit={victim.internal_id}, foreign=False)
    with zipfile.ZipFile(base) as z:
        assert DC.COMMENTS_PART not in z.namelist()                         # no comments part at all
    rep = DC.comment_document(base, st, tmp_path / "no9-r1.docx", round_no=1)
    assert [u["id"] for u in rep.unanchored] == [victim.internal_id] and not rep.ok
    assert "quote not found" in rep.unanchored[0]["reason"] and rep.anchor_kind[victim.internal_id] == "fallback"
    assert {c.internal_id for c in gco_comments(tmp_path / "no9-r1.docx")} == {o.internal_id for o in st.observations}
    with zipfile.ZipFile(tmp_path / "no9-r1.docx") as z:
        names = z.namelist()
        rels = z.read(DC.DOC_RELS).decode()
    assert DC.COMMENTS_PART in names and DC.COMMENTS_EX_PART in names and 'Target="comments.xml"' in rels
    assert DC.package_structure_problems(tmp_path / "no9-r1.docx") == []


# ------------------------------------------------------------------------ the owner's return
def test_lifecycle_diff_returned_detects_silent_edit_deleted_comment_owner_comment(life, tmp_path):
    st = life.h.states[1]
    o_edit, o_del, o_cmt = obs(life, 0), obs(life, 4), obs(life, 6)
    q = quote_of(o_edit)
    old = q.split()[2]
    ret = simulate_return(life.r1, tmp_path / "ret.docx", edit=(q, old, "عبارة"), delete_ref=o_del.internal_id,
                          owner_comment_on=quote_of(o_cmt))
    d = DC.diff_returned(life.r1, ret, st)
    silent = d.silent_edits
    assert len(silent) == 1 and silent[0].kind == "modified" and "عبارة" in silent[0].after and old in silent[0].before
    assert d.missing_gco_comments == [o_del.internal_id]
    assert len(d.owner_comments) == 1
    oc = d.owner_comments[0]
    assert oc.author == OWNER and oc.text == "تم التعديل وفق الملاحظة." and oc.nearest_ref == o_cmt.internal_id
    assert [a for a in d.anchor_issues if a.internal_id == o_edit.internal_id and a.kind == "quote_missing"]
    s = d.summary()
    assert s["silent_edits"] == 1 and s["missing_gco_comments"] == [o_del.internal_id] and not s["clean"]
    assert json.loads(d.to_json())["owner_comments"][0]["author"] == OWNER
    assert DC.diff_returned(life.r1, life.r1, st).clean


def test_lifecycle_diff_returned_tracked_change_altered_comment_moved_anchor(life, tmp_path):
    st = life.h.states[1]
    o_ins, o_alt, o_move = obs(life, 8), obs(life, 9), obs(life, 6)
    ret = simulate_return(life.r1, tmp_path / "ret2.docx", add_ins_on=quote_of(o_ins), alter_ref=o_alt.internal_id,
                          move_ref=o_move.internal_id)
    d = DC.diff_returned(life.r1, ret, st)
    new_tc = [t for t in d.tracked_changes if not t.pre_existing]
    assert len(new_tc) == 1 and new_tc[0].kind == "ins" and new_tc[0].author == OWNER and "موافقة" in new_tc[0].excerpt
    assert len(new_tc[0].excerpt) <= 60 and sum(1 for t in d.tracked_changes if t.pre_existing) == 2
    tracked_mod = [c for c in d.paragraph_changes if c.tracked and c.kind == "modified"]
    assert tracked_mod and quote_of(o_ins)[:20] in tracked_mod[0].before
    assert all(quote_of(o_ins)[:20] not in c.before for c in d.silent_edits)      # explained by the w:ins
    assert [a.internal_id for a in d.altered_gco_comments] == [o_alt.internal_id]
    assert d.altered_gco_comments[0].line_index == 2 and d.altered_gco_comments[0].after.startswith("نص عُدّل")
    moved = [a for a in d.anchor_issues if a.kind == "anchor_moved"]
    assert [a.internal_id for a in moved] == [o_move.internal_id] and "نص بديل" in moved[0].before
    assert d.missing_gco_comments == [] and not d.owner_comments
    d2 = DC.diff_returned(life.r1, ret)                                             # without a state
    assert [a.kind for a in d2.anchor_issues] == ["range_text_changed"]


def test_lifecycle_closed_text_integrity(life, tmp_path):
    st = life.h.states[4]
    closed_r3 = [o for o in st.ordered() if o.value(3) == C]
    o_intact, o_changed, o_missing, o_moved = (next(o for o in closed_r3 if o.reader_no == n) for n in ("01", "02", "03", "07"))
    assert [o.value(2) for o in (o_intact, o_changed, o_missing, o_moved)] == [C, C, C, RE]
    q_changed = quote_of(o_changed)
    ret = simulate_return(life.r1, tmp_path / "ret3.docx", edit=(q_changed, q_changed.split()[1], "الجهتان"),
                          delete_para_of=quote_of(o_missing), move_ref=o_moved.internal_id)
    findings = {f.internal_id: f for f in DC.closed_text_integrity(st, ret, 4)}
    assert set(findings) == {o.internal_id for o in closed_r3}
    assert findings[o_intact.internal_id].status == DC.INTACT and findings[o_intact.internal_id].closed_in_round == 2
    f = findings[o_changed.internal_id]
    assert f.status == DC.CHANGED and "الجهتان" in f.excerpt and 0.5 <= f.ratio < 1
    assert findings[o_missing.internal_id].status == DC.MISSING
    f = findings[o_moved.internal_id]
    assert f.status == DC.MOVED_IDENTICAL and f.closed_in_round == 3 and f.paragraph_index not in f.anchored_paragraphs
    assert all(f.status == DC.INTACT for i, f in findings.items() if i not in
               {o_changed.internal_id, o_missing.internal_id, o_moved.internal_id})
    assert all(f.status == DC.INTACT for f in DC.closed_text_integrity(st, life.r1, 4))   # untouched file
    assert st.by_id()[o_changed.internal_id].value(3) == C                                # state never changed
    assert DC._location_hint(o_changed, 3) is None and DC.MOVED_LINE_RE.search(
        "الجولة الثانية: تغيّر موضع الملاحظة إلى ص.≈9، المادة (7) – البند (7-2)").group("where").startswith("ص.")


# ------------------------------------------------------------------------ round n: replies on the returned file
def test_lifecycle_round3_followups_are_native_replies(life):
    h, st = life.h, life.h.states[3]
    expected = {o.internal_id for o in st.ordered() if o.birth_round < 3 and (o.value(2) != C or o.value(3) != C)}
    born = {o.internal_id for o in st.ordered() if o.birth_round == 3}
    assert born and expected
    assert set(life.rep3.replies_added) == {f"{i} / R3-FU" for i in expected}
    assert set(life.rep3.roots_found) == expected and life.rep3.roots_missing == [] and life.rep3.unanchored == []
    assert set(life.rep3.new_comments) == born and life.rep3.ok
    comments = WS.extract_comments(life.r3)
    by_cid = {c.cid: c for c in comments}
    follow = [c for c in comments if c.follow_up_round == 3]
    assert {c.internal_id for c in follow} == expected
    r1_pids = {c.internal_id: c.para_id for c in gco_comments(life.r1) if c.follow_up_round is None}
    ex = etree.fromstring(parts_of(life.r3)[DC.COMMENTS_EX_PART])
    parent_of = {e.get(f"{{{WS.W15NS}}}paraId"): e.get(f"{{{WS.W15NS}}}paraIdParent") for e in ex}
    for c in follow:
        root = by_cid[c.parent_cid]
        assert root.internal_id == c.internal_id and root.follow_up_round is None and root.parent_cid is None
        assert parent_of[c.para_id] == r1_pids[c.internal_id]            # parent = the round-1 comment's paraId
        assert c.status == st.by_id()[c.internal_id].value(3)
        assert c.done == root.done == (c.status == C)
    for c in comments:
        if c.follow_up_round is None:
            assert c.parent_cid is None
    assert sum(1 for c in comments if c.author == OWNER) == 1                 # the foreign comment survives
    rep = WS.reconcile(life.r3, h.files[3])
    assert rep.ok, rep.problems
    assert rep.checks["G-SYNC-9 follow-ups threaded"] is True
    assert DC.check_package(life.r3, h.files[3], st) == []


def test_lifecycle_round3_reply_ranges_follow_the_root(life):
    doc = DC.DocText(etree.fromstring(parts_of(life.r3)[DC.DOC_PART]))
    comments = WS.extract_comments(life.r3)
    by_cid = {c.cid: c for c in comments}
    for c in comments:
        if c.follow_up_round == 3:
            assert doc.ranges[c.cid].complete and doc.ranges[c.cid].text == doc.ranges[c.parent_cid].text
    born = next(c for c in comments if c.follow_up_round is None and (c.internal_id or "").startswith("GCO-FIV-R3-"))
    assert "نص مضاف 1" in doc.ranges[born.cid].text
    a = DC.DocText(etree.fromstring(parts_of(life.r1)[DC.DOC_PART]))
    assert [p.text for p in a.paras] == [p.text for p in doc.paras]


def test_lifecycle_round3_rerun_is_idempotent(life, tmp_path):
    rep = DC.comment_document(life.r3, life.h.states[3], tmp_path / "r3b.docx", round_no=3)
    assert rep.replies_added == [] and rep.new_comments == [] and rep.ok
    assert set(rep.skipped) == set(life.rep3.replies_added) | set(life.rep3.new_comments)
    assert len(WS.extract_comments(tmp_path / "r3b.docx")) == len(WS.extract_comments(life.r3))


def test_lifecycle_round3_missing_root_is_reported(life, tmp_path):
    st = life.h.states[3]
    victim = next(o for o in st.ordered() if o.reader_no == "05")             # [O, RE, RE]: needs a follow-up
    ret = simulate_return(life.r1, tmp_path / "ret4.docx", delete_ref=victim.internal_id)
    rep = DC.comment_document(ret, st, tmp_path / "r3-missing.docx", round_no=3, prior_docx=life.r1)
    assert rep.roots_missing == [victim.internal_id] and not rep.ok
    assert any("removed by the owner" in w and victim.internal_id in w for w in rep.warnings)
    fu = next(c for c in gco_comments(tmp_path / "r3-missing.docx") if c.internal_id == victim.internal_id)
    assert fu.follow_up_round == 3 and fu.parent_cid is None and fu.status == RE
    rep_sync = WS.reconcile(tmp_path / "r3-missing.docx", life.h.files[3])
    assert rep_sync.checks["G-SYNC-9 follow-ups threaded"] is False
    assert any("duplicate" not in p and victim.internal_id in p for p in rep_sync.problems)


# ------------------------------------------------------------------------ package checks and the write guard
def test_lifecycle_check_package_names_duplicate_ref(life, tmp_path):
    h, st = life.h, life.h.states[1]
    parts = parts_of(life.r1)
    croot = etree.fromstring(parts[DC.COMMENTS_PART])
    droot = etree.fromstring(parts[DC.DOC_PART])
    victim = st.ordered()[3]
    cid = next(c.cid for c in gco_comments(life.r1) if c.internal_id == victim.internal_id)
    src = next(c for c in croot.findall(W("comment")) if c.get(W("id")) == cid)
    dup = etree.fromstring(etree.tostring(src))
    new_id = str(max(int(c.get(W("id"))) for c in croot.findall(W("comment"))) + 1)
    dup.set(W("id"), new_id)
    dup.find(W("p")).set(f"{{{WS.W14NS}}}paraId", "7ABCDEF0")
    croot.append(dup)
    first_run = next(r for r in droot.iter(W("r")) if r.find(W("t")) is not None)
    first_run.addprevious(DC._marker(W("commentRangeStart"), int(new_id)))
    first_run.addnext(DC._reference_run(int(new_id)))
    first_run.addnext(DC._marker(W("commentRangeEnd"), int(new_id)))
    parts[DC.COMMENTS_PART] = etree.tostring(croot, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts[DC.DOC_PART] = etree.tostring(droot, xml_declaration=True, encoding="UTF-8", standalone=True)
    bad = tmp_path / "dup.docx"
    with zipfile.ZipFile(bad, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    problems = DC.check_package(bad, h.files[1], st)
    assert problems and any("duplicate REF" in p and victim.internal_id in p for p in problems)
    assert any("one REF per row" in p for p in problems)


def test_lifecycle_check_package_structural_problems(life, tmp_path):
    parts = parts_of(life.r1)
    no_comments = tmp_path / "nocomments.docx"
    with zipfile.ZipFile(no_comments, "w") as z:
        for n, b in parts.items():
            if n != DC.COMMENTS_PART:
                z.writestr(n, b)
    probs = DC.package_structure_problems(no_comments)
    assert any("missing" in p and "comments" in p for p in probs)
    droot = etree.fromstring(parts[DC.DOC_PART])
    end = next(e for e in droot.iter(W("commentRangeEnd")))
    end.getparent().remove(end)
    parts[DC.DOC_PART] = etree.tostring(droot, xml_declaration=True, encoding="UTF-8", standalone=True)
    unpaired = tmp_path / "unpaired.docx"
    with zipfile.ZipFile(unpaired, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    assert any("unpaired" in p for p in DC.package_structure_problems(unpaired))
    exroot = etree.fromstring(parts[DC.COMMENTS_EX_PART])
    exroot[0].set(f"{{{WS.W15NS}}}paraIdParent", "0BADBAD0")
    parts[DC.COMMENTS_EX_PART] = etree.tostring(exroot, xml_declaration=True, encoding="UTF-8", standalone=True)
    badparent = tmp_path / "badparent.docx"
    with zipfile.ZipFile(badparent, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    assert any("does not resolve" in p for p in DC.package_structure_problems(badparent))
    assert DC.check_package(badparent, life.h.files[1], life.h.states[1])     # non-empty for the broken package


def test_lifecycle_atomic_output_leaves_destination_absent(life, tmp_path):
    dest = tmp_path / "guarded.docx"
    with pytest.raises(RuntimeError, match="boom"):
        with DC.atomic_output(dest, stage="unit") as tmp:
            tmp.write_bytes(b"partial")
            raise RuntimeError("boom")
    assert not dest.exists() and not list(tmp_path.glob(".guarded.docx.*.tmp"))
    rejected = DC.rejected_path(dest)
    assert rejected.exists()
    rj = json.loads(rejected.read_text(encoding="utf-8"))
    assert rj["reason"] == "RuntimeError: boom" and rj["stage"] == "unit" and rj["destination"] == str(dest)
    # a corrupt source through comment_document: same guarantee
    corrupt = tmp_path / "corrupt.docx"
    corrupt.write_text("not a zip")
    dest2 = tmp_path / "out" / "r1.docx"
    with pytest.raises(Exception):
        DC.comment_document(corrupt, life.h.states[1], dest2, round_no=1)
    assert not dest2.exists() and DC.rejected_path(dest2).exists()
    assert "BadZipFile" in json.loads(DC.rejected_path(dest2).read_text(encoding="utf-8"))["reason"]
    # a later success removes the stale rejection report
    DC.comment_document(life.base, life.h.states[1], dest2, round_no=1)
    assert dest2.exists() and not DC.rejected_path(dest2).exists()


def test_lifecycle_normalisation():
    n, m = DC.normalise("ﻻ  تترتب‏ عَلى")
    assert n == "لا تترتب على" and len(m) == len(n)
    assert DC.normalise("لأ لإ لآ")[0] == "لا لا لا" and DC.normalise("أحمد")[0] == "أحمد"
    assert DC._quote_parts("… نص أول … نص ثانٍ …") == ["نص أول", "نص ثانٍ"]
    assert DC._quote_parts("نص ...") == ["نص"]
    assert DC._passage_in_text("نص أول … ثانٍ", "هنا نص أول ثم نص ثانٍ هناك")
    assert not DC._passage_in_text("ثانٍ … أول", "هنا نص أول ثم نص ثانٍ هناك")
    assert DC._excerpt("ك" * 100).endswith("…") and len(DC._excerpt("ك" * 100)) == 60


def test_lifecycle_rows_for_round(life):
    st = life.h.states[5]
    follow, born = DC.rows_for_round(st, 5)
    assert {o.reader_no for o in born} == set()
    assert {o.reader_no for o in follow} == {o.reader_no for o in st.ordered() if o.birth_round < 5 and (o.value(4) != C or o.value(5) != C)}
    assert "02" in {o.reader_no for o in follow} and "01" not in {o.reader_no for o in follow}   # 02 reclosed in r5
    f1, b1 = DC.rows_for_round(life.h.states[1], 1)
    assert f1 == [] and len(b1) == 10


@needs_lo
def test_lifecycle_libreoffice_renders_final_docx(life, tmp_path):
    pdf = lo_engine.to_pdf(life.r3, tmp_path / "pdf")
    assert pdf is not None and pdf.stat().st_size > 1000
