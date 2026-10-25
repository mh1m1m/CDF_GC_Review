"""TEST 29 — safe ingestion and formula injection (15A.6, G-SUB-7; patches V712-005/026/056)."""
from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cdf_gco_crms.config import round_schema as S
from cdf_gco_crms.excel.register_writer import is_formula_like
from cdf_gco_crms.excel.workbook_builder import build_workbook
from cdf_gco_crms.ingestion.safe_open import IngestError, inspect_package, require_safe
from cdf_gco_crms.ingestion.workbook_reader import find_data_sheet, parse_workbook
from cdf_gco_crms.model.entities import GText, RoundCells
from cdf_gco_crms.regression import synthetic as SY
from cdf_gco_crms.rounds.round_expansion import NewObservation, build_round1
from cdf_gco_crms.rounds.round_parser import SchemaError

HOSTILE = ["=1+1", "=HYPERLINK(\"http://x\",\"x\")", "+SUM(A1)", "-2+3", "@SUM(1)", "=cmd|' /C calc'!A0"]


def _state():
    obs = [SY.topic_obs(i) for i in range(3)]
    obs[0] = NewObservation(HOSTILE[0], HOSTILE[1], GText.plain(HOSTILE[5]))
    return build_round1(SY.fixture_document("SEC"), obs, date="2026-10-23")


def test_formula_like_text_is_written_as_text_and_reopens_as_text(tmp_path):
    st = _state()
    st.observations[1].cells[1] = RoundCells(closure=S.OPEN, response=HOSTILE[2], impl=S.PH_STATUS)
    out = tmp_path / "s.xlsx"
    build_workbook(st, out)
    ws = find_data_sheet(load_workbook(out, rich_text=True))
    hits = [(c.coordinate, c.data_type) for row in ws.iter_rows(min_row=2) for c in row
            if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("="))]
    assert hits == [], hits
    pw = parse_workbook(out)
    o = pw.state.by_id()[st.observations[0].internal_id]
    assert str(o.location) == HOSTILE[0] and str(o.quote) == HOSTILE[1] and o.comment.text == HOSTILE[5]
    assert all(is_formula_like(h) for h in HOSTILE) and not is_formula_like("-") and not is_formula_like("نص")


def test_formula_planted_in_data_sheet_is_read_literally_and_flagged(tmp_path):
    st = _state()
    out = tmp_path / "p.xlsx"
    build_workbook(st, out)
    wb = load_workbook(out)
    ws = find_data_sheet(wb)
    ws["G2"] = "=1+1"                        # planted into a requester cell of the returned file
    wb.save(out)
    pw = parse_workbook(out)
    assert any(w.startswith("FORMULA_IN_DATA") and "G2" in w for w in pw.warnings)
    o = pw.state.by_id()[st.observations[0].internal_id]
    assert str(o.cells[1].response) == "=1+1"
    out2 = tmp_path / "p2.xlsx"
    build_workbook(pw.state, out2)
    assert find_data_sheet(load_workbook(out2))["G2"].data_type != "f"


def _zip_with(src: Path, dst: Path, extra: dict[str, bytes] | None = None, drop: tuple = (), rename: dict | None = None):
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            if i.filename in drop:
                continue
            zo.writestr((rename or {}).get(i.filename, i.filename), zi.read(i.filename))
        for n, b in (extra or {}).items():
            zo.writestr(n, b)


@pytest.fixture(scope="module")
def clean(tmp_path_factory):
    out = tmp_path_factory.mktemp("sec") / "clean.xlsx"
    build_workbook(_state(), out)
    assert require_safe(out).ok
    return out


def test_hostile_packages_fail_closed(clean, tmp_path):
    cases = {
        "macro": dict(extra={"xl/vbaProject.bin": b"\x00" * 10}),
        "external": dict(extra={"xl/externalLinks/externalLink1.xml": b"<x/>"}),
        "traversal": dict(extra={"../../etc/passwd": b"x"}),
        "absolute": dict(extra={"/tmp/evil": b"x"}),
        "nomain": dict(drop=("xl/workbook.xml",)),
        "noct": dict(drop=("[Content_Types].xml",)),
        "bomb": dict(extra={"xl/pad.xml": b"<a>" + b" " * (90 * 1024 * 1024) + b"</a>"}),
    }
    expected = {"macro": "MACRO_OR_BINARY_PART", "external": "EXTERNAL_LINK", "traversal": "PATH_TRAVERSAL",
                "absolute": "PATH_TRAVERSAL", "nomain": "NO_MAIN_PART", "noct": "NO_CONTENT_TYPES", "bomb": "COMPRESSION_BOMB"}
    for name, kw in cases.items():
        dst = tmp_path / f"{name}.xlsx"
        _zip_with(clean, dst, **kw)
        rep = inspect_package(dst)
        assert not rep.ok and expected[name] in {c for c, _ in rep.findings}, (name, rep.findings)
        with pytest.raises(SchemaError) as ei:
            parse_workbook(dst)
        assert ei.value.code == "UNSAFE_INPUT"
    # not a zip at all, and a wrong extension
    junk = tmp_path / "junk.xlsx"; junk.write_bytes(b"PK\x03\x04garbage")
    assert "BAD_ZIP" in {c for c, _ in inspect_package(junk).findings} or "NOT_OOXML" in {c for c, _ in inspect_package(junk).findings}
    xlsm = tmp_path / "m.xlsm"; shutil.copy(clean, xlsm)
    assert "BAD_EXTENSION" in {c for c, _ in inspect_package(xlsm).findings}
    with pytest.raises(IngestError):
        require_safe(xlsm)


def test_failed_ingestion_writes_nothing(clean, tmp_path):
    dst = tmp_path / "macro.xlsx"
    _zip_with(clean, dst, extra={"xl/vbaProject.bin": b"\x00"})
    before = sorted(p.name for p in tmp_path.iterdir())
    with pytest.raises(SchemaError):
        parse_workbook(dst)
    assert sorted(p.name for p in tmp_path.iterdir()) == before
