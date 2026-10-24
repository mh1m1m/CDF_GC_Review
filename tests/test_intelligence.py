"""Synthetic analytical, persistence, provenance and draft-export regressions."""
import hashlib
from io import BytesIO
import json
from pathlib import Path
import zipfile

import pytest
from hypothesis import given, strategies as st
from docx import Document
from openpyxl import load_workbook
from cdf_gco_crms.calibration.service import CalibrationService
from cdf_gco_crms.calibration.store import CalibrationStore, ReviewerContext, AccessDenied, CalibrationError
from cdf_gco_crms.intelligence.service import IntelligenceService, measure
from cdf_gco_crms.intelligence.exports import export_dashboard
from tools.seed_visual_demo import seed


@pytest.fixture
def dashboard(tmp_path):
    path=tmp_path/'pilot.sqlite';case_id=seed(path)
    store=CalibrationStore(path);service=IntelligenceService(CalibrationService(store))
    user=ReviewerContext('SYNTHETIC_ADMIN',frozenset({'CALIBRATION_REVIEWER','CALIBRATION_ADMIN'}),True)
    yield service,store,user,case_id
    store.close()


def test_actual_orchestrator_sources_and_unknown_alignment(dashboard):
    service,_,user,case=dashboard
    data=service.dashboard(case,user)
    assert data['trace_integrity']
    assert data['actual_live_inference_count']==0
    assert data['execution_sources']['MOCK_SCRIPT']>0
    assert data['execution_sources']['DETERMINISTIC_RULES']>0
    assert {i['stage'] for i in data['invocations']} >= {'DRAFTING','MAT-01','BLACK_TEAM','MAT-02','RED_ADVOCATE','MAT-03','ADJUDICATOR','FINAL_REVIEWER'}
    assert all(i['status']=='OK' for i in data['invocations'])
    assert not data['release_gate']['authorized'] and not data['release_gate']['current_authority']
    assert all(o['productive_cycles']<=6 for o in data['objections'])
    metrics={m['id']:m for m in data['metrics']}
    assert metrics['evidence_coverage']['value']==0 # anchor occurrence never inflates reference approval
    assert metrics['human_alignment']['value'] is None
    assert metrics['calibration']['numerator']==0


def test_case_grants_are_required(dashboard):
    service,_,_,case=dashboard
    stranger=ReviewerContext('OTHER',frozenset({'CALIBRATION_REVIEWER'}),True)
    with pytest.raises(AccessDenied): service.dashboard(case,stranger)


def test_snapshot_corruption_refused(dashboard):
    service,store,user,case=dashboard
    with store.db:store.db.execute("UPDATE executive_runs SET snapshot_json='{}' WHERE case_id=?",(case,))
    with pytest.raises(CalibrationError,match='INTEGRITY'):service.dashboard(case,user)


def test_forged_live_label_cannot_inflate_counter(dashboard):
    service,store,user,case=dashboard
    row=store.db.execute('SELECT snapshot_json FROM executive_runs WHERE case_id=?',(case,)).fetchone()
    raw=json.loads(row[0]);raw['trace_records'][0]['Execution_Source']='PROVIDER_API'
    raw['trace_records'][0]['Mode']='LIVE'
    # Even an updated snapshot checksum does not repair the actual trace chain.
    text=json.dumps(raw)
    with store.db:store.db.execute('UPDATE executive_runs SET snapshot_json=?,snapshot_hash=? WHERE case_id=?',
        (text,hashlib.sha256(text.encode()).hexdigest(),case))
    data=service.dashboard(case,user)
    assert not data['trace_integrity'] and data['actual_live_inference_count']==0


def test_latest_submitted_calibration_counts_distinct_observations(dashboard):
    service,store,user,case=dashboard
    digest=store._case(case)['evidence_hash']
    for _ in range(2):store.save_decision(case,'SYN-001','KEEP_UNCHANGED','Synthetic rationale',digest,user,status='SUBMITTED')
    data=service.dashboard(case,user)
    metric=next(m for m in data['metrics'] if m['id']=='calibration')
    assert metric['numerator']==1 and metric['denominator']==4 and metric['value']==25


@given(st.integers(min_value=1,max_value=10000),st.integers(min_value=0,max_value=10000))
def test_fraction_invariants(denominator,number):
    numerator=min(denominator,number)
    metric=measure('test','Synthetic',numerator,denominator,['test'],'No confidence')
    assert 0<=metric['value']<=100 and metric['confidence'] is None
    assert measure('test','Empty',0,0,[],'Unknown')['value'] is None


def test_exports_share_metrics_and_remain_drafts(dashboard):
    service,_,user,case=dashboard;data=service.dashboard(case,user)
    word=export_dashboard(data,'docx');doc=Document(BytesIO(word))
    assert doc.tables[0].rows[1].cells[1].text=='0 / 4'
    assert 'INTERNAL DRAFT' in doc.core_properties.category
    with zipfile.ZipFile(BytesIO(word)) as archive:assert any(n.startswith('word/media/') for n in archive.namelist())
    excel=export_dashboard(data,'xlsx');wb=load_workbook(BytesIO(excel),data_only=False)
    assert wb['Executive']['B3'].value==0 and wb['Executive']['C3'].value==4
    assert wb['Executive']['D3'].data_type=='f'
    assert wb['Executive']._charts and wb['Executive']._charts[0].tagname=='doughnutChart'
    assert 'INTERNAL DRAFT' in wb.properties.category
    assert wb['Provenance']['B2'].value=='NOT_AUTHORIZED'


def test_empty_case_does_not_invent_review_data(tmp_path):
    store=CalibrationStore(tmp_path/'empty.sqlite')
    user=ReviewerContext('SYN',frozenset({'CALIBRATION_ADMIN','CALIBRATION_REVIEWER'}),True)
    store.import_case('SYN-EMPTY','a'*64,[{'observation_id':'SYN-O'}],user)
    data=IntelligenceService(CalibrationService(store)).dashboard('SYN-EMPTY',user)
    assert data['run_id'] is None and data['invocations']==[]
    assert data['actual_live_inference_count']==0
    assert next(m for m in data['metrics'] if m['id']=='objections_resolved')['value'] is None
    store.close()


def test_pdf_and_excel_formulas_with_independent_engine(dashboard,tmp_path):
    from cdf_gco_crms.qa import lo_engine
    from pypdf import PdfReader
    if not lo_engine.available() or 'Dev' in lo_engine.version():
        pytest.skip('Stable LibreOffice required for executive export acceptance')
    service,_,user,case=dashboard;data=service.dashboard(case,user)
    pdf=export_dashboard(data,'pdf')
    reader=PdfReader(BytesIO(pdf));assert len(reader.pages)>0
    assert 'SYNTHETIC' in ''.join(page.extract_text() or '' for page in reader.pages)
    original=export_dashboard(data,'xlsx');path=tmp_path/'recalculated.xlsx';path.write_bytes(original)
    assert lo_engine.recalc(path,timeout=60).ok
    wb=load_workbook(path,data_only=True)
    assert wb['Executive']['D3'].value==0
    assert wb['Executive']['D4'].value==1
    human_index=next(i for i,m in enumerate(data['metrics'],3) if m['id']=='human_alignment')
    assert wb['Executive'].cell(human_index,4).value in ('',None)
    # Byte source remains immutable; recalculation was performed on its separate copy.
    assert hashlib.sha256(original).digest()!=hashlib.sha256(path.read_bytes()).digest()


def test_repeat_rate_uses_black_attempts_not_advisory_or_red_turns(dashboard):
    service,_,user,case=dashboard;data=service.dashboard(case,user)
    metric=next(m for m in data['metrics'] if m['id']=='repeat_objections')
    attempts=sum(sum(t['actor']=='BLACK' for t in o['turns'])+
                 sum(a['actor']=='BLACK' for a in o['unproductive_attempts']) for o in data['objections'])
    assert metric['denominator']==attempts==2
