"""Create a fictional case and execute the actual orchestrator in MOCK/rule modes.

No fixture claims genuine inference, human labels, institutional approval or source authority.
"""
import argparse
import hashlib
from pathlib import Path
import sys
import tempfile
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document as WordDocument
from cdf_gco_crms.agents import AgentRunner, MockAdapter, TraceWriter
from cdf_gco_crms.agents import contracts as C
from cdf_gco_crms.agents.mat_advisor import MatAdvisor
from cdf_gco_crms.agents.mat_profile import load_profile, synthetic_fixture_path
from cdf_gco_crms.agents.orchestrator import ReviewOrchestrator
from cdf_gco_crms.calibration.store import CalibrationStore, ReviewerContext
from cdf_gco_crms.calibration.service import CalibrationService
from cdf_gco_crms.intelligence.service import IntelligenceService
from cdf_gco_crms.model.entities import Document, GText
from cdf_gco_crms.rounds.round_expansion import build_round1, NewObservation
from cdf_gco_crms.references.synthetic_index import synthetic_reference_index
from tools.mat_dna.seed_calibration_demo import DEMO


def seed(db):
    store = CalibrationStore(db)
    user = ReviewerContext('SYNTHETIC_ADMIN', frozenset({'CALIBRATION_ADMIN','CALIBRATION_REVIEWER','DNA_RULE_APPROVER'}), True)
    calibration = CalibrationService(store)
    intelligence = IntelligenceService(calibration)
    case_id = 'SYN-COMMAND-001'
    try:
        if store.db.execute('SELECT 1 FROM cases WHERE case_id=?', (case_id,)).fetchone():
            return case_id  # Never replace or append to a reviewer's existing case.
        with tempfile.TemporaryDirectory(prefix='cdf-visual-synthetic-') as temp:
            root = Path(temp)
            docx = root / 'SYNTHETIC-agreement.docx'
            word = WordDocument()
            quotes = [f'Synthetic clause {i}: approval and control evidence must be documented.' for i in range(1,5)]
            for quote in quotes: word.add_paragraph(quote)
            word.save(docx)
            calibration.import_case(case_id, hashlib.sha256(docx.read_bytes()).hexdigest(), DEMO, user)
            state = build_round1(Document(name='اتفاقية تجريبية',code='SYN',owning_department='الحوكمة والامتثال',is_demo=True),
                [NewObservation(f'الفقرة {i+1}',quotes[i],GText.plain(item['original_text'])) for i,item in enumerate(DEMO)],
                date='2026-10-24', document_version='SYNTHETIC-v1')
            for observation, item in zip(state.observations, DEMO): observation.internal_id=item['observation_id']
            def refs(payload): return [e['id'] for e in payload.get('evidence',[]) if e.get('kind')=='DOCUMENT'][:1]
            def black(payload):
                objection = payload['observation']['internal_id'] in ('SYN-001','SYN-003') and not payload.get('prior_exchange')
                return {'material_objection':objection,'objection_kind':'UNCLEAR_ACTION' if objection else 'NO_MATERIAL_OBJECTION',
                    'challenged_claim':'نطاق التوصية','proposition':'تحتاج التوصية إلى إجراء محدد ونطاق يمكن التحقق منه.' if objection else '',
                    'material_consequence':'تعذر التحقق من التنفيذ' if objection else '', 'evidence_refs':refs(payload),
                    'existing_controls':[],'requested_correction':'تحديد النطاق','reasoning_summary':'نص اصطناعي لاختبار دورة النقاش.',
                    'evidence_limitations':'لا توجد سلطة مؤسسية معتمدة في هذه التجربة.','confidence':0.5}
            script={C.ROLE_DRAFTING:{'observations':[],'gaps_not_drafted':[],'reasoning_summary':'حفظ المرشحات التجريبية الحالية.'},
                C.ROLE_BLACK:black,
                C.ROLE_RED:lambda p:{'stance':'REVISE','objection_accepted':True,'reasoning':'يقبل النص التجريبي تضييق نطاق التوصية.',
                    'evidence_refs':refs(p),'revised_comment':'التوصية: تحديد الإجراء ونطاق التحقق قبل الاعتماد.',
                    'revised_quote':None,'route_to':None,'missing_evidence':[],'escalate':False,'new_argument':True},
                C.ROLE_ADJUDICATOR:lambda p:{'disposition':'UPHELD_WITH_REVISION','objection_valid':True,'rationale':'تعديل اصطناعي مقترح ينتظر المراجع البشري.',
                    'decisive_evidence':refs(p),'inadmissible_refs':[],'facts':[],'inferences':[],'unverified_assumptions':[],
                    'remedy_assessment':'تضييق النطاق','open_question':None,'escalate':False},
                C.ROLE_EDITORIAL:{'revised_comment':'التوصية: تحديد الإجراء ونطاق التحقق قبل الاعتماد.','finding_part':'إجراء غير محدد',
                    'recommendation_part':'تحديد الإجراء','proposed_title':'نطاق التحقق','changes':['تضييق النطاق'],'preserved_references':[],
                    'requires_human_approval':True,'house_style':{'finding_before_recommendation':True,'no_accusation':True,'action_labelled':True,'within_word_ceiling':True}},
                C.ROLE_FINAL:{'release_recommendation':'NOT_READY','findings':[],
                    'unresolved_critical_objections':[],'contradictions':[],'jurisdiction_consistent':True,'unsupported_or_excessive':[],
                    'summary':'لا يصرح هذا المسار بإصدار مؤسسي.'}}
            trace = TraceWriter()
            runner = AgentRunner(MockAdapter(script=script), trace)
            profile=load_profile(synthetic_fixture_path())
            engine = ReviewOrchestrator(synthetic_reference_index(),runner=runner,mat_profile=profile,gate_dir=root/'gate')
            # Deliberately exercise existing MAT deterministic fallback alongside mock debate.
            engine.mat = MatAdvisor(profile,None)
            execution = engine.execute(state,docx,draft=True)
            intelligence.capture(case_id,execution,trace.records,user,gate_dir=root/'gate')
            return case_id
    finally:
        store.close()


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--db',default='build/private/mat-calibration-dev.sqlite')
    args=parser.parse_args()
    print('Synthetic case ready:',seed(args.db))
