#!/usr/bin/env python3
"""Generic renderer: MAT DNA JSON deliverables → .xlsx workbooks (public code; the content it renders stays private).

Usage: python tools/mat_dna/dna_deliverables.py DELIVERABLES_DIR
Renders, when present: MAT_PROFESSIONAL_DNA_QUESTIONNAIRE.json → .xlsx, MAT_RECOMMENDATION_QUESTION_DISTRIBUTION.json → .xlsx,
MAT_OBSERVATION_EVOLUTION_REGISTER.json → .xlsx. Lists and dicts are serialised as JSON text in cells.
"""
from __future__ import annotations
import json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


def _cell(v):
    return json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v


def _sheet(ws, rows: list[dict]):
    if not rows:
        return
    keys = list(rows[0].keys())
    ws.append(keys)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(r.get(k)) for k in keys])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 28
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_view.rightToLeft = True


def render(d: Path) -> list[Path]:
    out = []
    q = d / "MAT_PROFESSIONAL_DNA_QUESTIONNAIRE.json"
    if q.exists():
        data = json.loads(q.read_text(encoding="utf-8"))
        wb = Workbook(); ws = wb.active; ws.title = "questionnaire"; _sheet(ws, data["questions"])
        meta = wb.create_sheet("meta"); _sheet(meta, [{k: v for k, v in data.items() if k != "questions"}])
        p = d / "MAT_PROFESSIONAL_DNA_QUESTIONNAIRE.xlsx"; wb.save(p); out.append(p)
    r = d / "MAT_RECOMMENDATION_QUESTION_DISTRIBUTION.json"
    if r.exists():
        data = json.loads(r.read_text(encoding="utf-8"))
        wb = Workbook(); ws = wb.active; ws.title = "distribution"
        _sheet(ws, [{"metric": k, "value": _cell(v)} for k, v in data.items()])
        p = d / "MAT_RECOMMENDATION_QUESTION_DISTRIBUTION.xlsx"; wb.save(p); out.append(p)
    e = d / "MAT_OBSERVATION_EVOLUTION_REGISTER.json"
    if e.exists():
        data = json.loads(e.read_text(encoding="utf-8"))
        wb = Workbook(); ws = wb.active; ws.title = "register"; _sheet(ws, data["rows"])
        p = d / "MAT_OBSERVATION_EVOLUTION_REGISTER.xlsx"; wb.save(p); out.append(p)
    return out


if __name__ == "__main__":
    for p in render(Path(sys.argv[1])):
        print(p)
