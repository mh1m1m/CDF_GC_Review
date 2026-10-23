#!/usr/bin/env python3
"""Packages the source and the complete delivery, then reopens and checks every deliverable."""
import json
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from cdf_gco_crms.config import workbook_settings as W  # noqa: E402

OUT = ROOT / "deliverables"
SRC_ZIP = OUT / "CDF-GCO-CRMS-v7.12-SOURCE.zip"
PKG_ZIP = OUT / "CDF-GCO-CRMS-v7.12-COMPLETE-PACKAGE.zip"
EXCLUDE = ("__pycache__", ".pyc", ".pytest_cache")


def add_tree(z, base: Path, arc_prefix: str):
    for p in sorted(base.rglob("*")):
        if p.is_file() and not any(x in str(p) for x in EXCLUDE):
            z.write(p, f"{arc_prefix}/{p.relative_to(base)}")


def source_zip():
    with zipfile.ZipFile(SRC_ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for name in ("main.py", "migrate_v710_to_v711.py", "migrate_to_v712.py", "README.md", "requirements.txt", "requirements-lock.txt"):
            z.write(ROOT / name, f"cdf-gco-crms-v7.12/{name}")
        for d in ("cdf_gco_crms", "tests", "tools", "assets", "sources"):
            add_tree(z, ROOT / d, f"cdf-gco-crms-v7.12/{d}")


def complete_zip():
    import shutil
    shutil.copy2(ROOT / "migrate_v710_to_v711.py", OUT / "migrate_v710_to_v711.py"); shutil.copy2(ROOT / "migrate_to_v712.py", OUT / "migrate_to_v712.py")
    with zipfile.ZipFile(PKG_ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob("*")):
            if p.is_file() and p != PKG_ZIP and not any(x in str(p) for x in EXCLUDE):
                z.write(p, f"CDF-GCO-CRMS-v7.12/{p.relative_to(OUT)}")


def verify():
    import hashlib
    from openpyxl import load_workbook
    from docx import Document
    h = lambda f: hashlib.sha256(f.read_bytes()).hexdigest()  # noqa: E731
    originals = {h(f) for f in list((ROOT / "sources").rglob("*.xlsx")) + list((OUT / "migration_demo").rglob("*.xlsx"))
                 if "backup" not in f.parts}
    rows = []
    for p in sorted(OUT.rglob("*")):
        if not p.is_file():
            continue
        ok, note = True, ""
        try:
            if "backup" in p.parts:
                ok = h(p) in originals
                note = "backup byte-identical to its source" if ok else "backup differs from every source"
            elif p.suffix == ".xlsx":
                wb = load_workbook(p)
                pr = wb.properties
                props = (zipfile.ZipFile(p).read("docProps/core.xml") + zipfile.ZipFile(p).read("docProps/app.xml")).decode().lower()
                ok = pr.creator == W.AUTHOR and pr.lastModifiedBy == W.AUTHOR and not any(
                    x in props for x in W.PROHIBITED_PROPERTY_VALUES) and W.COMPANY in props
                note = f"{len(wb.sheetnames)} sheets; title={pr.title}"
            elif p.suffix == ".docx":
                d = Document(p)
                cp = d.core_properties
                props = (zipfile.ZipFile(p).read("docProps/core.xml") + zipfile.ZipFile(p).read("docProps/app.xml")).decode().lower()
                ok = cp.author == W.AUTHOR and cp.last_modified_by == W.AUTHOR and not any(
                    x in props for x in W.PROHIBITED_PROPERTY_VALUES)
                note = f"{len(d.paragraphs)} paragraphs; title={cp.title}"
            elif p.suffix == ".zip":
                bad = zipfile.ZipFile(p).testzip()
                ok = bad is None
                note = f"{len(zipfile.ZipFile(p).namelist())} entries"
            elif p.suffix in (".md",):
                head = p.read_text(encoding="utf-8").splitlines()[0]
                ok = W.AUTHOR in head and W.COMPANY in head
                note = "metadata header line"
            elif p.suffix == ".json":
                json.loads(p.read_text(encoding="utf-8"))
                note = "valid JSON"
            elif p.suffix == ".txt":
                note = f"{len(p.read_text(encoding='utf-8').splitlines())} lines"
        except Exception as e:  # noqa: BLE001
            ok, note = False, f"{type(e).__name__}: {e}"
        rows.append((str(p.relative_to(OUT)), ok, note))
    return rows


if __name__ == "__main__":
    stage = sys.argv[1] if len(sys.argv) > 1 else "all"
    if stage in ("source", "all"):
        import shutil
        shutil.copy2(ROOT / "migrate_v710_to_v711.py", OUT / "migrate_v710_to_v711.py"); shutil.copy2(ROOT / "migrate_to_v712.py", OUT / "migrate_to_v712.py")
        source_zip()
        if stage == "source":
            print("source zip written")
            sys.exit(0)
    complete_zip()
    res = verify()
    bad = [r for r in res if not r[1]]
    for r in res:
        if "/" not in r[0] or r[0].endswith(".zip"):
            print(("OK  " if r[1] else "BAD ") + r[0], "|", r[2][:90])
    print(f"{len(res)} files checked, {len(bad)} problems")
    for b in bad:
        print("PROBLEM", b)
    (ROOT / "build" / "package_check.json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
