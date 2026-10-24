"""Synthetic stable-LibreOffice smoke checks. Never accepts institutional inputs."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import signal
import subprocess
import tempfile
import time
import zipfile


def run_bounded(command: list[str], timeout: float) -> subprocess.CompletedProcess:
    """Isolate and reap this invocation's process group, including on timeout."""
    process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                               text=True, start_new_session=True)
    try:
        stdout, stderr = process.communicate(timeout=timeout)
    except BaseException:
        # A launcher may exit while a child retains the pipe. Kill its group even
        # when the launcher has already exited; never use global pkill/killall.
        try:
            os.killpg(process.pid, signal.SIGTERM)
        except ProcessLookupError:
            pass
        try:
            process.communicate(timeout=2)
        except subprocess.TimeoutExpired:
            pass
        finally:
            try:
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
            process.communicate()
        raise
    if process.returncode:
        raise subprocess.CalledProcessError(process.returncode, command, output=stdout, stderr=stderr)
    return subprocess.CompletedProcess(command, process.returncode, stdout, stderr)


def stable_version(value: str) -> bool:
    return bool(re.match(r"^LibreOffice \d+\.\d+", value)) and not re.search(
        r"alpha|beta|nightly|dev|rc\d", value, re.I)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def create_presentation(path: Path) -> None:
    """Minimal editable ODP using only stdlib ZIP/XML; no new presentation SDK."""
    content = '''<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" office:version="1.2">
 <office:body><office:presentation><draw:page draw:name="Synthetic slide">
 <draw:frame svg:x="1cm" svg:y="1cm" svg:width="20cm" svg:height="5cm">
 <draw:text-box><text:p>CDF SYNTHETIC IMPRESS</text:p></draw:text-box></draw:frame>
 </draw:page></office:presentation></office:body></office:document-content>'''
    manifest = '''<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">
 <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.presentation"/>
 <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>'''
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.presentation", compress_type=zipfile.ZIP_STORED)
        archive.writestr("content.xml", content)
        archive.writestr("META-INF/manifest.xml", manifest)


def validate_pdf(path: Path, marker: str, *, arabic: bool = False) -> dict:
    from pypdf import PdfReader
    reader = PdfReader(path, strict=True)
    if reader.is_encrypted or not reader.pages:
        raise AssertionError("PDF is encrypted or has no pages")
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if marker not in text:
        raise AssertionError(f"PDF is missing readable synthetic marker: {marker}")
    if arabic and sum("\u0600" <= c <= "\u06ff" for c in text) < 6:
        raise AssertionError("PDF is missing readable Arabic text")
    return {"pages": len(reader.pages), "readable_marker": True, "arabic_checked": arabic}


def smoke(timeout: float = 45) -> dict:
    from docx import Document
    from openpyxl import Workbook, load_workbook

    result = {"status": "FAILED", "checks": [], "runtime": {}}
    executable = shutil.which("soffice")
    if not executable:
        result["error"] = "soffice is not installed or not on PATH"
        return result
    result["executable"] = executable
    try:
        version = run_bounded([executable, "--version"], min(timeout, 15)).stdout.strip()
        result["version"] = version
        if not stable_version(version):
            raise AssertionError("Stable LibreOffice required; development/prerelease builds refused")
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    for name in ("base-image.txt", "packages.tsv", "python-packages.txt"):
        path = Path("/opt/cdf-runtime") / name
        if path.is_file():
            result["runtime"][name] = path.read_text()
    with tempfile.TemporaryDirectory(prefix="cdf-lo-smoke-") as temp:
        root = Path(temp)
        source = root / "input"
        source.mkdir()
        book = Workbook()
        sheet = book.active
        sheet["A1"], sheet["A2"], sheet["A3"] = 2, 3, "=SUM(A1:A2)"
        book.save(source / "calc.xlsx")
        with (source / "calc.xlsx").open("rb") as stream:
            assert load_workbook(stream, data_only=True).active["A3"].value is None, "Input must have no cached result"
        doc = Document()
        doc.add_paragraph("CDF SYNTHETIC WRITER")
        paragraph = doc.add_paragraph("مراجعة حوكمة تجريبية")
        paragraph.runs[0].font.name = "Noto Sans Arabic"
        doc.save(source / "writer.docx")
        create_presentation(source / "impress.odp")
        for component, filename, target in (("Calc", "calc.xlsx", "xlsx:Calc MS Excel 2007 XML"),
                                             ("Writer", "writer.docx", "pdf:writer_pdf_Export"),
                                             ("Impress", "impress.odp", "pdf:impress_pdf_Export")):
            started = time.monotonic()
            entry = {"component": component, "status": "FAILED"}
            input_file = source / filename
            original_hash = digest(input_file)
            output = root / component / "output"
            output.mkdir(parents=True)
            profile = root / component / "profile"
            try:
                run_bounded([executable, f"-env:UserInstallation={profile.as_uri()}",
                             "--headless", "--norestore", "--convert-to", target,
                             "--outdir", str(output), str(input_file)], timeout)
                if component == "Calc":
                    recalculated = output / filename
                    value = load_workbook(recalculated, data_only=True).active["A3"].value
                    assert value == 5, f"Expected independently computed SUM=5, got {value!r}"
                    assert load_workbook(recalculated, data_only=False).active["A3"].data_type == "f", "Formula must survive"
                    entry["computed_sum"] = value
                else:
                    entry.update(validate_pdf(output / (input_file.stem + ".pdf"),
                                 "CDF SYNTHETIC " + component.upper(), arabic=component == "Writer"))
                assert digest(input_file) == original_hash, "Conversion modified the source"
                entry.update(status="PASSED", source_unchanged=True)
            except Exception as exc:
                entry["error"] = f"{type(exc).__name__}: {exc}"
            entry["duration_seconds"] = round(time.monotonic() - started, 3)
            result["checks"].append(entry)
    result["status"] = "PASSED" if all(x["status"] == "PASSED" for x in result["checks"]) else "FAILED"
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--timeout", type=float, default=45, help="Positive per-process timeout in seconds")
    parser.add_argument("--report", type=Path, default=Path("build/libreoffice/smoke.json"))
    args = parser.parse_args()
    if args.timeout <= 0:
        parser.error("--timeout must be positive")
    result = smoke(args.timeout)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    # Detailed package provenance stays in the report; stdout stays concise.
    print(json.dumps({k: v for k, v in result.items() if k != "runtime"}, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "PASSED" else 1


if __name__ == "__main__":
    raise SystemExit(main())
