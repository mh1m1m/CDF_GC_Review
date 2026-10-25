"""v7.10 or v7.11 to v7.12 migration of a requester workbook (protocol 10.4A.13).

A v7.11 source already carries a dashboard and a helper sheet; both are removed and rebuilt in the v7.12
design, the data sheet being kept exactly as in a v7.10 migration.

The original file is never opened for writing. The migration works on a copy:
the data sheet is kept exactly as it is (every value, rich-text run, fill,
dropdown and conditional format), guidance notes are added to its headers
where none exist, the helper sheet and the dashboard are added, the print
view "عرض الطباعة" is removed only after the dashboard exists, and the result
is saved under the same file name in another folder (09.1 stems unchanged).
No round block is appended by a migration.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import shutil
from dataclasses import asdict, dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..config import round_schema as S
from ..config import workbook_settings as W
from ..excel.workbook_builder import RIYADH, assemble, save_atomic, set_properties
from ..guidance import guidance_dictionary as G
from ..ingestion.workbook_reader import parse_workbook
from ..model.entities import DashboardViewState
from ..reconciliation.historical import _signature


class MigrationError(Exception):
    pass


@dataclass
class MigrationReport:
    source: str
    source_sha256: str
    backup: str
    output: str
    output_sha256: str = ""
    source_kind: str = ""
    data_sheet: str = ""
    rounds: int = 0
    observations: int = 0
    removed_sheets: list[str] = field(default_factory=list)
    kept_extra_sheets: list[str] = field(default_factory=list)
    notes_added: int = 0
    input_messages_added: int = 0
    cells_compared: int = 0
    cell_differences: list[str] = field(default_factory=list)
    validations_before: int = 0
    validations_after: int = 0
    cf_ranges_before: int = 0
    cf_ranges_after: int = 0
    ids_preserved: bool = False
    no_round_added: bool = False
    dashboard_first_active: bool = False
    print_view_absent: bool = False
    metadata_ok: bool = False
    not_recorded: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    status: str = "FAILED"

    def to_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False, indent=2)


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _sheet_snapshot(ws) -> dict[str, tuple]:
    snap = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            fill = c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type else None
            snap[c.coordinate] = (_signature(c.value), c.number_format, fill)
    return snap


def migrate(src, out_dir, *, backup_dir=None, overrides: dict | None = None, round_meta: dict | None = None,
            add_guidance: bool = True, view_state: DashboardViewState | None = None) -> MigrationReport:
    src = Path(src).resolve()
    out_dir = Path(out_dir).resolve()
    if not src.exists():
        raise MigrationError(f"source not found: {src}")
    dst = out_dir / src.name
    if dst == src:
        raise MigrationError("the output must not overwrite the source workbook")
    stamp = dt.datetime.now(RIYADH).strftime("%Y%m%dT%H%M%S")
    backup_dir = Path(backup_dir or (out_dir / "backup")).resolve()
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"{src.stem}.{stamp}.bak.xlsx"
    shutil.copy2(src, backup)
    rep = MigrationReport(str(src), _sha(src), str(backup), str(dst))
    if _sha(backup) != rep.source_sha256:
        raise MigrationError("backup copy differs from the source")

    # 1. validate the source (header-driven discovery, IDs, birth rounds)
    parsed = parse_workbook(src, overrides=overrides, round_meta=round_meta)
    rep.source_kind = parsed.kind
    rep.data_sheet = parsed.data_ws.title
    rep.rounds = parsed.schema.current_round
    rep.observations = len(parsed.state.observations)
    rep.warnings += parsed.warnings
    if parsed.unkeyed_rows:
        rep.warnings.append(f"rows without Internal_ID kept as owner input (SC-41): {parsed.unkeyed_rows}")
    st = parsed.state
    for r in st.rounds:
        if not r.date:
            rep.not_recorded.append(f"round {r.number} date")
    if not st.document.cover_version:
        rep.not_recorded.append("cover version")
    if not st.document.owning_department:
        rep.not_recorded.append("owning department")

    # 2. work on a copy
    work = out_dir / f".{src.stem}.migrating.xlsx"
    out_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, work)
    try:
        wb = load_workbook(work, rich_text=True)
        data_ws = wb[parsed.data_ws.title]
        before = _sheet_snapshot(data_ws)
        rep.validations_before = len(data_ws.data_validations.dataValidation)
        rep.cf_ranges_before = len(list(data_ws.conditional_formatting))

        # 3. additive guidance on the data sheet (notes on headers, input messages on current response cells)
        if add_guidance:
            R = parsed.schema.current_round
            for c, key in ((1, "reader_no"), (3, "location"), (4, "quote"), (5, "comment")):
                cell = data_ws.cell(row=1, column=c)
                if cell.comment is None:
                    cell.comment = Comment(G.fixed_note(key), G.NOTE_AUTHOR, width=320, height=110)
                    rep.notes_added += 1
            for r, roles in parsed.schema.blocks.items():
                for role, col in roles.items():
                    cell = data_ws.cell(row=1, column=col)
                    if cell.comment is None:
                        text = G.header_note(role, r, R)
                        cell.comment = Comment(text, G.NOTE_AUTHOR, width=320,
                                               height=min(380, 17 * (text.count("\n") + len(text) // 45 + 1) + 28))
                        rep.notes_added += 1
            col = get_column_letter(parsed.schema.col(R, S.ROLE_RESPONSE))
            rows = sorted(parsed.rows.values())
            if rows:
                t, p = G.input_message("response_current")
                dv = DataValidation(type=None, allow_blank=True, showInputMessage=True, showErrorMessage=False,
                                    promptTitle=t, prompt=p)
                open_rows = [parsed.rows[o.internal_id] for o in st.observations if o.value(R) != S.CLOSED]
                if open_rows:
                    dv.sqref = " ".join(f"{col}{x}" for x in open_rows)
                    data_ws.add_data_validation(dv)
                    rep.input_messages_added += 1
            if R >= 3:
                for r in range(1, R - 1):
                    for c in parsed.schema.blocks[r].values():
                        data_ws.column_dimensions[get_column_letter(c)].outline_level = 1
                data_ws.sheet_properties.outlinePr.summaryRight = False

        # 4. remove an earlier dashboard or helper (prototype) before building the engine's own
        for name in (W.DASHBOARD_SHEET, W.HELPER_SHEET):
            if name in wb.sheetnames:
                del wb[name]
                rep.removed_sheets.append(name)
        for nm in list(wb.defined_names.keys()):
            if nm.startswith("crms_"):
                del wb.defined_names[nm]

        # 5. dashboard first, then retire the print view (proposal PHASE 10)
        an, view, reg, _refs, _dash = assemble(wb, st, data_ws, parsed.schema, parsed.rows, view_state=view_state)
        if W.LEGACY_PRINT_VIEW_SHEET in wb.sheetnames:
            del wb[W.LEGACY_PRINT_VIEW_SHEET]
            rep.removed_sheets.append(W.LEGACY_PRINT_VIEW_SHEET)
        rep.kept_extra_sheets = [s for s in wb.sheetnames if s not in (W.DASHBOARD_SHEET, W.HELPER_SHEET,
                                                                          data_ws.title)]
        build_time = dt.datetime.now(RIYADH)
        set_properties(wb, st, build_time)
        save_atomic(wb, dst, reg.expected(view), build_time)      # cached values of the default view
    finally:
        if work.exists():
            work.unlink()

    # 6. reopen and verify
    _verify(rep, dst, before, parsed)
    rep.output_sha256 = _sha(dst)
    if _sha(src) != rep.source_sha256:
        rep.warnings.append("source changed during migration")
        rep.status = "FAILED"
    return rep


def _verify(rep: MigrationReport, dst: Path, before: dict, parsed):
    wb = load_workbook(dst, rich_text=True)
    ws = wb[rep.data_sheet]
    after = _sheet_snapshot(ws)
    diffs = []
    for coord, sig in before.items():
        rep.cells_compared += 1
        if after.get(coord) != sig:
            diffs.append(f"{coord}: {sig[0][0][:40]!r} != {after.get(coord, (('',),))[0][0][:40]!r}")
    rep.cell_differences = diffs[:50]
    rep.validations_after = len(ws.data_validations.dataValidation)
    rep.cf_ranges_after = len(list(ws.conditional_formatting))
    reparsed = parse_workbook(dst)
    rep.ids_preserved = (sorted(o.internal_id for o in reparsed.state.observations)
                         == sorted(o.internal_id for o in parsed.state.observations))
    rep.no_round_added = reparsed.schema.rounds == parsed.schema.rounds
    rep.dashboard_first_active = wb.sheetnames[0] == W.DASHBOARD_SHEET and wb.active.title == W.DASHBOARD_SHEET
    rep.print_view_absent = W.LEGACY_PRINT_VIEW_SHEET not in wb.sheetnames
    p = wb.properties
    rep.metadata_ok = (p.creator == W.AUTHOR and p.lastModifiedBy == W.AUTHOR
                       and p.title == W.workbook_title(parsed.state.document.name, parsed.schema.current_round))
    ok = (not diffs and rep.ids_preserved and rep.no_round_added and rep.dashboard_first_active
          and rep.print_view_absent and rep.metadata_ok and rep.validations_after >= rep.validations_before
          and rep.cf_ranges_after >= rep.cf_ranges_before)
    rep.status = "PASSED" if ok else "FAILED"
