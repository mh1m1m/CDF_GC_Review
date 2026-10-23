"""Writer of v7.10-format requester workbooks (regression fixtures for the migration, TEST 13).

Reproduces Art. 10.4 of v7.10: one data sheet named per round, followed by the
print-view sheet "عرض الطباعة" (columns A, C, D, E and the latest two round
blocks, copied as values with the same formats, a first row stating that
responses are entered in the data sheet only). No dashboard, no helper sheet,
no notes. Fixtures are synthetic and flagged as such in their keywords.
"""
from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ..config import round_schema as S
from ..config import workbook_settings as W
from ..excel import styles as st
from ..config import design_tokens as T
from ..excel.register_writer import write_data_sheet
from ..model.entities import ReviewState

PRINT_NOTE = "تُدخل الإفادات وحالات التطبيق في ورقة الملاحظات فقط؛ هذه الورقة للطباعة."


def write_v710(state: ReviewState, path, *, build_time=None):
    import datetime as dt
    from ..excel.workbook_builder import RIYADH, postprocess
    build_time = build_time or dt.datetime.now(RIYADH)
    wb = Workbook()
    wb.remove(wb.active)
    ws, schema = write_data_sheet(wb, state, outline=False, guidance=False, closed_guide_colour=T.PLACEHOLDER_GRAY)
    R = state.current_round
    pv = wb.create_sheet(W.LEGACY_PRINT_VIEW_SHEET)
    pv.sheet_view.rightToLeft = True
    pv.sheet_view.showGridLines = False
    cols = [1, 3, 4, 5]
    for r in range(max(1, R - 1), R + 1):
        cols += schema.block_columns(r)
    pv.cell(row=1, column=1, value=PRINT_NOTE).font = st.font(10, bold=True)
    pv.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    for j, c in enumerate(cols, start=1):
        pv.column_dimensions[get_column_letter(j)].width = ws.column_dimensions[get_column_letter(c)].width
        for i in range(1, ws.max_row + 1):
            src = ws.cell(row=i, column=c)
            dst = pv.cell(row=i + 1, column=j, value=src.value)
            if src.has_style:
                dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
                dst.alignment, dst.number_format = copy(src.alignment), src.number_format
    for i in range(1, ws.max_row + 1):
        pv.row_dimensions[i + 1].height = ws.row_dimensions[i].height
    pv.page_setup.orientation = "landscape"
    pv.page_setup.paperSize = pv.PAPERSIZE_A4
    pv.page_setup.fitToWidth = 1
    pv.page_setup.fitToHeight = 0
    pv.sheet_properties.pageSetUpPr.fitToPage = True
    pv.print_area = f"A1:{get_column_letter(len(cols))}{ws.max_row + 1}"
    wb.active = 0
    p = wb.properties
    p.creator = p.lastModifiedBy = W.AUTHOR
    p.title = W.workbook_title(state.document.name, R)
    p.keywords = "بيانات تجريبية – نموذج اختبار بصيغة v7.10"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    postprocess(path, None, build_time)
    return path
