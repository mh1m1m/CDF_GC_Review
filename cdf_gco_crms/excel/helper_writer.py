"""Hidden helper sheet "بيانات اللوحة" (protocol 10.4A.3).

Tables, each identified by the names in row 1 (never by position):
  meta.*    document block and the view state read from the dashboard selectors
  list.*    dropdown lists
  round.*   round map: header texts and the column of each role, found by MATCH
            on the data sheet's header row (header-driven, 10.3)
  obs.*     one row per observation (row found by Internal_ID) and its values
            for the selected view
  long.*    one row per observation and round: official value and derived
            indicators (D.10)
  trend.*   one row per round
  feed.*/bar.*  chart feeds
  q.*/t.*   order of the action queue and of the observation table

The sheet holds only allowlisted fields (10.4A.4, G-SHARE-3).
"""
from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from ..config import round_schema as S
from ..config import workbook_settings as W
from ..observations import text_derivation as TD
from .formula_registry import NA, Registry
from . import styles as st


def q(s: str) -> str:
    return '"' + s.replace('"', '""') + '"'


def sheet_ref(name: str) -> str:
    return "'" + name.replace("'", "''") + "'!"


HP = sheet_ref(W.HELPER_SHEET)
DP = sheet_ref(W.DASHBOARD_SHEET)


@dataclass
class Table:
    name: str
    first_col: int
    cols: dict[str, int]
    nrows: int

    def c(self, col: str) -> str:
        return get_column_letter(self.cols[col])

    def cell(self, col: str, i: int, *, abs_: bool = False, sheet: bool = False) -> str:
        letter = self.c(col)
        ref = f"${letter}${i + 1}" if abs_ else f"{letter}{i + 1}"
        return (HP if sheet else "") + ref

    def rng(self, col: str, *, sheet: bool = False) -> str:
        letter = self.c(col)
        return (HP if sheet else "") + f"${letter}$2:${letter}${max(self.nrows, 1) + 1}"


class HelperBook:
    def __init__(self, ws):
        self.ws = ws
        self.next_col = 1
        self.tables: dict[str, Table] = {}

    def table(self, name: str, columns: list[str], nrows: int) -> Table:
        cols = {}
        for i, cname in enumerate(columns):
            col = self.next_col + i
            cols[cname] = col
            cell = self.ws.cell(row=1, column=col, value=f"{name}.{cname}")
            cell.font = st.font(9, bold=True)
        t = Table(name, self.next_col, cols, nrows)
        self.next_col += len(columns) + 1
        self.tables[name] = t
        return t


@dataclass
class HelperRefs:
    meta_rows: dict[str, int] = field(default_factory=dict)
    tables: dict[str, Table] = field(default_factory=dict)

    def meta(self, key: str, *, sheet: bool = True) -> str:
        return (HP if sheet else "") + f"$B${self.meta_rows[key]}"


META_KEYS = ["doc_name", "cover_version", "owning_department", "current_round", "current_label",
             "data_sheet", "is_demo", "schema", "engine_version", "obs_count", "round_count",
             "selected_round", "compare_round", "scope", "selected_k", "latest_date", "queue_count",
             "table_count", "complete_filled", "complete_needed", "selected_label", "compare_label"]

LONG_COLS = ["k", "round", "id", "birth", "exists", "row", "ccol", "fcol", "rcol", "icol", "value", "prev",
             "closed", "closed_event", "ecb", "ever", "reopen_event", "reopen_state", "stalled", "change",
             "priority", "reason", "action", "feedback", "fb_short", "fb_cut", "cur_loc", "basis_static", "resp",
             "resp_filled", "lfr", "impl", "impl_filled", "lfi", "lcr", "invalid"]

OBS_COLS = ["k", "reader", "reader_label", "id", "birth", "row", "location", "comment", "summary", "comment_cut",
            "title_disp", "title", "title_status", "title_by", "title_on", "title_ref",
            "idx", "exists", "v", "prev_disp", "change", "priority", "reason", "action", "cur_loc", "idxC", "vC", "codeS",
            "codeC", "qkey", "include", "rank", "vR", "link"]

TREND_COLS = ["r", "label", "issued", "new", "closed", "cond", "req", "open", "invalid", "rate",
              "closed_event", "reopen_event", "stalled", "reopen_state", "ever"]

DETAIL_CUT = 300


def write_helper(wb, an, schema, sel, reg: Registry, *, data_sheet: str, data_last_row: int,
                 data_last_col: int, state) -> HelperRefs:
    ws = wb.create_sheet(W.HELPER_SHEET)
    ws.sheet_state = "hidden"
    ws.sheet_view.rightToLeft = True
    hb = HelperBook(ws)
    refs = HelperRefs()
    R = an.R
    N = len(an.obs)
    obs = an.obs
    DS = sheet_ref(data_sheet)
    max_row = data_last_row + W.CALC_ROW_MARGIN
    max_col = data_last_col + 40
    data_rng = f"{DS}$A$1:${get_column_letter(max_col)}${max_row}"
    hdr_rng = f"{DS}$A$1:${get_column_letter(max_col)}$1"
    id_rng = f"{DS}$B$1:$B${max_row}"

    # ------------------------------------------------------------------ meta
    meta = hb.table("meta", ["key", "value"], len(META_KEYS))
    for i, key in enumerate(META_KEYS, start=1):
        ws.cell(row=i + 1, column=meta.cols["key"], value=key)
        refs.meta_rows[key] = i + 1
    def mset(key, value):
        ws.cell(row=refs.meta_rows[key], column=meta.cols["value"], value=value)
    doc = state.document
    cur = state.round(R)
    mset("doc_name", doc.name)
    mset("cover_version", doc.cover_version or cur.document_version or S.NOT_RECORDED)
    mset("owning_department", doc.owning_department or S.NOT_RECORDED)
    mset("current_round", R)
    mset("current_label", S.round_label(R) if R else "")
    mset("data_sheet", data_sheet)
    mset("is_demo", "1" if doc.is_demo else "0")
    mset("schema", "CDF-GCO-CRMS/requester-workbook/7.11")
    from .. import ENGINE_VERSION
    mset("engine_version", ENGINE_VERSION)
    mset("obs_count", N)
    mset("round_count", R)
    mset("latest_date", cur.date or S.NOT_RECORDED)

    # ------------------------------------------------------------------ lists
    lists = hb.table("list", ["scope", "status4", "impl_owner"], max(len(S.SCOPES), 5))
    for i, v in enumerate(S.SCOPES, start=1):
        ws.cell(row=i + 1, column=lists.cols["scope"], value=v)
    for i, v in enumerate(S.CLOSURE_VALUES, start=1):
        ws.cell(row=i + 1, column=lists.cols["status4"], value=v)
    for i, v in enumerate(S.OWNER_IMPL_VALUES, start=1):
        ws.cell(row=i + 1, column=lists.cols["impl_owner"], value=v)
    status4 = f"${lists.c('status4')}$2:${lists.c('status4')}$5"
    impl_owner = f"${lists.c('impl_owner')}$2:${lists.c('impl_owner')}$6"
    scope_rng = f"${lists.c('scope')}$2:${lists.c('scope')}${len(S.SCOPES) + 1}"

    # ------------------------------------------------------------------ rounds
    rt = hb.table("round", ["no", "ordinal", "label", "date", "version", "h_closure", "h_feedback",
                            "h_response", "h_impl", "c_closure", "c_feedback", "c_response", "c_impl"], R)
    for r in range(1, R + 1):
        rr = state.round(r)
        row = r + 1
        vals = {"no": r, "ordinal": S.ordinal(r), "label": S.round_label(r),
                "date": rr.date or S.NOT_RECORDED, "version": rr.document_version or S.NOT_RECORDED}
        for k, v in vals.items():
            ws.cell(row=row, column=rt.cols[k], value=v)
        for role, hk, ck in ((S.ROLE_CLOSURE, "h_closure", "c_closure"), (S.ROLE_FEEDBACK, "h_feedback", "c_feedback"),
                             (S.ROLE_RESPONSE, "h_response", "c_response"), (S.ROLE_IMPL, "h_impl", "c_impl")):
            col = schema.col(r, role)
            if col is None:
                ws.cell(row=row, column=rt.cols[hk], value="")
                ws.cell(row=row, column=rt.cols[ck], value=0)
                continue
            header = schema.header_text[col]
            ws.cell(row=row, column=rt.cols[hk], value=header)
            reg.put(ws, f"{rt.c(ck)}{row}",
                    f"IFERROR(MATCH({rt.c(hk)}{row},{hdr_rng},0),{col})", col)

    # meta formulas that read the dashboard selectors
    rounds_rng = rt.rng("label")
    obs_t = hb.table("obs", OBS_COLS, N)
    readers_rng = obs_t.rng("reader_label")
    reg.put(ws, refs.meta("selected_round", sheet=False),
            f"IFERROR(MATCH({DP}{sel.sel_round},{rounds_rng},0),{refs.meta('current_round', sheet=False)})",
            lambda v: v.S)
    reg.put(ws, refs.meta("compare_round", sheet=False),
            f"IFERROR(MATCH({DP}{sel.sel_compare},{rounds_rng},0),"
            f"MAX(1,{refs.meta('selected_round', sheet=False)}-1))", lambda v: v.C)
    reg.put(ws, refs.meta("scope", sheet=False),
            f"IF(ISNA(MATCH({DP}{sel.sel_scope},{scope_rng},0)),{q(S.SCOPE_ALL)},{DP}{sel.sel_scope})",
            lambda v: v.scope)
    reg.put(ws, refs.meta("selected_k", sheet=False),
            f"IFERROR(MATCH({DP}{sel.sel_obs},{readers_rng},0),1)", lambda v: v.k)
    reg.put(ws, refs.meta("selected_label", sheet=False),
            f"INDEX({rounds_rng},{refs.meta('selected_round', sheet=False)})", lambda v: S.round_label(v.S))
    reg.put(ws, refs.meta("compare_label", sheet=False),
            f"INDEX({rounds_rng},{refs.meta('compare_round', sheet=False)})", lambda v: S.round_label(v.C))

    Sx = refs.meta("selected_round", sheet=False)
    Cx = refs.meta("compare_round", sheet=False)
    Rx = refs.meta("current_round", sheet=False)
    scope_x = refs.meta("scope", sheet=False)

    # ------------------------------------------------------------------ long
    lt = hb.table("long", LONG_COLS, N * R)
    L = lambda col, i: lt.cell(col, i)                       # noqa: E731
    i = 0
    for a in obs:
        for r in range(1, R + 1):
            i += 1
            x = a.long[r]
            n = i + 1                                         # sheet row
            def put(col, formula, expected):
                reg.put(ws, f"{lt.c(col)}{n}", formula, expected)
            def val(col, v):
                ws.cell(row=n, column=lt.cols[col], value=v)
            val("k", a.k); val("round", r); val("id", a.iid); val("birth", a.birth); val("exists", x.exists)
            put("row", f"INDEX({obs_t.rng('row')},{L('k', i)})", a.row)
            put("ccol", f"INDEX({rt.rng('c_closure')},{L('round', i)})", schema.col(r, S.ROLE_CLOSURE))
            if r == 1:
                val("fcol", 0)
            else:
                put("fcol", f"INDEX({rt.rng('c_feedback')},{L('round', i)})", schema.col(r, S.ROLE_FEEDBACK))
            put("rcol", f"INDEX({rt.rng('c_response')},{L('round', i)})", schema.col(r, S.ROLE_RESPONSE))
            put("icol", f"INDEX({rt.rng('c_impl')},{L('round', i)})", schema.col(r, S.ROLE_IMPL))
            ex, rr_, bb = L("exists", i), L("round", i), L("birth", i)
            v_, p_ = L("value", i), L("prev", i)
            put("value", f'IF({ex}=1,IFERROR(INDEX(crms_data,{L("row", i)},{L("ccol", i)})&"",""),"")', x.value)
            above = lambda col: lt.cell(col, i - 1)       # noqa: E731
            put("prev", f'IF(OR({rr_}=1,{rr_}<={bb}),"",{above("value")})', x.prev)
            put("closed", f'IF({v_}={q(S.CLOSED)},1,0)', x.closed)
            put("closed_event", f'IF(AND({v_}={q(S.CLOSED)},{p_}<>"",{p_}<>{q(S.CLOSED)}),1,0)', x.closed_event)
            put("ecb", f'IF({rr_}<={bb},0,IF(OR({above("ecb")}=1,{above("value")}={q(S.CLOSED)}),1,0))', x.ecb)
            put("ever", f'IF(OR({L("ecb", i)}=1,{L("closed", i)}=1),1,0)', x.ever)
            put("reopen_event", f'IF(AND({p_}={q(S.CLOSED)},{v_}<>"",{v_}<>{q(S.CLOSED)}),1,0)', x.reopen_event)
            put("reopen_state", f'IF(AND({v_}<>"",{v_}<>{q(S.CLOSED)},{L("ecb", i)}=1),1,0)', x.reopen_state)
            put("stalled", f'IF(AND({rr_}-1>{bb},OR({v_}={q(S.REQUIRES_EDIT)},{v_}={q(S.OPEN)}),'
                           f'OR({p_}={q(S.REQUIRES_EDIT)},{p_}={q(S.OPEN)})),1,0)', x.stalled)
            change_f = (f'IF({v_}="","",IF({rr_}={bb},{q(S.CHANGE_NEW)},IF({p_}={q(S.CLOSED)},'
                        f'IF({v_}={q(S.CLOSED)},{q(S.CHANGE_STAYED_CLOSED)},{q(S.CHANGE_REOPENED)}),'
                        f'IF({v_}={q(S.CLOSED)},{q(S.CHANGE_CLOSED)},IF({p_}={v_},{q(S.CHANGE_NO_PROGRESS)},'
                        f'IF({v_}={q(S.CLOSED_COND)},{q(S.CHANGE_COND)},IF({p_}={q(S.CLOSED_COND)},'
                        f'{q(S.CHANGE_COND_REGRESS)},{q(S.CHANGE_SHIFT)})))))))')
            put("change", change_f, x.change)
            put("priority", f'IF(OR({v_}="",{v_}={q(S.CLOSED)}),"",IF({L("reopen_state", i)}=1,1,'
                            f'IF({L("stalled", i)}=1,2,IF({v_}={q(S.CLOSED_COND)},3,'
                            f'IF({v_}={q(S.REQUIRES_EDIT)},4,5)))))', x.priority)
            pr = L("priority", i)
            put("reason", f'IF({pr}="","",CHOOSE({pr},{q(S.REASON_REOPENED)},{q(S.REASON_STALLED)},'
                          f'{q(S.REASON_COND)},{q(S.REASON_EDIT)},IF({rr_}={bb},{q(S.REASON_AWAIT_FIRST)},'
                          f'{q(S.REASON_OPEN)})))', x.reason)
            val("action", x.action)
            if r == 1:
                val("feedback", "")
            else:
                put("feedback", f'IF(OR({ex}=0,{rr_}=1),"",IFERROR(INDEX(crms_data,{L("row", i)},'
                                f'{L("fcol", i)})&"",""))', x.feedback)
            val("fb_short", TD.cut(x.feedback, 70) if r >= 2 else "")
            val("fb_cut", TD.cut(x.feedback, DETAIL_CUT) if r >= 2 else "")
            val("cur_loc", x.cur_loc)
            val("basis_static", x.basis)
            put("resp", f'IF({ex}=1,IFERROR(INDEX(crms_data,{L("row", i)},{L("rcol", i)})&"",""),"")', x.resp)
            rs = L("resp", i)
            pref = S.CLOSED_GUIDE_PREFIX
            put("resp_filled", f'IF(AND({ex}=1,{rs}<>"",NOT(EXACT({rs},{q(S.PH_TEXT)})),'
                               f'LEFT({rs},LEN({q(pref)}))<>{q(pref)}),1,0)', x.resp_filled)
            put("lfr", f'IF({L("resp_filled", i)}=1,{rr_},IF({rr_}=1,0,{above("lfr")}))', x.lfr)
            put("impl", f'IF({ex}=1,IFERROR(INDEX(crms_data,{L("row", i)},{L("icol", i)})&"",""),"")', x.impl)
            put("impl_filled", f'IF(AND({ex}=1,ISNUMBER(MATCH({L("impl", i)},{impl_owner},0))),1,0)',
                x.impl_filled)
            put("lfi", f'IF({L("impl_filled", i)}=1,{rr_},IF({rr_}=1,0,{above("lfi")}))', x.lfi)
            put("lcr", f'IF({v_}={q(S.CLOSED)},IF(AND({rr_}>1,{p_}={q(S.CLOSED)}),{above("lcr")},{rr_}),"")', x.lcr)
            put("invalid", f'IF(AND({ex}=1,ISNA(MATCH({v_},{status4},0))),1,0)', x.invalid)

    # ------------------------------------------------------------------ obs
    O = lambda col, k: obs_t.cell(col, k)                 # noqa: E731
    for a in obs:
        k = a.k
        n = k + 1
        def oput(col, formula, expected):
            reg.put(ws, f"{obs_t.c(col)}{n}", formula, expected)
        def oval(col, v):
            c = ws.cell(row=n, column=obs_t.cols[col], value=v)
            if col in ("reader", "reader_label"):
                c.number_format = "@"
        oval("k", k); oval("reader", a.reader); oval("reader_label", a.reader_label); oval("id", a.iid)
        oval("birth", a.birth); oval("summary", a.summary); oval("comment_cut", TD.cut(a.comment, DETAIL_CUT))
        oval("title_disp", a.title_disp); oval("title", a.title or None); oval("title_status", a.title_status)
        oval("title_by", a.title_by or None); oval("title_on", a.title_on or None); oval("title_ref", a.title_ref or None)
        oput("row", f"IFERROR(MATCH({O('id', k)},{id_rng},0),{a.row})", a.row)
        oput("location", f'IFERROR(INDEX(crms_data,{O("row", k)},3)&"","")', a.location)
        oput("comment", f'IFERROR(INDEX(crms_data,{O("row", k)},5)&"","")', a.comment)
        oput("idx", f"({O('k', k)}-1)*{Rx}+{Sx}", lambda v, k=k: (k - 1) * v.R + v.S)
        idx = O("idx", k)
        oput("exists", f"INDEX({lt.rng('exists')},{idx})", lambda v, a=a: a.long[v.S].exists)
        oput("v", f"INDEX({lt.rng('value')},{idx})", lambda v, a=a: a.long[v.S].value)
        oput("prev_disp", f'IF({O("exists", k)}=0,"",IF(OR({Sx}=1,{Sx}={O("birth", k)}),{q(S.PREV_NA)},'
                          f'INDEX({lt.rng("value")},{idx}-1)))', lambda v, a=a: v.prev_disp[a.k])
        for col in ("change", "priority", "reason", "action", "cur_loc"):
            tail = '&""' if col in ("action", "cur_loc") else ""          # static text: blank must read as ""
            oput(col, f"INDEX({lt.rng(col)},{idx}){tail}", lambda v, a=a, col=col: getattr(a.long[v.S], col))
        oput("idxC", f"({O('k', k)}-1)*{Rx}+{Cx}", lambda v, k=k: (k - 1) * v.R + v.C)
        oput("vC", f"INDEX({lt.rng('value')},{O('idxC', k)})", lambda v, a=a: a.long[v.C].value)
        oput("codeS", f'IF({O("v", k)}="",0,IFERROR(MATCH({O("v", k)},{status4},0),9))',
             lambda v, a=a: v._code(a.long[v.S].value))
        oput("codeC", f'IF({O("vC", k)}="",0,IFERROR(MATCH({O("vC", k)},{status4},0),9))',
             lambda v, a=a: v._code(a.long[v.C].value))
        oput("qkey", f'IF({O("priority", k)}="","",{O("priority", k)}*10000+{O("k", k)})',
             lambda v, a=a: v.qkey[a.k])
        ex, vv = O("exists", k), O("v", k)
        inc = (f'IF({ex}=0,0,IF({scope_x}={q(S.SCOPE_ALL)},1,IF({scope_x}={q(S.SCOPE_ACTION)},'
               f'IF({vv}<>{q(S.CLOSED)},1,0),IF({scope_x}={q(S.SCOPE_STALLED)},INDEX({lt.rng("stalled")},{idx}),'
               f'IF({scope_x}={q(S.SCOPE_REOPENED)},INDEX({lt.rng("reopen_state")},{idx}),'
               f'IF({scope_x}={q(S.SCOPE_NEW)},IF({O("birth", k)}={Sx},1,0),IF({vv}={scope_x},1,0)))))))')
        oput("include", inc, lambda v, a=a: v.include[a.k])
        first = obs_t.cell("include", 1, abs_=True)
        oput("rank", f'IF({O("include", k)}=1,COUNTIF({first}:{O("include", k)},1),"")',
             lambda v, a=a: v.rank[a.k])
        oput("vR", f"INDEX({lt.rng('value')},({O('k', k)}-1)*{Rx}+{Rx})", a.long[R].value)
        link_expected = _link(data_sheet, a.row,
                              schema.col(R, S.ROLE_RESPONSE) if a.long[R].value not in ("", S.CLOSED) else 1)
        oput("link", f'"#\'"&{refs.meta("data_sheet", sheet=False)}&"\'!"&ADDRESS({O("row", k)},'
                     f'IF(AND({O("vR", k)}<>"",{O("vR", k)}<>{q(S.CLOSED)}),INDEX({rt.rng("c_response")},{Rx}),1),4)',
             link_expected)

    # ------------------------------------------------------------------ queue / table order
    qt = hb.table("q", ["i", "key", "k"], N)
    tt = hb.table("t", ["i", "k"], N)
    for i in range(1, N + 1):
        n = i + 1
        ws.cell(row=n, column=qt.cols["i"], value=i)
        ws.cell(row=n, column=tt.cols["i"], value=i)
        reg.put(ws, f"{qt.c('key')}{n}", f'IFERROR(SMALL({obs_t.rng("qkey")},{qt.cell("i", i)}),"")',
                lambda v, i=i: (v.queue[i - 1].long[v.S].priority * 10000 + v.queue[i - 1].k)
                if i <= len(v.queue) else "")
        reg.put(ws, f"{qt.c('k')}{n}", f'IF({qt.cell("key", i)}="","",MATCH({qt.cell("key", i)},{obs_t.rng("qkey")},0))',
                lambda v, i=i: v.queue[i - 1].k if i <= len(v.queue) else "")
        reg.put(ws, f"{tt.c('k')}{n}", f'IFERROR(MATCH({tt.cell("i", i)},{obs_t.rng("rank")},0),"")',
                lambda v, i=i: v.table[i - 1].k if i <= len(v.table) else "")

    # ------------------------------------------------------------------ trend / feeds
    tr = hb.table("trend", TREND_COLS, R)
    lr = lambda col: lt.rng(col)                          # noqa: E731
    for r in range(1, R + 1):
        n = r + 1
        t = an.trend[r]
        ws.cell(row=n, column=tr.cols["r"], value=r)
        ws.cell(row=n, column=tr.cols["label"], value=S.round_label(r))
        rc = tr.cell("r", r)
        def tput(col, formula, expected):
            reg.put(ws, f"{tr.c(col)}{n}", formula, expected)
        tput("issued", f"COUNTIFS({lr('round')},{rc},{lr('exists')},1)", t.issued)
        tput("new", f"COUNTIF({obs_t.rng('birth')},{rc})", t.new)
        for col, val_ in (("closed", S.CLOSED), ("cond", S.CLOSED_COND), ("req", S.REQUIRES_EDIT), ("open", S.OPEN)):
            tput(col, f"COUNTIFS({lr('round')},{rc},{lr('value')},{q(val_)})", getattr(t, col))
        tput("invalid", f"SUMIFS({lr('invalid')},{lr('round')},{rc})", t.invalid)
        tput("rate", f'IF({tr.cell("issued", r)}=0,"",{tr.cell("closed", r)}/{tr.cell("issued", r)})', t.rate)
        ws.cell(row=n, column=tr.cols["rate"]).number_format = "0%"
        for col in ("closed_event", "reopen_event", "stalled", "reopen_state", "ever"):
            tput(col, f"SUMIFS({lr(col)},{lr('round')},{rc})", getattr(t, col))
    bar = hb.table("bar", ["label", "closed", "cond", "req", "open"], 1)
    reg.put(ws, f"{bar.c('label')}2", f"{refs.meta('selected_label', sheet=False)}", lambda v: S.round_label(v.S))
    for col in ("closed", "cond", "req", "open"):
        reg.put(ws, f"{bar.c(col)}2", f"INDEX({tr.rng(col)},{Sx})", lambda v, col=col: getattr(v.t, col))

    # meta aggregates
    reg.put(ws, refs.meta("queue_count", sheet=False), f"COUNT({obs_t.rng('qkey')})", lambda v: len(v.queue))
    reg.put(ws, refs.meta("table_count", sheet=False), f"COUNTIF({obs_t.rng('include')},1)", lambda v: len(v.table))
    reg.put(ws, refs.meta("complete_needed", sheet=False),
            f'SUMPRODUCT(({lr("round")}={Rx})*({lr("exists")}=1)*({lr("value")}<>{q(S.CLOSED)}))',
            lambda v: v.completeness()[1])
    reg.put(ws, refs.meta("complete_filled", sheet=False),
            f'SUMPRODUCT(({lr("round")}={Rx})*({lr("exists")}=1)*({lr("value")}<>{q(S.CLOSED)})*{lr("resp_filled")})',
            lambda v: v.completeness()[0])

    # ------------------------------------------------------------------ names
    def name(nm, ref):
        wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    name("crms_data", data_rng)
    name("crms_R", refs.meta("current_round"))
    name("crms_S", refs.meta("selected_round"))
    name("crms_C", refs.meta("compare_round"))
    name("crms_k", refs.meta("selected_k"))
    name(W.NAME_ROUNDS, rt.rng("label", sheet=True))
    name(W.NAME_SCOPES, HP + scope_rng)
    name(W.NAME_READERS, obs_t.rng("reader_label", sheet=True))

    for col in range(1, hb.next_col):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.protection.sheet = True
    refs.tables = hb.tables
    return refs


def _link(sheet: str, row: int, col: int) -> str:
    return f"#'{sheet}'!{get_column_letter(col)}{row}"
