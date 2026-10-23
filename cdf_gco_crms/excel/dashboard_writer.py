"""Executive dashboard "لوحة المتابعة التنفيذية" (protocol 10.4A; UX directive zones A–G).

The sheet is a formula-driven view of the data sheet through the hidden
helper sheet. It is read-only except for four selector cells; no response is
entered here (10.4A.11). Every formula is registered with its independently
computed expected value (formula_registry), so the build can write cached
values and the verifier can compare a recalculation for any view.

Layout: a twelve-column grid B..M (D.10), right-to-left. Zones in order:
A header and selectors, B KPI cards, C closure level (narrative, secondary
indicators, status distribution), D action queue, F selected observation
(detail and timeline), E round analysis (history charts, transition matrix),
the adaptive observations table, G user guide.
"""
from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.chart.title import Title
from openpyxl.chart.text import Text
from openpyxl.comments import Comment
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Protection
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

from ..analytics.dashboard_analytics import Analytics, View, fixed1
from ..config import design_tokens as T
from ..config import round_schema as S
from ..config import workbook_settings as W
from ..guidance import guidance_dictionary as G
from ..observations import text_derivation as TD
from . import styles as st
from .formula_registry import Registry
from .helper_writer import HP, HelperRefs, q, sheet_ref

GRID0 = W.GRID_FIRST_COL           # column B
NGRID = W.GRID_COLS                # 12
QUEUE_MAX = 10


def gc(i: int) -> int:
    """Worksheet column of grid position i (0 = column B, the right edge in RTL)."""
    return GRID0 + i


def L(c: int) -> str:
    return get_column_letter(c)


# fixed rows of the first screen (selector cells must be known before the helper is written)
ROW = {
    "sp0": 1, "fund": 2, "title": 3, "doc": 4, "info_l": 5, "info_v": 6, "banner": 7, "nav": 8,
    "sel_l": 9, "sel_v": 10, "hist": 11, "sp1": 12, "kpi_l": 13, "kpi_v": 14, "kpi_s": 15, "sp2": 16,
    "band0": 17,                      # rows 17..24: closure level band
    "sp3": 25, "q_title": 26, "q_note": 27, "q_head": 28,
}
BAND_ROWS = 8
HEIGHT = {"sp0": 6, "fund": 18, "title": 34, "doc": 22, "info_l": 15, "info_v": 20, "banner": 18, "nav": 17,
          "sel_l": 15, "sel_v": 22, "hist": 16, "sp1": 8, "kpi_l": 20, "kpi_v": 44, "kpi_s": 16, "sp2": 8,
          "sp3": 8, "q_title": 22, "q_note": 14, "q_head": 20}
BAND_HEIGHT = 15
SECTION_HEIGHT = 24


@dataclass
class Selectors:
    sel_round: str = f"${L(gc(0))}${ROW['sel_v']}"
    sel_compare: str = f"${L(gc(3))}${ROW['sel_v']}"
    sel_scope: str = f"${L(gc(6))}${ROW['sel_v']}"
    sel_obs: str = f"${L(gc(9))}${ROW['sel_v']}"


@dataclass
class DashInfo:
    sections: dict[str, int] = field(default_factory=dict)       # section -> first row
    last_row: int = 0
    queue_rows: list[int] = field(default_factory=list)
    table_rows: list[int] = field(default_factory=list)
    link_cells: list[str] = field(default_factory=list)          # HYPERLINK formula cells
    nav_cells: list[str] = field(default_factory=list)           # static internal hyperlinks
    charts: list[dict] = field(default_factory=list)
    selectors: Selectors = field(default_factory=Selectors)
    detail_rows: dict[str, int] = field(default_factory=dict)
    timeline_rows: list[int] = field(default_factory=list)
    matrix_origin: tuple[int, int] = (0, 0)
    notes: list[str] = field(default_factory=list)


# ------------------------------------------------------------------ helpers
def _note(text: str, width: int = 300) -> Comment:
    lines = sum(max(1, len(seg) // 45 + 1) for seg in text.split("\n"))
    c = Comment(text, G.NOTE_AUTHOR)
    c.width = width
    c.height = min(380, 17 * lines + 28)
    return c


def _merge(ws, r1: int, g1: int, r2: int | None = None, g2: int | None = None):
    r2 = r2 or r1
    g2 = g1 if g2 is None else g2
    c1, c2 = gc(g1), gc(g2)
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    return ws.cell(row=r1, column=c1)


def _fill_range(ws, r1, g1, r2, g2, color):
    for r in range(r1, r2 + 1):
        for g in range(g1, g2 + 1):
            ws.cell(row=r, column=gc(g)).fill = st.fill(color)


def _box(ws, r1, g1, r2, g2, color=T.CARD_BORDER, style="thin"):
    for r in range(r1, r2 + 1):
        for g in range(g1, g2 + 1):
            cell = ws.cell(row=r, column=gc(g))
            cell.border = st.border(
                left=style if g == g2 else None,      # RTL: grid end is the visual left
                right=style if g == g1 else None,
                top=style if r == r1 else None,
                bottom=style if r == r2 else None,
                color=color)


def _rule_below(ws, r, g1=0, g2=NGRID - 1, color=T.SECTION_RULE):
    for g in range(g1, g2 + 1):
        ws.cell(row=r, column=gc(g)).border = st.border(None, None, None, "medium", color=color)


def _card(ws, r1, g1, r2, g2, fill=T.CARD_FILL, color=T.CARD_BORDER):
    """A white card on the page surface: solid fill and a thin warm-grey frame (D.10)."""
    _fill_range(ws, r1, g1, r2, g2, fill)
    _box(ws, r1, g1, r2, g2, color=color)


def _marker(ws, r, color=T.ACCENT):
    """Powder-blue marker in the reading-edge margin of a section row (D.10 section marker)."""
    ws.cell(row=r, column=GRID0 - 1).fill = st.fill(color)


def _page_surface(ws, last_row: int):
    """Paint the page surface behind everything that has no fill of its own (D.10)."""
    for r in range(1, last_row + 1):
        for c in range(1, gc(NGRID) + 1):
            cell = ws.cell(row=r, column=c)
            if cell.fill is None or cell.fill.fill_type is None:
                cell.fill = st.fill(T.PAGE_SURFACE)


def _text(ws, r, g1, g2, value, *, size=T.SZ_BODY, bold=False, color=T.BRAND_INK, h="right", v="center",
          wrap=True, fill=None, r2=None, italic=False):
    cell = _merge(ws, r, g1, r2, g2)
    cell.value = value
    cell.font = st.font(size, bold=bold, color=color, italic=italic)
    cell.alignment = st.align(h, v, wrap)
    if fill:
        _fill_range(ws, r, g1, r2 or r, g2, fill)
    return cell


def _formula(ws, reg: Registry, r, g1, g2, formula, expected, *, size=T.SZ_BODY, bold=False,
             color=T.BRAND_INK, h="right", v="center", wrap=True, fill=None, r2=None, fmt=None):
    cell = _merge(ws, r, g1, r2, g2)
    reg.put(ws, cell.coordinate, formula, expected)
    cell.font = st.font(size, bold=bold, color=color)
    cell.alignment = st.align(h, v, wrap)
    if fmt:
        cell.number_format = fmt
    if fill:
        _fill_range(ws, r, g1, r2 or r, g2, fill)
    return cell


def _section(ws, r, title, *, g1=0, g2=NGRID - 1, height=SECTION_HEIGHT):
    ws.row_dimensions[r].height = height
    _text(ws, r, g1, g2, title, size=T.SZ_SECTION, bold=True, v="bottom", wrap=False)
    _rule_below(ws, r, g1, g2)
    if g1 == 0:
        _marker(ws, r)


def _status_rules(ws, rng: str, top_left: str, *, bold=True):
    """D.2 fills for the four official values plus grey for 'not raised' (G-DASH-10)."""
    for value in S.CLOSURE_VALUES:
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'{top_left}="{value}"'], fill=st.fill(T.STATUS_FILL[value]),
            font=st.font(T.SZ_BODY, bold=bold, color=T.BRAND_INK), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'OR({top_left}="{S.NOT_RAISED}",{top_left}="{S.PREV_NA}")'],
        font=st.font(T.SZ_NOTE, color=T.MUTED_TEXT), stopIfTrue=True))


def _rich_title(text: str, size: int = 10) -> Title:
    cp = CharacterProperties(sz=size * 100, b=True, solidFill=T.BRAND_INK,
                             latin=DFont(typeface=T.FONT), cs=DFont(typeface=T.FONT))
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    t = Title(tx=Text(rich=RichText(p=[para])), overlay=False)
    return t


def _txpr(size: int = 9) -> RichText:
    cp = CharacterProperties(sz=size * 100, solidFill=T.CHART_TEXT, latin=DFont(typeface=T.FONT),
                             cs=DFont(typeface=T.FONT))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _style_chart(ch, *, rtl_categories: bool = True, legend_pos: str = "b"):
    ch.style = 2
    ch.roundedCorners = False
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
        ax.txPr = _txpr(9)
    ch.y_axis.majorGridlines = None
    if rtl_categories:
        ch.x_axis.scaling.orientation = "maxMin"
    if ch.legend is not None:
        ch.legend.position = legend_pos
        ch.legend.txPr = _txpr(9)
    ch.graphical_properties = GraphicalProperties(ln=None)


def _brand_mark(series, fill_hex: str):
    """A chart mark filled with a brand colour and outlined in ink: the outline keeps the 3:1 boundary that the
    pastel fill alone would not reach on white (WCAG 1.4.11; D.10)."""
    series.graphicalProperties = GraphicalProperties(solidFill=fill_hex)
    series.graphicalProperties.line.solidFill = T.CHART_OUTLINE
    series.graphicalProperties.line.width = 9525


def _col_px(width_chars: float) -> int:
    return int(width_chars * 7 + 5)


def _pt_to_cm(pt: float) -> float:
    return pt / 72 * 2.54


def _px_to_cm(px: float) -> float:
    return px / 96 * 2.54


# ------------------------------------------------------------------ main
def write_dashboard(wb, an: Analytics, view: View, schema, refs: HelperRefs, reg: Registry, state, *,
                    data_sheet: str, sel: Selectors, logo_path: str | None = None) -> DashInfo:
    ws = wb.create_sheet(W.DASHBOARD_SHEET, 0)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = "FF" + T.ACCENT
    info = DashInfo(selectors=sel)
    t = refs.tables
    obs_t, lt, tr, rt, qt, tt, bar, lists = (t["obs"], t["long"], t["trend"], t["round"], t["q"], t["t"],
                                             t["bar"], t["list"])
    R = an.R
    N = len(an.obs)
    DS = sheet_ref(data_sheet)
    doc = state.document
    cur = state.round(R)
    Sx, Rx, Cx, Kx = "crms_S", "crms_R", "crms_C", "crms_k"
    ROUNDS = W.NAME_ROUNDS

    def M(key):
        return refs.meta(key)

    def TR(col):
        return f"INDEX({tr.rng(col, sheet=True)},{Sx})"

    # column widths
    ws.column_dimensions["A"].width = W.MARGIN_WIDTH
    for g in range(NGRID):
        ws.column_dimensions[L(gc(g))].width = W.GRID_WIDTH
    ws.column_dimensions[L(gc(NGRID))].width = W.MARGIN_WIDTH
    for key, h in HEIGHT.items():
        ws.row_dimensions[ROW[key]].height = h
    if not doc.is_demo:
        ws.row_dimensions[ROW["banner"]].height = 4

    # ============================================================ ZONE A — header
    # top rule in ink across the full width, then the identity band in Warm Gray edge to edge (D.10)
    for c in range(1, gc(NGRID) + 1):
        ws.cell(row=ROW["sp0"], column=c).fill = st.fill(T.TOP_RULE)
        for rr in range(ROW["fund"], ROW["doc"] + 1):
            ws.cell(row=rr, column=c).fill = st.fill(T.TITLE_BAND)
    tg = 4 if logo_path else 0          # text starts beside the logo (logo on the reading side)
    if logo_path:
        _merge(ws, ROW["fund"], 0, ROW["doc"], tg - 1)
        _add_logo(ws, logo_path, ROW, HEIGHT)
        info.notes.append("logo: supplied official horizontal logo, background made transparent, ink #101820 on the "
                          "Warm Gray band, aspect ratio kept")
    _text(ws, ROW["fund"], tg, NGRID - 1, f"{W.FUND_NAME}   |   {W.GCO_NAME}", size=T.SZ_SUBTITLE,
          color=T.MUTED_TEXT, v="bottom")
    _text(ws, ROW["title"], tg, NGRID - 1, W.DASHBOARD_TITLE, size=T.SZ_TITLE, bold=True, color=T.TITLE_TEXT)
    _text(ws, ROW["doc"], tg, NGRID - 1, doc.name, size=T.SZ_DOC, bold=True, color=T.TITLE_TEXT, v="top")

    info_items = [
        ("الجولة الحالية", S.round_label(R) if R else S.NOT_RECORDED),
        ("إصدار الوثيقة", doc.cover_version or cur.document_version or S.NOT_RECORDED),
        ("تاريخ آخر مراجعة", cur.date or S.NOT_RECORDED),
        ("الإدارة صاحبة الوثيقة", doc.owning_department or S.NOT_RECORDED),
    ]
    for i, (lab, val) in enumerate(info_items):
        _text(ws, ROW["info_l"], 2 * i, 2 * i + 1, lab, size=T.SZ_NOTE, color=T.MUTED_TEXT, v="bottom", fill=T.CARD_FILL)
        _text(ws, ROW["info_v"], 2 * i, 2 * i + 1, val, size=T.SZ_BODY, bold=True, fill=T.CARD_FILL)
    _text(ws, ROW["info_l"], 8, 11, "المطلوب من الإدارة صاحبة الطلب في الجولة الحالية", size=T.SZ_NOTE,
          color=T.MUTED_TEXT, v="bottom", fill=T.CARD_FILL)
    need, filled = M("complete_needed"), M("complete_filled")
    next_f = (f'IF({need}=0,"لا توجد ملاحظات تتطلب إفادة في الجولة الحالية",'
              f'"الإفادة عن "&{need}&" ملاحظة غير مغلقة في ورقة الملاحظات (أُدخلت "&{filled}&")")')

    def next_exp(v: View):
        f_, n_ = v.completeness()
        return ("لا توجد ملاحظات تتطلب إفادة في الجولة الحالية" if n_ == 0
                else f"الإفادة عن {n_} ملاحظة غير مغلقة في ورقة الملاحظات (أُدخلت {f_})")
    _formula(ws, reg, ROW["info_v"], 8, 11, next_f, next_exp, bold=True, color=T.STATUS_TEXT[S.OPEN], fill=T.CARD_FILL)
    for g1, g2 in ((0, 1), (2, 3), (4, 5), (6, 7), (8, 11)):
        _box(ws, ROW["info_l"], g1, ROW["info_v"], g2, color=T.CARD_BORDER)

    if doc.is_demo:
        _text(ws, ROW["banner"], 0, NGRID - 1, S.DEMO_BANNER, size=T.SZ_BODY, bold=True, color=T.BRAND_INK,
              h="center", fill=T.DEMO_BAND)

    # navigation (filled in after the sections are placed)
    nav_targets: list[tuple[str, str]] = []

    # selectors
    sel_specs = [
        (0, "الجولة المعروضة", "dash_round", sel.sel_round, f"={ROUNDS}", S.round_label(view.S)),
        (3, "جولة المقارنة", "dash_compare", sel.sel_compare, f"={ROUNDS}", S.round_label(view.C)),
        (6, "نطاق جدول الملاحظات", "dash_scope", sel.sel_scope, f"={W.NAME_SCOPES}", view.scope),
        (9, "الملاحظة المختارة", "dash_obs", sel.sel_obs, f"={W.NAME_READERS}", view.reader_label or ""),
    ]
    for g, lab, key, ref, src, default in sel_specs:
        lc = _text(ws, ROW["sel_l"], g, g + 2, lab, size=T.SZ_NOTE, bold=True, color=T.MUTED_TEXT, v="bottom")
        lc.comment = _note(G.entry(key).note)
        cell = _merge(ws, ROW["sel_v"], g, None, g + 2)
        cell.value = default
        cell.number_format = "@"
        cell.font = st.font(11, bold=True, color=T.BRAND_INK)
        cell.alignment = st.align("center", "center", False)
        _fill_range(ws, ROW["sel_v"], g, ROW["sel_v"], g + 2, T.SELECTOR_FILL)
        _box(ws, ROW["sel_v"], g, ROW["sel_v"], g + 2, color=T.SELECTOR_FRAME)
        for gg in range(g, g + 3):
            ws.cell(row=ROW["sel_v"], column=gc(gg)).protection = Protection(locked=False)
        title, prompt = G.input_message(key)
        dv = DataValidation(type="list", formula1=src.lstrip("="), allow_blank=False, showDropDown=False,
                            showInputMessage=True, showErrorMessage=True, errorStyle="stop",
                            promptTitle=title, prompt=prompt, errorTitle=title,
                            error="يرجى الاختيار من القائمة.")
        dv.sqref = cell.coordinate
        ws.add_data_validation(dv)

    for g in range(NGRID):                       # powder-blue divider between the header zone and the cards
        ws.cell(row=ROW["sp1"], column=gc(g)).fill = st.fill(T.ACCENT)
    hist_f = (f'IF({Sx}={Rx},"","تعرض اللوحة "&{M("selected_label")}&" (جولة سابقة)، والجولة الحالية هي "'
              f'&{M("current_label")}&".")')
    hc = _formula(ws, reg, ROW["hist"], 0, NGRID - 1, hist_f, lambda v: v.historical_band(), bold=True,
                  h="center", size=T.SZ_BODY)
    ws.conditional_formatting.add(hc.coordinate, FormulaRule(
        formula=[f'LEN({hc.coordinate})>0'], fill=st.fill(T.HISTORIC_BAND), stopIfTrue=True))

    # ============================================================ ZONE B — KPI cards
    cards = [
        ("إجمالي الملاحظات", TR("issued"), lambda v: v.t.issued, T.BRAND_INK, "issued", "0"),
        ("الملاحظات المغلقة", TR("closed"), lambda v: v.t.closed, T.STATUS_TEXT[S.CLOSED], "closed", "0"),
        ("مغلقة بشرط", TR("cond"), lambda v: v.t.cond, T.STATUS_TEXT[S.CLOSED_COND], "cond", "0"),
        ("يتطلب تعديل", TR("req"), lambda v: v.t.req, T.STATUS_TEXT[S.REQUIRES_EDIT], "req", "0"),
        ("مفتوحة", TR("open"), lambda v: v.t.open, T.STATUS_TEXT[S.OPEN], "open", "0"),
        ("نسبة الإغلاق", f'IF({TR("issued")}=0,"",{TR("closed")}/{TR("issued")})',
         lambda v: v.t.rate, T.BRAND_INK, "rate", "0.0%"),
    ]
    for i, (lab, f, exp, color, attr, fmt) in enumerate(cards):
        g = 2 * i
        # the card's cap carries the brand fill of the status it counts; value and comparison sit on white
        lc = _text(ws, ROW["kpi_l"], g, g + 1, lab, size=T.SZ_CARD_LABEL, bold=True, color=T.BRAND_INK,
                   h="center", v="center", fill=T.CARD_CAP[attr])
        if attr == "rate":
            lc.comment = _note(G.entry("kpi_rate").note)
        _formula(ws, reg, ROW["kpi_v"], g, g + 1, f, exp, size=T.SZ_CARD_VALUE, bold=True, color=color,
                 h="center", fmt=fmt, wrap=False, fill=T.CARD_FILL)
        if attr == "rate":
            prev = f'INDEX({tr.rng("rate", sheet=True)},{Sx}-1)'
            sub_f = (f'IF({Sx}=1,"لا توجد جولة سابقة",IF({prev}="","",'
                     f'"في الجولة السابقة: "&FIXED({prev}*100,1)&"%"))')
            sub_e = (lambda v: "لا توجد جولة سابقة" if v.S == 1 else
                     (f"في الجولة السابقة: {fixed1(v.prev_count('rate') * 100)}%" if v.prev_count("rate") != "" else ""))
        else:
            prev = f'INDEX({tr.rng(attr, sheet=True)},{Sx}-1)'
            sub_f = f'IF({Sx}=1,"لا توجد جولة سابقة","في الجولة السابقة: "&{prev})'
            sub_e = (lambda v, attr=attr: "لا توجد جولة سابقة" if v.S == 1
                     else f"في الجولة السابقة: {v.prev_count(attr)}")
        _formula(ws, reg, ROW["kpi_s"], g, g + 1, sub_f, sub_e, size=T.SZ_NOTE, color=T.MUTED_TEXT,
                 h="center", v="top", fill=T.CARD_FILL)
        _box(ws, ROW["kpi_l"], g, ROW["kpi_s"], g + 1)

    # ============================================================ ZONE C — closure level band
    b0 = ROW["band0"]
    for r in range(b0, b0 + BAND_ROWS):
        ws.row_dimensions[r].height = BAND_HEIGHT
    ws.row_dimensions[b0].height = SECTION_HEIGHT
    _text(ws, b0, 0, 5, "مستوى إغلاق الملاحظات", size=T.SZ_SECTION, bold=True, v="bottom", wrap=False)
    _rule_below(ws, b0, 0, 5)
    _marker(ws, b0)
    _card(ws, b0 + 1, 0, b0 + 3, 5)                       # narrative card
    issued, closed, cond, req, open_ = TR("issued"), TR("closed"), TR("cond"), TR("req"), TR("open")
    rate = TR("rate")
    lab = M("selected_label")
    n1 = (f'IF({issued}=0,"لا توجد ملاحظات مسجلة حتى الآن.",IF({Sx}=1,"حتى "&{lab}&": عدد الملاحظات الصادرة "'
          f'&{issued}&"، وجميعها بانتظار الإفادة الأولى من الإدارة صاحبة الطلب.","حتى "&{lab}&": مغلقة "&{closed}'
          f'&" من أصل "&{issued}&" ("&FIXED({rate}*100,1)&"%)؛ يتطلب تعديل "&{req}&"؛ مغلقة بشرط "&{cond}'
          f'&"؛ مفتوحة "&{open_}&"."))')
    _formula(ws, reg, b0 + 1, 0, 5, n1, lambda v: v.narrative1(), r2=b0 + 2, v="top", fill=T.CARD_FILL)
    n2 = (f'IF(OR({Sx}=1,{issued}=0),"","خلال "&{lab}&": أُغلقت "&{TR("closed_event")}&"، وأُعيد فتح "'
          f'&{TR("reopen_event")}&"، وأُضيفت "&{TR("new")}&"، ومتعثرة "&{TR("stalled")}&".")')
    _formula(ws, reg, b0 + 3, 0, 5, n2, lambda v: v.narrative2(), color=T.MUTED_TEXT, v="top", fill=T.CARD_FILL)
    sec = [
        ("المتعثرة", TR("stalled"), lambda v: v.t.stalled, None),
        ("المعاد فتحها", TR("reopen_state"), lambda v: v.t.reopen_state, None),
        ("استدامة الإغلاق", f'{closed}&" من "&{TR("ever")}&" شهدت إغلاقاً"',
         lambda v: f"{v.t.closed} من {v.t.ever} شهدت إغلاقاً", None),
        ("أُغلقت خلال الجولة", TR("closed_event"), lambda v: v.t.closed_event, None),
        ("جديدة في الجولة", TR("new"), lambda v: v.t.new, None),
        ("جاهزية المراجعة", None, None, "readiness"),
    ]
    for i, (lab_, f, exp, special) in enumerate(sec):
        rr = b0 + 4 + (i // 3) * 2
        g = (i % 3) * 2
        _fill_range(ws, rr, g, rr + 1, g + 1, T.CARD_FILL)
        lc = _text(ws, rr, g, g + 1, lab_, size=T.SZ_NOTE, bold=True, color=T.MUTED_TEXT, h="center", v="bottom")
        if special == "readiness":
            lc.comment = _note(G.entry("readiness").note)
            rf = (f'IF({issued}=0,"{S.READY_EMPTY}",IF({TR("invalid")}>0,"{S.READY_NO_DATA}",'
                  f'IF({open_}+{req}>0,"{S.READY_ACTION}",IF({cond}>0,"{S.READY_COND}","{S.READY_ALL_CLOSED}"))))')
            _formula(ws, reg, rr + 1, g, g + 1, rf, lambda v: v.readiness(), size=T.SZ_NOTE, bold=True, h="center")
        else:
            if lab_ == "المعاد فتحها":
                lc.comment = _note("ملاحظات سبق إغلاقها ثم أصبحت غير مغلقة، وتبقى ضمن هذا العدد حتى يُعاد إغلاقها. "
                                   "يبقى الإغلاق السابق محفوظاً في سجل الجولات.")
            if lab_ == "المتعثرة":
                lc.comment = _note("ملاحظات بقيت \"يتطلب تعديل\" أو \"مفتوحة\" في جولتين متتاليتين بعد جولة إصدارها.")
            if lab_ == "استدامة الإغلاق":
                lc.comment = _note("عدد الملاحظات المغلقة في الجولة المعروضة من بين كل الملاحظات التي أُغلقت في أي "
                                   "جولة حتى الجولة المعروضة. لا يُعد دليلاً على فعالية الضوابط بعد التنفيذ.")
            _formula(ws, reg, rr + 1, g, g + 1, f, exp, size=T.SZ_SECONDARY_VALUE if special is None and lab_ != "استدامة الإغلاق" else T.SZ_BODY,
                     bold=True, h="center", wrap=True)
        _box(ws, rr, g, rr + 1, g + 1)
    # status distribution chart (left half of the band)
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "percentStacked"
    ch.overlap = 100
    ch.gapWidth = 40
    for col, value in (("closed", S.CLOSED), ("cond", S.CLOSED_COND), ("req", S.REQUIRES_EDIT), ("open", S.OPEN)):
        ref = Reference(wb[W.HELPER_SHEET], min_col=bar.cols[col], min_row=1, max_row=2)
        ch.add_data(ref, titles_from_data=True)
        s = ch.series[-1]
        s.tx.strRef = None
        from openpyxl.chart.series import SeriesLabel
        s.tx = SeriesLabel(v=value)
        _brand_mark(s, T.STATUS_FILL[value])
        s.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False)
        s.dLbls.txPr = _txpr(9)
    ch.set_categories(Reference(wb[W.HELPER_SHEET], min_col=bar.cols["label"], min_row=2, max_row=2))
    ch.title = _rich_title("توزيع حالات الإغلاق في الجولة المعروضة", 10)
    _style_chart(ch, rtl_categories=False)
    ch.y_axis.number_format = "0%"
    ch.y_axis.scaling.orientation = "maxMin"     # RTL: the bar starts at the right edge
    ch.y_axis.delete = True
    ch.x_axis.delete = True
    width_px = 6 * _col_px(W.GRID_WIDTH)
    ch.width = _px_to_cm(width_px) - 0.2
    ch.height = _pt_to_cm(BAND_ROWS * BAND_HEIGHT) - 0.1
    ws.add_chart(ch, f"{L(gc(6))}{b0}")
    info.charts.append({"name": "status_distribution", "anchor": f"{L(gc(6))}{b0}", "kind": "bar/percentStacked",
                        "series": [S.CLOSED, S.CLOSED_COND, S.REQUIRES_EDIT, S.OPEN],
                        "source": f"bar.closed..bar.open (row 2) of {W.HELPER_SHEET}"})

    # ============================================================ ZONE D — action queue
    r = ROW["q_title"]
    info.sections["queue"] = r
    qc = refs.meta("queue_count")
    _formula(ws, reg, r, 0, NGRID - 1, f'"الملاحظات التي تتطلب إجراء ("&{qc}&")"',
             lambda v: f"الملاحظات التي تتطلب إجراء ({len(v.queue)})", size=T.SZ_SECTION, bold=True,
             v="bottom", wrap=False)
    _rule_below(ws, r)
    _marker(ws, r)
    qn = (f'IF({qc}>{QUEUE_MAX},"تُعرض أول {QUEUE_MAX} ملاحظات من أصل "&{qc}&" بحسب ترتيب المتابعة؛ القائمة الكاملة في '
          f'جدول الملاحظات بنطاق \'الملاحظات التي تتطلب إجراء\'. الترتيب لأغراض المتابعة وليس تصنيفاً للأهمية.",'
          f'"الترتيب لأغراض المتابعة وليس تصنيفاً للأهمية: المعاد فتحها، ثم المتعثرة، ثم المغلقة بشرط، ثم ما يتطلب '
          f'تعديلاً، ثم المفتوحة.")')

    def qn_exp(v: View):
        n = len(v.queue)
        if n > QUEUE_MAX:
            return (f"تُعرض أول {QUEUE_MAX} ملاحظات من أصل {n} بحسب ترتيب المتابعة؛ القائمة الكاملة في جدول "
                    f"الملاحظات بنطاق 'الملاحظات التي تتطلب إجراء'. الترتيب لأغراض المتابعة وليس تصنيفاً للأهمية.")
        return ("الترتيب لأغراض المتابعة وليس تصنيفاً للأهمية: المعاد فتحها، ثم المتعثرة، ثم المغلقة بشرط، "
                "ثم ما يتطلب تعديلاً، ثم المفتوحة.")
    _formula(ws, reg, ROW["q_note"], 0, NGRID - 1, qn, qn_exp, size=T.SZ_NOTE, color=T.MUTED_TEXT, v="top")
    qh = ROW["q_head"]
    q_cols = [("الملاحظة #", 0, 0), ("عنوان الملاحظة", 1, 3), ("الحالة", 4, 4), ("سبب الظهور في القائمة", 5, 6),
              ("الإجراء المطلوب", 7, 10), ("الانتقال", 11, 11)]
    for lab_, g1, g2 in q_cols:
        _text(ws, qh, g1, g2, lab_, size=T.SZ_BODY, bold=True, h="center", fill=T.ACCENT)
    QN = min(N, QUEUE_MAX) if N else 1
    q_height = 15
    for o in an.obs:
        q_height = max(q_height, 13 * _lines(o.title_disp, 3 * W.GRID_WIDTH) + 5)
        for rr_ in range(1, R + 1):
            x = o.long[rr_]
            q_height = max(q_height, 13 * max(_lines(x.action, 4 * W.GRID_WIDTH),
                                              _lines(x.reason, 2 * W.GRID_WIDTH)) + 5)
    q_height = min(q_height, 60)
    _fill_range(ws, qh + 1, 0, qh + QN, NGRID - 1, T.CARD_FILL)
    for i in range(1, QN + 1):
        rr = qh + i
        info.queue_rows.append(rr)
        ws.row_dimensions[rr].height = q_height
        kq = qt.cell("k", i, sheet=True)

        def qv(col, i=i):
            return lambda v: (getattr(v.queue[i - 1].long[v.S], col) if i <= len(v.queue) else "")
        first_empty = (f'IF({qc}=0,"لا توجد ملاحظات تتطلب إجراء في الجولة المعروضة.","")' if i == 1 else '""')

        def lbl_exp(v, i=i):
            if i <= len(v.queue):
                return v.queue[i - 1].reader_label
            return "لا توجد ملاحظات تتطلب إجراء في الجولة المعروضة." if (i == 1 and not v.queue) else ""
        if i == 1:
            _formula(ws, reg, rr, 0, 0, f'IF({kq}="",{first_empty},INDEX({obs_t.rng("reader_label", sheet=True)},{kq}))',
                     lbl_exp, bold=True, h="center", wrap=True)
        else:
            _formula(ws, reg, rr, 0, 0, f'IF({kq}="","",INDEX({obs_t.rng("reader_label", sheet=True)},{kq}))',
                     lbl_exp, bold=True, h="center")
        _formula(ws, reg, rr, 1, 3, f'IF({kq}="","",INDEX({obs_t.rng("title_disp", sheet=True)},{kq})&"")',
                 lambda v, i=i: (v.queue[i - 1].title_disp if i <= len(v.queue) else ""))
        _formula(ws, reg, rr, 4, 4, f'IF({kq}="","",INDEX({obs_t.rng("v", sheet=True)},{kq}))', qv("value"),
                 h="center", bold=True)
        _formula(ws, reg, rr, 5, 6, f'IF({kq}="","",INDEX({obs_t.rng("reason", sheet=True)},{kq}))', qv("reason"))
        _formula(ws, reg, rr, 7, 10, f'IF({kq}="","",INDEX({obs_t.rng("action", sheet=True)},{kq})&"")', qv("action"))
        lc = _formula(ws, reg, rr, 11, 11,
                      f'IF({kq}="","",HYPERLINK(INDEX({obs_t.rng("link", sheet=True)},{kq}),"فتح الصف"))',
                      lambda v, i=i: "فتح الصف" if i <= len(v.queue) else "", h="center",
                      color=T.LINK)
        lc.font = st.font(T.SZ_BODY, color=T.LINK, underline="single")
        info.link_cells.append(lc.coordinate)
        _status_rules(ws, f"{L(gc(4))}{qh + 1}:{L(gc(4))}{qh + QN}", f"${L(gc(4))}{qh + 1}")
        ws.conditional_formatting.add(f"{L(gc(5))}{qh + 1}:{L(gc(5))}{qh + QN}", FormulaRule(
            formula=[f'${L(gc(5))}{qh + 1}="{S.REASON_REOPENED}"'], fill=st.fill(T.FLAG_REOPENED), stopIfTrue=True))
        ws.conditional_formatting.add(f"{L(gc(5))}{qh + 1}:{L(gc(5))}{qh + QN}", FormulaRule(
            formula=[f'${L(gc(5))}{qh + 1}="{S.REASON_STALLED}"'], fill=st.fill(T.FLAG_STALLED), stopIfTrue=True))
        # zebra and row rules last: lower priority than the status and flag fills (LibreOffice applies only the
        # first matching rule, Excel stacks them by priority)
        rule_border = st.border(None, None, None, "thin", color=T.CARD_BORDER)
        whole = f"{L(gc(0))}{qh + 1}:{L(gc(NGRID - 1))}{qh + QN}"
        ws.conditional_formatting.add(whole, FormulaRule(
            formula=[f'AND(${L(gc(0))}{qh + 1}<>"",MOD(ROW()-{qh + 1},2)=1)'], fill=st.fill(T.ZEBRA),
            border=rule_border))
        ws.conditional_formatting.add(whole, FormulaRule(formula=[f'${L(gc(0))}{qh + 1}<>""'], border=rule_border))
    r = qh + QN + 1

    # ============================================================ ZONE F — selected observation
    r += 1
    ws.row_dimensions[r - 1].height = 8
    info.sections["detail"] = r
    _section(ws, r, "تفاصيل الملاحظة المختارة ومسارها عبر الجولات")
    r += 1
    k = Kx
    idx = f"({k}-1)*{Rx}+{Sx}"
    ex = f"INDEX({lt.rng('exists', sheet=True)},{idx})"
    head_f = (f'INDEX({obs_t.rng("reader_label", sheet=True)},{k})&"{S.SEP}"&IF({ex}=0,"{S.NOT_RAISED}",'
              f'INDEX({lt.rng("value", sheet=True)},{idx}))')
    ws.row_dimensions[r].height = 26
    noobs = f'{M("obs_count")}=0'
    head_f = f'IF({noobs},"{S.READY_EMPTY_DETAIL}",{head_f})'
    hcell = _formula(ws, reg, r, 0, NGRID - 1, head_f,
                     lambda v: v.detail().get("head", S.READY_EMPTY_DETAIL), size=T.SZ_DETAIL_HEAD, bold=True,
                     fill=T.CARD_FILL)
    for value in (S.CLOSED_COND, S.CLOSED, S.REQUIRES_EDIT, S.OPEN):
        suffix = f"{S.SEP}{value}"
        ws.conditional_formatting.add(f"{L(gc(0))}{r}:{L(gc(NGRID - 1))}{r}", FormulaRule(
            formula=[f'RIGHT(${L(gc(0))}${r},LEN("{suffix}"))="{suffix}"'],
            fill=st.fill(T.STATUS_FILL[value]), stopIfTrue=True))
    info.detail_rows["head"] = r
    r += 1
    lfr = f"INDEX({lt.rng('lfr', sheet=True)},{idx})"
    lfi = f"INDEX({lt.rng('lfi', sheet=True)},{idx})"
    lcr = f"INDEX({lt.rng('lcr', sheet=True)},{idx})"
    fields = [
        ("loc", "الموضع عند الإصدار", f'IF({ex}=0,"",INDEX({obs_t.rng("location", sheet=True)},{k}))', 1),
        ("cur_loc", "الموضع الحالي", f'IF({ex}=0,"",INDEX({lt.rng("cur_loc", sheet=True)},{idx})&"")', 1),
        ("comment", "الملاحظة كما صدرت", f'IF({ex}=0,"",INDEX({obs_t.rng("comment_cut", sheet=True)},{k})&"")', 3),
        ("feedback", "آخر مرئيات الحوكمة",
         f'IF({ex}=0,"",IF({Sx}=1,"لا توجد مرئيات متابعة في الجولة الأولى.",INDEX({lt.rng("fb_cut", sheet=True)},{idx})&""))', 3),
        ("action", "الإجراء المطلوب", f'IF({ex}=0,"",INDEX({lt.rng("action", sheet=True)},{idx})&"")', 1),
        ("resp", "آخر إفادة للإدارة",
         f'IF({ex}=0,"",IF({lfr}=0,"{S.NO_RESPONSE}",INDEX({lt.rng("resp", sheet=True)},({k}-1)*{Rx}+{lfr})'
         f'&" ("&INDEX({rt.rng("label", sheet=True)},{lfr})&")"))', 2),
        ("impl", "آخر حالة تطبيق أفادت بها الإدارة",
         f'IF({ex}=0,"",IF({lfi}=0,"{S.NO_RESPONSE}",INDEX({lt.rng("impl", sheet=True)},({k}-1)*{Rx}+{lfi})'
         f'&" ("&INDEX({rt.rng("label", sheet=True)},{lfi})&")"))', 1),
        ("closure", "الإغلاق",
         f'IF({ex}=0,"",IF({lcr}="","غير مغلقة في الجولة المعروضة","أُغلقت في "&INDEX({rt.rng("label", sheet=True)},'
         f'{lcr})&"{S.SEP}أساس الإغلاق: "&IF(INDEX({lt.rng("basis_static", sheet=True)},({k}-1)*{Rx}+{lcr})="",'
         f'"{S.NOT_RECORDED}",INDEX({lt.rng("basis_static", sheet=True)},({k}-1)*{Rx}+{lcr}))))', 1),
    ]
    for key, lab_, f, lines in fields:
        ws.row_dimensions[r].height = 14 * lines + 4
        _text(ws, r, 0, 1, lab_, size=T.SZ_NOTE, bold=True, color=T.BRAND_INK, v="top", fill=T.LABEL_FILL)
        _formula(ws, reg, r, 2, NGRID - 1, f'IF({noobs},"",{f})', lambda v, key=key: v.detail().get(key, ""), v="top",
                 fill=T.CARD_FILL)
        for g in range(NGRID):
            ws.cell(row=r, column=gc(g)).border = st.border(None, None, None, "thin", color=T.CARD_BORDER)
        info.detail_rows[key] = r
        r += 1
    # link to the full row
    lnk = _formula(ws, reg, r, 2, NGRID - 1,
                   f'IF({noobs},"",HYPERLINK(INDEX({obs_t.rng("link", sheet=True)},{k}),"عرض النص الكامل وسجل الجولات في ورقة الملاحظات"))',
                   lambda v: "عرض النص الكامل وسجل الجولات في ورقة الملاحظات" if v.sel else "")
    lnk.font = st.font(T.SZ_BODY, color=T.LINK, underline="single")
    info.link_cells.append(lnk.coordinate)
    info.detail_rows["link"] = r
    r += 2
    # timeline
    tl_cols = [("الجولة", 0, 1), ("حالة الإغلاق", 2, 3), ("التغير", 4, 5), ("حالة التطبيق كما أفادت الإدارة", 6, 7),
               ("مرئيات الحوكمة (مختصرة)", 8, 11)]
    for lab_, g1, g2 in tl_cols:
        _text(ws, r, g1, g2, lab_, size=T.SZ_BODY, bold=True, h="center", fill=T.ACCENT)
    ws.row_dimensions[r].height = 20
    info.detail_rows["timeline_head"] = r
    r += 1
    birth = f"INDEX({obs_t.rng('birth', sheet=True)},{k})"
    if R:
        _fill_range(ws, r, 0, r + R - 1, NGRID - 1, T.CARD_FILL)
    for i in range(1, R + 1):
        info.timeline_rows.append(r)
        ws.row_dimensions[r].height = 30
        pos = f"({k}-1)*{Rx}+{i}"
        _formula(ws, reg, r, 0, 1, f'IF(OR({noobs},{i}>{Sx}),"",INDEX({rt.rng("label", sheet=True)},{i}))',
                 lambda v, i=i: S.round_label(i) if (i <= v.S and v.sel) else "", bold=True, h="center")
        _formula(ws, reg, r, 2, 3, f'IF(OR({noobs},{i}>{Sx}),"",IF({i}<{birth},"{S.NOT_RAISED}",INDEX({lt.rng("value", sheet=True)},{pos})))',
                 lambda v, i=i: v.timeline(i)[0], h="center", bold=True)
        _formula(ws, reg, r, 4, 5, f'IF(OR({noobs},{i}>{Sx},{i}<{birth}),"",INDEX({lt.rng("change", sheet=True)},{pos}))',
                 lambda v, i=i: v.timeline(i)[1], h="center")
        _formula(ws, reg, r, 6, 7, f'IF(OR({noobs},{i}>{Sx},{i}<{birth}),"",IF(INDEX({lt.rng("impl_filled", sheet=True)},{pos})=1,'
                                   f'INDEX({lt.rng("impl", sheet=True)},{pos}),""))',
                 lambda v, i=i: v.timeline(i)[2], h="center")
        _formula(ws, reg, r, 8, 11, f'IF(OR({noobs},{i}>{Sx},{i}<{birth}),"",INDEX({lt.rng("fb_short", sheet=True)},{pos})&"")',
                 lambda v, i=i: v.timeline(i)[3], size=T.SZ_NOTE)
        for g in range(NGRID):
            ws.cell(row=r, column=gc(g)).border = st.border(None, None, None, "thin", color=T.CARD_BORDER)
        r += 1
    if R:
        a, b = info.timeline_rows[0], info.timeline_rows[-1]
        first = f"{L(gc(2))}{a}"
        ws.conditional_formatting.add(first, FormulaRule(       # 10.4: round-1 مفتوحة = awaiting the first response
            formula=[f'${first}="{S.OPEN}"'], fill=st.fill(T.ROUND1_OPEN_FILL),
            font=st.font(T.SZ_BODY, bold=True, color=T.BRAND_INK), stopIfTrue=True))
        _status_rules(ws, f"{L(gc(2))}{a}:{L(gc(2))}{b}", f"${L(gc(2))}{a}")
    r += 1

    # ============================================================ ZONE E — round analysis
    ws.row_dimensions[r - 1].height = 8
    info.sections["rounds"] = r
    _section(ws, r, "تطور الملاحظات عبر جولات المراجعة (جميع الجولات المسجلة)")
    r += 1
    chart_rows = 13
    for rr in range(r, r + chart_rows):
        ws.row_dimensions[rr].height = 15
    hs = wb[W.HELPER_SHEET]
    cats = Reference(hs, min_col=tr.cols["label"], min_row=2, max_row=1 + max(R, 1))
    # stacked composition by round
    c1 = BarChart()
    c1.type = "col"
    c1.grouping = "stacked"
    c1.overlap = 100
    c1.gapWidth = 60
    for col, value in (("closed", S.CLOSED), ("cond", S.CLOSED_COND), ("req", S.REQUIRES_EDIT), ("open", S.OPEN)):
        c1.add_data(Reference(hs, min_col=tr.cols[col], min_row=2, max_row=1 + max(R, 1)), titles_from_data=False)
        from openpyxl.chart.series import SeriesLabel
        c1.series[-1].tx = SeriesLabel(v=value)
        _brand_mark(c1.series[-1], T.STATUS_FILL[value])
    c1.set_categories(cats)
    c1.title = _rich_title("حالات الإغلاق في كل جولة (عدد الملاحظات الصادرة حتى الجولة)", 10)
    _style_chart(c1)
    c1.width = _px_to_cm(6 * _col_px(W.GRID_WIDTH)) - 0.2
    c1.height = _pt_to_cm(chart_rows * 15) - 0.1
    ws.add_chart(c1, f"{L(gc(0))}{r}")
    # closure rate line
    c2 = LineChart()
    c2.add_data(Reference(hs, min_col=tr.cols["rate"], min_row=2, max_row=1 + max(R, 1)), titles_from_data=False)
    from openpyxl.chart.series import SeriesLabel
    c2.series[0].tx = SeriesLabel(v="نسبة الإغلاق")
    c2.series[0].graphicalProperties.line.solidFill = T.RATE_LINE
    c2.series[0].graphicalProperties.line.width = 28575
    c2.series[0].marker.symbol = "circle"
    c2.series[0].marker.size = 7
    c2.series[0].marker.graphicalProperties = GraphicalProperties(solidFill=T.RATE_MARKER)
    c2.series[0].marker.graphicalProperties.line.solidFill = T.CHART_OUTLINE
    c2.series[0].dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False)
    c2.series[0].dLbls.numFmt = "0%"
    c2.series[0].dLbls.position = "t"
    c2.series[0].dLbls.txPr = _txpr(9)
    c2.set_categories(cats)
    c2.title = _rich_title("اتجاه نسبة الإغلاق", 10)
    _style_chart(c2)
    c2.legend = None
    c2.y_axis.number_format = "0%"
    c2.y_axis.scaling.min = 0
    c2.y_axis.scaling.max = 1
    c2.width = c1.width
    c2.height = c1.height
    ws.add_chart(c2, f"{L(gc(6))}{r}")
    info.charts.append({"name": "round_composition", "anchor": f"{L(gc(0))}{r}", "kind": "col/stacked",
                        "series": [S.CLOSED, S.CLOSED_COND, S.REQUIRES_EDIT, S.OPEN], "source": "trend.closed..open"})
    info.charts.append({"name": "closure_rate", "anchor": f"{L(gc(6))}{r}", "kind": "line",
                        "series": ["نسبة الإغلاق"], "source": "trend.rate"})
    r += chart_rows + 1
    for rr in range(r, r + chart_rows):
        ws.row_dimensions[rr].height = 15
    c3 = BarChart()
    c3.type = "col"
    c3.grouping = "clustered"
    c3.gapWidth = 80
    for col, name, color in (("new", "ملاحظات جديدة", T.MOVEMENT_CHART["new"]),
                             ("closed_event", "أُغلقت خلال الجولة", T.MOVEMENT_CHART["closed"]),
                             ("reopen_event", "أُعيد فتحها", T.MOVEMENT_CHART["reopened"])):
        c3.add_data(Reference(hs, min_col=tr.cols[col], min_row=2, max_row=1 + max(R, 1)), titles_from_data=False)
        c3.series[-1].tx = SeriesLabel(v=name)
        _brand_mark(c3.series[-1], color)
    c3.set_categories(cats)
    c3.title = _rich_title("حركة الملاحظات في كل جولة", 10)
    _style_chart(c3)
    c3.width = c1.width
    c3.height = c1.height
    ws.add_chart(c3, f"{L(gc(0))}{r}")
    info.charts.append({"name": "round_movement", "anchor": f"{L(gc(0))}{r}", "kind": "col/clustered",
                        "series": ["ملاحظات جديدة", "أُغلقت خلال الجولة", "أُعيد فتحها"],
                        "source": "trend.new, trend.closed_event, trend.reopen_event"})
    # transition matrix (right side of the second band is the visual left: grid 6..11)
    mr = r
    info.matrix_origin = (mr, 6)
    _formula(ws, reg, mr, 6, 11, f'"مقارنة جولتين: من "&{M("compare_label")}&" إلى "&{M("selected_label")}',
             lambda v: f"مقارنة جولتين: من {S.round_label(v.C)} إلى {S.round_label(v.S)}", bold=True, size=11)
    _rule_below(ws, mr, 6, 11)
    mr += 1
    ws.row_dimensions[mr].height = 30
    _formula(ws, reg, mr, 6, 7, f'"في "&{M("compare_label")}&" / في "&{M("selected_label")}',
             lambda v: f"في {S.round_label(v.C)} / في {S.round_label(v.S)}", size=T.SZ_NOTE, bold=True,
             h="center", fill=T.ACCENT)
    col_codes = [(1, S.CLOSED), (2, S.CLOSED_COND), (3, S.REQUIRES_EDIT), (4, S.OPEN)]
    for j, (code, value) in enumerate(col_codes):
        _text(ws, mr, 8 + j, 8 + j, value, size=T.SZ_NOTE, bold=True, h="center", fill=T.STATUS_FILL[value])
    row_codes = [(0, "لم تكن مطروحة"), (1, S.CLOSED), (2, S.CLOSED_COND), (3, S.REQUIRES_EDIT), (4, S.OPEN)]
    for i, (rc, rlabel) in enumerate(row_codes):
        rr = mr + 1 + i
        ws.row_dimensions[rr].height = 18
        _text(ws, rr, 6, 7, rlabel, size=T.SZ_NOTE, bold=True, h="center",
              fill=T.STATUS_FILL.get(rlabel, T.CLOSED_ROW_GRAY))
        for j, (cc, _) in enumerate(col_codes):
            _formula(ws, reg, rr, 8 + j, 8 + j,
                     f'COUNTIFS({obs_t.rng("codeC", sheet=True)},{rc},{obs_t.rng("codeS", sheet=True)},{cc})',
                     lambda v, rc=rc, cc=cc: v.matrix().get((rc, cc), 0), h="center", bold=(rc == cc),
                     fill=T.CARD_FILL)
            ws.cell(row=rr, column=gc(8 + j)).border = st.THIN
    nr = mr + 1 + len(row_codes)
    _text(ws, nr, 6, 11, "تبيّن المصفوفة عدد الملاحظات بحسب حالتها في الجولتين. القطر يمثل الملاحظات التي لم تتغير حالتها، "
          "وصف \"مغلقة\" في جولة المقارنة مع عمود غير \"مغلقة\" يمثل الملاحظات المعاد فتحها.",
          size=T.SZ_NOTE, color=T.MUTED_TEXT, v="top", r2=nr + 2)
    r += chart_rows + 1

    # ============================================================ adaptive observations table
    ws.row_dimensions[r].height = 8
    r += 1
    info.sections["table"] = r
    tc = refs.meta("table_count")
    _formula(ws, reg, r, 0, NGRID - 1,
             f'"جدول الملاحظات في "&{M("selected_label")}&"{S.SEP}"&{refs.meta("scope")}&" ("&{tc}&")"',
             lambda v: f"جدول الملاحظات في {S.round_label(v.S)}{S.SEP}{v.scope} ({len(v.table)})",
             size=T.SZ_SECTION, bold=True, v="bottom", wrap=False)
    _rule_below(ws, r)
    _marker(ws, r)
    ws.row_dimensions[r].height = SECTION_HEIGHT
    r += 1
    th = r
    ws.row_dimensions[th].height = 30
    prev_h = f'IF({Sx}=1,"الجولة السابقة","الحالة في "&INDEX({rt.rng("label", sheet=True)},{Sx}-1))'
    sel_h = f'"الحالة في "&{M("selected_label")}'
    t_cols = [("الملاحظة #", 0, 0, None), ("عنوان الملاحظة", 1, 4, None),
              (None, 5, 5, (prev_h, lambda v: "الجولة السابقة" if v.S == 1 else f"الحالة في {S.round_label(v.S - 1)}")),
              (None, 6, 6, (sel_h, lambda v: f"الحالة في {S.round_label(v.S)}")),
              ("التغير", 7, 7, None), ("الإجراء المطلوب", 8, 10, None), ("الانتقال", 11, 11, None)]
    for lab_, g1, g2, dyn in t_cols:
        if dyn:
            _formula(ws, reg, th, g1, g2, dyn[0], dyn[1], bold=True, h="center", fill=T.ACCENT)
        else:
            c = _text(ws, th, g1, g2, lab_, size=T.SZ_BODY, bold=True, h="center", fill=T.ACCENT)
            if lab_ == "التغير":
                c.comment = _note(G.entry("kpi_change").note)
    t_height = 15
    for o in an.obs:
        t_height = max(t_height, 13 * _lines(o.title_disp, 4 * W.GRID_WIDTH) + 5)
        for rr_ in range(1, R + 1):
            x = o.long[rr_]
            t_height = max(t_height, 13 * max(_lines(x.action, 3 * W.GRID_WIDTH), _lines(x.change, W.GRID_WIDTH)) + 5)
    t_height = min(t_height, 72)
    TN = max(N, 1)
    _fill_range(ws, th + 1, 0, th + TN, NGRID - 1, T.CARD_FILL)
    for i in range(1, TN + 1):
        rr = th + i
        info.table_rows.append(rr)
        ws.row_dimensions[rr].height = t_height
        kt = tt.cell("k", i, sheet=True)

        def tv(fn, i=i):
            return lambda v: (fn(v, v.table[i - 1]) if i <= len(v.table) else "")
        if i == 1:
            empty = (f'IF({tc}=0,"لا توجد ملاحظات ضمن النطاق المختار في الجولة المعروضة.","")')
            lab_f = f'IF({kt}="",{empty},INDEX({obs_t.rng("reader_label", sheet=True)},{kt}))'

            def lab_e(v):
                if v.table:
                    return v.table[0].reader_label
                return "لا توجد ملاحظات ضمن النطاق المختار في الجولة المعروضة."
        else:
            lab_f = f'IF({kt}="","",INDEX({obs_t.rng("reader_label", sheet=True)},{kt}))'
            lab_e = tv(lambda v, a: a.reader_label)
        _formula(ws, reg, rr, 0, 0, lab_f, lab_e, bold=True, h="center")
        _formula(ws, reg, rr, 1, 4, f'IF({kt}="","",INDEX({obs_t.rng("title_disp", sheet=True)},{kt})&"")',
                 tv(lambda v, a: a.title_disp))
        _formula(ws, reg, rr, 5, 5, f'IF({kt}="","",INDEX({obs_t.rng("prev_disp", sheet=True)},{kt}))',
                 tv(lambda v, a: v.prev_disp[a.k]), h="center")
        _formula(ws, reg, rr, 6, 6, f'IF({kt}="","",INDEX({obs_t.rng("v", sheet=True)},{kt}))',
                 tv(lambda v, a: a.long[v.S].value), h="center", bold=True)
        _formula(ws, reg, rr, 7, 7, f'IF({kt}="","",INDEX({obs_t.rng("change", sheet=True)},{kt}))',
                 tv(lambda v, a: a.long[v.S].change), h="center", size=T.SZ_NOTE)
        _formula(ws, reg, rr, 8, 10, f'IF({kt}="","",INDEX({obs_t.rng("action", sheet=True)},{kt})&"")',
                 tv(lambda v, a: a.long[v.S].action), size=T.SZ_NOTE)
        lc = _formula(ws, reg, rr, 11, 11,
                      f'IF({kt}="","",HYPERLINK(INDEX({obs_t.rng("link", sheet=True)},{kt}),"فتح الصف"))',
                      tv(lambda v, a: "فتح الصف"), h="center")
        lc.font = st.font(T.SZ_BODY, color=T.LINK, underline="single")
        info.link_cells.append(lc.coordinate)
    if N:
        a, b = th + 1, th + TN
        _status_rules(ws, f"{L(gc(5))}{a}:{L(gc(6))}{b}", f"{L(gc(5))}{a}")
        chg = f"{L(gc(7))}{a}:{L(gc(7))}{b}"
        for value, color in ((S.CHANGE_REOPENED, T.FLAG_REOPENED), (S.CHANGE_NO_PROGRESS, T.FLAG_STALLED),
                             (S.CHANGE_CLOSED, T.STATUS_FILL[S.CLOSED]), (S.CHANGE_COND_REGRESS, T.FLAG_STALLED)):
            ws.conditional_formatting.add(chg, FormulaRule(formula=[f'${L(gc(7))}{a}="{value}"'],
                                                           fill=st.fill(color), stopIfTrue=True))
        ws.conditional_formatting.add(f"{L(gc(0))}{a}:{L(gc(0))}{b}", FormulaRule(
            formula=[f'AND(${L(gc(0))}{a}<>"",${L(gc(0))}{a}={sel.sel_obs})'],
            fill=st.fill(T.SELECTED_ROW), stopIfTrue=True))
        rule_border = st.border(None, None, None, "thin", color=T.CARD_BORDER)
        whole = f"{L(gc(0))}{a}:{L(gc(NGRID - 1))}{b}"
        ws.conditional_formatting.add(whole, FormulaRule(
            formula=[f'AND(${L(gc(0))}{a}<>"",MOD(ROW()-{a},2)=1)'], fill=st.fill(T.ZEBRA), border=rule_border))
        ws.conditional_formatting.add(whole, FormulaRule(formula=[f'${L(gc(0))}{a}<>""'], border=rule_border))
    r = th + TN + 2

    # ============================================================ ZONE G — user guide
    info.sections["help"] = r
    _section(ws, r, "دليل الاستخدام")
    r += 1
    help_lines = [
        ("أين تُكتب الإفادة؟", "تُعبأ الإفادات وحالات التطبيق في ورقة الملاحظات فقط، في عمودي الجولة الحالية للملاحظات غير "
                               "المغلقة. لا تُدخل أي بيانات في هذه اللوحة، فهي للعرض والانتقال."),
        ("معاني حالات الإغلاق", f"{S.CLOSED}: لا يلزم إجراء. {S.CLOSED_COND}: يلزم استيفاء الشرط المذكور في مرئيات الحوكمة. "
                                f"{S.REQUIRES_EDIT}: يلزم استكمال التعديل المذكور في سطر \"المطلوب:\". {S.OPEN}: لم تُقدَّم "
                                "إفادة أو تعديل يمكن التحقق منه. في الجولة الأولى تعني \"مفتوحة\" أن الملاحظة بانتظار الإفادة الأولى."),
        ("حالة التطبيق", "القيمة التي تختارها الإدارة صاحبة الطلب إفادة منها، ولا تُغلق الملاحظة بذاتها؛ تحدد إدارة الحوكمة "
                         "والامتثال المؤسسي حالة الإغلاق بعد التحقق من النص."),
        ("اختيار الجولة والمقارنة", "تغيّر القوائم المنسدلة في أعلى اللوحة الجولة المعروضة وجولة المقارنة ونطاق الجدول والملاحظة "
                                    "المختارة. الاختيار للعرض فقط ولا يغيّر أي بيانات أو تقييم."),
        ("السجل التاريخي", "تُحفظ تقييمات الحوكمة وإفادات الإدارة في كل جولة في أعمدتها بورقة الملاحظات دون تعديل. يمكن طي "
                           "أعمدة الجولات الأقدم أو إظهارها من علامة (+) أعلى الأعمدة."),
        ("المؤشرات المشتقة", "التغير والمتعثرة والمعاد فتحها وجاهزية المراجعة أوصاف للعرض مستمدة من حالات الإغلاق، وليست "
                             "حالات إغلاق. ولا تعني نسبة الإغلاق اعتماد الوثيقة أو مستوى امتثالها."),
        ("الاعتراض على الإغلاق", "يُكتب الاعتراض على إغلاق ملاحظة في خلية الإفادة الخاصة بها في الجولة الحالية، وسيُعاد تقييمه."),
        ("مثال توضيحي", f"{G.EXAMPLE_TAG} \"عُدّل البند (5-2) بإضافة اشتراط موافقة صاحب الصلاحية قبل إطلاق أي مبادرة، "
                        "انظر الصفحة 4 من النسخة المحدثة.\""),
    ]
    for lab_, txt in help_lines:
        ws.row_dimensions[r].height = 14 * _lines(txt, 10 * W.GRID_WIDTH) + 6
        _text(ws, r, 0, 1, lab_, size=T.SZ_NOTE, bold=True, color=T.BRAND_INK, v="top", fill=T.HELP_FILL)
        _text(ws, r, 2, NGRID - 1, txt, size=T.SZ_NOTE, color=T.BRAND_INK, v="top", fill=T.HELP_FILL)
        r += 1
    info.last_row = r

    # ============================================================ navigation row
    data_ws = wb[data_sheet]
    resp_col = schema.col(R, S.ROLE_RESPONSE) if R else 1
    nav = [("الملاحظات التي تتطلب إجراء", f"'{W.DASHBOARD_SHEET}'!{L(gc(0))}{info.sections['queue']}"),
           ("تفاصيل الملاحظة", f"'{W.DASHBOARD_SHEET}'!{L(gc(0))}{info.sections['detail']}"),
           ("تحليل الجولات", f"'{W.DASHBOARD_SHEET}'!{L(gc(0))}{info.sections['rounds']}"),
           ("جدول الملاحظات", f"'{W.DASHBOARD_SHEET}'!{L(gc(0))}{info.sections['table']}"),
           ("دليل الاستخدام", f"'{W.DASHBOARD_SHEET}'!{L(gc(0))}{info.sections['help']}"),
           ("ورقة الملاحظات (الإدخال)", f"'{data_sheet}'!{L(resp_col)}2")]
    for i, (lab_, target) in enumerate(nav):
        c = _text(ws, ROW["nav"], 2 * i, 2 * i + 1, lab_, size=T.SZ_NOTE, h="center", color=T.LINK)
        c.font = st.font(T.SZ_NOTE, bold=True, color=T.LINK, underline="single")
        c.hyperlink = Hyperlink(ref=c.coordinate, location=target, display=lab_)
        info.nav_cells.append(c.coordinate)

    # page surface behind every cell that no zone painted (cards stay white)
    _page_surface(ws, info.last_row + 1)

    # ============================================================ sheet settings
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{L(gc(NGRID))}{info.last_row}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.sheet_view.selection[0].activeCell = sel.sel_round.replace("$", "")
    ws.sheet_view.selection[0].sqref = sel.sel_round.replace("$", "")
    return info


LOGO_WIDTH_PX = 300


def _add_logo(ws, path: str, ROW, HEIGHT):
    """Place the supplied logo, proportions kept, vertically centred in the title band (G-UX-02)."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import pixels_to_EMU
    img = XLImage(path)
    ratio = img.height / img.width
    w = LOGO_WIDTH_PX
    h = int(round(w * ratio))
    band_px = (HEIGHT["fund"] + HEIGHT["title"] + HEIGHT["doc"]) * 96 / 72
    off_y = max(0, int((band_px - h) / 2))
    img.width, img.height = w, h
    marker = AnchorMarker(col=GRID0 - 1, colOff=pixels_to_EMU(24), row=ROW["fund"] - 1, rowOff=pixels_to_EMU(off_y))
    img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
    ws.add_image(img)


def _lines(text: str, width_chars: float) -> int:
    if not text:
        return 1
    per = max(6, int(width_chars * 1.05))
    import math
    return sum(max(1, math.ceil(len(seg) / per)) for seg in str(text).split("\n"))
