"""Assembles the requester workbook: data sheet, hidden helper sheet, executive dashboard.

Two entry points:
* ``build_workbook`` renders a fresh workbook from the review state (rounds 1..n).
* ``assemble`` adds the helper and dashboard to a workbook whose data sheet
  already exists (used by the v7.10 migration, which keeps the data sheet as
  delivered).

Output is written atomically: a temporary file is produced, post-processed
(cached values of the selected view, app.xml without library names), reopened
and checked, and only then moved over the target (directive §10).
"""
from __future__ import annotations

import datetime as dt
import os
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from zoneinfo import ZoneInfo

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.properties import CalcProperties

from .. import ENGINE_VERSION, WORKBOOK_SCHEMA
from ..analytics.dashboard_analytics import Analytics, View
from ..config import round_schema as S
from ..config import design_tokens as T
from ..config import workbook_settings as W
from ..model.entities import DashboardViewState, ReviewState
from ..rounds.round_parser import RoundSchema
from .dashboard_writer import DashInfo, Selectors, write_dashboard
from .formula_registry import Registry
from .helper_writer import HelperRefs, write_helper
from .register_writer import write_data_sheet

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
RIYADH = ZoneInfo("Asia/Riyadh")
DEFAULT_LOGO = Path(__file__).resolve().parents[2] / "assets" / "cdf_logo_horizontal_transparent.png"


@dataclass
class BuildResult:
    path: Path
    state: ReviewState
    schema: RoundSchema
    analytics: Analytics
    view: View
    registry: Registry
    refs: HelperRefs
    dash: DashInfo
    data_sheet: str
    rows: dict[str, int]
    cached: bool
    build_time: dt.datetime
    warnings: list[str] = field(default_factory=list)


def default_view(an: Analytics, vs: DashboardViewState | None) -> View:
    if vs is None:
        return an.view(an.R, S.SCOPE_ALL, an.default_reader_label(), max(1, an.R - 1))
    return an.view(vs.selected_round, vs.scope, vs.selected_reader or an.default_reader_label(),
                   getattr(vs, "compare_round", None) or max(1, vs.selected_round - 1))


def assemble(wb, state: ReviewState, data_ws, schema: RoundSchema, rows: dict[str, int], *,
             view_state: DashboardViewState | None = None, logo_path=None):
    data_sheet = data_ws.title
    an = Analytics(state, rows, round_cols=schema.blocks, data_sheet=data_sheet)
    view = default_view(an, view_state)
    reg = Registry()
    sel = Selectors()
    last_row = max(rows.values()) if rows else 1
    refs = write_helper(wb, an, schema, sel, reg, data_sheet=data_sheet, data_last_row=max(last_row, data_ws.max_row),
                        data_last_col=schema.last_col, state=state)
    logo = None
    if logo_path is not False:
        p = Path(logo_path) if logo_path else DEFAULT_LOGO
        logo = str(p) if p.exists() else None
    dash = write_dashboard(wb, an, view, schema, refs, reg, state, data_sheet=data_sheet, sel=sel, logo_path=logo)
    # order: dashboard, data sheet, helper (hidden)
    order = [wb[W.DASHBOARD_SHEET], data_ws, wb[W.HELPER_SHEET]]
    others = [s for s in wb._sheets if s not in order]
    wb._sheets = order + others
    for s in wb.worksheets:
        s.sheet_view.tabSelected = False
    wb.active = 0
    wb[W.DASHBOARD_SHEET].sheet_view.tabSelected = True
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    return an, view, reg, refs, dash


def set_properties(wb, state: ReviewState, build_time: dt.datetime, *, working_copy: bool = False):
    p = wb.properties
    p.creator = W.AUTHOR
    p.lastModifiedBy = W.AUTHOR
    p.title = W.workbook_title(state.document.name, state.current_round)
    p.created = build_time.astimezone(dt.timezone.utc).replace(tzinfo=None)
    p.modified = p.created
    p.subject = None
    keywords = ["بيانات تجريبية"] if state.document.is_demo else []
    if working_copy:                          # release gate: an unauthorised build is marked in the package itself
        from ..release.authorization_gate import WORKING_COPY_KEYWORD, WORKING_COPY_MARK
        keywords.append(WORKING_COPY_KEYWORD)
    p.keywords = "؛ ".join(keywords) if keywords else None
    p.description = None
    p.category = WORKING_COPY_MARK if working_copy else None


def build_workbook(state: ReviewState, out_path, *, view_state: DashboardViewState | None = None,
                   write_cache: bool = True, build_time: dt.datetime | None = None, logo_path=None,
                   guidance: bool = True, working_copy: bool = False) -> BuildResult:
    build_time = build_time or dt.datetime.now(RIYADH)
    wb = Workbook()
    wb.remove(wb.active)
    data_ws, schema = write_data_sheet(wb, state, guidance=guidance)
    rows = {o.internal_id: 2 + i for i, o in enumerate(state.ordered())}
    an, view, reg, refs, dash = assemble(wb, state, data_ws, schema, rows, view_state=view_state,
                                         logo_path=logo_path)
    set_properties(wb, state, build_time, working_copy=working_copy)
    out_path = Path(out_path)
    save_atomic(wb, out_path, reg.expected(view) if write_cache else None, build_time)
    return BuildResult(out_path, state, schema, an, view, reg, refs, dash, data_ws.title, rows, write_cache, build_time)


# ------------------------------------------------------------------ saving
def save_atomic(wb, out_path: Path, cached: dict | None, build_time: dt.datetime, *, require_dashboard: bool = True):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(out_path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        postprocess(tmp, cached, build_time)
        chk = load_workbook(tmp)                  # must reopen before it replaces anything
        if require_dashboard and W.DASHBOARD_SHEET not in chk.sheetnames:
            raise RuntimeError("dashboard sheet missing after save")
        chk.close()
        os.replace(tmp, out_path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def _sheet_paths(z: zipfile.ZipFile) -> dict[str, str]:
    wbx = etree.fromstring(z.read("xl/workbook.xml"))
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rid_to_target = {r.get("Id"): r.get("Target") for r in rels}
    out = {}
    for sh in wbx.iter(f"{{{NS_MAIN}}}sheet"):
        target = rid_to_target[sh.get(f"{{{NS_REL}}}id")]
        target = target.lstrip("/")
        out[sh.get("name")] = target if target.startswith("xl/") else "xl/" + target
    return out


def _set_cached(root, values: dict[str, object]):
    for c in root.iter(f"{{{NS_MAIN}}}c"):
        f = c.find(f"{{{NS_MAIN}}}f")
        if f is None:
            continue
        coord = c.get("r")
        if coord not in values:
            continue
        val = values[coord]
        v = c.find(f"{{{NS_MAIN}}}v")
        if v is None:
            v = etree.SubElement(c, f"{{{NS_MAIN}}}v")
        if val is None or val == "":
            c.set("t", "str")
            v.text = ""
        elif isinstance(val, bool):
            c.set("t", "b")
            v.text = "1" if val else "0"
        elif isinstance(val, (int, float)):
            if "t" in c.attrib:
                del c.attrib["t"]
            v.text = repr(float(val)) if isinstance(val, float) else str(val)
        elif isinstance(val, str) and val.startswith("#") and val.endswith(("!", "?", "A")) and len(val) <= 7:
            c.set("t", "e")
            v.text = val
        else:
            c.set("t", "str")
            v.text = str(val)


APP_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           '<DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop>'
           f'<Company>{W.COMPANY}</Company><LinksUpToDate>false</LinksUpToDate>'
           '<SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged></Properties>')


def postprocess(path, cached: dict | None, build_time: dt.datetime):
    """Write cached results for the selected view and library-free app properties (09.4, D5)."""
    path = str(path)
    by_sheet: dict[str, dict[str, object]] = {}
    for (sheet, coord), val in (cached or {}).items():
        by_sheet.setdefault(sheet, {})[coord] = val
    tmp = path + ".pp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        paths = _sheet_paths(zin)
        rev = {v: k for k, v in paths.items()}
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = APP_XML.encode("utf-8")
            elif item.filename == "docProps/core.xml":
                data = _fix_core(data, build_time)
            elif re.fullmatch(r"xl/comments/comment\d+\.xml", item.filename):
                data = _note_runs(data)
            elif re.fullmatch(r"xl/drawings/commentsDrawing\d+\.vml", item.filename):
                data = _note_vml(data)
            elif item.filename in rev and rev[item.filename] in by_sheet:
                root = etree.fromstring(data)
                _set_cached(root, by_sheet[rev[item.filename]])
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zi = zipfile.ZipInfo(item.filename, date_time=item.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = item.external_attr
            zout.writestr(zi, data)
    os.replace(tmp, path)


NOTE_FONT_SIZE = 9
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _note_runs(data: bytes) -> bytes:
    """Notes in Calibri 9 (D.9, D.12): plain note text becomes one run with explicit run properties."""
    root = etree.fromstring(data)
    q = lambda t: f"{{{MAIN_NS}}}{t}"  # noqa: E731
    for text in root.iter(q("text")):
        plain = [c for c in text if c.tag == q("t")]
        if not plain or len(plain) != len(text):
            for rpr in text.iter(q("rPr")):
                for f in rpr.findall(q("rFont")):
                    f.set("val", T.FONT)
            continue
        value = "".join(t.text or "" for t in plain)
        for t in plain:
            text.remove(t)
        r = etree.SubElement(text, q("r"))
        rpr = etree.SubElement(r, q("rPr"))
        etree.SubElement(rpr, q("sz")).set("val", str(NOTE_FONT_SIZE))
        etree.SubElement(rpr, q("color")).set("rgb", "FF" + T.TEXT_BLACK)
        etree.SubElement(rpr, q("rFont")).set("val", T.FONT)
        etree.SubElement(rpr, q("family")).set("val", "2")
        t = etree.SubElement(r, q("t"))
        t.text = value
        if value != value.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _note_vml(data: bytes) -> bytes:
    """Note boxes right-aligned for Arabic text and filled with a design token instead of the library default."""
    s = data.decode("utf-8")
    s = s.replace("text-align:left", "text-align:right")
    s = re.sub(r'(fillcolor|color2)="#ffffe1"', lambda m: f'{m.group(1)}="#{T.NOTE_FILL}"', s, flags=re.I)
    return s.encode("utf-8")


def _fix_core(data: bytes, build_time: dt.datetime) -> bytes:
    root = etree.fromstring(data)
    ns = {"cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
          "dc": "http://purl.org/dc/elements/1.1/", "dcterms": "http://purl.org/dc/terms/",
          "xsi": "http://www.w3.org/2001/XMLSchema-instance"}
    stamp = build_time.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    for tag in ("created", "modified"):
        el = root.find(f"dcterms:{tag}", ns)
        if el is None:
            el = etree.SubElement(root, f"{{{ns['dcterms']}}}{tag}")
            el.set(f"{{{ns['xsi']}}}type", "dcterms:W3CDTF")
        el.text = stamp
    lm = root.find("cp:lastModifiedBy", ns)
    if lm is None:
        lm = etree.SubElement(root, f"{{{ns['cp']}}}lastModifiedBy")
    lm.text = W.AUTHOR
    cr = root.find("dc:creator", ns)
    if cr is None:
        cr = etree.SubElement(root, f"{{{ns['dc']}}}creator")
    cr.text = W.AUTHOR
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
