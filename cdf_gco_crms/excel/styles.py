"""Shared openpyxl style factories (Calibri everywhere, 09/10.4)."""
from __future__ import annotations

from functools import lru_cache

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ..config import design_tokens as T


@lru_cache(maxsize=None)
def font(size: float = T.SZ_BODY, bold: bool = False, color: str = T.TEXT_BLACK,
         italic: bool = False, underline: str | None = None) -> Font:
    return Font(name=T.FONT, size=size, bold=bold, italic=italic, color="FF" + color,
                underline=underline)


@lru_cache(maxsize=None)
def fill(color: str | None) -> PatternFill:
    if color is None:
        return PatternFill(fill_type=None)
    return PatternFill(fill_type="solid", start_color="FF" + color, end_color="FF" + color)


@lru_cache(maxsize=None)
def align(h: str = "right", v: str = "top", wrap: bool = True, indent: int = 0,
          rtl: bool = True) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent,
                     readingOrder=2 if rtl else 0)


@lru_cache(maxsize=None)
def side(style: str | None = "thin", color: str = T.BORDER) -> Side:
    if style is None:
        return Side(style=None)
    return Side(style=style, color="FF" + color)


@lru_cache(maxsize=None)
def border(left: str | None = "thin", right: str | None = "thin", top: str | None = "thin",
           bottom: str | None = "thin", color: str = T.BORDER) -> Border:
    return Border(left=side(left, color), right=side(right, color),
                  top=side(top, color), bottom=side(bottom, color))


THIN = border()
NONE = Border()
