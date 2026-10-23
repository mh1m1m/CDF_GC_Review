"""Renders the cumulative data sheet from the review state (protocol Art. 10).

The renderer is deterministic: the same state always gives the same sheet.
It writes rows in Reader_No order (SC-40), one fixed block A–E, the round-1
block F–H and one four-column block per later round, with the dropdowns,
placeholders and fills of 10.3, 10.4 and 10.8.
"""
from __future__ import annotations

import math

from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..config import design_tokens as T
from ..config import round_schema as S
from ..guidance import guidance_dictionary as G
from ..model.entities import GText, Observation, ReviewState, is_formula_like, plain_text, text_cell_value
from ..rounds.round_parser import RoundSchema, discover
from . import styles as st


def _cell_value(v):
    """Cell value of a model field. Text is always written as text (G-SUB-7: no formula may originate from
    a free-text field of the register); a GText keeps its bold runs."""
    if isinstance(v, GText):
        return v.to_cell_value()
    if is_formula_like(v):
        return text_cell_value(v)                                   # inline string, never a formula
    return v


def _header_value(role: str | None, text: str):
    """Requester headers carry a 9 pt regular guidance line (10.8)."""
    if role in S.REQUESTER_ROLES and "\n" in text:
        first, second = text.split("\n", 1)
        return CellRichText([
            TextBlock(InlineFont(rFont=T.FONT, sz=10, b=True), first),
            TextBlock(InlineFont(rFont=T.FONT, sz=9, b=False), "\n" + second),
        ])
    return text


def planned_layout(rounds: int) -> dict[int, dict[str, int]]:
    """Canonical column of each role; used only to *write*. Readers use discovery."""
    layout = {1: {S.ROLE_CLOSURE: 6, S.ROLE_RESPONSE: 7, S.ROLE_IMPL: 8}}
    c = 9
    for r in range(2, rounds + 1):
        layout[r] = {}
        for role in S.ROUNDN_ROLES:
            layout[r][role] = c
            c += 1
    return layout


def _estimate_lines(text: str, width: float) -> int:
    if not text:
        return 1
    per_line = max(8, int(width * 1.15))
    return sum(max(1, math.ceil(len(seg) / per_line)) for seg in text.split("\n"))


def _note(text: str):
    from openpyxl.comments import Comment
    lines = sum(max(1, len(seg) // 48 + 1) for seg in text.split("\n"))
    c = Comment(text, G.NOTE_AUTHOR)
    c.width = 320
    c.height = min(400, 18 * lines + 30)
    return c


def write_data_sheet(wb, state: ReviewState, *, outline: bool = True, title: str | None = None,
                     guidance: bool = True, doc_family: str = G.DOC_AGREEMENT,
                     closed_guide_colour: str = T.CLOSED_GUIDE_TEXT) -> tuple:
    R = state.current_round
    ws = wb.create_sheet(title=title or S.data_sheet_name(R))
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    layout = planned_layout(R)
    obs = state.ordered()
    last_row = 1 + len(obs)
    last_col = max(max(b.values()) for b in layout.values())

    # ---------------------------------------------------------------- headers
    hdr_font = st.font(10, bold=True)
    for c, h in enumerate(S.FIXED_HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = hdr_font
        cell.fill = st.fill(T.HEADER_CYAN)
        cell.alignment = st.align("center", "center")
        cell.border = st.THIN
        ws.column_dimensions[get_column_letter(c)].width = S.FIXED_WIDTHS[c]
        key = {1: "reader_no", 3: "location", 4: "quote", 5: "comment"}.get(c)
        if guidance and key:
            cell.comment = _note(G.fixed_note(key))
    ws.column_dimensions["B"].hidden = True
    for r, roles in layout.items():
        for i, (role, c) in enumerate(roles.items()):
            text = S.header_text(role, r)
            cell = ws.cell(row=1, column=c, value=_header_value(role, text))
            cell.font = hdr_font
            cell.fill = st.fill(T.LIGHT_GRAY if role in S.REQUESTER_ROLES else T.HEADER_CYAN)
            cell.alignment = st.align("center", "center")
            cell.border = st.border(left="double" if i == 0 else "thin")
            ws.column_dimensions[get_column_letter(c)].width = S.ROLE_WIDTH[role]
            if guidance:
                cell.comment = _note(G.header_note(role, r, R, doc_family))
    ws.row_dimensions[1].height = 36

    # ------------------------------------------------------------------- rows
    for idx, o in enumerate(obs):
        row = 2 + idx
        _write_fixed(ws, row, o)
        for r, roles in layout.items():
            _write_block(ws, row, o, r, roles, R)
        ws.row_dimensions[row].height = _row_height(o, layout, R)

    # ------------------------------------------------------------ validations
    closed_rows = [2 + i for i, o in enumerate(obs) if o.value(R) == S.CLOSED]
    add_validations(ws, layout, last_row, current_round=R, closed_rows=closed_rows,
                    guidance=guidance, doc_family=doc_family)
    add_placeholder_formats(ws, layout, R, last_row, closed_guide_colour)

    # ------------------------------------------------------------------ views
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}{max(last_row, 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.print_area = f"A1:{get_column_letter(last_col)}{max(last_row, 1)}"
    if outline and R >= 3:
        for r in range(1, R - 1):          # every block older than the latest two
            for c in layout[r].values():
                ws.column_dimensions[get_column_letter(c)].outline_level = 1
        ws.sheet_properties.outlinePr.summaryRight = False
    schema = discover(ws)
    return ws, schema


def _write_fixed(ws, row: int, o: Observation):
    a = ws.cell(row=row, column=1, value=o.reader_no)
    a.number_format = "@"
    a.font = st.font(10, bold=True)
    a.alignment = st.align("center")
    ws.cell(row=row, column=2, value=o.internal_id).font = st.font(10)
    ws.cell(row=row, column=3, value=_cell_value(o.location)).font = st.font(10)
    ws.cell(row=row, column=4, value=_cell_value(o.quote)).font = st.font(10)
    ws.cell(row=row, column=5, value=_cell_value(o.comment)).font = st.font(10)
    for c in range(1, 6):
        cell = ws.cell(row=row, column=c)
        cell.border = st.THIN
        if c != 1:
            cell.alignment = st.align("right")


def _write_block(ws, row: int, o: Observation, r: int, roles: dict[str, int], R: int):
    cells = o.cells.get(r)
    not_raised = r < o.birth_round
    value = "" if not_raised else plain_text(cells.closure if cells else None).strip()
    for i, (role, c) in enumerate(roles.items()):
        cell = ws.cell(row=row, column=c)
        cell.border = st.border(left="double" if i == 0 else "thin")
        raw = None
        if cells is not None:
            raw = getattr(cells, role)
        if role == S.ROLE_CLOSURE:
            raw = S.NOT_RAISED if not_raised else raw
        cell.value = _cell_value(raw)
        centred = role in (S.ROLE_CLOSURE, S.ROLE_IMPL)
        cell.alignment = st.align("center" if centred else "right")
        if not_raised:
            cell.fill = st.fill(T.CLOSED_ROW_GRAY)
            cell.font = st.font(9, color=T.CLOSED_ROW_TEXT)
            continue
        if role == S.ROLE_CLOSURE:
            fill = T.ROUND1_OPEN_FILL if (r == 1 and value == S.OPEN) else T.STATUS_FILL.get(value)
            cell.fill = st.fill(fill)
            cell.font = st.font(10, bold=True)
        elif role == S.ROLE_FEEDBACK:
            cell.font = st.font(10)
        else:  # requester columns
            if value == S.CLOSED:
                cell.fill = st.fill(T.CLOSED_ROW_GRAY)
                cell.font = st.font(9, color=T.CLOSED_ROW_TEXT) if role == S.ROLE_IMPL else st.font(10)
            elif r == R:
                cell.fill = st.fill(T.INPUT)
                cell.font = st.font(10)
            else:
                cell.fill = st.fill(T.LIGHT_GRAY)
                cell.font = st.font(10)


def _row_height(o: Observation, layout, R: int) -> float:
    texts = [(plain_text(o.quote), S.FIXED_WIDTHS[4]), (plain_text(o.comment), S.FIXED_WIDTHS[5])]
    for r in (R - 1, R):
        if r >= 2 and r in o.cells:
            texts.append((plain_text(o.cells[r].feedback), S.ROLE_WIDTH[S.ROLE_FEEDBACK]))
            texts.append((plain_text(o.cells[r].response), S.ROLE_WIDTH[S.ROLE_RESPONSE]))
    lines = max(_estimate_lines(t, w) for t, w in texts)
    return float(min(409, max(30, 13.5 * lines + 6)))


def _ranges(rows: list[int]) -> list[tuple[int, int]]:
    out: list[tuple[int, int]] = []
    for r in sorted(rows):
        if out and r == out[-1][1] + 1:
            out[-1] = (out[-1][0], r)
        else:
            out.append((r, r))
    return out


def add_validations(ws, layout: dict[int, dict[str, int]], last_row: int, *, current_round: int | None = None,
                    closed_rows: list[int] | None = None, guidance: bool = True,
                    doc_family: str = G.DOC_AGREEMENT):
    """Inline-list dropdowns on every data row of every block (10.3, G-SYNC-6) and input messages (10.4A.14)."""
    ws.data_validations.dataValidation = []
    if last_row < 2:
        last_row = 2
    R = current_round or max(layout)
    closure_cols = [get_column_letter(b[S.ROLE_CLOSURE]) for b in layout.values()]
    impl_cols = [get_column_letter(b[S.ROLE_IMPL]) for b in layout.values()]
    dv_c = DataValidation(type="list", formula1='"' + ",".join(S.CLOSURE_VALUES) + '"',
                          allow_blank=True, showDropDown=False, showErrorMessage=True,
                          errorStyle="stop")
    if guidance:
        t, p = G.input_message("closure")
        dv_c.showInputMessage = True
        dv_c.promptTitle, dv_c.prompt = t, p
    dv_c.sqref = " ".join(f"{c}2:{c}{last_row}" for c in closure_cols)
    dv_i = DataValidation(type="list", formula1='"' + ",".join(S.IMPL_VALUES) + '"',
                          allow_blank=True, showDropDown=False, showErrorMessage=True,
                          showInputMessage=True, errorStyle="stop",
                          promptTitle=S.DV_IMPL_PROMPT_TITLE, prompt=S.DV_IMPL_PROMPT,
                          errorTitle=S.DV_IMPL_ERROR_TITLE, error=S.DV_IMPL_ERROR)
    dv_i.sqref = " ".join(f"{c}2:{c}{last_row}" for c in impl_cols)
    dvs = [dv_c, dv_i]
    if guidance:
        # input messages on the current round's free-text response cells, by closure state
        col = get_column_letter(layout[R][S.ROLE_RESPONSE])
        closed = set(closed_rows or [])
        open_rows = [r for r in range(2, last_row + 1) if r not in closed]
        for key, rows in (("response_current", open_rows), ("response_closed", sorted(closed))):
            if not rows:
                continue
            t, p = G.input_message(key, doc_family)
            dv = DataValidation(type=None, allow_blank=True, showInputMessage=True,
                                showErrorMessage=False, promptTitle=t, prompt=p)
            dv.sqref = " ".join(f"{col}{a}:{col}{b}" for a, b in _ranges(rows))
            dvs.append(dv)
    for dv in dvs:
        if dv.formula1 and len(dv.formula1) > 257:
            raise ValueError("inline list exceeds 255 characters")
        ws.add_data_validation(dv)


def add_placeholder_formats(ws, layout, R: int, last_row: int, closed_guide_colour: str = T.CLOSED_GUIDE_TEXT):
    """Placeholder look by conditional format: grey italic only while the cell holds the guidance (10.8).

    The closed-row guidance sits on the #D9D9D9 fill, where #808080 measures 2.80:1; from v7.11 it is
    painted #595959 (4.96:1, AMB-153). The v7.10 fixture writer passes #808080 to reproduce v7.10 files.
    """
    if last_row < 2:
        last_row = 2
    ph_font = Font(name=T.FONT, size=10, italic=True, color="FF" + T.PLACEHOLDER_GRAY)
    closed_font = Font(name=T.FONT, size=10, italic=True, color="FF" + closed_guide_colour)
    for r, roles in layout.items():
        for role in S.REQUESTER_ROLES:
            col = get_column_letter(roles[role])
            rng = f"{col}2:{col}{last_row}"
            if role == S.ROLE_RESPONSE:
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'EXACT(${col}2,"{S.PH_TEXT}")'], font=ph_font, stopIfTrue=True))
                pref = S.CLOSED_GUIDE_PREFIX
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'LEFT(${col}2,LEN("{pref}"))="{pref}"'], font=closed_font, stopIfTrue=True))
            else:
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'EXACT(${col}2,"{S.PH_STATUS}")'], font=ph_font, stopIfTrue=True))
