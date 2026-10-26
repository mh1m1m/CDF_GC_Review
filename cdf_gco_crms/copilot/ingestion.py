"""Bounded supplemental extraction. Original bytes and anchors are immutable."""

import base64, hashlib, json, tempfile, uuid, zipfile
from io import BytesIO
from pathlib import Path
from ..calibration.store import CalibrationError, canonical, now_utc
from ..calibration.input_safety import check_ooxml_zip
from ..execution.service import Conflict


class IngestionService:
    def __init__(self, app):
        self.app = app
        app.store.db.executescript(
            """CREATE TABLE IF NOT EXISTS cp_documents(document_id TEXT PRIMARY KEY,case_id TEXT NOT NULL,sha256 TEXT NOT NULL,format TEXT NOT NULL,content BLOB NOT NULL,anchors_json TEXT NOT NULL,warnings_json TEXT NOT NULL,version INTEGER NOT NULL,at TEXT NOT NULL,UNIQUE(case_id,sha256,format));
        CREATE TABLE IF NOT EXISTS cp_corrections(event_id TEXT PRIMARY KEY,document_id TEXT NOT NULL REFERENCES cp_documents(document_id),anchor_id TEXT NOT NULL,version INTEGER NOT NULL,text TEXT NOT NULL,rationale TEXT NOT NULL,actor TEXT NOT NULL,at TEXT NOT NULL);"""
        )

    def extract(self, content, kind):
        if not 0 < len(content) <= 60000:
            raise CalibrationError("DOCUMENT_SIZE_LIMIT_60000_BYTES")
        anchors = []
        warnings = []

        def add(location, text):
            if text:
                anchors.append(
                    {
                        "anchor_id": "A" + str(len(anchors) + 1),
                        "location": location,
                        "text": str(text)[:12000],
                    }
                )
            if len(anchors) > 1000:
                raise CalibrationError("EXTRACTION_ANCHOR_LIMIT")

        if kind == "image":
            raise CalibrationError("IMAGE_OCR_REQUIRES_AUTHORIZED_EXTRACTION_PROVIDER")
        if kind == "text":
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError as exc:
                raise CalibrationError("UTF8_TEXT_REQUIRED") from exc
            for i, line in enumerate(text.splitlines(), 1):
                add("line " + str(i), line)
        elif kind in ("docx", "xlsx"):
            with tempfile.TemporaryDirectory(prefix="cdf-extract-") as temp:
                path = Path(temp) / ("source." + kind)
                path.write_bytes(content)
                check_ooxml_zip(path, max_compressed=60000)
                with zipfile.ZipFile(path) as package:
                    if any("vbaproject" in n.lower() for n in package.namelist()):
                        raise CalibrationError("ACTIVE_CONTENT_NOT_SUPPORTED")
                    from ..assurance.safe_xml import has_dtd

                    for name in package.namelist():
                        if name.lower().endswith((".xml", ".rels")) and has_dtd(package.read(name)):
                            raise CalibrationError("XML_ENTITY_DECLARATION_REJECTED")
                if kind == "docx":
                    from docx import Document

                    doc = Document(path)
                    for i, p in enumerate(doc.paragraphs, 1):
                        add("paragraph " + str(i), p.text)
                    for t, table in enumerate(doc.tables, 1):
                        for r, row in enumerate(table.rows, 1):
                            for c, cell in enumerate(row.cells, 1):
                                add(f"table {t}/row {r}/cell {c}", cell.text)
                else:
                    from openpyxl import load_workbook

                    book = load_workbook(
                        path, read_only=True, data_only=False, keep_links=False
                    )
                    try:
                        for sheet in book:
                            if sheet.max_row > 1000 or sheet.max_column > 100:
                                raise CalibrationError("WORKBOOK_DIMENSION_LIMIT")
                            for row in sheet:
                                for cell in row:
                                    if cell.value is not None:
                                        add(
                                            sheet.title + "!" + cell.coordinate,
                                            str(cell.value),
                                        )
                        warnings.append("FORMULAS_PRESERVED_AS_TEXT_NOT_EXECUTED")
                    finally:
                        book.close()
        elif kind == "pdf":
            if not content.startswith(b"%PDF-"):
                raise CalibrationError("PDF_SIGNATURE_INVALID")
            import pymupdf

            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                if pdf.needs_pass or len(pdf) > 100:
                    raise CalibrationError("PDF_ENCRYPTED_OR_PAGE_LIMIT")
                for i, page in enumerate(pdf, 1):
                    text = page.get_text()
                    add("page " + str(i), text)
                    if not text.strip():
                        warnings.append("PAGE_" + str(i) + "_REQUIRES_AUTHORIZED_OCR")
            warnings.append("PDF_READING_ORDER_REQUIRES_MANUAL_VERIFICATION")
        if not anchors:
            warnings.append("NO_EXTRACTABLE_TEXT")
        if sum(len(a["text"]) for a in anchors) > 100000:
            raise CalibrationError("EXTRACTION_TEXT_LIMIT")
        return anchors, warnings

    def dispatch(self, action, r, p):
        a = self.app
        a.authorize_case(r.case_id, p)
        if action == "ingest":
            if r.synthetic is not True:
                raise CalibrationError(
                    "PUBLIC_SYNTHETIC_CONTENT_ONLY_PRIVATE_ACTIVATION_DISABLED"
                )
            try:
                content = base64.b64decode(r.content_base64, validate=True)
            except ValueError as exc:
                raise CalibrationError("INVALID_BASE64") from exc
            digest = hashlib.sha256(content).hexdigest()
            old = a.store.db.execute(
                "SELECT document_id FROM cp_documents WHERE case_id=? AND sha256=? AND format=?",
                (r.case_id, digest, r.format),
            ).fetchone()
            if old:
                return {
                    "document_id": old[0],
                    "sha256": digest,
                    "duplicate": True,
                    "used_as_review_source": False,
                }
            anchors, warnings = self.extract(content, r.format)
            doc = "DOC-" + uuid.uuid4().hex
            a.store.db.execute(
                "INSERT INTO cp_documents VALUES (?,?,?,?,?,?,?,?,?)",
                (
                    doc,
                    r.case_id,
                    digest,
                    r.format,
                    content,
                    canonical(anchors),
                    canonical(warnings),
                    1,
                    now_utc(),
                ),
            )
            return {
                "document_id": doc,
                "case_id": r.case_id,
                "sha256": digest,
                "version": 1,
                "anchors": anchors,
                "warnings": warnings,
                "extraction_confidence": None,
                "quality": "STRUCTURAL_EXTRACTION_REQUIRES_HUMAN_VERIFICATION",
                "used_as_review_source": False,
                "scope": "SUPPLEMENTAL_SYNTHETIC_DOCUMENT_NOT_AUTOMATICALLY_REVIEWED",
            }
        rows = a.store.db.execute(
            "SELECT * FROM cp_documents WHERE case_id=? AND (CAST(? AS TEXT) IS NULL OR document_id=?) ORDER BY at",
            (r.case_id, r.document_id, r.document_id),
        ).fetchall()
        if action == "correct":
            if len(rows) != 1 or not r.document_id:
                raise CalibrationError("DOCUMENT_NOT_FOUND")
            row = rows[0]
            if str(row["version"]) != r.expected_version:
                raise Conflict("EXTRACTION_VERSION_CHANGED")
            if (
                r.anchor_id
                not in [x["anchor_id"] for x in json.loads(row["anchors_json"])]
                or not r.corrected_text
                or not r.rationale
            ):
                raise CalibrationError("CORRECTION_ANCHOR_TEXT_AND_RATIONALE_REQUIRED")
            version = row["version"] + 1
            event = uuid.uuid4().hex
            a.store.db.execute(
                "INSERT INTO cp_corrections VALUES (?,?,?,?,?,?,?,?)",
                (
                    event,
                    r.document_id,
                    r.anchor_id,
                    version,
                    r.corrected_text,
                    r.rationale,
                    p.reviewer.reviewer_id,
                    now_utc(),
                ),
            )
            a.store.db.execute(
                "UPDATE cp_documents SET version=? WHERE document_id=?",
                (version, r.document_id),
            )
            return {
                "document_id": r.document_id,
                "version": version,
                "event_id": event,
                "original_bytes_preserved": True,
                "source_verified": False,
            }
        return {
            "case_id": r.case_id,
            "documents": [
                {
                    "document_id": row["document_id"],
                    "sha256": row["sha256"],
                    "format": row["format"],
                    "version": row["version"],
                    "original_anchors": json.loads(row["anchors_json"]),
                    "warnings": json.loads(row["warnings_json"]),
                    "corrections": [
                        dict(c)
                        for c in a.store.db.execute(
                            "SELECT anchor_id,version,text,rationale,actor,at FROM cp_corrections WHERE document_id=? ORDER BY version",
                            (row["document_id"],),
                        )
                    ],
                    "used_as_review_source": False,
                }
                for row in rows[r.cursor : r.cursor + r.limit]
            ],
            "total": len(rows),
        }
