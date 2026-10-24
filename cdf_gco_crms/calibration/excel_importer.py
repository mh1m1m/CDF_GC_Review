"""Offline XLSX observation evolution importer (private workspaces only).

Do not upload real institutional registers to the public repository or a
public CI environment. The caller controls confidential source access.
"""
from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .store import CalibrationError, CalibrationStore, ReviewerContext

MAX_XLSX_BYTES = 20 * 1024 * 1024

FIELD_ALIASES = {
    "observation_id": ["observation_id", "Observation_ID", "internal_id", "Internal_ID"],
    "original_text": ["original_text", "Original_Comment", "Initial_Text", "Original_Text"],
    "reviewed_text": ["reviewed_text", "MAT_Comment", "Revised_Text", "Final_Text"],
    "observed_change": ["observed_change", "Change_Type", "Observed_Change_Type", "Evolution"],
    "original_class": ["original_class", "Original_Classification", "Original_Comment_Class"],
    "reviewed_class": ["reviewed_class", "MAT_Classification", "Final_Comment_Class"],
    "matching_method": ["matching_method", "Matching_Method"],
    "matching_confidence": ["matching_confidence", "Matching_Confidence"],
    "manual_review": ["manual_review", "Manual_Review"],
    "document_article": ["document_article", "Document_Article"],
    "document_anchor": ["document_anchor", "Document_Anchor"],
}


def _locate(columns: list[str], key: str):
    normalized = {str(h).strip().lower().replace(" ", "_"): i for i, h in enumerate(columns)}
    for alias in FIELD_ALIASES[key]:
        ix = normalized.get(alias.lower())
        if ix is not None:
            return ix
    return None


def _read_cell(row, ix):
    return row[ix] if ix is not None and ix < len(row) else None


def preview_xlsx(path: str | Path) -> dict:
    p=Path(path).expanduser()
    if p.is_symlink() or not p.is_file() or p.suffix.lower()!=".xlsx":
        raise CalibrationError("INVALID_EVOLUTION_XLSX")
    if not (0 < p.stat().st_size <= MAX_XLSX_BYTES):
        raise CalibrationError("EVOLUTION_XLSX_SIZE_OUT_OF_RANGE")
    with p.open("rb") as handle:
        if handle.read(4)!=b"PK\x03\x04":
            raise CalibrationError("INVALID_EVOLUTION_XLSX_SIGNATURE")
    source_hash=sha256(p.read_bytes()).hexdigest()
    wb=load_workbook(p, read_only=True, data_only=True)
    try:
        ws=next((wb[name] for name in ("Evolution","Observation Evolution","الملاحظات","Evolution Register") if name in wb.sheetnames),wb.active)
        iterator=ws.iter_rows(values_only=True)
        headers=next(iterator,None)
        if not headers:
            raise CalibrationError("NO_XLSX_HEADERS")
        header=[str(h or "").strip() for h in headers]
        loc={k:_locate(header,k) for k in FIELD_ALIASES}
        for required in ("observation_id","original_text","reviewed_text"):
            if loc[required] is None:
                raise CalibrationError("MISSING_EVOLUTION_COLUMN:"+required)
        result=[]; seen=set(); exceptions=[]
        for row_number,row in enumerate(iterator,2):
            oid=str(_read_cell(row,loc["observation_id"]) or "").strip()
            if not oid:continue
            if oid in seen:
                exceptions.append({"row":row_number,"code":"DUPLICATE_OBSERVATION_ID","observation_id":oid})
                continue
            seen.add(oid)
            original=str(_read_cell(row,loc["original_text"]) or "")
            revised=str(_read_cell(row,loc["reviewed_text"]) or "")
            raw_change=str(_read_cell(row,loc["observed_change"]) or "")
            if "REMOVED_FROM_VERSION" in raw_change:
                change="REMOVED_FROM_VERSION"
            elif "ADDED_IN_REVIEW" in raw_change:
                change="ADDED_IN_REVIEW"
            elif "EMPTIED" in raw_change:
                change="EMPTIED_IN_PLACE"
            elif original and original==revised:
                change="KEEP_UNCHANGED"
            else:
                change="REVISED"
            observed={
                "observation_id":oid,"original_text":original,"reviewed_text":revised,
                "observed_change":change,"raw_change":raw_change,
                "primary_type_original":_read_cell(row,loc["original_class"]),
                "primary_type_reviewed":_read_cell(row,loc["reviewed_class"]),
                "matching_method":_read_cell(row,loc["matching_method"]),
                "matching_confidence":_read_cell(row,loc["matching_confidence"]),
                "document_article":_read_cell(row,loc["document_article"]),
                "document_anchor":_read_cell(row,loc["document_anchor"]),
                "manual_review":bool(_read_cell(row,loc["manual_review"])),
                "material_categories":[],"documented_withdrawal":False,
                "attribution_status":"NOT_VERIFIED_FROM_SPREADSHEET",
            }
            if change=="REMOVED_FROM_VERSION":
                observed["manual_review"]=True
            result.append(observed)
        return {"source_sha256":source_hash,"observations":result,"exceptions":exceptions,
                "row_count":len(result),"withdrawal_decisions_verified":0,
                "evidence_status":"REGISTER_ONLY_NOT_OOXML_VERIFIED"}
    finally:
        wb.close()


def import_xlsx(store:CalibrationStore, case_id:str,path:str|Path,reviewer:ReviewerContext)->dict:
    reviewer.require("CALIBRATION_ADMIN")
    result=preview_xlsx(path)
    if result["exceptions"]:
        raise CalibrationError("EVOLUTION_RECONCILIATION_EXCEPTIONS")
    evidence_hash=store.import_case(case_id,result["source_sha256"],result["observations"],reviewer)
    return {"case_id":case_id,"source_sha256":result["source_sha256"],
            "evidence_hash":evidence_hash,"observations":result["row_count"],
            "evidence_status":"REGISTER_ONLY_NOT_OOXML_VERIFIED"}
