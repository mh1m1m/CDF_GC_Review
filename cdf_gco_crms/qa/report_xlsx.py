"""Plain tabular report workbooks (crosswalk, regression results, brand register, UX audit, matrices).

Calibri, header in the brand surface colour, frozen header, filter, wrapped
text, metadata per 09.4 (author, last modified by, company, title), no
library names in the properties.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

from ..config import design_tokens as T
from ..config import workbook_settings as W
from ..excel import styles as st
from ..excel.workbook_builder import RIYADH, save_atomic


@dataclass
class SheetSpec:
    name: str
    headers: list[str]
    rows: list[list]
    widths: list[float] | None = None
    rtl: bool = False
    status_col: str | None = None          # header name whose PASS/FAIL values get a fill
    notes: list[str] = field(default_factory=list)   # lines written above the table


STATUS_FILLS = {"PASS": T.STATUS_FILL["مغلقة"], "PASSED": T.STATUS_FILL["مغلقة"], "MANUAL_PASS": T.STATUS_FILL["مغلقة بشرط"],
                "FAIL": T.STATUS_FILL["مفتوحة"], "FAILED": T.STATUS_FILL["مفتوحة"], "NOT_VERIFIED": T.STATUS_FILL["يتطلب تعديل"],
                "NOT_RUN": T.STATUS_FILL["يتطلب تعديل"], "IMPLEMENTED": T.STATUS_FILL["مغلقة"],
                "PARTIAL": T.STATUS_FILL["يتطلب تعديل"], "NOT_IMPLEMENTED": T.STATUS_FILL["مفتوحة"],
                "UNCHANGED": "FFFFFF", "MODIFIED": T.STATUS_FILL["مغلقة بشرط"], "ADDED": T.STATUS_FILL["مغلقة"],
                "REPLACED": T.STATUS_FILL["يتطلب تعديل"], "MOVED": T.STATUS_FILL["مغلقة بشرط"],
                "REMOVED_WITH_JUSTIFICATION": T.STATUS_FILL["مفتوحة"]}


def write_report(path, title: str, sheets: list[SheetSpec], *, build_time: dt.datetime | None = None) -> Path:
    build_time = build_time or dt.datetime.now(RIYADH)
    wb = Workbook()
    wb.remove(wb.active)
    for spec in sheets:
        ws = wb.create_sheet(spec.name[:31])
        ws.sheet_view.rightToLeft = spec.rtl
        ws.sheet_view.showGridLines = False
        r0 = 1
        for line in spec.notes:
            c = ws.cell(row=r0, column=1, value=line)
            c.font = st.font(10, bold=(r0 == 1))
            c.alignment = st.align("right" if spec.rtl else "left", "top", False, rtl=spec.rtl)
            r0 += 1
        if spec.notes:
            r0 += 1
        for j, h in enumerate(spec.headers, start=1):
            c = ws.cell(row=r0, column=j, value=h)
            c.font = st.font(10, bold=True)
            c.fill = st.fill(T.BRAND_SURFACE)
            c.alignment = st.align("center", "center", True, rtl=spec.rtl)
            c.border = st.THIN
        for i, row in enumerate(spec.rows, start=r0 + 1):
            for j, v in enumerate(row, start=1):
                if isinstance(v, (list, tuple, set)):
                    v = "; ".join(map(str, v))
                if isinstance(v, str) and len(v) > 32000:
                    v = v[:32000] + "…"
                c = ws.cell(row=i, column=j, value=v)
                c.font = st.font(10)
                c.alignment = st.align("right" if spec.rtl else "left", "top", True, rtl=spec.rtl)
                c.border = st.THIN
        widths = spec.widths or [18] * len(spec.headers)
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        last = r0 + len(spec.rows)
        ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(spec.headers))}{max(last, r0)}"
        if spec.status_col and spec.status_col in spec.headers and spec.rows:
            col = get_column_letter(spec.headers.index(spec.status_col) + 1)
            for value, color in STATUS_FILLS.items():
                ws.conditional_formatting.add(f"{col}{r0 + 1}:{col}{last}", FormulaRule(
                    formula=[f'${col}{r0 + 1}="{value}"'], fill=st.fill(color), stopIfTrue=True))
    p = wb.properties
    p.creator = p.lastModifiedBy = W.AUTHOR
    p.title = title
    path = Path(path)
    save_atomic(wb, path, None, build_time, require_dashboard=False)
    return path
