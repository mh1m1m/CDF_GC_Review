"""UX gates G-UX-01…20 (UX directive §12), executed on the saved workbook."""
from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as gl

from ..branding import brand_tokens as BT
from ..branding.contrast import contrast_ratio
from ..config import design_tokens as T
from ..config import round_schema as S
from ..config import workbook_settings as W
from ..guidance import guidance_dictionary as G
from ..ingestion.workbook_reader import parse_workbook
from . import inspect as I
from .gates import FAIL, MANUAL, NV, PASS, GateContext, GateResult, _status


def token_colours() -> dict[str, str]:
    """hex -> family for every colour defined in design_tokens."""
    out = {}
    brand = {t.hex for t in BT.BRAND_TOKENS}
    for name in dir(T):
        val = getattr(T, name)
        vals = val.values() if isinstance(val, dict) else [val]
        for v in vals:
            if isinstance(v, str) and re.fullmatch(r"[0-9A-F]{6}", v):
                fam = ("brand" if v in brand else "semantic" if v in T.STATUS_FILL.values()
                       or v in T.STATUS_TEXT.values() or v in T.STATUS_CHART.values() else "protocol/complementary")
                out.setdefault(v, fam)
    return out


def used_colours(path) -> set[str]:
    cols = set()
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n == "xl/styles.xml" or n.startswith("xl/charts/") or n.startswith("xl/worksheets/sheet"):
                txt = z.read(n).decode("utf-8", "replace")
                # the legacy indexed palette is a definition written by the library, not a colour in use
                txt = re.sub(r"<indexedColors>.*?</indexedColors>", "", txt, flags=re.S)
                cols |= {m.upper()[-6:] for m in re.findall(r'rgb="([0-9A-Fa-f]{8})"', txt)}
                cols |= {m.upper() for m in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', txt)}
            elif n.startswith("xl/drawings/") and n.endswith(".vml"):
                txt = z.read(n).decode("utf-8", "replace")
                cols |= {m.upper() for m in re.findall(r'(?:fillcolor|color2|color)="#([0-9A-Fa-f]{6})"', txt)}
            elif n.startswith("xl/comments/"):
                txt = z.read(n).decode("utf-8", "replace")
                cols |= {m.upper()[-6:] for m in re.findall(r'rgb="([0-9A-Fa-f]{8})"', txt)}
    return cols


# Placeholder grey of 10.8 on the requester fills: kept by the protocol, recorded in D.10 (AMB-153).
RECORDED_CONTRAST_EXCEPTIONS = {(T.PLACEHOLDER_GRAY, T.INPUT), (T.PLACEHOLDER_GRAY, T.LIGHT_GRAY)}
_CF_EXACT = re.compile(r'^EXACT\(\$([A-Z]+)(\d+),"(.*)"\)$', re.S)
_CF_PREFIX = re.compile(r'^LEFT\(\$([A-Z]+)(\d+),LEN\("(.*)"\)\)="\3"$', re.S)


def cf_text_pairs(ws):
    """(coordinate, font colour, fill colour) for every cell that a conditional font rule of the data sheet
    paints: the predicates of 10.8 (exact placeholder text, closed-row guidance prefix) are evaluated in Python."""
    from openpyxl.utils.cell import range_boundaries
    out = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            font = rule.dxf.font if rule.dxf is not None else None
            if font is None or font.color is None or not isinstance(font.color.rgb, str) or not rule.formula:
                continue
            f = rule.formula[0]
            m = _CF_EXACT.match(f) or _CF_PREFIX.match(f)
            if not m:
                continue
            exact = _CF_EXACT.match(f) is not None
            text = m.group(3)
            fg = font.color.rgb[-6:].upper()
            for rng in str(cf.sqref).split():
                c1, r1, c2, r2 = range_boundaries(rng)
                for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                    for c in row:
                        v = c.value if isinstance(c.value, str) else None
                        if v is None or not (v == text if exact else v.startswith(text)):
                            continue
                        out.append((c.coordinate, fg, _effective(c)[1]))
    return out, RECORDED_CONTRAST_EXCEPTIONS


def _effective(cell):
    fg = (cell.font.color.rgb if cell.font and cell.font.color is not None and isinstance(cell.font.color.rgb, str)
          else "FF000000")[-6:]
    bg = (cell.fill.fgColor.rgb[-6:] if cell.fill is not None and cell.fill.fill_type == "solid"
          and isinstance(cell.fill.fgColor.rgb, str) else "FFFFFF")
    return fg.upper(), bg.upper()


def run_ux_gates(ctx: GateContext, dash_results: list[GateResult]) -> list[GateResult]:
    out: list[GateResult] = []
    x = Path(ctx.xlsx)
    inp = x.name
    by = {g.gate: g for g in dash_results}
    wb = load_workbook(x, rich_text=True)
    pkg = I.package(x)
    pw = parse_workbook(x)
    R = pw.state.current_round
    dash = wb[W.DASHBOARD_SHEET]
    data = wb[pw.data_ws.title]

    # 01 brand authenticity
    tokens = token_colours()
    used = used_colours(x)
    allowed = set(tokens)
    unknown = sorted(c for c in used if c not in allowed)
    claims_used = [c.hex for c in BT.UNVERIFIED_CLAIMS if c.hex in used]
    out.append(GateResult("G-UX-01", "Brand authenticity", "every colour in the file is a design token with a "
                          "recorded family; no unverified brand claim used", inp, "0 unknown, 0 claims",
                          f"{len(used)} colours used, unknown={unknown}, unverified claims used={claims_used}",
                          _status(not unknown and not claims_used),
                          "brand tokens: " + ", ".join(f"#{t.hex} {t.verification}" for t in BT.BRAND_TOKENS[:2])))

    # 02 logo integrity
    imgs = [i for i in pkg.images if i["sheet"] == W.DASHBOARD_SHEET]
    if imgs and ctx.logo_source and Path(ctx.logo_source).exists():
        from PIL import Image
        src = Image.open(ctx.logo_source)
        ratio_src = src.height / src.width
        im = imgs[0]
        ratio = im["cy"] / im["cx"]
        with zipfile.ZipFile(x) as z:
            emb = z.read(im["part"])
        from io import BytesIO
        emb_img = Image.open(BytesIO(emb))
        ok = abs(ratio - ratio_src) / ratio_src < 0.01 and im["col"] == 1 and im["row"] == 1 and \
            emb_img.size == src.size
        out.append(GateResult("G-UX-02", "Logo integrity", "embedded logo = supplied logo (pixels), aspect ratio "
                              "kept, placed in the header band", inp, f"ratio {ratio_src:.4f} ±1 %, anchor B2",
                              f"ratio {ratio:.4f}, anchor col {im['col']} row {im['row']}, size {emb_img.size}",
                              _status(ok), f"source sha256 {hashlib.sha256(Path(ctx.logo_source).read_bytes()).hexdigest()[:16]}…"))
    else:
        out.append(GateResult("G-UX-02", "Logo integrity", "logo present", inp, "logo", "absent", FAIL))

    # 03 typography
    fonts = set()
    for ws in wb.worksheets:
        fonts |= I.cell_fonts(ws)
    chart_fonts = set()
    with zipfile.ZipFile(x) as z:
        for n in z.namelist():
            if n.startswith("xl/charts/"):
                chart_fonts |= set(re.findall(r'(?:latin|cs) typeface="([^"]+)"', z.read(n).decode()))
    note_fonts, plain_notes = set(), 0
    with zipfile.ZipFile(x) as z:
        for n in z.namelist():
            if re.fullmatch(r"xl/comments/comment\d+\.xml", n):
                txt = z.read(n).decode()
                note_fonts |= set(re.findall(r'<rFont val="([^"]+)"', txt))
                plain_notes += len(re.findall(r"<text><t[ >]", txt))
    out.append(GateResult("G-UX-03", "Typography", "fonts of every cell, rich-text run, chart text and Note", inp,
                          "Calibri only", f"cells={sorted(fonts)}, charts={sorted(chart_fonts)}, "
                          f"notes={sorted(note_fonts)}, notes without run font={plain_notes}",
                          _status(fonts <= {"Calibri"} and chart_fonts <= {"Calibri"} and note_fonts <= {"Calibri"}
                                  and not plain_notes)))

    # 04 contrast of static text/background pairs
    worst = []
    checked = 0
    for ws in (dash, data):
        for row in ws.iter_rows():
            for c in row:
                if c.value in (None, ""):
                    continue
                fg, bg = _effective(c)
                ratio = contrast_ratio(fg, bg)
                large = (c.font.sz or 11) >= 18 or ((c.font.sz or 11) >= 14 and c.font.b)
                need = 3.0 if large else 4.5
                checked += 1
                if ratio < need:
                    worst.append(f"{ws.title}!{c.coordinate} #{fg} on #{bg} = {ratio:.2f}")
    cf_pairs = [(T.BRAND_INK, f) for f in T.STATUS_FILL.values()] + [(T.BRAND_INK, T.FLAG_REOPENED),
                                                                      (T.BRAND_INK, T.FLAG_STALLED),
                                                                      (T.BRAND_INK, T.SELECTED_ROW),
                                                                      (T.MUTED_TEXT, T.ZEBRA)]
    cf_bad = [f"#{a} on #{b}" for a, b in cf_pairs if contrast_ratio(a, b) < 4.5]
    painted, exceptions = cf_text_pairs(data)
    cf_cells_bad = [f"{data.title}!{c} #{a} on #{b} = {contrast_ratio(a, b):.2f}" for c, a, b in painted
                    if contrast_ratio(a, b) < 4.5 and (a, b) not in RECORDED_CONTRAST_EXCEPTIONS]
    exc_seen = sorted({f"#{a} on #{b} = {contrast_ratio(a, b):.2f}" for c, a, b in painted
                       if (a, b) in RECORDED_CONTRAST_EXCEPTIONS})
    out.append(GateResult("G-UX-04", "Contrast", "WCAG 2.2 AA on every static text cell, every conditional pair of "
                          "the dashboard and every data-sheet cell painted by a conditional font rule", inp,
                          "≥ 4.5:1 (≥ 3:1 large text); recorded exceptions only",
                          f"{checked} static cells, {len(worst)} below; dashboard pairs below: {cf_bad}; "
                          f"{len(painted)} conditionally painted cells, {len(cf_cells_bad)} below",
                          _status(not worst and not cf_bad and not cf_cells_bad),
                          "recorded exceptions (D.10, AMB-153): placeholder #808080 on the requester fills: "
                          + ", ".join(exc_seen)
                          + ("; " + "; ".join((worst + cf_cells_bad)[:3]) if worst or cf_cells_bad else "")))

    # 05 dashboard readability (reuses the automated part of G-DASH-18)
    g18 = by.get("G-DASH-18")
    out.append(GateResult("G-UX-05", "Dashboard readability", g18.test if g18 else "", inp,
                          g18.expected if g18 else "", g18.actual if g18 else "", g18.status if g18 else NV,
                          g18.evidence if g18 else ""))

    # 06 table readability
    nowrap = []
    vwb = load_workbook(x, data_only=True)
    for ws, rows in ((dash, range(1, dash.max_row + 1)), (data, range(2, data.max_row + 1))):
        vws = vwb[ws.title]
        spans = {}
        for m in ws.merged_cells.ranges:
            spans[(m.min_row, m.min_col)] = sum((ws.column_dimensions[gl(c)].width or 8.43)
                                                for c in range(m.min_col, m.max_col + 1))
        for r in rows:
            for c in ws[r]:
                v = vws.cell(row=c.row, column=c.column).value
                if not isinstance(v, str) or (c.alignment and c.alignment.wrap_text):
                    continue
                width = spans.get((c.row, c.column), ws.column_dimensions[gl(c.column)].width or 8.43)
                size = (c.font.sz or 10) / 10
                if len(v) * size > width * 1.15:          # an unwrapped text must fit its (merged) width
                    nowrap.append(f"{ws.title}!{c.coordinate}")
    capped = [r for r in range(2, data.max_row + 1) if (data.row_dimensions[r].height or 15) > 409]
    out.append(GateResult("G-UX-06", "Table readability", "every text that is not wrapped fits its (merged) width at "
                          "its font size; row heights within Excel's 409 pt", inp, "0 / 0", f"{len(nowrap)} unwrapped, {len(capped)} over 409 pt",
                          _status(not nowrap and not capped), "; ".join(nowrap[:3])))

    # 07 guidance completeness
    missing = []
    for c in range(1, pw.schema.last_col + 1):
        if c == 2:
            continue
        if data.cell(row=1, column=c).comment is None:
            missing.append(f"header {gl(c)}1")
    for coord in ("B9", "E9", "H9", "K9"):
        if dash[coord].comment is None:
            missing.append(f"selector note {coord}")
    dv_prompts = {str(dv.sqref): dv for ws in (dash, data) for dv in ws.data_validations.dataValidation}
    resp_col = gl(pw.schema.col(R, S.ROLE_RESPONSE))
    impl_col = gl(pw.schema.col(R, S.ROLE_IMPL))
    covered_resp = any(resp_col in k and dv.prompt for k, dv in dv_prompts.items())
    covered_impl = any(impl_col in k and dv.prompt == S.DV_IMPL_PROMPT for k, dv in dv_prompts.items())
    sel_prompts = sum(1 for k, dv in dv_prompts.items() if k in ("B10", "E10", "H10", "K10") and dv.prompt)
    integrity = G.validate_integrity()
    help_row = any(c.value == "دليل الاستخدام" for row in dash.iter_rows() for c in row)
    ok = not missing and covered_resp and covered_impl and sel_prompts == 4 and not integrity and help_row
    out.append(GateResult("G-UX-07", "Guidance completeness", "note on every visible header and selector; input "
                          "message on current response and status cells and selectors; dictionary integrity; help "
                          "area", inp, "all present; 0 integrity problems",
                          f"missing={missing[:4]}; response prompt={covered_resp}; status prompt={covered_impl}; "
                          f"selector prompts={sel_prompts}/4; dictionary problems={integrity}; help={help_row}",
                          _status(ok)))

    # 08 hover (notes attached to the right cells; platform behaviour documented)
    with zipfile.ZipFile(x) as z:
        comment_parts = [n for n in z.namelist() if re.match(r"xl/comments/comment\d+\.xml", n) or
                         re.match(r"xl/comments\d+\.xml", n)]
        refs = []
        for n in comment_parts:
            root = etree.fromstring(z.read(n))
            refs += [c.get("ref") for c in root.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}comment")]
    expected_refs = {f"{gl(c)}1" for c in range(1, pw.schema.last_col + 1) if c != 2}
    ok = expected_refs <= set(refs)
    out.append(GateResult("G-UX-08", "Hover functionality", "legacy Notes present at the header and selector "
                          "cells (comments part + VML); platform behaviour documented", inp,
                          f"{len(expected_refs)} data-sheet header notes + dashboard notes",
                          f"{len(refs)} notes in {len(comment_parts)} parts; all headers covered={ok}", _status(ok),
                          "Excel desktop shows Notes on hover; Excel for the web and mobile show them on tap or not at "
                          "all; data-validation messages appear on selection, not on hover (documented in the guide)"))

    # 09 input validation
    lists = {}
    for dv in data.data_validations.dataValidation:
        if dv.type == "list":
            lists[dv.formula1] = str(dv.sqref)
    closure_list = '"' + ",".join(S.CLOSURE_VALUES) + '"'
    impl_list = '"' + ",".join(S.IMPL_VALUES) + '"'
    last = data.max_row
    need_c = [f"{gl(pw.schema.col(r, S.ROLE_CLOSURE))}2:{gl(pw.schema.col(r, S.ROLE_CLOSURE))}{last}" for r in range(1, R + 1)]
    need_i = [f"{gl(pw.schema.col(r, S.ROLE_IMPL))}2:{gl(pw.schema.col(r, S.ROLE_IMPL))}{last}" for r in range(1, R + 1)]
    ok = all(n in lists.get(closure_list, "") for n in need_c) and all(n in lists.get(impl_list, "") for n in need_i)
    out.append(GateResult("G-UX-09", "Input validation", "exact inline lists on every data row of every closure "
                          "and requester-status column (10.3)", inp, f"{R} closure + {R} status ranges",
                          f"closure ranges={lists.get(closure_list, '')}; status ranges={lists.get(impl_list, '')}",
                          _status(ok)))

    for gid, name, src in (("G-UX-10", "Historical integrity", "G-DASH-02"), ("G-UX-11", "Formula integrity", "G-DASH-16"),
                           ("G-UX-12", "Navigation", "G-DASH-06"), ("G-UX-13", "Semantic accuracy", "G-DASH-19"),
                           ("G-UX-15", "Compatibility", "G-DASH-13")):
        g = by[src]
        status = g.status
        if gid == "G-UX-13":
            g10 = by["G-DASH-10"]
            status = PASS if (g.status == PASS and g10.status == PASS) else FAIL
        out.append(GateResult(gid, name, f"same check as {src}" + (" and G-DASH-10" if gid == "G-UX-13" else ""),
                              g.input, g.expected, g.actual, status, g.evidence))

    # 14 visual regression: charts do not cover filled cells + manual inspection
    overlaps = []
    for ch_anchor in _chart_rects(x, W.DASHBOARD_SHEET):
        c1, r1, c2, r2 = ch_anchor
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                v = dash.cell(row=r, column=c).value
                if v not in (None, ""):
                    overlaps.append(f"{gl(c)}{r}")
    man = ctx.manual.get("G-UX-14")
    st_ = (MANUAL if man else NV) if not overlaps else FAIL
    out.append(GateResult("G-UX-14", "Visual regression", "chart rectangles contain no cell content; rendered pages "
                          "inspected for overlap and clipping", inp, "0 overlaps; inspection recorded",
                          f"{len(overlaps)} cells under charts", st_, man[1] if man else "; ".join(overlaps[:5])))

    # 16 performance
    if ctx.perf:
        p = ctx.perf
        ok = p["build_seconds"] < 120 and p["file_mb"] < 25
        out.append(GateResult("G-UX-16", "Performance", "build of a large register (observations × rounds)",
                              f"{p['observations']} observations × {p['rounds']} rounds",
                              "build < 120 s, file < 25 MB, bounded ranges",
                              f"build {p['build_seconds']:.1f} s, {p['file_mb']:.1f} MB, {p['formulas']} formulas"
                              + (f", LibreOffice recalculation {p['recalc_seconds']:.1f} s" if p.get('recalc_seconds') else ""),
                              _status(ok)))
    else:
        out.append(GateResult("G-UX-16", "Performance", "not run", inp, "", "", NV))

    # 17 information security
    texts = I.all_text_parts(x)
    sheet_txt = " ".join(v for k, v in texts.items() if k.startswith("xl/worksheets/") or k.startswith("xl/sharedStrings")
                         or k.startswith("xl/comments") or k.startswith("xl/charts"))
    patterns = {
        "severity codes": r"\b(CRITICAL|EDITORIAL)\b|Severity",
        "integrity labels": r"\bI-\d\d\b|INTEGRITY_|سلامة الردود",
        "closure-path codes": r"VERIFIED_TEXT|GCO_REASSESSMENT|ROUTED_RESIDUAL|RISK_ACCEPTED",
        "hashes": r"\b[0-9a-f]{64}\b",
        "Arabic severity words": r"(?<![\w])(حرجة|عالية الخطورة)(?![\w])",
    }
    hits = {k: len(re.findall(p, sheet_txt)) for k, p in patterns.items()}
    vis_ids = [c.coordinate for row in load_workbook(x, data_only=True)[W.DASHBOARD_SHEET].iter_rows() for c in row
               if isinstance(c.value, str) and re.search(r"GCO-[A-Z]{2,4}-R\d+", c.value)]
    vwb17 = load_workbook(x, data_only=True)
    code = pw.state.document.code
    code_cells = [f"{ws.title}!{c.coordinate}" for ws in vwb17.worksheets for row in ws.iter_rows() for c in row
                  if isinstance(c.value, str) and c.value.strip() == code]
    banner = any(c.value == S.DEMO_BANNER for row in vwb17[W.DASHBOARD_SHEET].iter_rows() for c in row)
    demo_ok = banner == bool(pw.state.document.is_demo)
    ok = not any(hits.values()) and not vis_ids and not code_cells and demo_ok
    out.append(GateResult("G-UX-17", "Information security", "allowlist scan of every sheet, note and chart part; "
                          "no Internal_ID visible on the dashboard; no document code as a value (D1); demonstration "
                          "band present exactly when the data are synthetic (AMB-154)", inp, "0 hits",
                          f"{hits}; Internal_IDs visible on dashboard: {len(vis_ids)}; document code cells: "
                          f"{code_cells[:3]}; synthetic={bool(pw.state.document.is_demo)}, band={banner}", _status(ok),
                          "Internal_IDs exist only in hidden column B and the hidden helper sheet (D4)"))

    # 18 user experience
    vals = load_workbook(x, data_only=True)[W.DASHBOARD_SHEET]
    next_label = dash["J5"].value
    next_val = vals["J6"].value
    nav_input = [c for c in dash[8] if c.hyperlink is not None and c.hyperlink.location
                 and c.hyperlink.location.startswith(f"'{pw.data_ws.title}'!{resp_col}")]
    ok = bool(next_label) and bool(next_val) and bool(nav_input) and help_row
    out.append(GateResult("G-UX-18", "User experience", "the next required action is stated on the first screen; "
                          "one click reaches the current-round input column; help present", inp,
                          "stated; link; help", f"'{next_label}': '{next_val}'; input link={bool(nav_input)}",
                          _status(ok)))

    # 19 design consistency
    sizes = {(dash.cell(row=14, column=c).font.sz, dash.cell(row=13, column=c).font.sz) for c in (2, 4, 6, 8, 10, 12)}
    merged = {str(m) for m in dash.merged_cells.ranges}
    cards = all(f"{gl(c)}14:{gl(c + 1)}14" in merged for c in (2, 4, 6, 8, 10, 12))
    out.append(GateResult("G-UX-19", "Design consistency", "six KPI cards with identical size and type scale",
                          inp, "1 size pair; 6 equal merges", f"size pairs={sizes}; equal cards={cards}",
                          _status(len(sizes) == 1 and cards)))

    # 20 workbook integrity
    with zipfile.ZipFile(x) as z:
        missing_cache = 0
        total_f = 0
        for n in z.namelist():
            if n.startswith("xl/worksheets/sheet"):
                root = etree.fromstring(z.read(n))
                for c in root.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"):
                    if c.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f") is not None:
                        total_f += 1
                        if c.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v") is None:
                            missing_cache += 1
    xsd = by["G-DASH-13"].actual.split(";")[0]
    ok = missing_cache == 0 and "XSD PASS" in xsd
    out.append(GateResult("G-UX-20", "Workbook integrity", "reopen; every formula carries a cached value; XSD",
                          inp, "0 formulas without cache; XSD PASS", f"{total_f} formulas, {missing_cache} without "
                          f"cache; {xsd}", _status(ok)))
    return sorted(out, key=lambda g: g.gate)


def _chart_rects(path, sheet_name):
    """(col1, row1, col2, row2) 1-based cell rectangles covered by each chart anchor on a sheet."""
    pkg = I.package(path)
    out = []
    wb = load_workbook(path)
    ws = wb[sheet_name]
    with zipfile.ZipFile(path) as z:
        part = [p for n, s, p in pkg.sheets if n == sheet_name][0]
        rels = I._rels(z, part)
        for rid, (typ, target) in rels.items():
            if typ != "drawing":
                continue
            root = etree.fromstring(z.read(target))
            for anchor in root:
                if anchor.find(".//c:chart", I.NS) is None:
                    continue
                frm = anchor.find("xdr:from", I.NS)
                ext = anchor.find("xdr:ext", I.NS)
                c0 = int(frm.find("xdr:col", I.NS).text) + 1
                r0 = int(frm.find("xdr:row", I.NS).text) + 1
                w_px = int(ext.get("cx")) / 9525
                h_px = int(ext.get("cy")) / 9525
                c, acc = c0, 0.0
                while acc + (ws.column_dimensions[gl(c)].width or 8.43) * 7 + 5 <= w_px + 2:
                    acc += (ws.column_dimensions[gl(c)].width or 8.43) * 7 + 5
                    c += 1
                r, acc = r0, 0.0
                while acc + (ws.row_dimensions[r].height or 15) * 96 / 72 <= h_px + 2:
                    acc += (ws.row_dimensions[r].height or 15) * 96 / 72
                    r += 1
                out.append((c0, r0, max(c0, c - 1), max(r0, r - 1)))
    return out
