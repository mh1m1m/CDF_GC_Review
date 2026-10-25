"""Read-back inspection of a delivered workbook (directive §11.3).

Every check here reads the saved file (openpyxl for cells and styles, the
OOXML parts for charts, drawings, notes and properties); nothing is taken
from the in-memory build.
"""
from __future__ import annotations

import os
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook

from ..config import workbook_settings as W

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    "dc": "http://purl.org/dc/elements/1.1/",
    "ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
}


@dataclass
class ChartInfo:
    part: str
    kinds: list[str]
    grouping: str | None
    series_refs: list[str]
    series_names: list[str]
    category_refs: list[str]
    sheet: str | None = None


@dataclass
class PackageInfo:
    path: Path
    parts: list[str]
    sheets: list[tuple[str, str, str]]          # (name, state, part)
    active_tab: int
    charts: list[ChartInfo] = field(default_factory=list)
    images: list[dict] = field(default_factory=list)
    core: dict = field(default_factory=dict)
    app: dict = field(default_factory=dict)
    calc_full_on_load: bool = False


def _rels(z, part):
    d, f = os.path.split(part)
    rp = f"{d}/_rels/{f}.rels"
    if rp not in z.namelist():
        return {}
    root = etree.fromstring(z.read(rp))
    out = {}
    for r in root:
        t = r.get("Target")
        if not t.startswith("/"):
            t = os.path.normpath(os.path.join(d, t)).replace("\\", "/")
        else:
            t = t.lstrip("/")
        out[r.get("Id")] = (r.get("Type").rsplit("/", 1)[-1], t)
    return out


def package(path) -> PackageInfo:
    path = Path(path)
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        wbx = etree.fromstring(z.read("xl/workbook.xml"))
        wrels = _rels(z, "xl/workbook.xml")
        sheets = []
        for sh in wbx.findall("m:sheets/m:sheet", NS):
            rid = sh.get(f"{{{NS['r']}}}id")
            sheets.append((sh.get("name"), sh.get("state") or "visible", wrels[rid][1]))
        bv = wbx.find("m:bookViews/m:workbookView", NS)
        active = int(bv.get("activeTab") or 0) if bv is not None else 0
        calc = wbx.find("m:calcPr", NS)
        info = PackageInfo(path, names, sheets, active,
                           calc_full_on_load=calc is not None and calc.get("fullCalcOnLoad") in ("1", "true"))
        for name, _, part in sheets:
            for rid, (typ, target) in _rels(z, part).items():
                if typ != "drawing":
                    continue
                drel = _rels(z, target)
                droot = etree.fromstring(z.read(target))
                for anchor in list(droot):
                    for gf in anchor.iter(f"{{{NS['c']}}}chart"):
                        cpart = drel[gf.get(f"{{{NS['r']}}}id")][1]
                        info.charts.append(_chart(z, cpart, name))
                    for pic in anchor.iter(f"{{{NS['xdr']}}}pic"):
                        blip = pic.find(".//a:blip", NS)
                        ext = anchor.find("xdr:ext", NS)
                        frm = anchor.find("xdr:from", NS)
                        img_part = drel[blip.get(f"{{{NS['r']}}}embed")][1]
                        info.images.append({"sheet": name, "part": img_part,
                                            "cx": int(ext.get("cx")) if ext is not None else None,
                                            "cy": int(ext.get("cy")) if ext is not None else None,
                                            "col": int(frm.find("xdr:col", NS).text),
                                            "row": int(frm.find("xdr:row", NS).text),
                                            "bytes": len(z.read(img_part))})
        core = etree.fromstring(z.read("docProps/core.xml"))
        for tag, ns in (("creator", "dc"), ("title", "dc"), ("lastModifiedBy", "cp"), ("keywords", "cp")):
            el = core.find(f"{ns}:{tag}", NS)
            info.core[tag] = el.text if el is not None else None
        app = etree.fromstring(z.read("docProps/app.xml"))
        for el in app:
            info.app[etree.QName(el).localname] = el.text
    return info


def _chart(z, part, sheet) -> ChartInfo:
    root = etree.fromstring(z.read(part))
    plot = root.find(".//c:plotArea", NS)
    kinds, grouping = [], None
    for el in plot:
        ln = etree.QName(el).localname
        if ln.endswith("Chart"):
            kinds.append(ln)
            g = el.find("c:grouping", NS)
            if g is not None:
                grouping = g.get("val")
            bd = el.find("c:barDir", NS)
            if bd is not None:
                kinds[-1] += f"/{bd.get('val')}"
    sref, names, cref = [], [], []
    for ser in plot.iter(f"{{{NS['c']}}}ser"):
        v = ser.find("c:val/c:numRef/c:f", NS)
        sref.append(v.text if v is not None else "")
        tx = ser.find("c:tx/c:v", NS)
        if tx is None:
            tx = ser.find("c:tx/c:strRef/c:f", NS)
        names.append(tx.text if tx is not None else "")
        c = ser.find("c:cat/c:numRef/c:f", NS)
        if c is None:
            c = ser.find("c:cat/c:strRef/c:f", NS)
        cref.append(c.text if c is not None else "")
    return ChartInfo(part, kinds, grouping, sref, names, cref, sheet)


def all_text_parts(path) -> dict[str, str]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n).decode("utf-8", "replace") for n in z.namelist()
                if n.endswith((".xml", ".rels", ".vml"))}


def xsd_validate(path, schema_dir: str | None = None) -> tuple[str, list[str]]:
    """Validate worksheet, workbook, styles, drawing and chart parts against the ISO 29500 transitional XSDs.

    Returns (status, messages); status is NOT_RUN when no schema directory is available.
    """
    schema_dir = schema_dir or os.environ.get("CRMS_OOXML_XSD_DIR") or \
        "/mnt/skills/public/xlsx/scripts/office/schemas/ISO-IEC29500-4_2016"
    sml = Path(schema_dir) / "sml.xsd"
    if not sml.exists():
        return "NOT_RUN", [f"schema directory not available: {schema_dir}"]
    schemas = {}
    for key, fn in (("sml", "sml.xsd"), ("chart", "dml-chart.xsd"), ("drawing", "dml-spreadsheetDrawing.xsd")):
        try:
            schemas[key] = etree.XMLSchema(etree.parse(str(Path(schema_dir) / fn)))
        except Exception as e:  # noqa: BLE001
            return "NOT_RUN", [f"cannot load {fn}: {e}"]
    msgs = []
    checked = 0
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("xl/worksheets/sheet") or n in ("xl/workbook.xml", "xl/styles.xml"):
                key = "sml"
            elif n.startswith("xl/charts/chart"):
                key = "chart"
            elif n.startswith("xl/drawings/drawing") and n.endswith(".xml"):
                key = "drawing"
            else:
                continue
            doc = etree.fromstring(z.read(n))
            checked += 1
            if not schemas[key].validate(doc):
                for err in list(schemas[key].error_log):
                    # Excel itself writes xml:space="preserve" on <t>, and protocol 10.9 requires it; the ISO
                    # transitional XSD does not declare the attribute. Known schema gap, not a defect.
                    if "XML/1998/namespace}space" in err.message and "}t'" in err.message:
                        continue
                    msgs.append(f"{n}: line {err.line}: {err.message}")
                    if len(msgs) > 40:
                        break
    return ("PASS" if not msgs else "FAIL"), [f"{checked} parts validated"] + msgs


def cell_fonts(ws) -> set[str]:
    out = set()
    for row in ws.iter_rows():
        for c in row:
            if c.has_style and c.font is not None and c.font.name:
                out.add(c.font.name)
            v = c.value
            if hasattr(v, "__iter__") and not isinstance(v, str):
                for part in v:
                    f = getattr(part, "font", None)
                    if f is not None and getattr(f, "rFont", None):
                        out.add(f.rFont)
    return out


def min_font_size(ws) -> float:
    m = 99.0
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, "") and c.font is not None and c.font.sz:
                m = min(m, float(c.font.sz))
    return m


def hyperlink_target(value: str) -> tuple[str, str] | None:
    """'#'sheet'!G12' -> (sheet, G12)."""
    m = re.match(r"^#'(?P<s>(?:[^']|'')+)'!(?P<c>\$?[A-Z]+\$?\d+)$", value or "")
    if not m:
        return None
    return m.group("s").replace("''", "'"), m.group("c").replace("$", "")
