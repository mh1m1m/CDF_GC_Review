"""CLI -> independent MCP SDK client -> rendered React -> actual document engines.

No paid credentials, private files, direct MCP function calls, or fabricated
remote-client receipts. A fresh private temporary SQLite database is shared by
all processes; only synthetic, credential-free evidence leaves that directory.
"""

from __future__ import annotations

import argparse
import asyncio
import base64
from contextlib import contextmanager
import hashlib
from io import BytesIO
import json
import math
import os
from pathlib import Path
import re
import secrets
import socket
import subprocess
import sys
import tempfile
import time
import zipfile

import httpx
from mcp import ClientSession
from mcp.client.streamable_http import streamable_http_client

ROOT = Path(__file__).resolve().parents[2]
CASE = "SYN-COOPERATION-001"
TERMINAL = {
    "COMPLETED",
    "WAITING_FOR_HUMAN",
    "FAILED",
    "CANCELLED",
    "BLOCKED_AUTHORIZATION",
}


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2))


def free_port():
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


@contextmanager
def serving(command, env, health, log):
    with log.open("w") as output:
        process = subprocess.Popen(
            command,
            cwd=ROOT,
            env=env,
            stdin=subprocess.DEVNULL,
            stdout=output,
            stderr=output,
            start_new_session=True,
        )
        try:
            for _ in range(150):
                try:
                    if httpx.get(health, timeout=1).status_code == 200:
                        break
                except httpx.HTTPError:
                    pass
                if process.poll() is not None:
                    raise RuntimeError("ACCEPTANCE_SERVER_EXITED: " + log.name)
                time.sleep(0.1)
            else:
                raise RuntimeError("ACCEPTANCE_SERVER_START_TIMEOUT")
            yield process
        finally:
            from ..execution.jobs import kill_tree

            kill_tree(process.pid)
            process.wait(timeout=5)


def stable_fields(view):
    return {
        key: view[key]
        for key in (
            "case_id",
            "run_id",
            "observation_ids",
            "agent_status",
            "evidence_references",
            "calibration_version",
            "release_gate",
        )
    } | {
        "metrics": [
            {
                "id": m["id"],
                "numerator": m["numerator"],
                "denominator": m["denominator"],
            }
            for m in view["metrics"]
        ],
        "artifact_ids": [a["artifact_id"] for a in view["artifacts"]],
    }


class Acceptance:
    def __init__(self, output, env):
        self.output, self.env = output, env
        self.commands, self.invocations, self.security = [], [], []
        self.artifacts = {}

    def cli(self, *args, expected=0):
        command = [sys.executable, "-m", "cdf_gco_crms.cli", *args, "--json"]
        result = subprocess.run(
            command, cwd=ROOT, env=self.env, capture_output=True, text=True, timeout=260
        )
        record = {"command": command, "exit_status": result.returncode}
        try:
            payload = json.loads(result.stdout)
        except ValueError:
            raise AssertionError("CLI_NOT_JSON: " + str(args)) from None
        record["response"] = payload
        self.commands.append(record)
        assert result.returncode == expected, record
        return payload.get("data", payload)

    def phase_cli(self, exports):
        self.cli("init")
        created = self.cli(
            "case",
            "create",
            "--case",
            CASE,
            "--source",
            "synthetic",
            "--idempotency-key",
            "accept-create-001",
        )
        self.source_hash = created["source_hash"]
        self.cli(
            "case", "import", "--case", CASE
        )  # creation atomically imported the fixed DOCX
        self.review = self.cli(
            "review",
            "run",
            "--case",
            CASE,
            "--executor",
            "deterministic",
            "--wait",
            "--timeout",
            "120",
            "--idempotency-key",
            "accept-review-001",
        )
        assert self.review["job"]["status"] == "WAITING_FOR_HUMAN"
        self.cli("review", "observations", "--case", CASE)
        self.debate = self.cli(
            "debate",
            "run",
            "--case",
            CASE,
            "--executor",
            "mock",
            "--wait",
            "--timeout",
            "120",
            "--idempotency-key",
            "accept-debate-001",
        )
        assert self.debate["job"]["status"] == "WAITING_FOR_HUMAN"
        self.trace = self.cli("debate", "trace", "--case", CASE)
        self.authority = self.cli(
            "authority", "check", "--case", CASE, "--query", "تعاون"
        )
        self.organization = self.cli(
            "organization", "check", "--case", CASE, "--query", "تعاون"
        )
        assert self.authority["authorities"] and self.organization["responsibilities"]
        self.cli("calibration", "list", "--case", CASE)
        view = self.cli("case", "status", "--case", CASE)
        self.calibration = self.cli(
            "calibration",
            "save",
            "--case",
            CASE,
            "--observation",
            view["observation_ids"][0],
            "--decision",
            "KEEP_UNCHANGED",
            "--rationale",
            "Synthetic local reviewer retained the source-grounded finding; no institutional signature.",
            "--evidence-hash",
            view["evidence_hash"],
            "--idempotency-key",
            "accept-calibrate-001",
        )
        self.cli("gate", "--case", CASE)
        for export in exports:
            kind = {"word": "docx", "excel": "xlsx", "pdf": "pdf"}[export]
            job = self.cli(
                "report",
                export,
                "--case",
                CASE,
                "--wait",
                "--timeout",
                "180",
                "--idempotency-key",
                "accept-export-" + export,
            )
            assert job["job"]["status"] == "COMPLETED", job
            aid = job["result"]["artifact_id"]
            destination = (
                self.output
                / "synthetic-reports"
                / ("INTERNAL_DRAFT__executive." + kind)
            )
            metadata = self.cli(
                "artifact",
                "get",
                "--case",
                CASE,
                "--artifact",
                aid,
                "--output",
                str(destination),
            )
            self.artifacts[kind] = metadata
        self.cli_view = self.cli("case", "status", "--case", CASE)
        self.canonical = stable_fields(self.cli_view)
        assert (
            len(self.canonical["observation_ids"]) == 8
            and not self.canonical["release_gate"]["authorized"]
        )
        assert self.cli_view["actual_live_inference_count"] == 0
        assert len({a["stage"] for a in self.canonical["agent_status"]}) >= 8

    async def phase_mcp(self, url):
        async with httpx.AsyncClient() as client:
            invalid = await client.post(
                url,
                headers={"Authorization": "Bearer invalid-synthetic-credential"},
                json={},
            )
            assert invalid.status_code == 401
            self.security.append(
                {
                    "test": "MCP_INVALID_AUTH",
                    "status": "PASS",
                    "http_status": invalid.status_code,
                }
            )
        async with httpx.AsyncClient(
            headers={
                "Authorization": "Bearer " + self.env["CDF_CALIBRATION_DEV_TOKEN"]
            },
            timeout=60,
        ) as authenticated_http, streamable_http_client(
            url, http_client=authenticated_http
        ) as (
            read,
            send,
            _,
        ):
            async with ClientSession(read, send) as client:
                initialized = await client.initialize()
                discovery = await client.list_tools()
                from ..execution.contracts import TOOLS

                assert set(TOOLS) == {t.name for t in discovery.tools}
                assert all(t.inputSchema and t.outputSchema for t in discovery.tools)
                write(
                    self.output / "mcp-tool-discovery.json",
                    {
                        "initialization": initialized.model_dump(mode="json"),
                        "tools": [t.model_dump(mode="json") for t in discovery.tools],
                    },
                )

                async def call(tool, arguments, error=False):
                    response = await client.call_tool(tool, {"request": arguments})
                    payload = response.structuredContent
                    self.invocations.append(
                        {
                            "tool": tool,
                            "arguments": arguments,
                            "is_error": bool(response.isError),
                            "response": payload,
                        }
                    )
                    assert bool(response.isError) == error, self.invocations[-1]
                    if error:
                        return payload
                    assert payload["ok"]
                    return payload["data"]

                view = await call("cdf_get_case_status", {"case_id": CASE})
                self.mcp_fields = stable_fields(view)
                assert self.mcp_fields == self.canonical
                created = await call("cdf_create_case", {"case_id": CASE, "source": "synthetic-cooperation-v1", "idempotency_key": "accept-create-001"})
                assert created["replayed"] and created["source_hash"] == self.source_hash
                for tool in (
                    "cdf_get_observations",
                    "cdf_get_debate_trace",
                    "cdf_get_dashboard_metrics",
                    "cdf_get_release_gate_status",
                    "cdf_get_calibration_history",
                    "cdf_check_delegated_authority",
                    "cdf_get_department_responsibilities",
                    "cdf_get_reference_status",
                ):
                    await call(tool, {"case_id": CASE})
                # Replay the accepted CLI operation; one job ID, one backend attempt, no provider replay.
                replay = await call(
                    "cdf_review_document",
                    {
                        "case_id": CASE,
                        "executor": "deterministic",
                        "timeout": 120,
                        "idempotency_key": "accept-review-001",
                    },
                )
                assert (
                    replay["replayed"]
                    and replay["job_id"] == self.review["job"]["job_id"]
                    and replay["attempts"] == 1
                )
                exported = (
                    await call(
                        "cdf_generate_excel_dashboard",
                        {
                            "case_id": CASE,
                            "idempotency_key": "accept-export-excel",
                            "timeout": 180,
                        },
                    )
                    if "xlsx" in self.artifacts
                    else None
                )
                if exported:
                    polled = await call(
                        "cdf_get_job_status",
                        {"case_id": CASE, "job_id": exported["job_id"]},
                    )
                    result = await call(
                        "cdf_get_job_result",
                        {"case_id": CASE, "job_id": exported["job_id"]},
                    )
                    assert (
                        polled["status"] == "COMPLETED"
                        and result["result"]["artifact_id"]
                        == self.artifacts["xlsx"]["artifact_id"]
                    )
                # A distinct MCP mutation is a draft calibration decision, excluded from submitted metric version.
                await call(
                    "cdf_calibrate_observation",
                    {
                        "case_id": CASE,
                        "observation_id": self.canonical["observation_ids"][1],
                        "decision": "KEEP_UNCHANGED",
                        "rationale": "Synthetic MCP draft only; no submitted human approval.",
                        "status": "DRAFT",
                        "expected_evidence_hash": view["evidence_hash"],
                        "idempotency_key": "accept-mcp-draft-001",
                    },
                )
                denied = await call(
                    "cdf_get_case_status", {"case_id": "SYN-NOT-GRANTED"}, error=True
                )
                self.security.append(
                    {
                        "test": "MCP_CASE_SCOPE",
                        "status": "PASS",
                        "error": denied["error"],
                    }
                )
                for label, tool, args in [
                    (
                        "MCP_FORGED_IDENTITY",
                        "cdf_get_case_status",
                        {"case_id": CASE, "reviewer_id": "FORGED"},
                    ),
                    (
                        "MCP_ISSUE_WITHOUT_AUTHORIZATION",
                        "cdf_generate_word_report",
                        {"case_id": CASE, "issue": True},
                    ),
                    (
                        "MCP_MATERIAL_WITHDRAWAL",
                        "cdf_calibrate_observation",
                        {
                            "case_id": CASE,
                            "observation_id": self.canonical["observation_ids"][-1],
                            "decision": "WITHDRAW",
                            "rationale": "Synthetic unsupported withdrawal.",
                            "expected_evidence_hash": view["evidence_hash"],
                        },
                    ),
                ]:
                    response = await client.call_tool(tool, {"request": args})
                    assert response.isError, label
                    self.security.append({"test": label, "status": "PASS"})
                for kind, meta in self.artifacts.items():
                    await call(
                        "cdf_get_artifact_status",
                        {"case_id": CASE, "artifact_id": meta["artifact_id"]},
                    )
                    content, offset = bytearray(), 0
                    while True:
                        part = await call(
                            "cdf_get_artifact",
                            {
                                "case_id": CASE,
                                "artifact_id": meta["artifact_id"],
                                "cursor": offset,
                            },
                        )
                        content.extend(base64.b64decode(part["content_base64"]))
                        # Avoid duplicating document bytes in the human-readable invocation log.
                        self.invocations[-1]["response"]["data"][
                            "content_base64"
                        ] = "[OMITTED; SHA-256 VERIFIED]"
                        if part["next_offset"] is None:
                            break
                        offset = part["next_offset"]
                    source = (
                        self.output
                        / "synthetic-reports"
                        / ("INTERNAL_DRAFT__executive." + kind)
                    )
                    assert bytes(content) == source.read_bytes()
                    assert hashlib.sha256(content).hexdigest() == meta["sha256"]
                # Refresh all interfaces after the actual MCP mutation.
                self.mcp_fields = stable_fields(
                    await call("cdf_get_case_status", {"case_id": CASE})
                )
        self.cli_view = self.cli("case", "status", "--case", CASE)
        self.canonical = stable_fields(self.cli_view)
        assert self.mcp_fields == self.canonical

    def phase_react(self, workspace):
        api_port, web_port = free_port(), free_port()
        env = {
            **self.env,
            "CDF_API_ORIGIN": f"http://127.0.0.1:{api_port}",
            "CDF_WEB_ORIGIN": f"http://127.0.0.1:{web_port}",
            "CDF_WEB_PORT": str(web_port),
            "CDF_UNIFIED_CASE_ID": CASE,
            "CDF_UNIFIED_REPORT_DIR": str(self.output),
            "PLAYWRIGHT_JUNIT_OUTPUT_FILE": str(self.output / "test-results.xml"),
        }
        write(self.output / "expected-interface.json", self.canonical)
        write(self.output / "expected-artifacts.json", self.artifacts)
        with serving(
            [
                sys.executable,
                "-m",
                "uvicorn",
                "apps.api.main:app",
                "--host",
                "127.0.0.1",
                "--port",
                str(api_port),
            ],
            env,
            env["CDF_API_ORIGIN"] + "/api/health",
            workspace / "api.log",
        ):
            command = [
                "node",
                "node_modules/@playwright/test/cli.js",
                "test",
                "e2e/unified.spec.ts",
                "--reporter=json,junit",
                "--workers=1",
            ]
            result = subprocess.run(
                command,
                cwd=ROOT / "apps/web",
                env=env,
                capture_output=True,
                text=True,
                timeout=180,
            )
            (self.output / "browser-results.json").write_text(result.stdout)
            (self.output / "browser-stderr.log").write_text(result.stderr.replace(self.env["CDF_CALIBRATION_DEV_TOKEN"], "[REDACTED]"))
            self.commands.append({"command": command, "exit_status": result.returncode})
            assert (
                result.returncode == 0
            ), "REACT_ACCEPTANCE_FAILED (see browser-results.json)"
            report = json.loads(result.stdout)
            assert (
                report["stats"]["unexpected"] == 0
                and report["stats"]["skipped"] == 0
                and report["stats"]["expected"] >= 1
            )
        return json.loads((self.output / "react-interface.json").read_text())

    def validate_exports(self, exports):
        from docx import Document
        from lxml import etree
        from openpyxl import load_workbook
        import pymupdf
        from ..qa import lo_engine

        results = {}
        reports = self.output / "synthetic-reports"
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        if "word" in exports:
            content = (reports / "INTERNAL_DRAFT__executive.docx").read_bytes()
            with zipfile.ZipFile(BytesIO(content)) as package:
                root = etree.fromstring(
                    package.read("word/document.xml"),
                    etree.XMLParser(resolve_entities=False, no_network=True),
                )
                text = " ".join(root.xpath("//w:t/text()", namespaces=ns))
                assert (
                    CASE in text
                    and self.canonical["run_id"] in text
                    and all(oid in text for oid in self.canonical["observation_ids"])
                )
                assert (
                    "التوصية:" in text
                    and "التصنيف المادي" in text
                    and "SYN-DOA-COOP" in text
                    and "SYN-ORG-1" in text
                )
                assert root.xpath("//w:bidi", namespaces=ns) and root.xpath(
                    "//w:tbl", namespaces=ns
                )
                assert "INTERNAL DRAFT" in text and "NOT AUTHORIZED FOR ISSUE" in text
            assert len(Document(BytesIO(content)).tables) >= 3
            results["word"] = {
                "status": "PASS",
                "editable_text": True,
                "rtl": True,
                "all_observation_ids": True,
            }
        if "excel" in exports:
            path = reports / "INTERNAL_DRAFT__executive.xlsx"
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            workbook = load_workbook(path, data_only=False)
            assert workbook["Executive"]["B1"].value == CASE
            assert [
                r[0].value for r in list(workbook["Observations"].rows)[1:]
            ] == self.canonical["observation_ids"]
            assert (
                workbook["Observations"].max_row == 9
                and len(workbook["Executive"]._charts) >= 1
            )
            assert all(
                workbook["Executive"].cell(i, 4).data_type == "f" for i in range(3, 11)
            )
            assert all(workbook["Observations"].cell(i, 6).value for i in range(2, 10))
            assert self.canonical["run_id"] in [
                str(c.value) for row in workbook["Provenance"] for c in row
            ]
            assert lo_engine.available() and not re.search(
                r"alpha|beta|dev|nightly", lo_engine.version(), re.I
            ), "STABLE_CALC_REQUIRED"
            with tempfile.TemporaryDirectory(prefix="cdf-accept-calc-") as temp:
                copy = Path(temp) / "recalculated.xlsx"
                copy.write_bytes(path.read_bytes())
                recalculated = lo_engine.recalc(copy, timeout=120)
                assert recalculated.ok, recalculated.message
                data = load_workbook(copy, data_only=True)
                fractions = []
                for index, metric in enumerate(self.canonical["metrics"], 3):
                    actual = data["Executive"].cell(index, 4).value
                    expected = (
                        metric["numerator"] / metric["denominator"]
                        if metric["numerator"] is not None and metric["denominator"]
                        else None
                    )
                    assert (
                        actual in ("", None)
                        if expected is None
                        else math.isclose(actual, expected, abs_tol=1e-9)
                    )
                    fractions.append(
                        {"id": metric["id"], "actual": actual, "expected": expected}
                    )
                preview = lo_engine.to_pdf(path, Path(temp) / "preview", timeout=120)
                assert preview is not None
                with pymupdf.open(preview) as pdf:
                    pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.3, 1.3)).save(
                        self.output / "screenshots" / "excel-preview.png"
                    )
            assert hashlib.sha256(path.read_bytes()).hexdigest() == before
            results["excel"] = {
                "status": "PASS",
                "engine": lo_engine.version(),
                "source_unchanged": True,
                "native_charts": True,
                "formula_fractions": fractions,
                "limitation": "Calc validates formulas; Microsoft Excel not executed. Original x14 sparkline may be dropped by other editors.",
            }
        if "pdf" in exports:
            path = reports / "INTERNAL_DRAFT__executive.pdf"
            assert path.read_bytes().startswith(b"%PDF-")
            with pymupdf.open(path) as pdf:
                text = "\n".join(p.get_text() for p in pdf)
                # PDF bidi extraction may split ASCII identifiers; normalize whitespace.
                assert CASE in re.sub(r"\s+", "", text)
                assert self.canonical["run_id"] in re.sub(r"\s+", "", text)
                assert pdf.page_count >= 2
                pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.3, 1.3)).save(
                    self.output / "screenshots" / "word-preview.png"
                )
                results["pdf"] = {
                    "status": "PASS",
                    "pages": pdf.page_count,
                    "case_and_run_present": True,
                    "rendered": True,
                }
        write(self.output / "export-validation.json", results)
        return results


def scan_evidence(output, token):
    """Heuristic leakage gate: source-safe fixture, no secret bytes/session hash in logs/documents."""
    from docx import Document
    from openpyxl import load_workbook

    patterns = [
        token.encode(),
        hashlib.sha256(token.encode()).hexdigest().encode(),
        b"-----BEGIN PRIVATE KEY",
        b"ghp_",
        b"sk-ant-",
    ]
    scanned = []
    for path in output.rglob("*"):
        if not path.is_file():
            continue
        contents = path.read_bytes()
        assert not any(p in contents for p in patterns), "EVIDENCE_CREDENTIAL_LEAK"
        if path.suffix in (".docx", ".xlsx"):
            with zipfile.ZipFile(path) as package:
                for item in package.namelist():
                    if item.endswith((".xml", ".rels")):
                        assert not any(
                            p in package.read(item) for p in patterns
                        ), "DOCUMENT_CREDENTIAL_LEAK"
        scanned.append(
            {
                "path": str(path.relative_to(output)),
                "sha256": hashlib.sha256(contents).hexdigest(),
                "bytes": len(contents),
            }
        )
    return scanned


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--mode", choices=["public-synthetic"], default="public-synthetic"
    )
    parser.add_argument("--interfaces", default="cli,mcp,react")
    parser.add_argument("--exports", default="word,excel,pdf")
    parser.add_argument(
        "--report-dir", type=Path, default=Path("build/unified-acceptance")
    )
    args = parser.parse_args(argv)
    interfaces, exports = args.interfaces.split(","), (
        args.exports.split(",") if args.exports else []
    )
    if (
        set(interfaces) - {"cli", "mcp", "react"}
        or "cli" not in interfaces
        or set(exports) - {"word", "excel", "pdf"}
    ):
        parser.error(
            "CLI is the prerequisite; supported interfaces cli,mcp,react and exports word,excel,pdf"
        )
    output = args.report_dir.resolve()
    if output.exists() and any(output.iterdir()):
        parser.error("Use a new empty report directory to prevent stale evidence")
    output.mkdir(parents=True, exist_ok=True)
    (output / "screenshots").mkdir()
    summary = {
        "local_synthetic_acceptance": "FAIL",
        "case_id": CASE,
        "chatgpt_connection": "BLOCKED_EXTERNAL",
        "claude_connection": "BLOCKED_EXTERNAL",
        "verified_live_inference": "NOT_RUN",
        "private_institutional_activation": "PENDING_AUTHORIZATION",
    }
    acceptance = None
    try:
        with tempfile.TemporaryDirectory(prefix="cdf-unified-") as temp:
            workspace = Path(temp)
            env = {
                **os.environ,
                "CDF_OPERATION_MODE": "PUBLIC_SYNTHETIC_DEMO",
                "CDF_CALIBRATION_MODE": "SYNTHETIC_LOCAL_DEV",
                "CDF_CALIBRATION_DB": str(workspace / "synthetic.sqlite"),
                "CDF_CALIBRATION_DEV_TOKEN": secrets.token_urlsafe(36),
                "CDF_AI_PROVIDER": "",
                "CDF_ENABLE_SUPERVISED_CLI": "0",
                "CDF_ENABLE_PROVIDER_API": "0",
            }
            # Never propagate optional provider credentials to deterministic subprocesses.
            for name in (
                "OPENAI_API_KEY",
                "ANTHROPIC_API_KEY",
                "CDF_PUBLIC_UI_DIRECTORY",
            ):
                env.pop(name, None)
            acceptance = Acceptance(output, env)
            acceptance.phase_cli(exports)
            matrix = {"cli": acceptance.canonical}
            if "mcp" in interfaces:
                port = free_port()
                with serving(
                    [
                        sys.executable,
                        "-m",
                        "cdf_gco_crms.interfaces.mcp_server",
                        "--port",
                        str(port),
                    ],
                    env,
                    f"http://127.0.0.1:{port}/health",
                    workspace / "mcp.log",
                ):
                    asyncio.run(acceptance.phase_mcp(f"http://127.0.0.1:{port}/mcp"))
                matrix.update(cli=acceptance.canonical, mcp=acceptance.mcp_fields)
            if "react" in interfaces:
                matrix["react"] = acceptance.phase_react(workspace)
                assert (
                    matrix["react"] == acceptance.canonical
                ), "CROSS_INTERFACE_MISMATCH"
            validated = acceptance.validate_exports(exports)
            write(
                output / "interface-matrix.json",
                {
                    "status": "PASS",
                    "fields": list(acceptance.canonical),
                    "interfaces": matrix,
                },
            )
            write(
                output / "provider-provenance.json",
                {
                    "deterministic_review": acceptance.review,
                    "explicit_mock_debate": acceptance.debate,
                    "trace": acceptance.trace,
                },
            )
            summary.update(
                local_synthetic_acceptance="PASS",
                source_hash=acceptance.source_hash,
                run_id=acceptance.canonical["run_id"],
                observation_ids=acceptance.canonical["observation_ids"],
                calibration_event=acceptance.calibration["event_hash"],
                artifacts=acceptance.artifacts,
                interfaces=interfaces,
                export_validation=validated,
                actual_provider_api_inference_count=0,
            )
            write(output / "command-results.json", acceptance.commands)
            write(output / "mcp-invocations.json", acceptance.invocations)
            write(output / "security-results.json", acceptance.security)
            write(output / "execution-summary.json", summary)
            manifest = scan_evidence(output, env["CDF_CALIBRATION_DEV_TOKEN"])
            write(
                output / "manifest.json",
                {
                    "schema": "CDF-UNIFIED-ACCEPTANCE-v1",
                    "synthetic": True,
                    "commit": os.getenv("CDF_ACCEPTANCE_COMMIT", "UNCOMMITTED_SNAPSHOT"),
                    "files": manifest,
                    "limitations": [
                        "NO_EXTERNAL_CLIENT_CONNECTION_PROVEN",
                        "NO_LIVE_INFERENCE_IN_THIS_RUN",
                        "HEURISTIC_LEAKAGE_SCAN_NOT_DLP_CERTIFICATION",
                    ],
                },
            )
        print(json.dumps(summary, ensure_ascii=False))
        return 0
    except Exception as exc:
        summary["error"] = str(exc)
        write(output / "execution-summary.json", summary)
        if acceptance:
            write(output / "command-results.json", acceptance.commands)
            write(output / "mcp-invocations.json", acceptance.invocations)
            write(output / "security-results.json", acceptance.security)
        raise


if __name__ == "__main__":
    raise SystemExit(main())
