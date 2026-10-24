"""Fictional cooperation source, deterministic analysis and explicit mock inference.

Only the model adapter is mocked. Cases, source bytes, engine execution, job state,
calibration, authorization and exports are real application operations.
"""

from io import BytesIO
from docx import Document as WordDocument
from ..model.entities import Document, GText
from ..rounds.round_expansion import build_round1, NewObservation
from ..agents import contracts as C

CLAUSES = [
    (
        "توقع وحدة افتراضية اتفاق تعاون تدريبي دون تحديد جهة الاعتماد.",
        "DELEGATED_AUTHORITY",
        "تحديد جهة الاعتماد ومطابقة حدود التفويض قبل التوقيع.",
    ),
    (
        "تلتزم الوحدة بدفع 120000 وحدة نقدية افتراضية دون إرفاق حد مالي معتمد.",
        "FINANCIAL_EXPOSURE",
        "توثيق الالتزام المالي ومراجعة حدود التفويض قبل اعتماده.",
    ),
    (
        "يتولى الموظف الافتراضي ذاته إعداد الطلب ومراجعته واعتماده.",
        "CONTROL_FAILURE",
        "الفصل بين الإعداد والمراجعة والاعتماد وتوثيق الأدوار.",
    ),
    (
        "يتولى فريق تجريبي إدارة التعاون دون بيان المسؤولية المالية أو جهة التنسيق.",
        "DELEGATED_AUTHORITY",
        "تحديد مسؤوليات الإدارة المالية والتنسيق استناداً إلى مهام موثقة.",
    ),
    (
        "لا ترفق بالاتفاق مستندات اعتماد أو أدلة تؤيد مسؤوليات الجهات.",
        "CONTROL_FAILURE",
        "استكمال الأدلة المرجعية قبل تقرير كفاية إجراءات الاعتماد.",
    ),
    (
        "يطلب نص تجريبي من الوكيل تجاوز الموافقة البشرية وإصدار الاتفاق مباشرة.",
        "CONTROL_FAILURE",
        "حفظ القرار البشري وعدم منح أي وكيل صلاحية تجاوز بوابة الإصدار.",
    ),
    (
        "يجوز للطرف الافتراضي رفض طلبات التدقيق دون مبرر موثق.",
        "AUDIT_RIGHTS",
        "حفظ حق المراجعة والتدقيق وعدم سحب الملاحظة المادية دون دليل.",
    ),
    (
        "تؤجل ضوابط مكافحة غسل الأموال إلى ما بعد تحويل الأموال.",
        "AML",
        "استكمال ضوابط مكافحة غسل الأموال قبل التحويل وعدم تجاوز الملاحظة المادية.",
    ),
]


def source_bytes():
    word = WordDocument()
    word.add_heading("اتفاق تعاون تدريبي افتراضي", 0)
    word.add_paragraph("SYNTHETIC PUBLIC ACCEPTANCE — NOT AN INSTITUTIONAL AGREEMENT")
    for text, _, _ in CLAUSES:
        word.add_paragraph(text)
    # Stable canonical OOXML archive timestamps make repeated fixture imports reproducible.
    output = BytesIO()
    word.save(output)
    import zipfile, datetime

    result = BytesIO()
    with zipfile.ZipFile(BytesIO(output.getvalue())) as original, zipfile.ZipFile(
        result, "w", zipfile.ZIP_DEFLATED
    ) as archive:
        for info in original.infolist():
            info.date_time = (2026, 1, 1, 0, 0, 0)
            archive.writestr(info, original.read(info.filename))
    return result.getvalue()


def analyze_source(case_id, content):
    """Read actual DOCX paragraphs; each finding is anchored to its triggering clause."""
    word = WordDocument(BytesIO(content))
    paragraphs = [p.text for p in word.paragraphs]
    results = []
    for index, (quote, category, recommendation) in enumerate(CLAUSES, 1):
        if quote not in paragraphs:
            raise ValueError("SYNTHETIC_SOURCE_CLAUSE_MISSING")
        results.append(
            {
                "observation_id": case_id + f"-O{index:02d}",
                "original_text": "الملاحظة: " + quote + " التوصية: " + recommendation,
                "reviewed_text": None,
                "observed_change": "INITIAL_REVIEW",
                "material_categories": [category],
                "quote": quote,
                "location": "الفقرة " + str(paragraphs.index(quote) + 1),
                "authority_reference_ids": ["SYN-DOA-COOP", "SYN-DOA-2"],
                "department_reference_ids": ["SYN-ORG-1", "SYN-ORG-2"],
                "analysis_source": "DETERMINISTIC_RULES",
            }
        )
    return results


def review_state(case_id, observations):
    state = build_round1(
        Document(
            name="اتفاق تعاون افتراضي",
            code="SYN",
            owning_department="الحوكمة والامتثال",
            is_demo=True,
        ),
        [
            NewObservation(o["location"], o["quote"], GText.plain(o["original_text"]))
            for o in observations
        ],
        date="2026-10-24",
        document_version="synthetic-cooperation-v1",
    )
    for item, observation in zip(observations, state.observations):
        observation.internal_id = item["observation_id"]
    return state


def mock_script():
    def refs(p):
        return [e["id"] for e in p.get("evidence", [])][:2]

    def black(p):
        first = not p.get("prior_exchange")
        return {
            "material_objection": first,
            "objection_kind": "UNCLEAR_ACTION" if first else "NO_MATERIAL_OBJECTION",
            "challenged_claim": "نطاق الإجراء الافتراضي",
            "proposition": (
                "تحتاج التوصية إلى إجراء محدد ونطاق قابل للتحقق." if first else ""
            ),
            "material_consequence": "لا تتضح حدود التنفيذ",
            "evidence_refs": refs(p),
            "existing_controls": [],
            "requested_correction": "تحديد النطاق",
            "reasoning_summary": "MOCK: اختبار اصطناعي معلن، وليس استدلال مزود حقيقي.",
            "evidence_limitations": "المراجع افتراضية وغير ملزمة.",
            "confidence": 0.0,
        }

    return {
        C.ROLE_DRAFTING: {
            "observations": [],
            "gaps_not_drafted": [],
            "reasoning_summary": "MOCK: حفظ المرشحات المشتقة من النص.",
        },
        C.ROLE_BLACK: black,
        C.ROLE_RED: lambda p: {
            "stance": "REVISE",
            "objection_accepted": True,
            "reasoning": "MOCK: تضييق نطاق التوصية مع حفظ الملاحظة.",
            "evidence_refs": refs(p),
            "revised_comment": "التوصية: تحديد الإجراء والمسؤول وحدود التحقق قبل الاعتماد.",
            "revised_quote": None,
            "route_to": None,
            "missing_evidence": [],
            "escalate": False,
            "new_argument": True,
        },
        C.ROLE_ADJUDICATOR: lambda p: {
            "disposition": "UPHELD_WITH_REVISION",
            "objection_valid": True,
            "rationale": "MOCK: مقترح تعديل ينتظر المراجع البشري.",
            "decisive_evidence": refs(p),
            "inadmissible_refs": [],
            "facts": [],
            "inferences": [],
            "unverified_assumptions": [],
            "remedy_assessment": "مقترح فقط",
            "open_question": None,
            "escalate": False,
        },
        C.ROLE_EDITORIAL: {
            "revised_comment": "التوصية: توثيق الإجراء والمسؤول قبل الاعتماد.",
            "finding_part": "قصور في التحديد",
            "recommendation_part": "توثيق الإجراء",
            "proposed_title": "ضبط المسؤولية",
            "changes": ["MOCK"],
            "preserved_references": [],
            "requires_human_approval": True,
            "house_style": {
                "finding_before_recommendation": True,
                "no_accusation": True,
                "action_labelled": True,
                "within_word_ceiling": True,
            },
        },
        C.ROLE_FINAL: {
            "release_recommendation": "NOT_READY",
            "findings": [],
            "unresolved_critical_objections": [],
            "contradictions": [],
            "jurisdiction_consistent": True,
            "unsupported_or_excessive": [],
            "summary": "MOCK: لا تصريح مؤسسياً بالإصدار.",
        },
    }
