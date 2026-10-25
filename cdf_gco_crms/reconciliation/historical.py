"""Historical preservation and return-cycle intake (10.6, SC-38, SC-40, SC-59, SC-60, G-DASH-02/11).

Cells are compared by (Internal_ID, column key), never by coordinates, so a
re-sorted or re-rendered sheet is compared correctly.
"""
from __future__ import annotations

import copy
from dataclasses import dataclass, field

from openpyxl.cell.rich_text import CellRichText, TextBlock

from ..config import round_schema as S
from ..model.entities import GText, ReviewState, plain_text

FIXED_KEYS = ("A", "C", "D", "E")
GCO_ROLE_KEYS = (S.ROLE_CLOSURE, S.ROLE_FEEDBACK)
OWNER_ROLE_KEYS = (S.ROLE_RESPONSE, S.ROLE_IMPL)


def _signature(value) -> tuple:
    """(plain text, bold-run structure) of a cell value."""
    if isinstance(value, GText):
        return value.text, tuple((s.text, s.bold) for s in value.segments)
    if isinstance(value, CellRichText):
        runs = []
        for p in value:
            if isinstance(p, TextBlock):
                runs.append((p.text, bool(p.font.b)))
            else:
                runs.append((str(p), False))
        g = GText(runs)
        return g.text, tuple((s.text, s.bold) for s in g.segments)
    t = plain_text(value)
    return t, ((t, False),) if t else ()


def snapshot(state: ReviewState) -> dict[tuple[str, str], tuple]:
    snap: dict[tuple[str, str], tuple] = {}
    for o in state.observations:
        snap[(o.internal_id, "A")] = (o.reader_no, ())
        snap[(o.internal_id, "C")] = _signature(o.location)
        snap[(o.internal_id, "D")] = _signature(o.quote)
        snap[(o.internal_id, "E")] = _signature(o.comment)
        for r, c in o.cells.items():
            for role in S.ROUNDN_ROLES:
                if r == 1 and role == S.ROLE_FEEDBACK:
                    continue
                snap[(o.internal_id, f"R{r}.{role}")] = _signature(getattr(c, role))
    return snap


@dataclass
class Difference:
    internal_id: str
    key: str
    before: str
    after: str
    kind: str          # TEXT | FORMAT | MISSING | ADDED


def compare(before: dict, after: dict, *, rounds: set[int] | None = None,
            include_fixed: bool = True, ignore_format: bool = False,
            owner_rounds: set[int] | None = None) -> list[Difference]:
    """Differences for the fixed block and the given rounds. ``owner_rounds`` limits the owner columns
    compared (the latest delivered round's owner cells legitimately change when the owner answers)."""
    diffs: list[Difference] = []
    def in_scope(key: str) -> bool:
        if key in FIXED_KEYS:
            return include_fixed
        r = int(key[1:].split(".")[0])
        if owner_rounds is not None and key.endswith((".response", ".impl")) and r not in owner_rounds:
            return False
        return rounds is None or r in rounds
    for k, v in before.items():
        if not in_scope(k[1]):
            continue
        if k not in after:
            diffs.append(Difference(k[0], k[1], v[0], "", "MISSING"))
            continue
        w = after[k]
        if v[0] != w[0]:
            diffs.append(Difference(k[0], k[1], v[0], w[0], "TEXT"))
        elif not ignore_format and v[1] != w[1]:
            diffs.append(Difference(k[0], k[1], v[0], w[0], "FORMAT"))
    return diffs


@dataclass
class IntakeReport:
    gco_cell_edits: list[Difference] = field(default_factory=list)      # restored (SC-38, SC-59)
    late_owner_edits: list[Difference] = field(default_factory=list)    # kept, logged (10.6)
    format_losses: list[Difference] = field(default_factory=list)       # restored (SC-60)
    missing_rows: list[str] = field(default_factory=list)               # restored, no response (SC-40)
    owner_inputs: int = 0
    notes: list[str] = field(default_factory=list)


def _ws_norm(s: str) -> str:
    return " ".join(s.split())


def merge_returned(delivered: ReviewState, returned: ReviewState) -> tuple[ReviewState, IntakeReport]:
    """GCO cells from the delivered copy, owner cells from the returned copy (G-VER-4)."""
    rep = IntakeReport()
    out = copy.deepcopy(delivered)
    ret = returned.by_id()
    R = delivered.current_round
    for o in out.observations:
        r_o = ret.get(o.internal_id)
        if r_o is None:
            rep.missing_rows.append(o.internal_id)
            continue
        for key, mine, theirs in (("C", o.location, r_o.location), ("D", o.quote, r_o.quote),
                                  ("E", o.comment, r_o.comment)):
            a, b = _signature(mine), _signature(theirs)
            if a[0] != b[0]:
                kind = "FORMAT" if _ws_norm(a[0]) == _ws_norm(b[0]) else "TEXT"
                d = Difference(o.internal_id, key, a[0], b[0], kind)
                (rep.format_losses if kind == "FORMAT" else rep.gco_cell_edits).append(d)
        for r, cells in o.cells.items():
            rc = r_o.cells.get(r)
            if rc is None:
                continue
            for role in GCO_ROLE_KEYS:
                if r == 1 and role == S.ROLE_FEEDBACK:
                    continue
                a, b = _signature(getattr(cells, role)), _signature(getattr(rc, role))
                if a[0] != b[0]:
                    kind = "FORMAT" if _ws_norm(a[0]) == _ws_norm(b[0]) else "TEXT"
                    d = Difference(o.internal_id, f"R{r}.{role}", a[0], b[0], kind)
                    (rep.format_losses if kind == "FORMAT" else rep.gco_cell_edits).append(d)
            for role in OWNER_ROLE_KEYS:
                a, b = _signature(getattr(cells, role)), _signature(getattr(rc, role))
                if r == R:
                    if b[0] not in ("", S.PH_TEXT, S.PH_STATUS) and not b[0].startswith(S.CLOSED_GUIDE_PREFIX) \
                            and b[0] != S.IMPL_PRIOR_CLOSED:
                        rep.owner_inputs += 1
                elif a[0] != b[0]:
                    rep.late_owner_edits.append(Difference(o.internal_id, f"R{r}.{role}", a[0], b[0], "TEXT"))
                setattr(cells, role, getattr(rc, role))           # owner text kept as returned
    extra = set(ret) - set(out.by_id())
    if extra:
        rep.notes.append(f"rows with unknown Internal_ID in the returned workbook: {sorted(extra)}")
    return out, rep


@dataclass
class ViewSheetEdit:
    sheet: str
    coordinate: str
    delivered: object
    returned: object


def view_sheet_edits(delivered_path, returned_path) -> tuple[list[ViewSheetEdit], list[str]]:
    """Differences on the dashboard and the helper sheet between the delivered and the returned workbook (10.6,
    15A.3(e), SC-61 to SC-64). Selector cells (the only unlocked cells, SC-62) are interface preferences and are
    skipped. Nothing found here is input: text typed here is logged as Channel X input and both sheets are rebuilt.
    Returns the edits and the sheets missing from the returned workbook."""
    from openpyxl import load_workbook
    from ..config import workbook_settings as W
    a = load_workbook(delivered_path)
    b = load_workbook(returned_path)
    edits, missing = [], []
    for name in (W.DASHBOARD_SHEET, W.HELPER_SHEET):
        if name not in a.sheetnames:
            continue
        if name not in b.sheetnames:
            missing.append(name)
            continue
        wa, wb_ = a[name], b[name]
        rows = max(wa.max_row, wb_.max_row)
        cols = max(wa.max_column, wb_.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                ca, cb = wa.cell(row=r, column=c), wb_.cell(row=r, column=c)
                if name == W.DASHBOARD_SHEET and ca.protection is not None and ca.protection.locked is False:
                    continue
                va, vb = ca.value, cb.value
                if str(va if va is not None else "") != str(vb if vb is not None else ""):
                    edits.append(ViewSheetEdit(name, ca.coordinate, va, vb))
    return edits, missing
